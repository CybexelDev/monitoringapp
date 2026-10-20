from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib import messages
from django.contrib.auth.models import User  
from .models import *
from datetime import datetime
from django.views.decorators.csrf import csrf_exempt
from datetime import date
import openpyxl
from django.views.decorators.http import (
    require_GET,
    require_POST,
    require_http_methods,
)
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.contrib.auth.hashers import check_password, make_password
import random
from django.core.mail import send_mail
from datetime import time, timedelta
from django.utils.dateparse import parse_datetime,parse_time
from django.utils.timezone import localtime, make_aware, is_naive
from django.views.decorators.cache import never_cache
from django.utils.cache import add_never_cache_headers
from .models import AnnouncementRecipient
from django.utils.dateparse import parse_time
from django.http import JsonResponse
from django.db import transaction
from django.db.models.functions import Lower, Trim
from openpyxl.utils import get_column_letter
from django.utils.dateparse import parse_time
from django.db import IntegrityError
from django.db.models import Q
from chat.models import Conversation
import calendar


@never_cache
def index(request):
    return render(request,'index.html')




def admin_logout(request):
    logout(request)
    request.session.flush()
    return redirect("index")

@never_cache
@require_POST
def delete_department(request, pk):
    from django.http import HttpResponseForbidden, JsonResponse
    from django.db.models.deletion import ProtectedError, RestrictedError

    if _admin_notepad_owner(request) is None:
        return HttpResponseForbidden("Admin access required.")
    ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    item = get_object_or_404(Department, pk=pk)
    try:
        item.delete()
    except (ProtectedError, RestrictedError):
        message = "This department is still in use and cannot be deleted."
        if ajax:
            return JsonResponse({"ok": False, "message": message}, status=409)
        messages.error(request, message)
        return redirect("admin_dashboard")
    message = "Department deleted successfully."
    if ajax:
        return JsonResponse({"ok": True, "message": message})
    messages.success(request, message)
    return redirect("admin_dashboard")


@never_cache
@require_POST
def delete_team(request, pk):
    from django.http import HttpResponseForbidden, JsonResponse
    from django.db.models.deletion import ProtectedError, RestrictedError

    if _admin_notepad_owner(request) is None:
        return HttpResponseForbidden("Admin access required.")
    ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    item = get_object_or_404(Team, pk=pk)
    try:
        item.delete()
    except (ProtectedError, RestrictedError):
        message = "This team is still in use and cannot be deleted."
        if ajax:
            return JsonResponse({"ok": False, "message": message}, status=409)
        messages.error(request, message)
        return redirect("admin_dashboard")
    message = "Team deleted successfully."
    if ajax:
        return JsonResponse({"ok": True, "message": message})
    messages.success(request, message)
    return redirect("admin_dashboard")


@never_cache
@require_GET
def admin_reports(request):
    from django.http import HttpResponseForbidden
    if _admin_notepad_owner(request) is None:
        return HttpResponseForbidden("Admin access required.")
    from django.utils.dateparse import parse_date

    report_tab = request.GET.get("tab", "leads")
    if report_tab not in {"leads", "members"}:
        report_tab = "leads"
    role = "Team Lead" if report_tab == "leads" else "Team Member"
    show_all = request.GET.get("all") == "1"
    export = request.GET.get("export") == "1"
    filter_date = None
    raw_date = request.GET.get("date", "").strip()
    if raw_date and raw_date != "None":
        try:
            selected_date = parse_date(raw_date)
        except ValueError:
            selected_date = None
        if selected_date:
            filter_date = selected_date.isoformat()
        else:
            messages.error(request, "Choose a valid report date.")

    morning_reports = MorningReport.objects.filter(
        user__job_Position__iexact=role
    ).select_related("user").order_by("-created_at", "-id")
    evening_reports = EveningReport.objects.filter(
        user__job_Position__iexact=role
    ).select_related("user").order_by("-created_at", "-id")

    # Filter a local calendar day with a timestamp range. This avoids
    # applying a date-conversion function to every stored timestamp.
    day = parse_date(filter_date) if filter_date else (None if show_all else timezone.localdate())
    if day is not None:
        from datetime import datetime as report_datetime, time as report_time, timedelta as report_delta
        from django.conf import settings as report_settings

        start = report_datetime.combine(day, report_time.min)
        end = report_datetime.combine(day + report_delta(days=1), report_time.min)
        if report_settings.USE_TZ:
            tz = timezone.get_current_timezone()
            start = timezone.make_aware(start, tz)
            end = timezone.make_aware(end, tz)
        morning_reports = morning_reports.filter(created_at__gte=start, created_at__lt=end)
        evening_reports = evening_reports.filter(created_at__gte=start, created_at__lt=end)

    # Export filtered reports to Excel
    if export:
        wb = Workbook()
        ws = wb.active
        ws.title = "Team Reports"

        headers = ["Report Type", "Date", "Time", "Name", "Team", "Department", "Report", "Status"]
        ws.append(headers)

        def add_report_to_sheet(report_list, report_type):
            for report in report_list:
                local_dt = timezone.localtime(report.created_at)
                ws.append([
                    report_type,
                    local_dt.strftime("%Y-%m-%d"),
                    local_dt.strftime("%I:%M %p"),
                    getattr(report.user, "name", report.user.username),
                    report.team,
                    report.department,
                    report.report_text,
                    getattr(report, "status", ""),
                ])

        add_report_to_sheet(morning_reports, "Morning")
        add_report_to_sheet(evening_reports, "Evening")

        for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        column_widths = [15, 15, 12, 20, 15, 20, 50, 15]
        for i, width in enumerate(column_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = 'attachment; filename="team_reports.xlsx"'
        wb.save(response)
        return response

    context = {
        'report_tab': report_tab,
        'morning_reports': morning_reports,
        'evening_reports': evening_reports,
        'show_all': show_all,
        'filter_date': filter_date,
    }
    return render(request, 'admin_panel/admin_reports.html', context)


@never_cache
@require_http_methods(["GET", "POST"])
def admin_profile(request):
    from django.http import HttpResponseForbidden
    if _admin_notepad_owner(request) is None:
        return HttpResponseForbidden("Admin access required.")
    user = _admin_chat_user(request)
    if user is None:
        return HttpResponseForbidden("A Management employee login is required for this profile.")

    if request.method == "POST":
        action = request.POST.get("action")

        # AJAX Profile Image Update
        if action == "edit_profile" and request.FILES.get("profile_image"):
            user.profile_image = request.FILES["profile_image"]
            user.save()
            return JsonResponse({
                "success": True,
                "image_url": user.profile_image.url
            })

        # Normal form update for other details
        elif action == "edit_profile":
            user.name = request.POST.get("name")
            user.email = request.POST.get("email")
            user.phone = request.POST.get("phone")
            user.work_location = request.POST.get("work_location", user.work_location)
            user.save()
            messages.success(request, "Profile updated successfully!")
            return redirect("admin_profile")

    # Render page
    response = render(request, "admin_panel/admin_profile.html", {"user": user})
    response["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response["Pragma"] = "no-cache"
    response["Expires"] = "0"
    return response

from django.views.decorators.http import require_http_methods, require_POST

def _admin_managed_users(request):
    """Admin portal access; only Team Leads and Team Members are managed here."""
    from django.db.models import Q
    allowed = request.user.is_authenticated and request.user.is_superuser
    if not allowed:
        session_user = User.objects.filter(pk=request.session.get("user_id")).first()
        allowed = bool(session_user and session_user.status == "active" and
                       session_user.job_Position.strip().lower() == "management")
    if not allowed:
        return None
    return User.objects.filter(Q(job_Position__iexact="Team Lead") |
                               Q(job_Position__iexact="Team Member"))

from django.contrib import messages
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Q
from django.http import JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET, require_POST, require_http_methods
from .models import Announcement, AnnouncementRecipient, User as MonitoringUser

def _announcement_role_user(request, role):
    if request.session.get("position") != role:
        return None
    user_id = request.session.get("user_id")
    if not user_id:
        return None
    return MonitoringUser.objects.filter(
        pk=user_id,
        job_Position__iexact={"team_lead": "Team Lead", "team_member": "Team Member"}[role],
    ).first()


def _announcement_inbox(user):
    return AnnouncementRecipient.objects.filter(
        recipient=user, announcement__is_active=True,
    ).filter(
        Q(announcement__is_admin_announcement=True)
        | Q(announcement__created_by__job_Position__iexact="Team Lead")
    ).select_related("announcement", "announcement__created_by", "announcement__admin_created_by")


@never_cache
@require_POST
def teamlead_admin_announcement_seen(request, pk):
    user = _announcement_role_user(request, "team_lead")
    if user is None:
        return JsonResponse({"ok": False, "message": "Please sign in again."}, status=403)
    rows = _announcement_inbox(user).filter(announcement__is_admin_announcement=True)
    row = get_object_or_404(rows, announcement_id=pk)
    rows.filter(pk=row.pk, read_at__isnull=True).update(read_at=timezone.now())
    return JsonResponse({"ok": True})

from .models import Designation
from django.db import transaction, IntegrityError
from django.http import JsonResponse


def _department_designation_options():
    options = {}
    for item in Designation.objects.order_by("name", "id"):
        options.setdefault(str(item.department_id), []).append(item.name)
    return options


def _valid_department_designation(department_id, designation):
    try:
        department_id = int(department_id)
    except (ValueError, TypeError):
        return False
    return Designation.objects.filter(
        department_id=department_id, name=designation
    ).exists()


@never_cache
@require_http_methods(["GET", "POST"])
def admin_dashboard(request):
    if _admin_notepad_owner(request) is None:
        return redirect("admin_login")
    ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"

    def finish(message, ok=True):
        if ajax:
            return JsonResponse({"ok": ok, "message": message}, status=200 if ok else 400)
        (messages.success if ok else messages.error)(request, message)
        return redirect("admin_dashboard")

    if request.method == "POST":
        names = [value.strip() for value in request.POST.getlist("designation_names") if value.strip()]
        # Discard duplicate names within this form, ignoring case.
        names = list({name.casefold(): name for name in names}.values())
        if any(len(name) > 100 for name in names):
            return finish("Designation names must be 100 characters or fewer.", False)
        try:
            with transaction.atomic():
                if "edit_department" in request.POST:
                    try:
                        department_id = int(request.POST.get("department_id", ""))
                    except (ValueError, TypeError):
                        return finish("Choose a valid department.", False)
                    department = Department.objects.select_for_update().filter(pk=department_id).first()
                    if department is None:
                        return finish("Department not found.", False)
                    department_name = request.POST.get("department_name", "").strip()
                    if not department_name or len(department_name) > 100:
                        return finish("Enter a department name with up to 100 characters.", False)
                    if Department.objects.exclude(pk=department.pk).filter(name__iexact=department_name).exists():
                        return finish("This department name already exists.", False)
                    raw_names = request.POST.getlist("designation_names")
                    raw_ids = request.POST.getlist("designation_ids")
                    if len(raw_names) != len(raw_ids) or not raw_names:
                        return finish("Add at least one designation.", False)
                    existing = {str(item.pk): item for item in department.designations.select_for_update()}
                    changes = []; seen_names = set(); seen_ids = set()
                    for item_id, raw_name in zip(raw_ids, raw_names):
                        name = raw_name.strip()
                        if not name or len(name) > 100 or name.casefold() in seen_names:
                            return finish("Use unique designation names, each up to 100 characters.", False)
                        if item_id and (item_id not in existing or item_id in seen_ids):
                            return finish("Invalid designation selection.", False)
                        seen_names.add(name.casefold()); seen_ids.add(item_id)
                        changes.append((item_id, name))
                    removed = [item for key, item in existing.items() if key not in seen_ids]
                    for item in removed:
                        if User.objects.filter(department=department, designation__iexact=item.name).exists():
                            return finish(f"{item.name} is assigned to employees. Reassign them before removing it.", False)
                    # Temporarily rename retained rows to permit swapping two designation names.
                    from uuid import uuid4
                    for item_id, name in changes:
                        if item_id:
                            Designation.objects.filter(pk=existing[item_id].pk).update(name="tmp-" + uuid4().hex)
                    for item in removed:
                        item.delete()
                    # Match employees against original values before renaming to avoid swap collisions.
                    employee_changes = []
                    for item_id, name in changes:
                        if item_id:
                            item = existing[item_id]
                            employee_ids = list(User.objects.filter(department=department, designation=item.name).values_list("pk", flat=True))
                            employee_changes.append((employee_ids, name))
                            item.name = name; item.save(update_fields=["name"])
                        else:
                            Designation.objects.create(department=department, name=name)
                    for employee_ids, name in employee_changes:
                        User.objects.filter(pk__in=employee_ids).update(designation=name)
                    department.name = department_name; department.save(update_fields=["name"])
                    return finish("Department updated successfully.")
                if "add_department" in request.POST:
                    name = request.POST.get("department_name", "").strip()
                    if not name or len(name) > 100:
                        return finish("Enter a department name with up to 100 characters.", False)
                    if not names:
                        return finish("Add at least one designation.", False)
                    if Department.objects.filter(name__iexact=name).exists():
                        return finish("This department already exists. Add designations to its existing row.", False)
                    department = Department.objects.create(name=name)
                    for designation in names:
                        Designation.objects.create(department=department, name=designation)
                    return finish("Department and designations added successfully.")
                if "add_designation" in request.POST:
                    try:
                        department_id = int(request.POST.get("department_id", ""))
                    except (ValueError, TypeError):
                        return finish("Choose a valid department.", False)
                    department = Department.objects.filter(pk=department_id).first()
                    if department is None or not names:
                        return finish("Choose a department and enter a designation.", False)
                    for name in names:
                        if not Designation.objects.filter(department=department, name__iexact=name).exists():
                            Designation.objects.create(department=department, name=name)
                    return finish("Designations saved successfully.")
                if "add_team" in request.POST:
                    name = request.POST.get("team_name", "").strip()
                    if not name or len(name) > 100:
                        return finish("Enter a team name with up to 100 characters.", False)
                    if Team.objects.filter(name__iexact=name).exists():
                        return finish("This team already exists.", False)
                    Team.objects.create(name=name)
                    return finish("Team added successfully.")
                return finish("Invalid dashboard action.", False)
        except IntegrityError:
            return finish("This department, team or designation already exists.", False)

    from django.core.paginator import Paginator
    from django.db.models import Count, Q

    managed = _admin_managed_users(request)
    if managed is None:
        return redirect("admin_login")
    all_teams = list(Team.objects.order_by("name"))
    department_ids = set(Department.objects.values_list("pk", flat=True))
    team_ids = {team.pk for team in all_teams}

    def selected_id(name, allowed):
        try:
            value = int(request.GET.get(name, ""))
        except (ValueError, TypeError):
            return None
        return value if value in allowed else None

    department_id = selected_id("department", department_ids)
    team_id = selected_id("team", team_ids)
    department_teams = {}
    for dept_id, assigned_team in managed.exclude(department_id=None).exclude(team_id=None).values_list("department_id", "team_id").distinct():
        department_teams.setdefault(str(dept_id), []).append(assigned_team)
    if department_id and team_id not in department_teams.get(str(department_id), []):
        team_id = None
    employees = managed
    if department_id:
        employees = employees.filter(department_id=department_id)
    if team_id:
        employees = employees.filter(team_id=team_id)
    employee_counts = employees.aggregate(
        total=Count("pk"),
        leads=Count("pk", filter=Q(job_Position__iexact="Team Lead")),
        members=Count("pk", filter=Q(job_Position__iexact="Team Member")),
        active=Count("pk", filter=Q(status__iexact="active")),
        inactive=Count("pk", filter=Q(status__iexact="inactive")),
    )
    employee_page = Paginator(employees.select_related("department", "team").order_by("name", "pk"), 25).get_page(request.GET.get("page"))

    departments = list(Department.objects.prefetch_related("designations").order_by("name"))
    return render(request, "admin_panel/admin_dashboard.html", {
        "departments": departments,
        "department_editor_items": {str(dept.pk): {"name": dept.name, "designations": [{"id": item.pk, "name": item.name} for item in dept.designations.all()]} for dept in departments},
        "teams": all_teams,
        "filter_department": str(department_id) if department_id else "",
        "filter_team": str(team_id) if team_id else "",
        "department_team_options": department_teams,
        "employee_page": employee_page,
        "employee_counts": employee_counts,
    })


@never_cache
@require_http_methods(["GET", "POST"])
def admin_usermanagement(request):
    from django.http import HttpResponseForbidden
    managed_users = _admin_managed_users(request)
    if managed_users is None:
        return HttpResponseForbidden("Admin access required.")

    if request.method == "GET" and "designation_department" in request.GET:
        from django.http import JsonResponse
        from .models import Designation
        try:
            department_id = int(request.GET.get("designation_department", ""))
        except (TypeError, ValueError):
            return JsonResponse({"ok": False, "message": "Invalid department."}, status=400)
        if not Department.objects.filter(pk=department_id).exists():
            return JsonResponse({"ok": False, "message": "Department not found."}, status=404)
        names = list(Designation.objects.filter(
            department_id=department_id
        ).order_by("name").values_list("name", flat=True))
        return JsonResponse({"ok": True, "items": names})

    users = managed_users.select_related("department", "team").order_by("-id")

    # Hide currently logged-in management/admin user
    current_user_id = request.session.get("user_id")

    if current_user_id:
        users = users.exclude(id=current_user_id)
    departments = Department.objects.all()
    teams = Team.objects.all()
    job_positions = (
        managed_users.exclude(job_Position__isnull=True)
        .exclude(job_Position__exact="")
        .values_list('job_Position', flat=True)
        .distinct()
        .order_by('job_Position')
    )

    if request.method == "POST":
        if request.POST.get("job_Position", "") not in {"Team Lead", "Team Member"}:
            messages.error(request, "Choose Team Lead or Team Member.")
            return redirect("admin_usermanagement")
        date_value = request.POST.get("joining_date", "")
        if date_value:
            try:
                datetime.strptime(date_value, "%Y-%m-%d")
            except ValueError:
                messages.error(request, "Enter a valid joining date.")
                return redirect("admin_usermanagement")
        user_id = request.POST.get("id")  # for edit (None for new user)
        name = request.POST.get("name")
        employee_id = request.POST.get("employee_id")
        email = request.POST.get("email")
        phone = request.POST.get("phone")
        department_id = request.POST.get("department")
        team_id = request.POST.get("team")
        job_Position = request.POST.get("job_Position")
        designation = request.POST.get("designation")
        work_location = request.POST.get("work_location")
        username = request.POST.get("username")
        password = request.POST.get("password")
        status = request.POST.get("status")
        image = request.FILES.get("profile_image")
        joining_date = request.POST.get("joining_date")

        designation = (designation or "").strip()
        if not _valid_department_designation(department_id, designation):
            messages.error(request, "Choose a designation belonging to the selected department.")
            return redirect("admin_usermanagement")

        # Foreign key fetch
        department = Department.objects.get(id=department_id) if department_id else None
        team = Team.objects.get(id=team_id) if team_id else None

        # Helper for uniqueness validation
        def check_unique(field_name, value):
            qs = User.objects.filter(**{field_name: value})
            if user_id:
                qs = qs.exclude(id=user_id)
            return qs.exists()

        # ========== EDIT EXISTING USER ==========
        if user_id:
            user = get_object_or_404(managed_users, id=user_id)

            # Uniqueness checks
            if check_unique("employee_id", employee_id):
                messages.error(request, "⚠️ Employee ID already exists for another user.")
                return redirect("admin_usermanagement")
            if check_unique("username", username):
                messages.error(request, "⚠️ Username already exists for another user.")
                return redirect("admin_usermanagement")
            if check_unique("email", email):
                messages.error(request, "⚠️ Email already exists for another user.")
                return redirect("admin_usermanagement")

            # Update details
            user.name = name
            user.employee_id = employee_id
            user.email = email
            user.phone = phone
            user.department = department
            user.team = team
            user.job_Position = job_Position
            user.designation = designation
            user.work_location = work_location
            user.username = username
            user.status = status

            if password:
                user.password = make_password(password)
            if image:
                user.profile_image = image
            if joining_date:
                try:
                    user.joining_date = datetime.strptime(joining_date, "%Y-%m-%d").date()
                except ValueError:
                    messages.error(request, "❌ Invalid date format! Use YYYY-MM-DD.")
                    return redirect("admin_usermanagement")

            try:
                user.save()
                messages.success(request, "✅ User updated successfully!")
            except IntegrityError as e:
                messages.error(request, f"⚠️ Cannot update user. Unique field conflict ({str(e)})")
                return redirect("admin_usermanagement")

        # ========== CREATE NEW USER ==========
        else:
            if check_unique("employee_id", employee_id):
                messages.error(request, "⚠️ Employee ID already exists. Please choose another one.")
                return redirect("admin_usermanagement")
            if check_unique("username", username):
                messages.error(request, "⚠️ Username already exists. Please choose another one.")
                return redirect("admin_usermanagement")
            if check_unique("email", email):
                messages.error(request, "⚠️ Email already exists. Please choose another one.")
                return redirect("admin_usermanagement")

            try:
                User.objects.create(
                    name=name,
                    employee_id=employee_id,
                    email=email,
                    phone=phone,
                    department=department,
                    team=team,
                    job_Position=job_Position,
                    designation=designation,
                    work_location=work_location,
                    username=username,
                    password=make_password(password) if password else "",
                    status=status,
                    profile_image=image,
                    joining_date=datetime.strptime(joining_date, "%Y-%m-%d").date() if joining_date else None,
                )
                messages.success(request, "✅ User created successfully!")
            except IntegrityError as e:
                messages.error(request, f"⚠️ Cannot create user. Unique field conflict ({str(e)})")
                return redirect("admin_usermanagement")

        return redirect("admin_usermanagement")

    # =======================================================
    # ===============   GET REQUEST (FILTERS)   =============
    # =======================================================

    search = request.GET.get("search", "").strip()
    department_filter = request.GET.get("department", "")
    team_filter = request.GET.get("team", "")
    position_filter = request.GET.get("position", "")

    if search:
        users = users.filter(
            Q(name__icontains=search) |
            Q(email__icontains=search) |
            Q(employee_id__icontains=search)
        )

    if department_filter:
        users = users.filter(department__name=department_filter)

    if team_filter:
        users = users.filter(team__name=team_filter)

    if position_filter:
        users = users.filter(job_Position=position_filter)

    context = {
        "users": users,
        "departments": departments,
        "teams": teams,
        "job_positions": job_positions,
        "designation_options": _department_designation_options(),
    }
    return render(request, "admin_panel/admin_usermanagement.html", context)


from django.views.decorators.http import require_POST

@never_cache
@require_POST
def edit_user(request):
    from django.http import HttpResponseForbidden, JsonResponse

    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"

    def finish(message, ok=False, status=400):
        if is_ajax:
            payload = {"ok": ok, "message": message}
            if ok:
                payload["id"] = user.pk
            return JsonResponse(payload, status=200 if ok else status)
        if ok:
            messages.success(request, message)
        else:
            messages.error(request, message)
        return redirect("admin_usermanagement")
    managed_users = _admin_managed_users(request)
    if managed_users is None:
        if is_ajax:
            return JsonResponse({"ok": False, "message": "Admin access required."}, status=403)
        return HttpResponseForbidden("Admin access required.")

    if request.method == "POST":
        if request.POST.get("edit_job_position", "") not in {"Team Lead", "Team Member"}:
            return finish("Choose Team Lead or Team Member.")
        date_value = request.POST.get("edit_joining_date", "")
        if date_value:
            try:
                datetime.strptime(date_value, "%Y-%m-%d")
            except ValueError:
                return finish("Enter a valid joining date.")
        user_id = request.POST.get("id")
        user = get_object_or_404(managed_users, id=user_id)

        # Get form values
        employee_id = request.POST.get("edit_emp_id")
        username = request.POST.get("edit_username")
        email = request.POST.get("edit_email")
        name = request.POST.get("edit_name")
        phone = request.POST.get("edit_phone")
        job_position = request.POST.get("edit_job_position")
        designation = request.POST.get("edit_designation")
        work_location = request.POST.get("edit_work_location")
        status = request.POST.get("edit_status")
        password = request.POST.get("edit_password")
        joining_date = request.POST.get("edit_joining_date")
        dept_id = request.POST.get("edit_department")
        team_id = request.POST.get("edit_team")

        designation = (designation or "").strip()
        if not _valid_department_designation(dept_id, designation):
            return finish("Choose a designation belonging to the selected department.")

        # ========== Check uniqueness before saving ==========
        if User.objects.filter(employee_id=employee_id).exclude(id=user.id).exists():
            return finish("⚠️ Employee ID already exists for another user.")

        if User.objects.filter(username=username).exclude(id=user.id).exists():
            return finish("⚠️ Username already exists for another user.")

        if User.objects.filter(email=email).exclude(id=user.id).exists():
            return finish("⚠️ Email already exists for another user.")

        # ========== Update fields ==========
        user.employee_id = employee_id
        user.username = username
        user.email = email
        user.name = name
        user.phone = phone
        user.job_Position = job_position
        user.designation = designation
        user.work_location = work_location
        user.status = status

        if password:
            user.password = make_password(password)

        # Department & Team
        try:
            user.department = Department.objects.get(id=dept_id) if dept_id else None
            user.team = Team.objects.get(id=team_id) if team_id else None
        except (Department.DoesNotExist, Team.DoesNotExist, ValueError, TypeError):
            return finish("Choose a valid department and team.")

        # Joining Date
        if joining_date:
            try:
                user.joining_date = datetime.strptime(joining_date, "%Y-%m-%d").date()
            except ValueError:
                return finish("❌ Invalid date format! Use YYYY-MM-DD.")

        # Profile Image
        if request.FILES.get("edit_profile_upload"):
            user.profile_image = request.FILES["edit_profile_upload"]

        # ========== Save with try-except ==========
        try:
            user.save()

        except IntegrityError as e:
            return finish("Cannot update user: employee ID, username or email is already in use.")

        return finish("User updated successfully!", ok=True)

    return redirect("admin_usermanagement")


@never_cache
@require_POST
def delete_user(request, id):
    from django.http import HttpResponseForbidden
    managed_users = _admin_managed_users(request)
    if managed_users is None:
        return HttpResponseForbidden("Admin access required.")

    user = get_object_or_404(managed_users, id=id)
    user.delete()
    messages.success(request, "Employee deleted successfully.")
    return redirect('admin_usermanagement') 


from datetime import datetime, timedelta
from urllib.parse import urlsplit
from django.http import JsonResponse
from django.db import transaction
from django.db.models import Q
from django.contrib import messages
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_time
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods
from .models import GoogleMeeting, Department, User


from urllib.parse import urlencode
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponseBadRequest, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_POST, require_http_methods
from .models import AdminNotepad, User as MonitoringUser


def _admin_notepad_owner(request):
    auth_user = request.user
    if auth_user.is_authenticated and auth_user.is_active and auth_user.is_superuser:
        return {"auth_owner": auth_user}
    if request.session.get("position") != "management":
        return None
    user_id = request.session.get("user_id")
    if not user_id:
        return None
    user = MonitoringUser.objects.filter(
        pk=user_id, job_Position__iexact="Management"
    ).first()
    if user is None:
        return None
    return {"management_owner": user}


@never_cache
@require_http_methods(["GET", "POST"])
def admin_notepad(request):
    owner = _admin_notepad_owner(request)
    if owner is None:
        return redirect("admin_login")
    owned = AdminNotepad.objects.filter(**owner)
    params = request.POST if request.method == "POST" else request.GET
    query = params.get("q", "").strip()[:200]
    sort = params.get("sort", "updated")
    orderings = {"updated": ("-updated_at", "-id"), "created": ("-created_at", "-id"), "title": ("title", "id")}
    if sort not in orderings:
        sort = "updated"
    notes = owned
    if query:
        notes = notes.filter(Q(title__icontains=query) | Q(content__icontains=query))
    page_obj = Paginator(notes.order_by(*orderings[sort]), 4).get_page(params.get("page"))
    selected_id = request.POST.get("note_id") if request.method == "POST" else None
    note = None
    if selected_id:
        try:
            selected_id = int(selected_id)
        except (ValueError, TypeError):
            return HttpResponseBadRequest("Invalid note ID.")
        note = get_object_or_404(owned, pk=selected_id)
    error = ""
    title = note.title if note else ""
    content = (note.content or "") if note else ""
    if request.method == "POST":
        title = request.POST.get("title", "").strip() or "Untitled"
        content = request.POST.get("content", "")
        if len(title) > 255:
            error = "Title must be 255 characters or fewer."
        elif not content.strip():
            error = "Write some content before saving the note."
        else:
            if note:
                note.title, note.content = title, content
                note.save()
                if request.headers.get("X-Requested-With") == "XMLHttpRequest":
                    return JsonResponse({
                        "ok": True, "id": note.pk, "title": note.title,
                        "content": note.content or "",
                        "updated_at": timezone.localtime(note.updated_at).strftime("%d %b %Y, %I:%M %p"),
                    })
                messages.success(request, "Note updated successfully.")
            else:
                note = AdminNotepad.objects.create(**owner, title=title, content=content)
                messages.success(request, "Note created successfully.")
            return redirect(reverse("admin_notepad") + "?" + urlencode({
                "q": query if selected_id else "", "sort": sort if selected_id else "updated", "page": page_obj.number if selected_id else 1,
            }))
    if error and selected_id and request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return JsonResponse({"ok": False, "message": error}, status=400)
    response = render(request, "admin_panel/admin_notepad.html", {
        "note": None, "page_obj": page_obj, "query": query, "sort": sort,
        "form_title": title if not selected_id else "", "form_content": content if not selected_id else "", "note_error": error,
        "total_notes": owned.count(),
    })
    if error:
        response.status_code = 400
    return response

@never_cache
@require_POST
def admin_notepad_delete(request, pk):
    owner = _admin_notepad_owner(request)
    if owner is None:
        return redirect("admin_login")
    note = get_object_or_404(AdminNotepad, pk=pk, **owner)
    note.delete()
    messages.success(request, "Note deleted successfully.")
    sort = request.POST.get("sort", "updated")
    if sort not in {"updated", "created", "title"}:
        sort = "updated"
    return redirect(reverse("admin_notepad") + "?" + urlencode({
        "q": request.POST.get("q", "").strip()[:200],
        "sort": sort, "page": request.POST.get("page", "1"),
    }))



# Add imports and helpers to monitoringapp/views.py ONCE.
from pathlib import Path
from types import SimpleNamespace
from django.http import FileResponse, Http404
from django.core.exceptions import ValidationError
from django.core.validators import URLValidator
from django.shortcuts import get_object_or_404, redirect, render
from django.contrib import messages
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_POST, require_http_methods
from django.urls import reverse
from .models import AdminKnowledge, Knowledge, User as MonitoringUser


def _repository_role_user(request, role):
    if request.session.get("position") != role:
        return None
    user_id = request.session.get("user_id")
    if not user_id:
        return None
    db_role = {"team_lead": "Team Lead", "team_member": "Team Member"}[role]
    return MonitoringUser.objects.filter(pk=user_id, job_Position__iexact=db_role).first()


def _repository_card(resource, source):
    is_admin = source == "admin"
    uploader = (resource.auth_owner or resource.management_owner) if is_admin else resource.user
    label = "Admin" if is_admin else resource.user.job_Position
    file = None
    if resource.file:
        file = SimpleNamespace(
            name=resource.file.name,
            url=reverse("admin_repository_file", args=[resource.pk]) if is_admin else resource.file.url,
        )
    return SimpleNamespace(
        id=resource.pk, card_key=f"{source}-{resource.pk}",
        title=resource.title, description=resource.description or "",
        link=resource.link or "", file=file, created_at=resource.created_at,
        user=uploader, owner_username=uploader.username if uploader else "Admin",
        source_label=label, is_admin_resource=is_admin,
    )


@never_cache
@require_http_methods(["GET"])
def admin_repository_file(request, pk):
    if _admin_notepad_owner(request) is None and _repository_role_user(request, "team_lead") is None:
        raise Http404("Resource not found.")
    resource = get_object_or_404(AdminKnowledge, pk=pk)
    if not resource.file:
        raise Http404("File not found.")
    try:
        stream = resource.file.open("rb")
    except (FileNotFoundError, OSError):
        raise Http404("File not found.")
    return FileResponse(stream, as_attachment=True, filename=Path(resource.file.name).name)

@never_cache
@require_http_methods(["GET", "POST"])
def admin_repository(request):
    owner = _admin_notepad_owner(request)
    if owner is None:
        return redirect("admin_login")

    error = ""
    if request.method == "POST":
        title = request.POST.get("title", "").strip()
        description = request.POST.get("description", "").strip()
        link = request.POST.get("link", "").strip()
        file = request.FILES.get("file")

        if not title or len(title) > 255:
            error = "Enter a title with a maximum of 255 characters."
        elif link:
            try:
                URLValidator(schemes=["http", "https"])(link)
                if len(link) > 200:
                    raise ValidationError("Link is too long.")
            except ValidationError:
                error = "Enter a valid HTTP/HTTPS link with a maximum of 200 characters."

        if not error:
            AdminKnowledge.objects.create(
                **owner, title=title, description=description,
                link=link, file=file,
            )
            messages.success(request, "Resource added successfully.")
            return redirect("admin_repository")

    admin_items = AdminKnowledge.objects.select_related("auth_owner", "management_owner")
    team_items = Knowledge.objects.select_related("user", "department").filter(
        user__job_Position__iexact="Team Lead"
    ) | Knowledge.objects.select_related("user", "department").filter(
        user__job_Position__iexact="Team Member"
    )
    knowledge_items = [_repository_card(item, "admin") for item in admin_items]
    knowledge_items += [_repository_card(item, "team") for item in team_items]
    knowledge_items.sort(key=lambda item: (item.created_at, item.card_key), reverse=True)
    return render(request, "admin_panel/admin_repository.html", {
        "knowledge_items": knowledge_items,
        "repository_error": error,
        "form_title": request.POST.get("title", "") if error else "",
        "form_description": request.POST.get("description", "") if error else "",
        "form_link": request.POST.get("link", "") if error else "",
    }, status=400 if error else 200)

@never_cache
@require_POST
def admin_repository_delete(request, pk):
    if _admin_notepad_owner(request) is None:
        return redirect("admin_login")
    resource = get_object_or_404(AdminKnowledge, pk=pk)
    resource.delete()
    messages.success(request, "Resource deleted successfully.")
    return redirect("admin_repository")



@never_cache
@require_http_methods(["GET", "POST"])
def teamlead_repository(request):
    user = _repository_role_user(request, "team_lead")
    if user is None:
        return redirect("index")

    if request.method == "POST":
        title = request.POST.get("title", "").strip()
        description = request.POST.get("description")
        link = request.POST.get("link")
        file = request.FILES.get("file")

        if not title or len(title) > 255:
            messages.error(request, "Enter a title with a maximum of 255 characters.")
            return redirect("teamlead_repository")
        if user.department_id is None:
            messages.error(request, "Set your department before adding a resource.")
            return redirect("teamlead_repository")
        if link:
            try:
                URLValidator(schemes=["http", "https"])(link)
            except ValidationError:
                messages.error(request, "Enter a valid HTTP/HTTPS link.")
                return redirect("teamlead_repository")
        Knowledge.objects.create(
            department=user.department,
            user=user,
            title=title,
            description=description,
            link=link,
            file=file
        )

        messages.success(request, "Resource added successfully.")
        return redirect("teamlead_repository")

    team_items = Knowledge.objects.filter(department=user.department).select_related("user") if user.department_id else Knowledge.objects.none()
    knowledge_items = [_repository_card(item, "team") for item in team_items]
    knowledge_items += [_repository_card(item, "admin") for item in AdminKnowledge.objects.select_related("auth_owner", "management_owner")]
    knowledge_items.sort(key=lambda item: (item.created_at, item.card_key), reverse=True)

    response = render(request, "team_lead/teamlead_repository.html", {"knowledge_items": knowledge_items, "user": user})
    
    # Extra cache protection
    response["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response["Pragma"] = "no-cache"
    response["Expires"] = "0"
    
    return response

@never_cache
@require_POST
def teamlead_repository_delete(request, pk):
    user = _repository_role_user(request, "team_lead")
    if user is None:
        return redirect("index")
    resource = get_object_or_404(Knowledge, pk=pk, department=user.department)
    resource.delete()
    messages.success(request, "Resource deleted successfully.")
    return redirect("teamlead_repository")


@never_cache
def teammember_repository(request):
    user = _repository_role_user(request, "team_member")
    if user is None:
        return redirect("index")

    department = user.department if hasattr(user, "department") and user.department else Department.objects.first()

    if request.method == "POST":
        title = request.POST.get("title")
        description = request.POST.get("description")
        link = request.POST.get("link")
        file = request.FILES.get("file")

        if not (title or "").strip():
            messages.error(request, "Please enter a resource title.")
            return redirect("teammember_repository")

        # ✅ Make sure department is not None before saving
        if not department:
            messages.error(request, "No department found for this user or in database.")
            return redirect("teammember_repository")

        Knowledge.objects.create(
            department=department,
            user=user,
            title=title,
            description=description,
            link=link,
            file=file
        )

        messages.success(request, "Resource added successfully.")
        return redirect("teammember_repository")

    # ✅ Fetch repository items department-wise
    knowledge_items = Knowledge.objects.filter(department=department).order_by("-created_at")

    return render(request, "team_member/teammember_repository.html", {
        "knowledge_items": knowledge_items
    })

@require_POST
@never_cache
def teammember_repository_delete(request, pk):
    user = _repository_role_user(request, "team_member")
    if user is None:
        return redirect("index")

    department = user.department if hasattr(user, "department") and user.department else Department.objects.first()
    resource = get_object_or_404(Knowledge, id=pk, department=department)

    # Delete only after the confirmation form submits.
    resource.delete()
    messages.success(request, "Resource deleted successfully.")
    return redirect("teammember_repository")

from chat.models import Conversation
from django.db import transaction
from django.db.models import Q
from django.http import JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods
from .models import User, Group, GroupMember, GroupMessage


def _admin_chat_user(request):
    position = str(request.session.get("position", "")).strip().lower().replace(" ", "_")
    if position != "management" or not request.session.get("user_id"):
        return None
    # Chat identity is the same custom User used by Lead/Member messages.
    return User.objects.filter(
        pk=request.session["user_id"], job_Position__iexact="Management"
    ).first()


@never_cache
@require_http_methods(["GET", "POST"])
def admin_chat(request):
    current_user = _admin_chat_user(request)
    if current_user is None:
        return redirect("index")

    if request.method == "POST":
        action = request.POST.get("action")

        if action == "create_group":
            return _member_chat_create_group(
                request,
                current_user,
                group_view_name="admin_group_chat",
            )

        return _member_chat_error("Invalid action.")

    context = _member_chat_context(current_user)
    context["user_role"] = "management"

    return render(
        request,
        "admin_panel/admin_chat.html",
        context,
    )


@never_cache
@require_http_methods(["GET", "POST"])
def admin_chat_room(request, user_id):
    from .member_chat_actions import chat_action

    current_user = _admin_chat_user(request)
    if current_user is None:
        return redirect("index")

    other_user = get_object_or_404(
        User,
        pk=user_id,
    )

    if current_user.pk == other_user.pk:
        return redirect("admin_chat")

    # Preserve existing conversations in either user ordering.
    room = (
        Conversation.objects.filter(
            Q(user1=current_user, user2=other_user)
            | Q(user1=other_user, user2=current_user)
        )
        .order_by("pk")
        .first()
    )

    if room is None:
        first, second = sorted(
            [current_user, other_user],
            key=lambda person: person.pk,
        )

        room, _ = Conversation.objects.get_or_create(
            user1=first,
            user2=second,
        )

    if request.method == "POST":
        action = request.POST.get("action")

        if action == "create_group":
            return _member_chat_create_group(
                request,
                current_user,
                group_view_name="admin_group_chat",
            )

        return chat_action(
            request,
            current_user,
            room,
            group=False,
        )

    messages_qs = (
        room.messages
        .exclude(hidden_for=current_user)
        .select_related("sender")
        .order_by("created_at", "pk")
    )

    # Used by the chat UI to reconcile messages without refreshing.
    if request.GET.get("messages") == "1":
        return JsonResponse({
            "ok": True,
            "items": [
                _member_chat_payload(message)
                for message in messages_qs
            ],
        })

    context = _member_chat_context(current_user)

    context.update({
        "user_role": "management",
        "room": room,
        "other_user": other_user,
        "messages": messages_qs,
    })

    return render(
        request,
        "admin_panel/admin_chat_room.html",
        context,
    )


@never_cache
@require_http_methods(["GET", "POST"])
def admin_group_chat(request, group_id):
    import logging

    from asgiref.sync import async_to_sync
    from channels.layers import get_channel_layer

    from .chat_receipts import broadcast_chat_activity    
    from .member_chat_actions import chat_action

    current_user = _admin_chat_user(request)
    if current_user is None:
        return redirect("index")
    group = get_object_or_404(Group, pk=group_id)

    membership = get_object_or_404(
        GroupMember,
        group=group,
        user=current_user,
    )

    can_manage = (
        group.created_by_id == current_user.pk
        or membership.role == "admin"
    )

    def member_state():
        current_membership = GroupMember.objects.filter(
            group=group,
            user=current_user,
        ).first()

        permitted = bool(
            current_membership
            and (
                group.created_by_id == current_user.pk
                or current_membership.role == "admin"
            )
        )

        entries = list(
            GroupMember.objects.filter(group=group)
            .select_related("user")
            .order_by("user__name", "pk")
        )

        available_users = []

        if permitted:
            available_users = [
                {
                    "id": person.pk,
                    "name": person.name,
                }
                for person in User.objects.exclude(
                    group_memberships__group=group
                ).order_by("name", "pk")
            ]

        return {
            "ok": True,
            "can_manage": permitted,
            "count": len(entries),
            "members": [
                {
                    "id": entry.user_id,
                    "name": entry.user.name,
                    "role": entry.role,
                    "is_creator": (
                        entry.user_id == group.created_by_id
                    ),
                    "is_self": (
                        entry.user_id == current_user.pk
                    ),
                    "profile": (
                        entry.user.profile_image.url
                        if entry.user.profile_image
                        else None
                    ),
                }
                for entry in entries
            ],
            "all_users": available_users,
        }

    def notify_members():
        layer = get_channel_layer()

        if layer:
            try:
                async_to_sync(layer.group_send)(
                    f"group_{group.pk}",
                    {"type": "chat_change"},
                )
            except Exception:
                logging.getLogger(__name__).exception(
                    "Group membership update delivery failed."
                )

        async_to_sync(broadcast_chat_activity)(
            True,
            group.pk,
        )

    if request.method == "POST":
        action = request.POST.get("action", "")

        if action == "create_group":
            return _member_chat_create_group(
                request,
                current_user,
                group_view_name="admin_group_chat",
            )

        if action in {
            "edit_message",
            "delete_for_me",
            "delete_for_everyone",
            "clear_chat",
        }:
            return chat_action(
                request,
                current_user,
                group,
                group=True,
            )

        if action not in {
            "add_member",
            "remove_member",
            "make_admin",
            "remove_admin",
        }:
            return _member_chat_error("Invalid action.")

        if not can_manage:
            return _member_chat_error(
                "Only the group creator or an admin "
                "can manage members.",
                403,
            )

        try:
            target_id = int(
                request.POST.get("user_id", "")
            )
        except (ValueError, TypeError):
            return _member_chat_error(
                "Select a valid member."
            )

        with transaction.atomic():
            # Serialize membership changes for this group.
            locked_group = (
                Group.objects.select_for_update()
                .get(pk=group.pk)
            )

            membership = get_object_or_404(
                GroupMember,
                group=group,
                user=current_user,
            )

            if (
                locked_group.created_by_id != current_user.pk
                and membership.role != "admin"
            ):
                return _member_chat_error(
                    "You no longer have permission "
                    "to manage this group.",
                    403,
                )

            if action == "add_member":
                target_user = get_object_or_404(
                    User,
                    pk=target_id,
                )

                _, created = (
                    GroupMember.objects.get_or_create(
                        group=group,
                        user=target_user,
                    )
                )

                notice = (
                    f"{target_user.name} added to the group."
                    if created
                    else f"{target_user.name} is already a member."
                )

            else:
                target = get_object_or_404(
                    GroupMember.objects.select_related("user"),
                    group=group,
                    user_id=target_id,
                )

                if target_id == locked_group.created_by_id:
                    return _member_chat_error(
                        "The group creator cannot be "
                        "removed or demoted."
                    )

                if target_id == current_user.pk:
                    return _member_chat_error(
                        "You cannot remove yourself or "
                        "change your own admin role here."
                    )

                if action == "remove_member":
                    member_count = (
                        GroupMember.objects.filter(
                            group=group
                        ).count()
                    )

                    if member_count <= 2:
                        return _member_chat_error(
                            "Keep at least two people "
                            "in the group."
                        )

                    name = target.user.name
                    target.delete()

                    notice = (
                        f"{name} removed from the group."
                    )

                else:
                    target.role = (
                        "admin"
                        if action == "make_admin"
                        else "member"
                    )

                    target.save(update_fields=["role"])

                    notice = (
                        f"{target.user.name} is now an admin."
                        if action == "make_admin"
                        else (
                            "Admin access removed for "
                            f"{target.user.name}."
                        )
                    )

            transaction.on_commit(notify_members)

        return JsonResponse({
            **member_state(),
            "message": notice,
        })

    if request.GET.get("members") == "1":
        return JsonResponse(member_state())

    messages_qs = (
        GroupMessage.objects.filter(group=group)
        .exclude(hidden_for=current_user)
        .select_related("sender")
        .order_by("timestamp", "pk")
    )

    if request.GET.get("messages") == "1":
        return JsonResponse({
            "ok": True,
            "items": [
                _member_chat_payload(message, True)
                for message in messages_qs
            ],
        })

    context = _member_chat_context(current_user)

    context.update({
        "user_role": "management",
        "group": group,
        "messages": messages_qs,
        "members": (
            GroupMember.objects.filter(group=group)
            .select_related("user")
            .order_by("user__name", "pk")
        ),
        "all_users": (
            User.objects.exclude(
                group_memberships__group=group
            ).order_by("name", "pk")
        ),
        "can_manage_group": can_manage,
    })

    return render(
        request,
        "admin_panel/admin_group_chat.html",
        context,
    )



def add_contact(request):
    if request.method == "POST":
        name = request.POST.get("name")
        phone = request.POST.get("phone")

        # prevent duplicate phone numbers
        if ExtraContact.objects.filter(phone=phone).exists():
            messages.error(request, "Contact with this phone already exists!")
        else:
            ExtraContact.objects.create(name=name, phone=phone)
            messages.success(request, "Contact added successfully!")

        return redirect("admin_chat")  # back to chat page

    return redirect("admin_chat")

from .models import AdminNotification
from django.http import JsonResponse
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET, require_POST
from django.shortcuts import redirect, get_object_or_404, render


def _admin_notification_owned(request):
    owner = _admin_notepad_owner(request)
    if owner is None:
        return None
    _process_admin_reminders(owner)
    return AdminNotification.objects.filter(**owner)

def _create_admin_notification(owner, *, kind, title, message="", url="", event_key=None):
    """Call after a successful Admin operation; owner comes from _admin_notepad_owner.
    A stable event_key avoids duplicates on repeated processing.
    """
    if owner is None:
        return None
    values = dict(owner, kind=kind, title=title, message=message, url=url)
    if event_key:
        owner_id = next(iter(owner.values())).pk
        key = f"admin:{next(iter(owner))}:{owner_id}:{event_key}"
        if len(key) > 180:
            raise ValueError("Notification event key is too long.")
        return AdminNotification.objects.get_or_create(event_key=key, defaults=values)[0]
    return AdminNotification.objects.create(**values)


def _admin_notification_count(owned):
    return owned.filter(is_read=False, is_archived=False).count()


@never_cache
@require_GET
def admin_notifications(request):
    owned = _admin_notification_owned(request)
    if owned is None:
        if request.GET.get("format") == "json":
            return JsonResponse({"ok": False, "message": "Please sign in again."}, status=403)
        return redirect("admin_login")
    context = {"items": owned.order_by("-created_at", "-id"), "unread_count": _admin_notification_count(owned)}
    if request.GET.get("format") == "json":
        data = {"ok": True, "unread_count": context["unread_count"]}
        if request.GET.get("counts") != "1":
            data["html"] = render_to_string("admin_panel/admin_notifications.html", context, request=request)
        return JsonResponse(data)
    return render(request, "admin_panel/admin_notifications.html", context)


@never_cache
@require_POST
def admin_notifications_read_all(request):
    owned = _admin_notification_owned(request)
    if owned is None:
        return JsonResponse({"ok": False, "message": "Admin access denied."}, status=403)
    owned.filter(is_read=False, is_archived=False).update(is_read=True)
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return JsonResponse({"ok": True, "message": "All notifications marked as read.", "unread_count": _admin_notification_count(owned)})
    return redirect("admin_notifications")


@never_cache
@require_POST
def admin_notification_open(request, notification_id):
    owned = _admin_notification_owned(request)
    if owned is None:
        return JsonResponse({"ok": False, "message": "Admin access denied."}, status=403)
    item = get_object_or_404(owned, pk=notification_id)
    action = request.POST.get("action", "open")
    changes = {"mark_read": {"is_read": True}, "mark_unread": {"is_read": False}, "archive": {"is_archived": True}, "restore": {"is_archived": False}, "open": {"is_read": True}}
    if action not in changes:
        return JsonResponse({"ok": False, "message": "Invalid notification action."}, status=400)
    owned.filter(pk=item.pk).update(**changes[action])
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return JsonResponse({"ok": True, "message": "Notification updated.", "unread_count": _admin_notification_count(owned)})
    return redirect("admin_notifications")

import calendar
from datetime import date, datetime, time
from django.db import transaction
from django.http import JsonResponse
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods
from .models import AdminReminder, AdminNotification


def _process_admin_reminders(owner=None):
    now = timezone.now()
    due = AdminReminder.objects.filter(is_completed=False, notified_at__isnull=True, remind_at__lte=now)
    if owner is not None:
        due = due.filter(**owner)
    created = 0
    for reminder_id in due.values_list("pk", flat=True).iterator():
        with transaction.atomic():
            reminder = AdminReminder.objects.select_for_update().filter(pk=reminder_id, is_completed=False, notified_at__isnull=True, remind_at__lte=now).first()
            if reminder is None:
                continue
            recipient = {"management_owner_id": reminder.management_owner_id} if reminder.management_owner_id else {"auth_owner_id": reminder.auth_owner_id}
            event = timezone.localtime(reminder.event_at).strftime("%d %b %Y, %I:%M %p")
            notification, added = AdminNotification.objects.get_or_create(
                event_key=f"admin-reminder:{reminder.pk}:{reminder.schedule_version}",
                defaults=dict(recipient, kind="reminder", title=reminder.title, message=f"Event: {event}\n{reminder.description}".strip(), url=""),
            )
            reminder.notification = notification
            reminder.notified_at = now
            reminder.save(update_fields=["notification", "notified_at", "updated_at"])
            created += int(added)
    return created


def _admin_reminder_item(reminder):
    event = timezone.localtime(reminder.event_at)
    remind = timezone.localtime(reminder.remind_at)
    return {"id": reminder.pk, "title": reminder.title, "description": reminder.description,
            "date": event.strftime("%Y-%m-%d"), "event_at": event.strftime("%Y-%m-%dT%H:%M"),
            "remind_at": remind.strftime("%Y-%m-%dT%H:%M"), "is_completed": reminder.is_completed}


def _admin_reminder_time(value):
    try:
        parsed = parse_datetime(value or "")
        if parsed and timezone.is_naive(parsed):
            parsed = timezone.make_aware(parsed, timezone.get_current_timezone())
        return parsed
    except (TypeError, ValueError, OverflowError):
        return None


@never_cache
@require_http_methods(["GET", "POST"])
def admin_reminders(request):
    owner = _admin_notepad_owner(request)
    ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    if owner is None:
        if ajax or request.GET.get("format") == "json":
            return JsonResponse({"ok": False, "message": "Admin access required."}, status=403)
        return redirect("admin_login")
    reminders = AdminReminder.objects.filter(**owner)

    def failure(message, status=400):
        if ajax:
            return JsonResponse({"ok": False, "message": message}, status=status)
        messages.error(request, message)
        return redirect("admin_reminders")

    def success(message, reminder=None, deleted_id=None):
        if ajax:
            return JsonResponse({"ok": True, "message": message, "id": reminder.pk if reminder else deleted_id,
                                 "item": _admin_reminder_item(reminder) if reminder else None,
                                 "is_completed": reminder.is_completed if reminder else None})
        messages.success(request, message)
        return redirect("admin_reminders")

    if request.method == "POST":
        action = request.POST.get("action", "")
        if action not in {"create", "update", "delete", "complete", "reopen"}:
            return failure("Invalid reminder action.")
        with transaction.atomic():
            reminder = None
            if action != "create":
                try:
                    reminder_id = int(request.POST.get("reminder_id", ""))
                except (TypeError, ValueError):
                    return failure("Invalid reminder ID.")
                reminder = reminders.select_for_update().filter(pk=reminder_id).first()
                if reminder is None:
                    return failure("Reminder not found.", 404)
            if action == "delete":
                deleted_id = reminder.pk
                if reminder.notification_id:
                    AdminNotification.objects.filter(pk=reminder.notification_id, **owner).update(is_archived=True)
                reminder.delete()
                return success("Reminder deleted.", deleted_id=deleted_id)
            if action in {"complete", "reopen"}:
                reminder.is_completed = action == "complete"
                reminder.save(update_fields=["is_completed", "updated_at"])
                if reminder.notification_id and reminder.is_completed:
                    AdminNotification.objects.filter(pk=reminder.notification_id, **owner).update(is_read=True)
                # Reopening keeps an already-delivered alert delivered; reschedule to get a new alert.
                return success("Reminder completed." if reminder.is_completed else "Reminder reopened.", reminder)
            title = request.POST.get("title", "").strip()
            description = request.POST.get("description", "").strip()
            event_at = _admin_reminder_time(request.POST.get("event_at"))
            remind_at = _admin_reminder_time(request.POST.get("remind_at"))
            if not title or len(title) > 180:
                return failure("Enter a title with a maximum of 180 characters.")
            if event_at is None or remind_at is None:
                return failure("Enter valid event and reminder dates and times.")
            if remind_at > event_at:
                return failure("Reminder must be on or before the event.")
            changed = reminder is None or reminder.event_at != event_at or reminder.remind_at != remind_at
            if changed and remind_at <= timezone.now():
                return failure("Choose a future reminder time.")
            if reminder is None:
                reminder = AdminReminder(**owner)
            elif changed:
                if reminder.notification_id:
                    AdminNotification.objects.filter(pk=reminder.notification_id, **owner).update(is_archived=True)
                reminder.schedule_version += 1
            if changed:
                reminder.notified_at = None
                reminder.notification = None
                reminder.is_completed = False
            reminder.title, reminder.description = title, description
            reminder.event_at, reminder.remind_at = event_at, remind_at
            reminder.save()
            return success("Reminder created." if action == "create" else "Reminder updated.", reminder)

    _process_admin_reminders(owner)
    today = timezone.localdate()
    try:
        selected = date.fromisoformat(request.GET.get("month", today.strftime("%Y-%m")) + "-01")
        next_month = selected.replace(year=selected.year + 1, month=1) if selected.month == 12 else selected.replace(month=selected.month + 1)
    except (TypeError, ValueError):
        selected = today.replace(day=1)
        next_month = selected.replace(year=selected.year + 1, month=1) if selected.month == 12 else selected.replace(month=selected.month + 1)
    tz = timezone.get_current_timezone()
    start = timezone.make_aware(datetime.combine(selected, time.min), tz)
    end = timezone.make_aware(datetime.combine(next_month, time.min), tz)
    monthly = reminders.filter(event_at__gte=start, event_at__lt=end).order_by("event_at", "id")
    context = {"selected_month": selected.strftime("%Y-%m"), "month_label": selected.strftime("%B %Y"),
               "calendar_weeks": calendar.Calendar(firstweekday=0).monthdayscalendar(selected.year, selected.month),
               "calendar_items": [_admin_reminder_item(r) for r in monthly], "today": today.isoformat(), "reminders": monthly}
    if request.GET.get("format") == "json":
        return JsonResponse({"ok": True, "items": context["calendar_items"], "selected_month": context["selected_month"],
                             "month_label": context["month_label"], "calendar_weeks": context["calendar_weeks"], "today": context["today"]})
    return render(request, "admin_panel/admin_reminders.html", context)



@never_cache
def login_view(request):
    # Already logged in → redirect
    if request.session.get("user_id") and request.session.get("position"):
        position = request.session["position"]
        if position == "team_lead":
            return redirect("teamlead_dashboard")
        elif str(position).strip().lower().replace(" ", "_") in {"accounts", "accounts_team"}:
            return redirect("accounts_dashboard")
        elif position == "management":
            return redirect("admin_dashboard")
        else:
            return redirect("teammember_dashboard")

    if request.method == "POST":
        username = request.POST.get("username", "").strip()
        password = request.POST.get("password", "").strip()

        # Fetch user from DB
        try:
            user = User.objects.get(username=username)
        except User.DoesNotExist:
            messages.error(request, "Invalid username or password")
            return render(request, "user_login.html")

        # Check hashed password
        if not check_password(password, user.password):
            messages.error(request, "Invalid username or password")
            return render(request, "user_login.html")

        # ✅ Update user status and last login time (normalize status explicitly)
        now = timezone.now()
        user.status = "active"                # canonical value (no spaces, lowercase)
        user.last_login_time = now
        user.last_activity = now

        # if model happens to have is_active (boolean), keep it in sync
        if hasattr(user, "is_active"):
            try:
                setattr(user, "is_active", True)
            except Exception:
                pass

        user.save()

        # Normalize DB position
        db_position = user.job_Position.strip().lower().replace(" ", "_")

        # Save session
        request.session["user_id"] = user.id
        request.session["position"] = db_position
        request.session["login_time"] = str(now)

        # Redirect based on DB position only
        if db_position == "team_lead":
            return redirect("teamlead_dashboard")
        elif db_position in {"accounts", "accounts_team"}:
            return redirect("accounts_dashboard")
        elif db_position == "management":
            return redirect("admin_dashboard")
        else:
            return redirect("teammember_dashboard")

    return render(request, "user_login.html")



def get_logged_in_user_api(request):
    user_id = request.session.get("user_id")
    if not user_id:
        return JsonResponse({"name": None, "profile_image_url": None})
    user = User.objects.filter(id=user_id).first()
    if user is None:
        return JsonResponse({"name": None, "profile_image_url": None})
    return JsonResponse({
        "name": user.name,
        "username": user.username,
        "job_position": user.job_Position,
        "profile_image_url": user.profile_image.url if user.profile_image else None,
    })




@never_cache
def teamlead_edit_member(request):

    # -----------------------------------------
    # TEAM LEAD ACCESS CHECK
    # -----------------------------------------
    if (
        not request.session.get("user_id")
        or request.session.get("position") != "team_lead"
    ):
        return redirect("index")

    try:
        team_lead = User.objects.get(
            id=request.session["user_id"]
        )
    except User.DoesNotExist:
        request.session.flush()
        return redirect("index")


    # -----------------------------------------
    # ONLY POST ALLOWED
    # -----------------------------------------
    if request.method != "POST":
        return redirect("teamlead_dashboard")


    member_id = request.POST.get("id")

    if not member_id:
        messages.error(request, "Member ID is missing.")
        return redirect("teamlead_dashboard")


    # -----------------------------------------
    # SECURITY:
    # ONLY TEAM MEMBERS FROM TL'S DEPARTMENT
    # -----------------------------------------
    user = get_object_or_404(
        User,
        id=member_id,
        department=team_lead.department,
        job_Position__iexact="Team Member",
    )


    # -----------------------------------------
    # FORM VALUES
    # -----------------------------------------
    employee_id = request.POST.get("employee_id", "").strip()
    username = request.POST.get("username", "").strip()
    email = request.POST.get("email", "").strip()
    name = request.POST.get("name", "").strip()
    phone = request.POST.get("phone", "").strip()

    job_position = request.POST.get(
        "job_position",
        "Team Member"
    ).strip()

    designation = request.POST.get(
        "designation",
        ""
    ).strip()

    work_location = request.POST.get(
        "work_location",
        ""
    ).strip()

    status = request.POST.get(
        "status",
        ""
    ).strip()

    password = request.POST.get(
        "password",
        ""
    ).strip()

    joining_date = request.POST.get(
        "joining_date",
        ""
    ).strip()

    department_name = request.POST.get(
        "department",
        ""
    ).strip()

    team_name = request.POST.get(
        "team",
        ""
    ).strip()


    # -----------------------------------------
    # UNIQUENESS CHECKS
    # -----------------------------------------
    if User.objects.filter(
        employee_id=employee_id
    ).exclude(id=user.id).exists():

        messages.error(
            request,
            "Employee ID already exists for another user."
        )

        return redirect("teamlead_dashboard")


    if User.objects.filter(
        username=username
    ).exclude(id=user.id).exists():

        messages.error(
            request,
            "Username already exists for another user."
        )

        return redirect("teamlead_dashboard")


    if User.objects.filter(
        email=email
    ).exclude(id=user.id).exists():

        messages.error(
            request,
            "Email already exists for another user."
        )

        return redirect("teamlead_dashboard")


    # -----------------------------------------
    # UPDATE BASIC DETAILS
    # -----------------------------------------
    user.employee_id = employee_id
    user.username = username
    user.email = email
    user.name = name
    user.phone = phone

    user.designation = designation
    user.work_location = work_location
    user.status = status


    # -----------------------------------------
    # KEEP JOB POSITION AS TEAM MEMBER
    # -----------------------------------------
    user.job_Position = "Team Member"


    # -----------------------------------------
    # DEPARTMENT
    # -----------------------------------------
    if department_name:

        try:
            department = Department.objects.get(
                name__iexact=department_name
            )

            # Team Lead can only keep member
            # inside their own department.
            if department.id != team_lead.department_id:
                messages.error(
                    request,
                    "You can only edit members from your department."
                )

                return redirect("teamlead_dashboard")

            user.department = department

        except Department.DoesNotExist:

            messages.error(
                request,
                "Selected department does not exist."
            )

            return redirect("teamlead_dashboard")


    # -----------------------------------------
    # TEAM
    # -----------------------------------------
    if team_name:

        try:
            team = Team.objects.get(
                name__iexact=team_name
            )

            user.team = team

        except Team.DoesNotExist:

            messages.error(
                request,
                "Selected team does not exist."
            )

            return redirect("teamlead_dashboard")

    else:

        user.team = None


    # -----------------------------------------
    # JOINING DATE
    # -----------------------------------------
    if joining_date:

        try:

            user.joining_date = datetime.strptime(
                joining_date,
                "%Y-%m-%d"
            ).date()

        except ValueError:

            messages.error(
                request,
                "Invalid joining date."
            )

            return redirect("teamlead_dashboard")

    else:

        user.joining_date = None


    # -----------------------------------------
    # PASSWORD
    # Blank = Keep Existing Password
    # -----------------------------------------
    if password:

        user.password = make_password(
            password
        )


    # -----------------------------------------
    # PROFILE IMAGE
    # -----------------------------------------
    if request.FILES.get("profile_image"):

        user.profile_image = request.FILES[
            "profile_image"
        ]


    # -----------------------------------------
    # SAVE
    # -----------------------------------------
    try:

        user.save()

        messages.success(
            request,
            "Team member updated successfully."
        )

    except IntegrityError:

        messages.error(
            request,
            "Unable to update member because of a unique field conflict."
        )

        return redirect("teamlead_dashboard")


    return redirect("teamlead_dashboard")


@never_cache
def teamlead_dashboard(request):
    # 🔒 Restrict access
    if not request.session.get("user_id") or request.session.get("position") != "team_lead":
        return redirect("index")

    try:
        team_lead = User.objects.get(id=request.session['user_id'])
    except User.DoesNotExist:
        request.session.flush()
        return redirect("index")

        # Save default Morning and Evening times from an effective date.
    if (
        request.method == "POST"
        and request.POST.get("action") == "save_default_report_schedule"
    ):
        from django.utils import timezone
        from django.utils.dateparse import parse_date

        day_kind = request.POST.get("day_kind", "regular")
        try:
            effective_from = parse_date(
                request.POST.get("effective_from", "")
            )
        except ValueError:
            effective_from = None

        if day_kind not in ("regular", "saturday") or effective_from is None:
            messages.error(request, "Select a valid schedule and start date.")
            return redirect("teamlead_dashboard")

        if effective_from < timezone.localdate():
            messages.error(request, "Start date must be today or later.")
            return redirect("teamlead_dashboard")

        timings = []
        for report_type in ("morning", "evening"):
            start_time = parse_time(
                request.POST.get(f"{report_type}_start", "")
            )
            end_time = parse_time(
                request.POST.get(f"{report_type}_end", "")
            )

            if not start_time or not end_time or start_time >= end_time:
                messages.error(
                    request,
                    f"Enter valid {report_type} start and end times."
                )
                return redirect("teamlead_dashboard")

            timings.append((report_type, start_time, end_time))

        with transaction.atomic():
            for report_type, start_time, end_time in timings:
                ReportDefaultSchedule.objects.update_or_create(
                    team_lead=team_lead,
                    effective_from=effective_from,
                    day_kind=day_kind,
                    report_type=report_type,
                    defaults={
                        "start_time": start_time,
                        "end_time": end_time,
                        "is_active": True,
                    },
                )

        messages.success(
            request,
            f"Default report times saved from {effective_from:%d %B %Y}."
        )
        return redirect(
            f"{request.path}?month={effective_from:%Y-%m}"
        )

    # Save morning/evening windows for one calendar date.
    if request.method == "POST" and request.POST.get("action") == "save_date_report_schedule":
        from django.utils.dateparse import parse_date

        try:
            schedule_date = parse_date(request.POST.get("schedule_date", ""))
        except ValueError:
            schedule_date = None

        if schedule_date is None:
            messages.error(request, "Select a valid report date.")
            return redirect("teamlead_dashboard")

        changes = []

        for report_type in ("morning", "evening"):
            enabled = request.POST.get(f"{report_type}_enabled") == "on"
            start_value = request.POST.get(f"{report_type}_start", "")
            end_value = request.POST.get(f"{report_type}_end", "")

            if enabled:
                start_time = parse_time(start_value)
                end_time = parse_time(end_value)

                if not start_time or not end_time or start_time >= end_time:
                    messages.error(
                        request,
                        f"Enter a valid {report_type} start and end time."
                    )
                    return redirect(
                        f"{request.path}?month={schedule_date:%Y-%m}"
                    )

                changes.append((report_type, True, start_time, end_time))
            else:
                changes.append((report_type, False, None, None))

        with transaction.atomic():
            for report_type, enabled, start_time, end_time in changes:
                if enabled:
                    ReportDateSchedule.objects.update_or_create(
                        team_lead=team_lead,
                        date=schedule_date,
                        report_type=report_type,
                        defaults={
                            "start_time": start_time,
                            "end_time": end_time,
                            "is_active": True,
                        },
                    )
                else:
                    date_setting = ReportDateSchedule.objects.filter(
                        team_lead=team_lead,
                        date=schedule_date,
                        report_type=report_type,
                    ).first()

                    if date_setting is not None:
                        date_setting.is_active = False
                        date_setting.save(update_fields=["is_active"])
                        continue

                    # Save an inactive date override so the default
                    # does not apply to this date.
                    start_value = request.POST.get(
                        f"{report_type}_start", ""
                    )
                    end_value = request.POST.get(
                        f"{report_type}_end", ""
                    )
                    disabled_start = parse_time(start_value)
                    disabled_end = parse_time(end_value)

                    if not disabled_start or not disabled_end:
                        kinds = (
                            ("saturday", "regular")
                            if schedule_date.weekday() == 5
                            else ("regular",)
                        )
                        for kind in kinds:
                            default_setting = (
                                ReportDefaultSchedule.objects.filter(
                                    team_lead=team_lead,
                                    effective_from__lte=schedule_date,
                                    day_kind=kind,
                                    report_type=report_type,
                                )
                                .order_by("-effective_from", "-id")
                                .first()
                            )
                            if default_setting is not None:
                                disabled_start = default_setting.start_time
                                disabled_end = default_setting.end_time
                                break

                    if disabled_start and disabled_end:
                        ReportDateSchedule.objects.create(
                            team_lead=team_lead,
                            date=schedule_date,
                            report_type=report_type,
                            start_time=disabled_start,
                            end_time=disabled_end,
                            is_active=False,
                        )

        messages.success(
            request,
            f"Report times saved for {schedule_date:%d %B %Y}."
        )
        return redirect(f"{request.path}?month={schedule_date:%Y-%m}")
    if request.method == "POST" and request.POST.get("action") == "mark_sundays_holiday":
        month_value = request.POST.get("month", "")

        try:
            first_day = date.fromisoformat(f"{month_value}-01")
        except ValueError:
            messages.error(request, "Select a valid month.")
            return redirect("teamlead_dashboard")

        last_day = calendar.monthrange(first_day.year, first_day.month)[1]

        with transaction.atomic():
            for number in range(1, last_day + 1):
                holiday_date = date(first_day.year, first_day.month, number)
                if holiday_date.weekday() == 6:
                    TeamReportLeave.objects.update_or_create(
                        team_lead=team_lead,
                        date=holiday_date,
                        defaults={
                            "morning_leave": True,
                            "evening_leave": True,
                            "reason": "Sunday holiday",
                        },
                    )

        messages.success(request, "All Sundays marked as Holiday.")
        return redirect(f"{request.path}?month={month_value}")

    # Save one-date leave for the team lead and all members of the team.
    if request.method == "POST" and request.POST.get("action") == "save_team_report_leave":
        from django.utils import timezone
        from django.utils.dateparse import parse_date

        leave_date = parse_date(request.POST.get("leave_date", ""))
        if leave_date is None or leave_date < timezone.localdate():
            messages.error(request, "Select a valid date today or later.")
            return redirect("teamlead_dashboard")

        morning_leave = request.POST.get("morning_leave") == "on"
        evening_leave = request.POST.get("evening_leave") == "on"
        if not morning_leave and not evening_leave:
            # Clearing both selections removes any previous leave on this date.
            TeamReportLeave.objects.filter(team_lead=team_lead, date=leave_date).delete()
            messages.success(request, f"Report leave cleared for {leave_date}.")
        else:
            TeamReportLeave.objects.update_or_create(
                team_lead=team_lead,
                date=leave_date,
                defaults={
                    "morning_leave": morning_leave,
                    "evening_leave": evening_leave,
                    "reason": request.POST.get("reason", "").strip()[:255],
                },
            )
            messages.success(request, f"Report leave saved for {leave_date}.")

        return redirect(f"{request.path}?month={leave_date:%Y-%m}")
    # =========================================================
    # ⚙️ REPORT SCHEDULE CONFIGURATION
    # =========================================================

    report_days = [
        (0, "Monday"),
        (1, "Tuesday"),
        (2, "Wednesday"),
        (3, "Thursday"),
        (4, "Friday"),
        (5, "Saturday"),
    ]

    # ---------------------------------------------------------
    # SAVE REPORT SCHEDULE FROM DASHBOARD MODAL
    # ---------------------------------------------------------
    if (
        request.method == "POST"
        and request.POST.get("action") == "update_report_schedule"
    ):

        for day_value, day_name in report_days:

            # =========================
            # MORNING
            # =========================

            morning_start = request.POST.get(
                f"morning_start_{day_value}"
            )

            morning_end = request.POST.get(
                f"morning_end_{day_value}"
            )

            morning_enabled = request.POST.get(
                f"morning_enabled_{day_value}"
            )

            morning_setting = ReportTimeSetting.objects.filter(
                team_lead=team_lead,
                day_of_week=day_value,
                report_type="morning",
            ).first()

            if (
                morning_enabled == "on"
                and morning_start
                and morning_end
            ):

                ReportTimeSetting.objects.update_or_create(
                    team_lead=team_lead,
                    day_of_week=day_value,
                    report_type="morning",
                    defaults={
                        "start_time": morning_start,
                        "end_time": morning_end,
                        "is_active": True,
                    },
                )

            elif morning_setting:

                morning_setting.is_active = False

                morning_setting.save(
                    update_fields=["is_active"]
                )

            # =========================
            # EVENING
            # =========================

            evening_start = request.POST.get(
                f"evening_start_{day_value}"
            )

            evening_end = request.POST.get(
                f"evening_end_{day_value}"
            )

            evening_enabled = request.POST.get(
                f"evening_enabled_{day_value}"
            )

            evening_setting = ReportTimeSetting.objects.filter(
                team_lead=team_lead,
                day_of_week=day_value,
                report_type="evening",
            ).first()

            if (
                evening_enabled == "on"
                and evening_start
                and evening_end
            ):

                ReportTimeSetting.objects.update_or_create(
                    team_lead=team_lead,
                    day_of_week=day_value,
                    report_type="evening",
                    defaults={
                        "start_time": evening_start,
                        "end_time": evening_end,
                        "is_active": True,
                    },
                )

            elif evening_setting:

                evening_setting.is_active = False

                evening_setting.save(
                    update_fields=["is_active"]
                )

            messages.success(request, "Report schedule updated successfully.")
        return redirect("teamlead_dashboard")

      


    
        # ---------------------------------
    # 📢 Team Lead Announcements
    # ---------------------------------
# ---------------------------------
# 📢 Team Lead Announcements
# ---------------------------------
    if request.method == "POST" and request.POST.get("announcement_submit") == "1":

        announcement_message = request.POST.get("message", "").strip()
        send_to = request.POST.get("send_to", "team")
        selected_member_ids = request.POST.getlist("recipient_ids")

        if not announcement_message:
            messages.error(request, "Please enter an announcement message.")
            return redirect("teamlead_dashboard")

        # Create announcement
        announcement = Announcement.objects.create(
            title="Team Lead Announcement",
            message=announcement_message,
            created_by=team_lead
        )

        # Only Team Members under this Team Lead's department
        eligible_members = User.objects.filter(
            department=team_lead.department,
            job_Position__iexact="Team Member"
        ).exclude(
            id=team_lead.id
        )

        # ---------------------------------
        # Send to Entire Team
        # ---------------------------------
        if send_to == "team":

            recipients = [
                AnnouncementRecipient(
                    announcement=announcement,
                    recipient=member
                )
                for member in eligible_members
            ]

            if recipients:
                AnnouncementRecipient.objects.bulk_create(recipients)

            messages.success(
                request,
                "Announcement sent to the entire team."
            )

        # ---------------------------------
        # Send to Selected Members
        # ---------------------------------
        elif send_to == "selected":

            selected_members = eligible_members.filter(
                id__in=selected_member_ids
            )

            if not selected_members.exists():
                announcement.delete()
                messages.error(
                    request,
                    "Please select at least one team member."
                )
                return redirect("teamlead_dashboard")

            recipients = [
                AnnouncementRecipient(
                    announcement=announcement,
                    recipient=member
                )
                for member in selected_members
            ]

            AnnouncementRecipient.objects.bulk_create(recipients)

            messages.success(
                request,
                "Announcement sent to selected members."
            )


    # --- existing code for login_time, report times, etc. ---

    # ------------------------------
    # 🧩 Team Members
    # ------------------------------

    team_members_qs = User.objects.filter(
        department=team_lead.department,
        job_Position__iexact="Team Member"
    ).exclude(
        id=team_lead.id
    )

    # Get filters from GET request
    search = request.GET.get("search", "").strip()
    team_filter = request.GET.get("team", "")
    position_filter = request.GET.get("position", "")
    department_filter = request.GET.get("department", "")

    # Apply filters dynamically
    if search:
        team_members_qs = team_members_qs.filter(
            Q(name__icontains=search) | Q(email__icontains=search)
        )
    if team_filter:
        team_members_qs = team_members_qs.filter(team__name__iexact=team_filter)
    if position_filter:
        team_members_qs = team_members_qs.filter(job_Position__iexact=position_filter)
    if department_filter:
        team_members_qs = team_members_qs.filter(department__name__iexact=department_filter)

    # Dropdown data
    teams = User.objects.filter(team__isnull=False).values_list('team__name', flat=True).distinct()
    positions = User.objects.filter(job_Position__isnull=False).values_list('job_Position', flat=True).distinct()
    departments = User.objects.filter(department__isnull=False).values_list('department__name', flat=True).distinct()

    # --- Keep your existing stats and report context below ---
    annotated = team_members_qs.annotate(status_clean=Lower(Trim('status')))
    active_members = annotated.filter(status_clean='active').count()
    inactive_members = annotated.filter(status_clean='inactive').count()
    total_users = team_members_qs.count()

        # =========================================================
    # 🕒 CURRENT DAY REPORT SCHEDULE
    # =========================================================

    today = localtime().date()
    current_day = today.weekday()

    today_leave = TeamReportLeave.objects.filter(
        team_lead=team_lead,
        date=today,
    ).first()

    morning_on_leave = bool(today_leave and today_leave.morning_leave)
    evening_on_leave = bool(today_leave and today_leave.evening_leave)

    def get_today_schedule(report_type, on_leave):
        if on_leave:
            return None

        # A date-specific setting has first priority.
        date_setting = ReportDateSchedule.objects.filter(
            team_lead=team_lead,
            date=today,
            report_type=report_type,
        ).first()

        if date_setting is not None:
            return date_setting if date_setting.is_active else None

        # Saturday timing, if saved; otherwise regular timing.
        kinds = (
            ("saturday", "regular")
            if today.weekday() == 5
            else ("regular",)
        )

        for kind in kinds:
            default_setting = (
                ReportDefaultSchedule.objects.filter(
                    team_lead=team_lead,
                    effective_from__lte=today,
                    day_kind=kind,
                    report_type=report_type,
                )
                .order_by("-effective_from", "-id")
                .first()
            )
            if default_setting is not None:
                return default_setting if default_setting.is_active else None

        return None

    current_morning = get_today_schedule(
        "morning", morning_on_leave
    )
    current_evening = get_today_schedule(
        "evening", evening_on_leave
    )

    morning_start = (
        current_morning.start_time if current_morning else None
    )
    morning_end = (
        current_morning.end_time if current_morning else None
    )
    evening_start = (
        current_evening.start_time if current_evening else None
    )
    evening_end = (
        current_evening.end_time if current_evening else None
    )

        # =========================================================
    # 📝 TEAM LEAD REPORT SUBMISSION
    # =========================================================

    morning_allowed = (
        is_within_time_range(
            morning_start,
            morning_end
        )
        if morning_start and morning_end
        else False
    )

    evening_allowed = (
        is_within_time_range(
            evening_start,
            evening_end
        )
        if evening_start and evening_end
        else False
    )

    # ---------------------------------------------------------
    # HANDLE MORNING / EVENING REPORT SUBMISSION
    # ---------------------------------------------------------

    if request.method == "POST":

        now_time = localtime().time()

        # =====================================================
        # 🌅 MORNING REPORT
        # =====================================================

        if "morning_submit" in request.POST:

            if not current_morning:

                messages.error(
                    request,
                    "Morning report submission is not available today."
                )

            elif is_within_time_range(
                morning_start,
                morning_end,
                now_time
            ):

                report_text = request.POST.get(
                    "morning_report",
                    ""
                ).strip()

                status = request.POST.get(
                    "morning_status"
                )

                if report_text and status:

                    MorningReport.objects.create(
                        user=team_lead,
                        department=(
                            str(team_lead.department.name)
                            if team_lead.department
                            else "Unassigned"
                        ),
                        team=(
                            str(team_lead.team.name)
                            if team_lead.team
                            else "Unassigned"
                        ),
                        report_text=report_text,
                        status=status,
                    )
                    messages.success(request, "Morning report submitted successfully.")

                    return redirect(
                        "teamlead_dashboard"
                    )

                else:

                    messages.error(
                        request,
                        "Please fill in all fields before submitting."
                    )

            else:

                messages.error(
                    request,
                    f"You can only submit morning reports between "
                    f"{morning_start.strftime('%I:%M %p')} and "
                    f"{morning_end.strftime('%I:%M %p')}."
                )

        # =====================================================
        # 🌇 EVENING REPORT
        # =====================================================

        elif "evening_submit" in request.POST:

            if not current_evening:

                messages.error(
                    request,
                    "Evening report submission is not available today."
                )

            elif is_within_time_range(
                evening_start,
                evening_end,
                now_time
            ):

                report_text = request.POST.get(
                    "evening_report",
                    ""
                ).strip()

                status = request.POST.get(
                    "evening_status"
                )

                if report_text and status:

                    EveningReport.objects.create(
                        user=team_lead,
                        department=(
                            str(team_lead.department.name)
                            if team_lead.department
                            else "Unassigned"
                        ),
                        team=(
                            str(team_lead.team.name)
                            if team_lead.team
                            else "Unassigned"
                        ),
                        report_text=report_text,
                        status=status,
                    )
                    messages.success(request, "Evening report submitted successfully.")
                    return redirect(
                        "teamlead_dashboard"
                    )

                else:

                    messages.error(
                        request,
                        "Please fill in all fields before submitting."
                    )

            else:

                messages.error(
                    request,
                    f"You can only submit evening reports between "
                    f"{evening_start.strftime('%I:%M %p')} and "
                    f"{evening_end.strftime('%I:%M %p')}."
                )

    # =========================================================
    # 📊 TEAM LEAD REPORTS — LAST 24 HOURS
    # =========================================================

    report_cutoff = now() - timedelta(hours=24)

    morning_reports = MorningReport.objects.filter(
        user=team_lead,
        created_at__gte=report_cutoff
    ).values(
        "report_text",
        "status",
        "created_at"
    )

    for report in morning_reports:
        report["type"] = "Morning"

    evening_reports = EveningReport.objects.filter(
        user=team_lead,
        created_at__gte=report_cutoff
    ).values(
        "report_text",
        "status",
        "created_at"
    )

    for report in evening_reports:
        report["type"] = "Evening"

    all_reports = sorted(
        list(morning_reports) + list(evening_reports),
        key=lambda report: report["created_at"],
        reverse=True
    )


    # =========================================================
    # 📅 ALL SAVED REPORT SCHEDULES
    # =========================================================

    schedule_qs = ReportTimeSetting.objects.filter(
        team_lead=team_lead
    ).order_by(
        "day_of_week",
        "report_type"
    )

    # report_schedule = {}

    # for day_value, day_name in report_days:

    #     morning = schedule_qs.filter(
    #         day_of_week=day_value,
    #         report_type="morning",
    #     ).first()

    #     evening = schedule_qs.filter(
    #         day_of_week=day_value,
    #         report_type="evening",
    #     ).first()

    #     report_schedule[day_value] = {
    #         "day_name": day_name,
    #         "morning": morning,
    #         "evening": evening,
    #     }

    report_schedule = []

    for day_value, day_name in report_days:

        morning = schedule_qs.filter(
            day_of_week=day_value,
            report_type="morning",
        ).first()

        evening = schedule_qs.filter(
            day_of_week=day_value,
            report_type="evening",
        ).first()

        report_schedule.append({
            "day_value": day_value,
            "day_name": day_name,
            "morning": morning,
            "evening": evening,
        })
            # Dates to show in the report schedule modal
    selected_month = request.GET.get("month", "")
    try:
        year, month = map(int, selected_month.split("-"))
        if not 1 <= month <= 12:
            raise ValueError
        first_date = date(year, month, 1)
    except (ValueError, TypeError):
        first_date = localtime().date().replace(day=1)

    year, month = first_date.year, first_date.month
    days_in_month = calendar.monthrange(year, month)[1]

    saved_schedules = {
        (item.date, item.report_type): item
        for item in ReportDateSchedule.objects.filter(
            team_lead=team_lead,
            date__year=year,
            date__month=month,
        )
    }

    saved_leaves = {
        item.date: item
        for item in TeamReportLeave.objects.filter(
            team_lead=team_lead,
            date__year=year,
            date__month=month,
        )
    }

    last_date = date(year, month, days_in_month)
    default_schedules = list(
        ReportDefaultSchedule.objects.filter(
            team_lead=team_lead,
            effective_from__lte=last_date,
        ).order_by("-effective_from", "-id")
    )

    def timing_for(report_date, report_type):
        # A saved date setting takes priority, even when disabled.
        date_setting = saved_schedules.get((report_date, report_type))
        if date_setting is not None:
            return date_setting, "date"

        # Saturday can have its own default. Otherwise use regular timing.
        kinds = (
            ("saturday", "regular")
            if report_date.weekday() == 5
            else ("regular",)
        )

        for kind in kinds:
            for setting in default_schedules:
                if (
                    setting.effective_from <= report_date
                    and setting.day_kind == kind
                    and setting.report_type == report_type
                ):
                    return setting, kind

        return None, None

    date_report_schedule = []
    for day_number in range(1, days_in_month + 1):
        report_date = date(year, month, day_number)
        morning, morning_source = timing_for(report_date, "morning")
        evening, evening_source = timing_for(report_date, "evening")

        date_report_schedule.append({
            "date": report_date,
            "weekday": report_date.strftime("%A"),
            "is_sunday": report_date.weekday() == 6,
            "morning": morning,
            "evening": evening,
            "morning_source": morning_source,
            "evening_source": evening_source,
            "leave": saved_leaves.get(report_date),
        })

        # My tasks and assigned tasks currently in progress
        ongoing_tasks = Task.objects.filter(
        created_by=team_lead,
        status="in_progress",
    ).select_related(
        "assigned_to"
    ).order_by("-created_at", "-id")

    return render(
        request,
        "team_lead/teamlead_dashboard.html",
        {
            "team_lead": team_lead,
            "ongoing_tasks": ongoing_tasks,
            "team_members": team_members_qs,
            "teams": teams,
            "positions": positions,
            "departments": departments,

            "total_users": total_users,
            "active_members": active_members,
            "inactive_members": inactive_members,

            "morning_start": morning_start,
            "morning_end": morning_end,
            "evening_start": evening_start,
            "evening_end": evening_end,
            "morning_allowed": morning_allowed,
            "evening_allowed": evening_allowed,
            "all_reports": all_reports,
            "morning_on_leave": morning_on_leave,
            "evening_on_leave": evening_on_leave,

            "report_schedule": report_schedule,
            "date_report_schedule": date_report_schedule,
            "selected_month": first_date.strftime("%Y-%m"),
            "selected_month_label": first_date.strftime("%B %Y"),
            "current_day": current_day,
        }
    )


def is_within_time_range(start_time, end_time, now=None):
    now = now or localtime().time()
    return start_time <= now <= end_time

# =========================================================
# 🎥 TEAM LEAD GOOGLE MEET
# =========================================================

@never_cache
def teamlead_google_meet(request):

    # =====================================================
    # 🔒 TEAM LEAD ONLY
    # =====================================================
    if (
        not request.session.get("user_id")
        or request.session.get("position") != "team_lead"
    ):
        return redirect("index")

    # =====================================================
    # 👤 LOGGED-IN TEAM LEAD
    # =====================================================
    team_lead = get_object_or_404(
        User,
        id=request.session["user_id"]
    )

    # =====================================================
    # 👥 TEAM MEMBERS
    # =====================================================
    team_members = User.objects.filter(
        department=team_lead.department,
        job_Position__iexact="Team Member"
    ).exclude(
        id=team_lead.id
    ).order_by("name")

    # =====================================================
    # 📋 ELIGIBLE MEMBERS
    # =====================================================
    eligible_members = User.objects.filter(
        department=team_lead.department,
        job_Position__iexact="Team Member"
    ).exclude(
        id=team_lead.id
    )

    # =====================================================
    # 📩 POST
    # =====================================================
    if request.method == "POST":

        action = request.POST.get(
            "action",
            "create_meeting"
        ).strip()

        # =================================================
        # COMMON FORM DATA
        # =================================================
        purpose = request.POST.get(
            "purpose",
            ""
        ).strip()

        meeting_date = request.POST.get(
            "date",
            ""
        ).strip()

        meeting_time = request.POST.get(
            "time",
            ""
        ).strip()

        meeting_link = request.POST.get(
            "meeting_link",
            ""
        ).strip()

        participant_type = request.POST.get(
            "participant_type",
            "all"
        ).strip()

        selected_member_ids = request.POST.getlist(
            "participants"
        )

        # =================================================
        # REQUIRED VALIDATION
        # =================================================
        if not purpose:
            messages.error(
                request,
                "Please enter the meeting purpose."
            )
            return redirect("teamlead_google_meet")

        if not meeting_date:
            messages.error(
                request,
                "Please select a meeting date."
            )
            return redirect("teamlead_google_meet")

        if not meeting_time:
            messages.error(
                request,
                "Please select a meeting time."
            )
            return redirect("teamlead_google_meet")

        if not meeting_link:
            messages.error(
                request,
                "Please enter the Google Meet link."
            )
            return redirect("teamlead_google_meet")

        # =================================================
        # DATE VALIDATION
        # =================================================
        try:
            parsed_date = datetime.strptime(
                meeting_date,
                "%Y-%m-%d"
            ).date()

        except ValueError:
            messages.error(
                request,
                "Invalid meeting date."
            )
            return redirect("teamlead_google_meet")

        # =================================================
        # TIME VALIDATION
        # =================================================
        try:
            parsed_time = datetime.strptime(
                meeting_time,
                "%H:%M"
            ).time()

        except ValueError:
            messages.error(
                request,
                "Invalid meeting time."
            )
            return redirect("teamlead_google_meet")

        # =================================================
        # GOOGLE MEET LINK VALIDATION
        # =================================================
        if not (
            meeting_link.startswith(
                "https://meet.google.com/"
            )
            or meeting_link.startswith(
                "http://meet.google.com/"
            )
        ):
            messages.error(
                request,
                "Please enter a valid Google Meet link."
            )
            return redirect("teamlead_google_meet")

        # =================================================
        # 👥 PARTICIPANTS VALIDATION
        # =================================================
        if participant_type == "all":

            final_members = eligible_members

        elif participant_type == "selected":

            if not selected_member_ids:

                messages.error(
                    request,
                    "Please select at least one team member."
                )

                return redirect(
                    "teamlead_google_meet"
                )

            final_members = eligible_members.filter(
                id__in=selected_member_ids
            )

            if not final_members.exists():

                messages.error(
                    request,
                    "Selected team members are not valid."
                )

                return redirect(
                    "teamlead_google_meet"
                )

        else:

            messages.error(
                request,
                "Invalid participant selection."
            )

            return redirect(
                "teamlead_google_meet"
            )

        # =================================================
        # ✏️ EDIT MEETING
        # =================================================
        if action == "edit_meeting":

            meeting_id = request.POST.get(
                "meeting_id"
            )

            if not meeting_id:

                messages.error(
                    request,
                    "Meeting ID is missing."
                )

                return redirect(
                    "teamlead_google_meet"
                )

            meeting = get_object_or_404(
                GoogleMeeting,
                id=meeting_id,
                created_by=team_lead
            )
            # ---------------------------------------------
            # PREVENT EDITING ENDED MEETINGS
            # ---------------------------------------------
            if meeting.meeting_status == "ended":
                messages.error(
                    request,
                    "This meeting has already ended and can no longer be edited."
                )
                return redirect("teamlead_google_meet")

            # ---------------------------------------------
            # UPDATE
            # ---------------------------------------------
            meeting.purpose = purpose
            meeting.date = parsed_date
            meeting.time = parsed_time
            meeting.meeting_link = meeting_link

            meeting.save()

            # ---------------------------------------------
            # UPDATE PARTICIPANTS
            # ---------------------------------------------
            meeting.participants.set(
                final_members
            )

            messages.success(
                request,
                "Google Meet updated successfully."
            )

            return redirect(
                "teamlead_google_meet"
            )

        # =================================================
        # ➕ CREATE MEETING
        # =================================================
        meeting = GoogleMeeting.objects.create(
            created_by=team_lead,
            date=parsed_date,
            time=parsed_time,
            meeting_link=meeting_link,
            purpose=purpose,
        )

        meeting.participants.set(
            final_members
        )

        messages.success(
            request,
            "Google Meet scheduled successfully."
        )

        return redirect(
            "teamlead_google_meet"
        )

    # =====================================================
    # 📅 MEETINGS
    # =====================================================
    meetings = GoogleMeeting.objects.filter(
        Q(created_by=team_lead) | Q(participants=team_lead)
    ).distinct().prefetch_related(
        "participants"
    ).order_by(
        "date",
        "time"
    )

    # =====================================================
    # CONTEXT
    # =====================================================
    context = {
        "team_lead": team_lead,
        "team_members": team_members,
        "meetings": meetings,
    }

    return render(
        request,
        "team_lead/teamlead_google_meet.html",
        context
    )


@never_cache
def teamlead_announcements(request):
    # 🔒 Restrict access
    if (
        not request.session.get("user_id")
        or request.session.get("position") != "team_lead"
    ):
        return redirect("index")

    try:
        team_lead = User.objects.get(
            id=request.session["user_id"]
        )
    except User.DoesNotExist:
        request.session.flush()
        return redirect("index")

    if team_lead.job_Position.lower() != "team lead":
        return redirect("index")

    # =========================================================
    # 📢 ANNOUNCEMENT ACTIONS
    # =========================================================

    if request.method == "POST":

        action = request.POST.get("action", "").strip()

        # =====================================================
        # ✏️ EDIT ANNOUNCEMENT
        # =====================================================

        if action == "edit_announcement":

            announcement_id = request.POST.get(
                "announcement_id"
            )

            edited_message = request.POST.get(
                "message", ""
            ).strip()

            if not edited_message:
                messages.error(
                    request,
                    "Announcement message cannot be empty."
                )
                return redirect(
                    "teamlead_announcements"
                )

            try:
                announcement = Announcement.objects.get(
                    id=announcement_id,
                    created_by=team_lead,
                    is_active=True
                )
            except Announcement.DoesNotExist:
                messages.error(
                    request,
                    "Announcement not found."
                )
                return redirect(
                    "teamlead_announcements"
                )

            # Update the SAME announcement.
            # All existing recipients will see the
            # updated message.
            announcement.message = edited_message
            announcement.save()

            messages.success(
                request,
                "Announcement updated successfully."
            )

            return redirect(
                "teamlead_announcements"
            )


        # =====================================================
        # 🗑️ DELETE ANNOUNCEMENT
        # =====================================================

        elif action == "delete_announcement":

            announcement_id = request.POST.get(
                "announcement_id"
            )

            try:
                announcement = Announcement.objects.get(
                    id=announcement_id,
                    created_by=team_lead,
                    is_active=True
                )
            except Announcement.DoesNotExist:
                messages.error(
                    request,
                    "Announcement not found."
                )
                return redirect(
                    "teamlead_announcements"
                )

            # Soft delete.
            # Recipient records stay safe, but the
            # announcement will no longer be displayed.
            announcement.is_active = False
            announcement.save(
                update_fields=["is_active"]
            )

            messages.success(
                request,
                "Announcement deleted successfully."
            )

            return redirect(
                "teamlead_announcements"
            )


        # =====================================================
        # 📢 CREATE ANNOUNCEMENT
        # =====================================================

        elif action == "create_announcement":

            announcement_message = request.POST.get(
                "message", ""
            ).strip()

            send_to = request.POST.get(
                "send_to",
                "team"
            )

            selected_member_ids = request.POST.getlist(
                "recipient_ids"
            )

            # Empty message validation
            if not announcement_message:
                messages.error(
                    request,
                    "Please enter an announcement message."
                )
                return redirect(
                    "teamlead_announcements"
                )

            # -------------------------------------------------
            # Create announcement
            # -------------------------------------------------

            announcement = Announcement.objects.create(
                title="Team Lead Announcement",
                message=announcement_message,
                created_by=team_lead
            )

            # -------------------------------------------------
            # Only Team Members under this Team Lead's
            # department
            # -------------------------------------------------

            eligible_members = User.objects.filter(
                department=team_lead.department,
                job_Position__iexact="Team Member"
            ).exclude(
                id=team_lead.id
            )

            # -------------------------------------------------
            # Send to Entire Team
            # -------------------------------------------------

            if send_to == "team":

                recipients = [
                    AnnouncementRecipient(
                        announcement=announcement,
                        recipient=member
                    )
                    for member in eligible_members
                ]

                if recipients:
                    AnnouncementRecipient.objects.bulk_create(
                        recipients
                    )

                messages.success(
                    request,
                    "Announcement sent to the entire team."
                )


            # -------------------------------------------------
            # Send to Selected Members
            # -------------------------------------------------

            elif send_to == "selected":

                selected_members = eligible_members.filter(
                    id__in=selected_member_ids
                )

                if not selected_members.exists():

                    announcement.delete()

                    messages.error(
                        request,
                        "Please select at least one team member."
                    )

                    return redirect(
                        "teamlead_announcements"
                    )

                recipients = [
                    AnnouncementRecipient(
                        announcement=announcement,
                        recipient=member
                    )
                    for member in selected_members
                ]

                AnnouncementRecipient.objects.bulk_create(
                    recipients
                )

                messages.success(
                    request,
                    "Announcement sent to selected members."
                )

            return redirect(
                "teamlead_announcements"
            )


    # =========================================================
    # 📢 PREVIOUS ANNOUNCEMENTS
    # =========================================================

    announcements = Announcement.objects.filter(
    created_by=team_lead,
    is_active=True
).prefetch_related(
    "recipients__recipient"
).order_by(
    "-created_at"
)[:20]


    # =========================================================
    # 👥 TEAM MEMBERS
    # =========================================================

    team_members = User.objects.filter(
        department=team_lead.department,
        job_Position__iexact="Team Member"
    ).exclude(
        id=team_lead.id
    ).order_by("name")


    response = render(
        request,
        "team_lead/teamlead_announcements.html",
        {
            "announcements": announcements,
            "team_members": team_members,
            "team_lead": team_lead,
            "admin_received": _announcement_inbox(team_lead).filter(announcement__is_admin_announcement=True).order_by("-announcement__created_at", "-pk"),
        }
    )

    if request.method == "GET" and request.GET.get("format") == "json":
        return JsonResponse({"ok": True, "html": response.content.decode(response.charset)})
    return response

@never_cache
def teamlead_chat(request):
    user_id = request.session.get("user_id")

    position = (
        str(request.session.get("position", ""))
        .strip()
        .lower()
        .replace(" ", "_")
    )

    if not user_id or position != "team_lead":
        return redirect("index")

    current_user = get_object_or_404(User, pk=user_id)

    if request.method == "POST":
        action = request.POST.get("action")

        if action == "create_group":
            return _member_chat_create_group(
                request,
                current_user,
                group_view_name="teamlead_group_chat_view",
            )

        return _member_chat_error("Invalid action.")

    context = _member_chat_context(current_user)
    context["user_role"] = "team_lead"

    return render(
        request,
        "team_lead/teamlead_chat.html",
        context,
    )


@never_cache
def teamlead_chat_room(request, user_id):
    from .member_chat_actions import chat_action

    session_user_id = request.session.get("user_id")

    position = (
        str(request.session.get("position", ""))
        .strip()
        .lower()
        .replace(" ", "_")
    )

    if not session_user_id or position != "team_lead":
        return redirect("index")

    current_user = get_object_or_404(
        User,
        pk=session_user_id,
    )

    other_user = get_object_or_404(
        User,
        pk=user_id,
    )

    if current_user.pk == other_user.pk:
        return redirect("teamlead_chat")

    # Preserve existing conversations in either user ordering.
    room = (
        Conversation.objects.filter(
            Q(user1=current_user, user2=other_user)
            | Q(user1=other_user, user2=current_user)
        )
        .order_by("pk")
        .first()
    )

    if room is None:
        first, second = sorted(
            [current_user, other_user],
            key=lambda person: person.pk,
        )

        room, _ = Conversation.objects.get_or_create(
            user1=first,
            user2=second,
        )

    if request.method == "POST":
        action = request.POST.get("action")

        if action == "create_group":
            return _member_chat_create_group(
                request,
                current_user,
                group_view_name="teamlead_group_chat_view",
            )

        return chat_action(
            request,
            current_user,
            room,
            group=False,
        )

    messages_qs = (
        room.messages
        .exclude(hidden_for=current_user)
        .select_related("sender")
        .order_by("created_at", "pk")
    )

    # Used by the chat UI to reconcile messages without refreshing.
    if request.GET.get("messages") == "1":
        return JsonResponse({
            "ok": True,
            "items": [
                _member_chat_payload(message)
                for message in messages_qs
            ],
        })

    context = _member_chat_context(current_user)

    context.update({
        "user_role": "team_lead",
        "room": room,
        "other_user": other_user,
        "messages": messages_qs,
    })

    return render(
        request,
        "team_lead/teamlead_chat_room.html",
        context,
    )

@never_cache
def teamlead_group_chat_view(request, group_id):
    import logging

    from asgiref.sync import async_to_sync
    from channels.layers import get_channel_layer

    from .chat_receipts import broadcast_chat_activity
    from .member_chat_actions import chat_action

    user_id = request.session.get("user_id")

    position = (
        str(request.session.get("position", ""))
        .strip()
        .lower()
        .replace(" ", "_")
    )

    if not user_id or position != "team_lead":
        return redirect("index")

    current_user = get_object_or_404(User, pk=user_id)
    group = get_object_or_404(Group, pk=group_id)

    membership = get_object_or_404(
        GroupMember,
        group=group,
        user=current_user,
    )

    can_manage = (
        group.created_by_id == current_user.pk
        or membership.role == "admin"
    )

    def member_state():
        current_membership = GroupMember.objects.filter(
            group=group,
            user=current_user,
        ).first()

        permitted = bool(
            current_membership
            and (
                group.created_by_id == current_user.pk
                or current_membership.role == "admin"
            )
        )

        entries = list(
            GroupMember.objects.filter(group=group)
            .select_related("user")
            .order_by("user__name", "pk")
        )

        available_users = []

        if permitted:
            available_users = [
                {
                    "id": person.pk,
                    "name": person.name,
                }
                for person in User.objects.exclude(
                    group_memberships__group=group
                ).order_by("name", "pk")
            ]

        return {
            "ok": True,
            "can_manage": permitted,
            "count": len(entries),
            "members": [
                {
                    "id": entry.user_id,
                    "name": entry.user.name,
                    "role": entry.role,
                    "is_creator": (
                        entry.user_id == group.created_by_id
                    ),
                    "is_self": (
                        entry.user_id == current_user.pk
                    ),
                    "profile": (
                        entry.user.profile_image.url
                        if entry.user.profile_image
                        else None
                    ),
                }
                for entry in entries
            ],
            "all_users": available_users,
        }

    def notify_members():
        layer = get_channel_layer()

        if layer:
            try:
                async_to_sync(layer.group_send)(
                    f"group_{group.pk}",
                    {"type": "chat_change"},
                )
            except Exception:
                logging.getLogger(__name__).exception(
                    "Group membership update delivery failed."
                )

        async_to_sync(broadcast_chat_activity)(
            True,
            group.pk,
        )

    if request.method == "POST":
        action = request.POST.get("action", "")

        if action == "create_group":
            return _member_chat_create_group(
                request,
                current_user,
                group_view_name="teamlead_group_chat_view",
            )

        if action in {
            "edit_message",
            "delete_for_me",
            "delete_for_everyone",
            "clear_chat",
        }:
            return chat_action(
                request,
                current_user,
                group,
                group=True,
            )

        if action not in {
            "add_member",
            "remove_member",
            "make_admin",
            "remove_admin",
        }:
            return _member_chat_error("Invalid action.")

        if not can_manage:
            return _member_chat_error(
                "Only the group creator or an admin "
                "can manage members.",
                403,
            )

        try:
            target_id = int(
                request.POST.get("user_id", "")
            )
        except (ValueError, TypeError):
            return _member_chat_error(
                "Select a valid member."
            )

        with transaction.atomic():
            # Serialize membership changes for this group.
            locked_group = (
                Group.objects.select_for_update()
                .get(pk=group.pk)
            )

            membership = get_object_or_404(
                GroupMember,
                group=group,
                user=current_user,
            )

            if (
                locked_group.created_by_id != current_user.pk
                and membership.role != "admin"
            ):
                return _member_chat_error(
                    "You no longer have permission "
                    "to manage this group.",
                    403,
                )

            if action == "add_member":
                target_user = get_object_or_404(
                    User,
                    pk=target_id,
                )

                _, created = (
                    GroupMember.objects.get_or_create(
                        group=group,
                        user=target_user,
                    )
                )

                notice = (
                    f"{target_user.name} added to the group."
                    if created
                    else f"{target_user.name} is already a member."
                )

            else:
                target = get_object_or_404(
                    GroupMember.objects.select_related("user"),
                    group=group,
                    user_id=target_id,
                )

                if target_id == locked_group.created_by_id:
                    return _member_chat_error(
                        "The group creator cannot be "
                        "removed or demoted."
                    )

                if target_id == current_user.pk:
                    return _member_chat_error(
                        "You cannot remove yourself or "
                        "change your own admin role here."
                    )

                if action == "remove_member":
                    member_count = (
                        GroupMember.objects.filter(
                            group=group
                        ).count()
                    )

                    if member_count <= 2:
                        return _member_chat_error(
                            "Keep at least two people "
                            "in the group."
                        )

                    name = target.user.name
                    target.delete()

                    notice = (
                        f"{name} removed from the group."
                    )

                else:
                    target.role = (
                        "admin"
                        if action == "make_admin"
                        else "member"
                    )

                    target.save(update_fields=["role"])

                    notice = (
                        f"{target.user.name} is now an admin."
                        if action == "make_admin"
                        else (
                            "Admin access removed for "
                            f"{target.user.name}."
                        )
                    )

            transaction.on_commit(notify_members)

        return JsonResponse({
            **member_state(),
            "message": notice,
        })

    if request.GET.get("members") == "1":
        return JsonResponse(member_state())

    messages_qs = (
        GroupMessage.objects.filter(group=group)
        .exclude(hidden_for=current_user)
        .select_related("sender")
        .order_by("timestamp", "pk")
    )

    if request.GET.get("messages") == "1":
        return JsonResponse({
            "ok": True,
            "items": [
                _member_chat_payload(message, True)
                for message in messages_qs
            ],
        })

    context = _member_chat_context(current_user)

    context.update({
        "user_role": "team_lead",
        "group": group,
        "messages": messages_qs,
        "members": (
            GroupMember.objects.filter(group=group)
            .select_related("user")
            .order_by("user__name", "pk")
        ),
        "all_users": (
            User.objects.exclude(
                group_memberships__group=group
            ).order_by("name", "pk")
        ),
        "can_manage_group": can_manage,
    })

    return render(
        request,
        "team_lead/teamlead_group_chat.html",
        context,
    )



def teammember_chat(request):
    # Check session-based login
    user_id = request.session.get("user_id")
    if not user_id:
        return redirect("login_view")

    current_user = get_object_or_404(User, id=user_id)

    # Users list excluding current user
    users = list(User.objects.exclude(id=current_user.id))
    users.sort(key=lambda u: u.name.lower())

    extra_contacts = ExtraContact.objects.all()
    groups = Group.objects.filter(memberships__user=current_user).distinct()

    # Handle Create Group from modal
    if request.method == "POST" and request.POST.get("action") == "create_group":
        group_name = request.POST.get("group_name")
        member_ids = request.POST.getlist("members")  # list of user ids

        if group_name:
            with transaction.atomic():
                # ✅ Create group with required created_by field
                group = Group.objects.create(name=group_name, created_by=current_user)

                # Add creator as group member
                GroupMember.objects.create(group=group, user=current_user)

                # Add other selected members
                for uid in member_ids:
                    user = User.objects.get(id=uid)
                    GroupMember.objects.get_or_create(group=group, user=user)

            messages.success(request, f"Group '{group_name}' created successfully!")
            return redirect("teammember_chat")
        else:
            messages.error(request, "Please provide a group name.")

    context = {
        "current_user": current_user,
        "users": users,
        "extra_contacts": extra_contacts,
        "groups": groups,
        "role": "teammember",
    }

    return render(request, "team_member/teammember_chat.html", context)

def get_or_create_room(user1, user2):
    
    if user1.id > user2.id:
        user1, user2 = user2, user1
    room, created = ChatRoom.objects.get_or_create(user1=user1, user2=user2)
    return room


def chat_room(request, user_id):
    current_user_id = request.session.get("user_id")
    if not current_user_id:
        return redirect("login_view")

    current_user = get_object_or_404(User, id=current_user_id)
    other_user = get_object_or_404(User, id=user_id)
    room = get_or_create_room(current_user, other_user)
    users = User.objects.exclude(id=current_user.id).order_by('name')
    extra_contacts = User.objects.filter(status='inactive').exclude(id=current_user.id)
    messages_list = room.messages.order_by("timestamp")
    groups = Group.objects.filter(memberships__user=current_user).distinct()

    # ✅ Normalize role safely
    user_role = request.session.get("position", "").strip().lower().replace(" ", "_")

    # Handle group creation
    if request.method == "POST" and request.POST.get("action") == "create_group":
        group_name = request.POST.get("group_name")
        member_ids = request.POST.getlist("members")

        if group_name:
            with transaction.atomic():
                group = Group.objects.create(name=group_name, created_by=current_user)
                GroupMember.objects.create(group=group, user=current_user)
                for uid in member_ids:
                    user = User.objects.get(id=uid)
                    GroupMember.objects.get_or_create(group=group, user=user)

            messages.success(request, f"Group '{group_name}' created successfully!")
            return redirect("chat_room", user_id=user_id)
        else:
            messages.error(request, "Please provide a group name.")

    context = {
        "room": room,
        "current_user": current_user,
        "other_user": other_user,
        "messages": messages_list,
        "users": users,
        "extra_contacts": extra_contacts,
        "groups": groups,
        "user_role": user_role,   # ✅ correct key used in template
    }

    return render(request, "chat_room.html", context)


def group_chat_view(request, group_id):
    user_id = request.session.get("user_id")
    if not user_id:
        return redirect("login_view")

    current_user = get_object_or_404(User, id=user_id)

    # ✅ Correct role logic
    user_role = request.session.get("position", "").strip().lower().replace(" ", "_")

    group = get_object_or_404(Group, id=group_id)
    messages_qs = GroupMessage.objects.filter(group=group).order_by('timestamp')
    members = GroupMember.objects.filter(group=group).select_related('user')
    all_users = User.objects.exclude(id__in=[m.user.id for m in members])

    users = list(User.objects.exclude(id=current_user.id))
    users.sort(key=lambda u: u.name.lower())
    extra_contacts = ExtraContact.objects.all()
    groups = Group.objects.filter(memberships__user=current_user).distinct()

    if request.method == "POST":
        action = request.POST.get("action")
        if action == "add_member":
            new_user_id = request.POST.get("user_id")
            new_user = get_object_or_404(User, id=new_user_id)
            GroupMember.objects.get_or_create(group=group, user=new_user)
            messages.success(request, f"{new_user.name} added to {group.name}.")
            return redirect('group_chat_view', group_id=group.id)

        elif action == "remove_member":
            rem_user_id = request.POST.get("user_id")
            member = get_object_or_404(GroupMember, group=group, user_id=rem_user_id)
            member.delete()
            messages.warning(request, "Member removed successfully.")
            return redirect('group_chat_view', group_id=group.id)

    # AJAX partial
    if request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.GET.get('partial') == 'true':
        return render(request, 'partials/group_chat_partial.html', {
            "group": group,
            "messages": messages_qs,
            "members": members,
            "all_users": all_users,
            "current_user": current_user,
            "users": users,
            "extra_contacts": extra_contacts,
            "groups": groups,
            "user_role": user_role,  # ✅ pass correct role
        })

    # Full page render
    context = {
        "group": group,
        "messages": messages_qs,
        "members": members,
        "all_users": all_users,
        "current_user": current_user,
        "users": users,
        "extra_contacts": extra_contacts,
        "groups": groups,
        "user_role": user_role,  # ✅ correct value
    }
    return render(request, 'group_chat.html', context)


from .models import ReportDateSchedule, ReportDefaultSchedule, TeamReportLeave

@never_cache
def teammember_dashboard(request):
    if (
        not request.session.get("user_id")
        or request.session.get("position") != "team_member"
    ):
        return redirect("login_view")

    # ✅ Fetch logged-in team member
    try:
        team_member = User.objects.get(
            id=request.session["user_id"]
        )
    except (User.DoesNotExist, ValueError, TypeError):
        request.session.flush()
        return redirect("login_view")

    team_name = team_member.team

    # =========================================================
    # 🕒 LOGIN TIME
    # =========================================================
    login_time_str = request.session.get("login_time")
    try:
        login_time = parse_datetime(login_time_str) if login_time_str else None
    except (ValueError, TypeError):
        login_time = None

    if login_time:
        if is_naive(login_time):
            login_time = make_aware(login_time)
        login_time = localtime(login_time)

    # =========================================================
    # 📅 TODAY'S DAY
    # Monday = 0
    # Tuesday = 1
    # ...
    # Saturday = 5
    # Sunday = 6
    # =========================================================
    current_datetime = localtime()
    today = current_datetime.date()
    current_day = today.weekday()

    # Use the lead of this member's team.
    team_lead = None
    if team_member.team_id:
        team_lead = User.objects.filter(
            team_id=team_member.team_id,
            job_Position__iexact="Team Lead",
        ).first()

    today_leave = (
        TeamReportLeave.objects.filter(team_lead=team_lead, date=today).first()
        if team_lead else None
    )
    morning_on_leave = bool(today_leave and today_leave.morning_leave)
    evening_on_leave = bool(today_leave and today_leave.evening_leave)

    def get_today_schedule(report_type, on_leave):
        # This follows the same precedence as the Team Lead dashboard.
        if team_lead is None or on_leave:
            return None
        date_setting = ReportDateSchedule.objects.filter(
            team_lead=team_lead, date=today, report_type=report_type,
        ).first()
        if date_setting is not None:
            # An explicitly disabled date must not fall back to default timings.
            return date_setting if date_setting.is_active else None
        kinds = ("saturday", "regular") if today.weekday() == 5 else ("regular",)
        for kind in kinds:
            default_setting = ReportDefaultSchedule.objects.filter(
                team_lead=team_lead, effective_from__lte=today,
                day_kind=kind, report_type=report_type,
            ).order_by("-effective_from", "-id").first()
            if default_setting is not None:
                return default_setting if default_setting.is_active else None
        return None

    morning_setting = get_today_schedule("morning", morning_on_leave)
    evening_setting = get_today_schedule("evening", evening_on_leave)
    morning_start = morning_setting.start_time if morning_setting else None
    morning_end = morning_setting.end_time if morning_setting else None
    evening_start = evening_setting.start_time if evening_setting else None
    evening_end = evening_setting.end_time if evening_setting else None

    # =========================================================
    # 📝 HANDLE REPORT SUBMISSIONS
    # =========================================================
    if request.method == "POST":
        now_time = localtime().time()

        # =====================================================
        # 🌅 MORNING REPORT
        # =====================================================
        if "morning_submit" in request.POST:

            # No schedule configured for today
            if morning_on_leave:
                messages.warning(request, "Morning report is marked as leave/holiday today.")
            elif not morning_setting:
                messages.error(
                    request,
                    "Morning report submission is not available today."
                )

            elif is_within_time_range(
                morning_start,
                morning_end,
                now_time
            ):
                report_text = request.POST.get("morning_report", "").strip()
                status = request.POST.get("morning_status")

                if report_text and status in {"Pending", "In Progress", "Completed"}:
                    MorningReport.objects.create(
                        user=team_member,
                        department=(
                            str(team_member.department.name)
                            if team_member.department
                            else "Unassigned"
                        ),
                        team=(
                            str(team_member.team.name)
                            if team_member.team
                            else "Unassigned"
                        ),
                        report_text=report_text,
                        status=status,
                    )

                    messages.success(
                        request,
                        "Morning report submitted successfully."
                    )

                    return redirect("teammember_dashboard")

                else:
                    messages.error(
                        request,
                        "Please fill in all fields before submitting."
                    )

            else:
                messages.error(
                    request,
                    f"You can only submit morning reports between "
                    f"{morning_start.strftime('%I:%M %p')} and "
                    f"{morning_end.strftime('%I:%M %p')}."
                )

        # =====================================================
        # 🌇 EVENING REPORT
        # =====================================================
        elif "evening_submit" in request.POST:

            # No schedule configured for today
            if evening_on_leave:
                messages.warning(request, "Evening report is marked as leave/holiday today.")
            elif not evening_setting:
                messages.error(
                    request,
                    "Evening report submission is not available today."
                )

            elif is_within_time_range(
                evening_start,
                evening_end,
                now_time
            ):
                report_text = request.POST.get("evening_report", "").strip()
                status = request.POST.get("evening_status")

                if report_text and status in {"Pending", "In Progress", "Completed"}:
                    EveningReport.objects.create(
                        user=team_member,
                        department=(
                            str(team_member.department.name)
                            if team_member.department
                            else "Unassigned"
                        ),
                        team=(
                            str(team_member.team.name)
                            if team_member.team
                            else "Unassigned"
                        ),
                        report_text=report_text,
                        status=status,
                    )

                    messages.success(
                        request,
                        "Evening report submitted successfully."
                    )

                    return redirect("teammember_dashboard")

                else:
                    messages.error(
                        request,
                        "Please fill in all fields before submitting."
                    )

            else:
                messages.error(
                    request,
                    f"You can only submit evening reports between "
                    f"{evening_start.strftime('%I:%M %p')} and "
                    f"{evening_end.strftime('%I:%M %p')}."
                )

    if request.method == "POST":
        return redirect("teammember_dashboard")

    # =========================================================
    # 📊 FETCH REPORTS FROM LAST 24 HOURS
    # =========================================================
    report_cutoff = now() - timedelta(hours=24)

    morning_reports = list(MorningReport.objects.filter(
        user=team_member,
        created_at__gte=report_cutoff
    ).values(
        "report_text",
        "status",
        "created_at"
    ))

    for r in morning_reports:
        r["type"] = "Morning"

    evening_reports = list(EveningReport.objects.filter(
        user=team_member,
        created_at__gte=report_cutoff
    ).values(
        "report_text",
        "status",
        "created_at"
    ))

    for r in evening_reports:
        r["type"] = "Evening"

    # =========================================================
    # 🔄 COMBINE + SORT REPORTS
    # =========================================================
    all_reports = sorted(
        list(morning_reports) + list(evening_reports),
        key=lambda x: x["created_at"],
        reverse=True
    )

    # =========================================================
    # 📢 TEAM MEMBER ANNOUNCEMENTS
    # =========================================================
    announcements = Announcement.objects.filter(
        recipients__recipient=team_member,
        is_active=True,
        created_at__gte=now() - timedelta(hours=12)
    ).distinct().order_by("-created_at")

    # Seen is recorded only when the member selects an announcement card.
    # Loading the dashboard does not change read receipts.

    # =========================================================
    # 🧭 RENDER DASHBOARD
    # =========================================================
    return render(
        request,
        "team_member/teammember_dashboard.html",
        {
            "user": team_member,
            "team_name": team_name,
            "announcements": announcements,

            # Report availability
            "morning_allowed": (
                is_within_time_range(
                    morning_start,
                    morning_end
                )
                if morning_start and morning_end
                else False
            ),

            "evening_allowed": (
                is_within_time_range(
                    evening_start,
                    evening_end
                )
                if evening_start and evening_end
                else False
            ),

            # Login
            "login_time": login_time,

            # Reports
            "all_reports": all_reports,

            # Morning timing
            "morning_start": morning_start,
            "morning_end": morning_end,

            # Evening timing
            "evening_start": evening_start,
            "evening_end": evening_end,

            # Optional extra context
            "current_day": current_day,
            "team_lead": team_lead,
            "morning_on_leave": morning_on_leave,
            "evening_on_leave": evening_on_leave,
        },
    )


@never_cache
def teamlead_reports(request):
    """Team lead report list, daily holiday rows, and matching Excel export."""
    if (
        not request.session.get("user_id")
        or request.session.get("position") != "team_lead"
    ):
        return redirect("index")

    team_lead = get_object_or_404(User, id=request.session["user_id"])
    report_tab = request.GET.get("tab", "members")
    if report_tab not in ("my", "members"):
        report_tab = "members"

    show_all = request.GET.get("all") == "1"
    filter_date = None
    filter_date_str = request.GET.get("date")
    if filter_date_str:
        try:
            filter_date = datetime.strptime(filter_date_str, "%Y-%m-%d").date()
        except ValueError:
            messages.error(request, "Invalid date format.")
            return redirect("teamlead_reports")

    report_date = filter_date or localtime().date()
    if report_tab == "my":
        report_filters = {"user": team_lead}
    elif team_lead.team_id:
        report_filters = {
            "user__team_id": team_lead.team_id,
            "user__job_Position__iexact": "Team Member",
        }
    else:
        # A lead without a team cannot see unrelated unassigned users' reports.
        report_filters = {"user__id__in": []}

    morning_reports = MorningReport.objects.filter(
        **report_filters
    ).select_related("user").order_by("-created_at")
    evening_reports = EveningReport.objects.filter(
        **report_filters
    ).select_related("user").order_by("-created_at")

    report_day_leave = None
    holiday_users = []
    if not show_all:
        morning_reports = morning_reports.filter(created_at__date=report_date)
        evening_reports = evening_reports.filter(created_at__date=report_date)
        report_day_leave = TeamReportLeave.objects.filter(
            team_lead=team_lead, date=report_date
        ).first()
        if report_day_leave and (
            report_day_leave.morning_leave or report_day_leave.evening_leave
        ):
            if report_tab == "my":
                holiday_users = [team_lead]
            elif team_lead.team_id:
                holiday_users = list(
                    User.objects.filter(
                        team_id=team_lead.team_id,
                        job_Position__iexact="Team Member",
                    ).order_by("name")
                )

    if "export" in request.GET:
        wb = Workbook()
        ws = wb.active
        ws.title = "My Reports" if report_tab == "my" else "Team Members Reports"
        ws.append([
            "Date", "Time", "Report Type", "Name",
            "Team", "Department", "Report", "Status",
        ])

        for reports, report_type in (
            (morning_reports, "Morning"),
            (evening_reports, "Evening"),
        ):
            for report in reports:
                report_time = localtime(report.created_at)
                ws.append([
                    report_time.strftime("%Y-%m-%d"),
                    report_time.strftime("%I:%M %p"),
                    report_type,
                    str(report.user.name),
                    str(report.team or ""),
                    str(report.department or ""),
                    str(report.report_text or ""),
                    str(report.status or ""),
                ])

        if report_day_leave:
            for report_type, is_holiday in (
                ("Morning", report_day_leave.morning_leave),
                ("Evening", report_day_leave.evening_leave),
            ):
                if is_holiday:
                    for person in holiday_users:
                        ws.append([
                            report_date.strftime("%Y-%m-%d"),
                            "",
                            report_type,
                            str(person.name),
                            str(person.team.name) if person.team else "",
                            str(person.department.name) if person.department else "",
                            report_day_leave.reason or "Holiday",
                            "Holiday",
                        ])

        for cell in ws[1]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column, width in enumerate(
            [15, 12, 14, 22, 20, 22, 50, 16], start=1
        ):
            ws.column_dimensions[get_column_letter(column)].width = width

        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )
        filename = (
            "my_reports.xlsx" if report_tab == "my"
            else "team_members_reports.xlsx"
        )
        response["Content-Disposition"] = (
            f'attachment; filename="{filename}"'
        )
        wb.save(response)
        return response

    return render(
        request,
        "team_lead/teamlead_reports.html",
        {
            "morning_reports": morning_reports,
            "evening_reports": evening_reports,
            "show_all": show_all,
            "filter_date": filter_date,
            "report_tab": report_tab,
            "report_date": report_date,
            "report_day_leave": report_day_leave,
            "holiday_users": holiday_users,
        },
    )


@never_cache
def teamlead_project_assigning(request):
    # 🔒 Prevent KeyError: check session before accessing user
    if not request.session.get("user_id") or request.session.get("position") != "team_lead":
        return redirect("index")

    team_lead = get_object_or_404(User, id=request.session["user_id"])
    departments = Department.objects.all()

    team_members = User.objects.filter(
        team=team_lead.team,
        job_Position__icontains="team member"
    ).exclude(id=team_lead.id)

    if request.method == "POST":
        # First create the project without files/images
        project = ProjectAssign.objects.create(
            team=team_lead.team,
            department_id=request.POST.get("department"),
            assign_to_id=request.POST.get("assign_to"),
            assigned_by=team_lead,
            work_name=request.POST.get("work_name"),
            work_type=request.POST.get("work_type"),
            category=request.POST.get("category"),
            description=request.POST.get("description"),
            deadline=request.POST.get("deadline"),
            additional_notes=request.POST.get("additional_notes"),
            color_preference=request.POST.get("color_preference"),
            content_example=request.POST.get("content_example"),
            priority=request.POST.get("priority"),
        )

        # Handle multiple uploaded files
        for file in request.FILES.getlist("upload_file[]"):
            ProjectFile.objects.create(project=project, file=file)

        # Handle multiple uploaded images
        for image in request.FILES.getlist("upload_image[]"):
            ProjectImage.objects.create(project=project, image=image)

        messages.success(request, "Project assigned successfully.")

        return redirect("teamlead_project_assigning")

    projects = ProjectAssign.objects.filter(team=team_lead.team)

    return render(request, "team_lead/teamlead_project_assigning.html", {
        "departments": departments,
        "team_lead": team_lead,
        "team_members": team_members,
        "projects": projects,
    })

def project_assign_edit(request, pk):
    project = get_object_or_404(ProjectAssign, id=pk)
    departments = Department.objects.all()
    team_members = User.objects.filter(
        team=project.team, job_Position__iexact="Team Member"
    ).exclude(id=project.assigned_by.id)

    if request.method == "POST":
        project.department_id = request.POST.get("department")
        project.assign_to_id = request.POST.get("assign_to")
        project.work_name = request.POST.get("work_name")
        project.work_type = request.POST.get("work_type")
        project.category = request.POST.get("category")
        project.description = request.POST.get("description")

        deadline = request.POST.get("deadline")
        project.deadline = deadline if deadline else None

        project.additional_notes = request.POST.get("additional_notes")
        project.color_preference = request.POST.get("color_preference")
        project.content_example = request.POST.get("content_example")
        project.priority = request.POST.get("priority")
        project.save()

        # ✅ Save new files (without removing old ones)
        for file in request.FILES.getlist("upload_file[]"):
            ProjectFile.objects.create(project=project, file=file)

        # ✅ Save new images (without removing old ones)
        for image in request.FILES.getlist("upload_image[]"):
            ProjectImage.objects.create(project=project, image=image)

        messages.success(request, "Project updated successfully.")

        return redirect("teamlead_project_assigning")

    return render(request, "team_lead/teamlead_project_assigning.html", {
        "project": project,
        "departments": departments,
        "team_members": team_members,
    })

def project_assign_delete(request, pk):
    project = get_object_or_404(ProjectAssign, id=pk)
    project.delete()
    messages.success(request, "Project deleted successfully.")
    return redirect("teamlead_project_assigning")

def update_project_status(request, pk):
    user_id = request.session.get("user_id")  
    if not user_id:
        messages.error(request, "You must be logged in to update status.")
        return redirect("index")  # or your login page

    user = get_object_or_404(User, id=user_id)
    project = get_object_or_404(ProjectAssign, pk=pk, assign_to=user)

    if request.method == "POST":
        new_status = request.POST.get("status")
        if new_status in dict(ProjectAssign.STATUS_CHOICES):
            project.status = new_status
            project.save()
            messages.success(request, f"Project status updated to {new_status}")
        else:
            messages.error(request, "Invalid status selected.")

    return redirect("teammember_project")
from urllib.parse import urlencode
from django.db.models import Q
from django.core.paginator import Paginator
from django.http import HttpResponseBadRequest, JsonResponse
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST


@never_cache
def teamlead_notepad(request):
    user_id = request.session.get("user_id")
    if not user_id or request.session.get("position") != "team_lead":
        return redirect("index")
    user = get_object_or_404(User, pk=user_id)
    owned = Notepad.objects.filter(user=user)
    params = request.POST if request.method == "POST" else request.GET
    query = params.get("q", "").strip()[:200]
    sort = params.get("sort", "updated")
    orderings = {"updated": ("-updated_at", "-id"), "created": ("-created_at", "-id"), "title": ("title", "id")}
    if sort not in orderings:
        sort = "updated"
    notes = owned
    if query:
        notes = notes.filter(Q(title__icontains=query) | Q(content__icontains=query))
    page_obj = Paginator(notes.order_by(*orderings[sort]), 4).get_page(params.get("page"))
    selected_id = request.POST.get("note_id") if request.method == "POST" else None
    note = None
    if selected_id:
        try:
            selected_id = int(selected_id)
        except (ValueError, TypeError):
            return HttpResponseBadRequest("Invalid note ID.")
        note = get_object_or_404(owned, pk=selected_id)
    error = ""
    title = note.title if note else ""
    content = (note.content or "") if note else ""
    if request.method == "POST":
        title = request.POST.get("title", "").strip() or "Untitled"
        content = request.POST.get("content", "")
        if len(title) > 255:
            error = "Title must be 255 characters or fewer."
        elif not content.strip():
            error = "Write some content before saving the note."
        else:
            if note:
                note.title, note.content = title, content
                note.save()
                if request.headers.get("X-Requested-With") == "XMLHttpRequest":
                    return JsonResponse({
                        "ok": True, "id": note.pk, "title": note.title,
                        "content": note.content or "",
                        "updated_at": timezone.localtime(note.updated_at).strftime("%d %b %Y, %I:%M %p"),
                    })
                messages.success(request, "Note updated successfully.")
            else:
                note = Notepad.objects.create(user=user, title=title, content=content)
                messages.success(request, "Note created successfully.")
            return redirect(reverse("teamlead_notepad") + "?" + urlencode({
                "q": query if selected_id else "", "sort": sort if selected_id else "updated", "page": page_obj.number if selected_id else 1,
            }))
    if error and selected_id and request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return JsonResponse({"ok": False, "message": error}, status=400)
    response = render(request, "team_lead/teamlead_notepad.html", {
        "note": None, "page_obj": page_obj, "query": query, "sort": sort,
        "form_title": title if not selected_id else "", "form_content": content if not selected_id else "", "note_error": error,
        "total_notes": owned.count(),
    })
    if error:
        response.status_code = 400
    return response


@never_cache
@require_POST
def teamlead_notepad_delete(request, pk):
    user_id = request.session.get("user_id")
    if not user_id or request.session.get("position") != "team_lead":
        return redirect("index")
    user = get_object_or_404(User, pk=user_id)
    note = get_object_or_404(Notepad, pk=pk, user=user)
    note.delete()
    messages.success(request, "Note deleted successfully.")
    sort = request.POST.get("sort", "updated")
    if sort not in {"updated", "created", "title"}:
        sort = "updated"
    return redirect(reverse("teamlead_notepad") + "?" + urlencode({
        "q": request.POST.get("q", "").strip()[:200],
        "sort": sort, "page": request.POST.get("page", "1"),
    }))


from urllib.parse import urlencode
from django.db.models import Q
from django.core.paginator import Paginator
from django.http import HttpResponseBadRequest, JsonResponse
from django.urls import reverse
from django.shortcuts import get_object_or_404, redirect, render
from django.contrib import messages
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_POST

@never_cache
def teammember_notepad(request):
    user_id = request.session.get("user_id")
    if not user_id or request.session.get("position") != "team_member":
        return redirect("index")
    user = get_object_or_404(User, pk=user_id)
    owned = Notepad.objects.filter(user=user)
    params = request.POST if request.method == "POST" else request.GET
    query = params.get("q", "").strip()[:200]
    sort = params.get("sort", "updated")
    orderings = {"updated": ("-updated_at", "-id"), "created": ("-created_at", "-id"), "title": ("title", "id")}
    if sort not in orderings:
        sort = "updated"
    notes = owned
    if query:
        notes = notes.filter(Q(title__icontains=query) | Q(content__icontains=query))
    page_obj = Paginator(notes.order_by(*orderings[sort]), 4).get_page(params.get("page"))
    selected_id = request.POST.get("note_id") if request.method == "POST" else None
    note = None
    if selected_id:
        try:
            selected_id = int(selected_id)
        except (ValueError, TypeError):
            return HttpResponseBadRequest("Invalid note ID.")
        note = get_object_or_404(owned, pk=selected_id)
    error = ""
    title = note.title if note else ""
    content = (note.content or "") if note else ""
    if request.method == "POST":
        title = request.POST.get("title", "").strip() or "Untitled"
        content = request.POST.get("content", "")
        if len(title) > 255:
            error = "Title must be 255 characters or fewer."
        elif not content.strip():
            error = "Write some content before saving the note."
        else:
            if note:
                note.title, note.content = title, content
                note.save()
                if request.headers.get("X-Requested-With") == "XMLHttpRequest":
                    return JsonResponse({
                        "ok": True, "id": note.pk, "title": note.title,
                        "content": note.content or "",
                        "updated_at": timezone.localtime(note.updated_at).strftime("%d %b %Y, %I:%M %p"),
                    })
                messages.success(request, "Note updated successfully.")
            else:
                note = Notepad.objects.create(user=user, title=title, content=content)
                messages.success(request, "Note created successfully.")
            return redirect(reverse("teammember_notepad") + "?" + urlencode({
                "q": query if selected_id else "", "sort": sort if selected_id else "updated", "page": page_obj.number if selected_id else 1,
            }))
    if error and selected_id and request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return JsonResponse({"ok": False, "message": error}, status=400)
    response = render(request, "team_member/teammember_notepad.html", {
        "note": None, "page_obj": page_obj, "query": query, "sort": sort,
        "form_title": title if not selected_id else "", "form_content": content if not selected_id else "", "note_error": error,
        "total_notes": owned.count(),
    })
    if error:
        response.status_code = 400
    return response


@never_cache
@require_POST
def teammember_notepad_delete(request, pk):
    user_id = request.session.get("user_id")
    if not user_id or request.session.get("position") != "team_member":
        return redirect("index")
    user = get_object_or_404(User, pk=user_id)
    note = get_object_or_404(Notepad, pk=pk, user=user)
    note.delete()
    messages.success(request, "Note deleted successfully.")
    sort = request.POST.get("sort", "updated")
    if sort not in {"updated", "created", "title"}:
        sort = "updated"
    return redirect(reverse("teammember_notepad") + "?" + urlencode({
        "q": request.POST.get("q", "").strip()[:200],
        "sort": sort, "page": request.POST.get("page", "1"),
    }))


from django.shortcuts import redirect, get_object_or_404
from django.views.decorators.cache import never_cache

@never_cache
def teamlead_profile(request):
    user_id = request.session.get("user_id")
    if not user_id:
        return redirect("index")

    user = get_object_or_404(User, id=user_id)

    if request.method == "POST":
        action = request.POST.get("action")

        # AJAX Profile Image Update
        if action == "edit_profile" and request.FILES.get("profile_image"):
            user.profile_image = request.FILES["profile_image"]
            user.save()
            return JsonResponse({
                "success": True,
                "image_url": user.profile_image.url
            })

        # Normal form update for other details
        elif action == "edit_profile":
            user.name = request.POST.get("name")
            user.email = request.POST.get("email")
            user.phone = request.POST.get("phone")
            user.work_location = request.POST.get("work_location")
            user.save()
            messages.success(request, "Profile updated successfully!")
            return redirect("teamlead_profile")

    # Render page
    response = render(request, "team_lead/teamlead_profile.html", {"user": user})
    response["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response["Pragma"] = "no-cache"
    response["Expires"] = "0"
    return response

@never_cache
@require_GET
def teammember_announcements(request):
    """Show announcements addressed to this member, separated by sender."""
    user = _announcement_role_user(request, "team_member")
    if user is None:
        return redirect("index")
    recipients = _announcement_inbox(user)
    source = request.GET.get("source", "lead")
    if source not in {"lead", "admin"}:
        source = "lead"
    lead_rows = recipients.filter(announcement__is_admin_announcement=False)
    admin_rows = recipients.filter(announcement__is_admin_announcement=True)
    selected = admin_rows if source == "admin" else lead_rows
    page_obj = Paginator(
        selected.order_by("-announcement__created_at", "-pk"), 10
    ).get_page(request.GET.get("page"))
    return render(request, "team_member/teammember_announcements.html", {
        "user": user,
        "page_obj": page_obj,
        "unread_count": recipients.filter(read_at__isnull=True).count(),
        "announcement_source": source,
        "lead_announcement_count": lead_rows.count(),
        "admin_announcement_count": admin_rows.count(),
    })


@never_cache
@require_POST
def teammember_announcement_seen(request, pk):
    """Opening a member's own announcement records its first read time."""
    member = _announcement_role_user(request, "team_member")
    if member is None:
        return JsonResponse({"ok": False, "message": "Please sign in again."}, status=403)
    user_id = member.pk
    recipients = _announcement_inbox(member)
    row = get_object_or_404(
        recipients.select_related("announcement", "announcement__created_by"),
        announcement_id=pk,
    )
    # Keep the first read timestamp; repeated views must not overwrite it.
    first_seen = recipients.filter(announcement_id=pk, read_at__isnull=True).update(read_at=timezone.now())
    if first_seen and not row.announcement.is_admin_announcement:
        member = get_object_or_404(User, pk=user_id)
        TeamLeadNotification.objects.get_or_create(
            event_key=f"announcement-seen:{pk}:{user_id}",
            defaults={
                "recipient": row.announcement.created_by, "kind": "announcement",
                "title": "Announcement seen",
                "message": f"{member.name} saw your announcement: {row.announcement.title}",
                "url": reverse("teamlead_announcements"),
            },
        )
    announcement = row.announcement
    author_name = announcement.author_name
    created = timezone.localtime(announcement.created_at).strftime("%d %b %Y, %I:%M %p")
    return JsonResponse({
        "ok": True,
        "title": announcement.title or "Team Announcement",
        "message": announcement.message,
        "meta": f"Posted by {author_name} · {created}",
        "unread_count": recipients.filter(read_at__isnull=True).count(),
    })


from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET


@never_cache
@require_GET
def teammember_google_meet(request):
    user_id = request.session.get("user_id")
    if not user_id or request.session.get("position") != "team_member":
        return redirect("index")

    user = get_object_or_404(User, pk=user_id)

    meetings = GoogleMeeting.objects.filter(
        participants=user,
    ).select_related(
        "created_by",
    ).prefetch_related(
        "participants",
    ).distinct().order_by(
        "-date", "-time", "-id",
    )
    return render(request, "team_member/teammember_google_meet.html", {
        "user": user,
        "meetings": meetings,
    })


from datetime import datetime
from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.timezone import localtime
from django.views.decorators.cache import never_cache
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter


@never_cache
def teammember_reports(request):
    """Current member reports, team holidays, and matching Excel export."""
    if (
        not request.session.get("user_id")
        or request.session.get("position") != "team_member"
    ):
        return redirect("index")

    user = get_object_or_404(User, id=request.session["user_id"])
    report_tab = "my"
    # Resolve the same team lead used by the existing Member dashboard.
    team_lead = None
    if user.team_id:
        team_lead = User.objects.filter(
            team_id=user.team_id, job_Position__iexact="Team Lead",
        ).first()

    show_all = request.GET.get("all") == "1" and not request.GET.get("date")
    filter_date = None
    filter_date_str = request.GET.get("date")
    if filter_date_str:
        try:
            filter_date = datetime.strptime(filter_date_str, "%Y-%m-%d").date()
        except ValueError:
            messages.error(request, "Invalid date format.")
            return redirect("teammember_reports")

    report_date = filter_date or localtime().date()
    report_filters = {"user": user}

    morning_reports = MorningReport.objects.filter(
        **report_filters
    ).select_related("user").order_by("-created_at")
    evening_reports = EveningReport.objects.filter(
        **report_filters
    ).select_related("user").order_by("-created_at")

    report_day_leave = None
    holiday_users = []
    if not show_all:
        morning_reports = morning_reports.filter(created_at__date=report_date)
        evening_reports = evening_reports.filter(created_at__date=report_date)
        if team_lead:
            report_day_leave = TeamReportLeave.objects.filter(
                team_lead=team_lead, date=report_date
            ).first()
        if report_day_leave and (
            report_day_leave.morning_leave or report_day_leave.evening_leave
        ):
            holiday_users = [user]

    if "export" in request.GET:
        wb = Workbook()
        ws = wb.active
        ws.title = "My Reports" if report_tab == "my" else "Team Members Reports"
        ws.append([
            "Date", "Time", "Report Type", "Name",
            "Team", "Department", "Report", "Status",
        ])

        for reports, report_type in (
            (morning_reports, "Morning"),
            (evening_reports, "Evening"),
        ):
            for report in reports:
                report_time = localtime(report.created_at)
                ws.append([
                    report_time.strftime("%Y-%m-%d"),
                    report_time.strftime("%I:%M %p"),
                    report_type,
                    str(report.user.name),
                    str(report.team or ""),
                    str(report.department or ""),
                    str(report.report_text or ""),
                    str(report.status or ""),
                ])

        if report_day_leave:
            for report_type, is_holiday in (
                ("Morning", report_day_leave.morning_leave),
                ("Evening", report_day_leave.evening_leave),
            ):
                if is_holiday:
                    for person in holiday_users:
                        ws.append([
                            report_date.strftime("%Y-%m-%d"),
                            "",
                            report_type,
                            str(person.name),
                            str(person.team.name) if person.team else "",
                            str(person.department.name) if person.department else "",
                            report_day_leave.reason or "Holiday",
                            "Holiday",
                        ])

        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.value = "'" + cell.value

        for cell in ws[1]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column, width in enumerate(
            [15, 12, 14, 22, 20, 22, 50, 16], start=1
        ):
            ws.column_dimensions[get_column_letter(column)].width = width

        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )
        filename = (
            "my_reports.xlsx" if report_tab == "my"
            else "team_members_reports.xlsx"
        )
        response["Content-Disposition"] = (
            f'attachment; filename="{filename}"'
        )
        wb.save(response)
        return response

    return render(
        request,
        "team_member/teammember_reports.html",
        {
            "user": user,
            "morning_reports": morning_reports,
            "evening_reports": evening_reports,
            "show_all": show_all,
            "filter_date": filter_date,
            "report_tab": report_tab,
            "report_date": report_date,
            "report_day_leave": report_day_leave,
            "holiday_users": holiday_users,
        },
    )



@never_cache
@require_POST
def teammember_project_seen(request, pk):
    if request.session.get("position") != "team_member":
        return JsonResponse({"ok": False}, status=403)

    user_id = request.session.get("user_id")
    user = User.objects.filter(
        pk=user_id,
        job_Position__iexact="Team Member"
    ).first()

    if user is None:
        return JsonResponse({"ok": False}, status=403)

    project = get_object_or_404(
        ProjectAssign.objects.select_related("assigned_by"),
        pk=pk,
        assign_to=user,
    )

    with transaction.atomic():
        updated = ProjectAssign.objects.filter(
            pk=project.pk,
            seen_at__isnull=True,
        ).update(seen_at=timezone.now())

        if updated and project.assigned_by_id != user.pk:
            team_lead = project.assigned_by

            if (
                team_lead
                and (team_lead.job_Position or "").strip().lower()
                == "team lead"
            ):
                TeamLeadNotification.objects.create(
                    recipient=team_lead,
                    kind="project",
                    title="Project Viewed",
                    message=(
                        f"{user.name} viewed the assigned project "
                        f"'{project.work_name}'."
                    ),
                    url=reverse("teamlead_project_assigning"),
                )

    return JsonResponse({
        "ok": True,
        "first_view": bool(updated),
    })



from django.contrib import messages
from django.db.models import Q
from django.http import HttpResponseBadRequest
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_POST


@never_cache
def teammember_task(request):
    user_id = request.session.get("user_id")
    if not user_id or request.session.get("position") != "team_member":
        return redirect("index")
    user = get_object_or_404(User, pk=user_id)

    if request.method == "POST":
        title = request.POST.get("title", "").strip()
        description = request.POST.get("description", "").strip()
        if not title:
            messages.error(request, "Please enter a task title.")
            return redirect("teammember_task")
        Task.objects.create(
            title=title,
            description=description,
            assigned_to=user,
            created_by=user,
            status="pending",
            progress=0,
        )
        messages.success(request, "Your task has been created successfully.")
        return redirect("teammember_task")

    # Include older personal tasks that have no assignee.
    personal_filter = Q(created_by=user) & (
        Q(assigned_to=user) | Q(assigned_to__isnull=True)
    )
    received_filter = Q(assigned_to=user) & ~Q(created_by=user)
    tasks = Task.objects.filter(personal_filter | received_filter).select_related(
        "created_by", "assigned_to"
    ).order_by("-created_at", "-id")
    my_tasks = tasks.filter(personal_filter)
    assigned_tasks = tasks.filter(received_filter)
    return render(request, "team_member/teammember_task.html", {
        "user": user,
        "tasks": tasks,
        "my_tasks": my_tasks,
        "assigned_tasks": assigned_tasks,
        "total_tasks": tasks.count(),
        "pending_tasks": tasks.filter(status="pending").count(),
        "completed_tasks": tasks.filter(status="completed").count(),
    })


@never_cache
@require_POST
def delete_task(request, task_id):
    user_id = request.session.get("user_id")
    if not user_id or request.session.get("position") != "team_member":
        return redirect("index")
    user = get_object_or_404(User, pk=user_id)
    # A member cannot delete a task assigned by their team lead.
    personal_tasks = Task.objects.filter(created_by=user).filter(
        Q(assigned_to=user) | Q(assigned_to__isnull=True)
    )
    task = get_object_or_404(personal_tasks, pk=task_id)
    task.delete()
    messages.success(request, "Task deleted successfully.")
    return redirect("teammember_task")


@never_cache
def teamlead_task(request):

    user_id = request.session.get("user_id")

    if not user_id or request.session.get("position") != "team_lead":
        return redirect("index")

    user = get_object_or_404(User, id=user_id)


    # =========================================================
    # TEAM MEMBERS
    # =========================================================

    team_members = User.objects.filter(
        team=user.team
    ).exclude(
        id=user.id
    ).order_by("name")


    # =========================================================
    # MY TASKS
    # =========================================================

    my_tasks = Task.objects.filter(assigned_to=user).order_by("-created_at")


    # =========================================================
    # ASSIGNED TASKS
    # =========================================================

    assigned_tasks = Task.objects.filter(
        created_by=user
    ).exclude(
        assigned_to=user
    ).select_related(
        "assigned_to"
    ).order_by("-created_at")


    # =========================================================
    # ALL TASKS
    # Used for summary cards
    # =========================================================

    all_tasks = Task.objects.filter(Q(created_by=user) | Q(assigned_to=user))


    total_tasks = all_tasks.count()


    pending_tasks = all_tasks.filter(
        status="pending"
    ).count()


    completed_tasks = all_tasks.filter(
        status="completed"
    ).count()


    # =========================================================
    # CREATE / ASSIGN TASK
    # =========================================================

    if request.method == "POST":

        title = request.POST.get(
            "title",
            ""
        ).strip()

        description = request.POST.get(
            "description",
            ""
        ).strip()

        task_mode = request.POST.get(
            "task_mode",
            "personal"
        )

        assigned_to_id = request.POST.get(
            "assigned_to"
        )


        # =====================================================
        # TITLE VALIDATION
        # =====================================================

        if not title:

            messages.error(
                request,
                "Please enter a task title."
            )

            return redirect("teamlead_task")


        # =====================================================
        # MY TASK
        # =====================================================

        if task_mode == "personal":

            Task.objects.create(
                title=title,
                description=description,
                assigned_to=user,
                created_by=user,
                status="pending",
                progress=0
            )

            messages.success(
                request,
                "Your task has been created successfully."
            )

            return redirect("teamlead_task")


        # =====================================================
        # ASSIGN TASK
        # =====================================================

        if task_mode == "assign":

            if not assigned_to_id:

                messages.error(
                    request,
                    "Please select a team member."
                )

                return redirect("teamlead_task")


            # -------------------------------------------------
            # ONLY CURRENT TEAM MEMBERS ARE ALLOWED
            # -------------------------------------------------

            assigned_user = get_object_or_404(
                User,
                id=assigned_to_id,
                team=user.team
            )


            Task.objects.create(
                title=title,
                description=description,
                assigned_to=assigned_user,
                created_by=user,
                status="pending",
                progress=0
            )


            messages.success(
                request,
                f"Task assigned to {assigned_user.name} successfully."
            )

            return redirect("teamlead_task")


        # =====================================================
        # INVALID TASK MODE
        # =====================================================

        messages.error(
            request,
            "Invalid task request."
        )

        return redirect("teamlead_task")


    # =========================================================
    # PAGE
    # =========================================================

    response = render(
        request,
        "team_lead/teamlead_task.html",
        {
            "my_tasks": my_tasks,
            "assigned_tasks": assigned_tasks,
            "team_members": team_members,
            "team_lead": user,

            # Summary cards
            "total_tasks": total_tasks,
            "pending_tasks": pending_tasks,
            "completed_tasks": completed_tasks,
        }
    )


    response["Cache-Control"] = (
        "no-store, no-cache, "
        "must-revalidate, max-age=0"
    )

    response["Pragma"] = "no-cache"

    response["Expires"] = "0"


    return response

@never_cache
@require_POST
def update_task_teamlead(request, task_id):
    user_id = request.session.get("user_id")
    if not user_id or request.session.get("position") != "team_lead":
        return redirect("index")
    user = get_object_or_404(User, pk=user_id)
    task = get_object_or_404(Task.objects.filter(Q(created_by=user) | Q(assigned_to=user)), pk=task_id)
    status = request.POST.get("status", "")
    progress = {"pending": 0, "in_progress": 50, "completed": 100}
    if status not in progress:
        messages.error(request, "Invalid task status.")
    else:
        task.status, task.progress = status, progress[status]
        task.save(update_fields=["status", "progress"])
        messages.success(request, "Task status updated successfully.")
    return redirect("teamlead_task")



def delete_task_teamlead(request, task_id):
    user_id = request.session.get("user_id")
    if not user_id:
        return redirect("login_view")

    user = User.objects.get(id=user_id)
    task = get_object_or_404(Task, id=task_id, created_by=user)

    if request.method == "POST":
        task.delete()
        messages.success(request, "Task deleted successfully.")

    return redirect("teamlead_task")



@never_cache
def teamlead_logout(request):
    if "user_id" in request.session:
        try:
            user = User.objects.get(id=request.session["user_id"])
            user.status = "inactive"
            user.last_logout_time = timezone.now()
            user.save()
        except User.DoesNotExist:
            pass

    request.session.flush()
    response = redirect("index")
    response["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response["Pragma"] = "no-cache"
    response["Expires"] = "0"
    return response

@never_cache
def teammember_logout(request):
    if "user_id" in request.session:
        try:
            user = User.objects.get(id=request.session["user_id"])
            user.status = "inactive"
            user.last_logout_time = timezone.now()

            if hasattr(user, "is_active"):
                try:
                    setattr(user, "is_active", False)
                except Exception:
                    pass

            user.save()
        except User.DoesNotExist:
            pass

    request.session.flush()

    response = redirect("index")
    response["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response["Pragma"] = "no-cache"
    response["Expires"] = "0"
    return response


def update_inactive_users():
    # Heartbeats from several tabs must not all run the same bulk write.
    from django.core.cache import cache

    if not cache.add("monitoring:inactive-presence-sweep:v1", True, timeout=30):
        return
    cutoff_time = timezone.now() - timedelta(minutes=10)
    try:
        User.objects.filter(
            status="active", last_activity__lt=cutoff_time
        ).update(status="inactive")
    except Exception:
        cache.delete("monitoring:inactive-presence-sweep:v1")
        raise

@csrf_exempt
def update_activity(request):
    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid request"}, status=400)

    user_id = request.session.get("user_id")
    if not user_id:
        return JsonResponse({"success": False, "message": "Not logged in"}, status=401)


    updated = User.objects.filter(pk=user_id).update(
        last_activity=timezone.now(), status="active"
    )
    if not updated:
        return JsonResponse({"success": False, "message": "User not found"}, status=404)

    update_inactive_users()
    return JsonResponse({"success": True, "status": "active"})




def get_user_status(request, user_id):
    from datetime import timedelta
    from django.http import JsonResponse
    from django.utils import timezone

    user = User.objects.filter(pk=user_id).values("status", "last_activity").first()
    if user is None:
        return JsonResponse({"success": False, "message": "User not found"}, status=404)

    cutoff = timezone.now() - timedelta(minutes=10)
    last_activity = user["last_activity"]
    active = (
        str(user["status"] or "").strip().lower() == "active"
        and last_activity is not None
        and last_activity > cutoff
    )
    response = JsonResponse({
        "success": True,
        "status": "active" if active else "inactive",
    })
    response["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    return response



from django.views.decorators.http import require_POST
from django.utils.http import url_has_allowed_host_and_scheme


def _notification_teamlead(request):
    if (
        not request.session.get("user_id")
        or request.session.get("position") != "team_lead"
    ):
        return None

    return User.objects.filter(
        pk=request.session["user_id"]
    ).first()

@never_cache
def teamlead_notifications(request):
    team_lead = _notification_teamlead(request)

    if team_lead is None:
        return redirect("login_view")

    notifications = TeamLeadNotification.objects.filter(
        recipient=team_lead
    ).order_by("-created_at", "-id")

    return render(
        request,
        "team_lead/teamlead_notifications.html",
        {
            "items": notifications,
            "unread_count": notifications.filter(
                is_read=False,
                is_archived=False,
            ).count(),
        },
    )


@require_POST
def teamlead_notifications_read_all(request):
    team_lead = _notification_teamlead(request)

    if team_lead is None:
        return redirect("login_view")

    TeamLeadNotification.objects.filter(
        recipient=team_lead,
        is_read=False,
        is_archived=False,
    ).update(is_read=True)

    return redirect("teamlead_notifications")


@require_POST
def teamlead_notification_open(request, notification_id):
    team_lead = _notification_teamlead(request)

    if team_lead is None:
        return redirect("login_view")

    notification = get_object_or_404(
        TeamLeadNotification,
        pk=notification_id,
        recipient=team_lead,
    )

    action = request.POST.get("action", "open")

    changes = {
        "mark_read": {"is_read": True},
        "mark_unread": {"is_read": False},
        "archive": {"is_archived": True},
        "restore": {"is_archived": False},
    }

    if action in changes:
        TeamLeadNotification.objects.filter(
            pk=notification.pk,
            recipient=team_lead,
        ).update(**changes[action])

        return redirect("teamlead_notifications")

    if action != "open":
        return redirect("teamlead_notifications")

    if not notification.is_read:
        notification.is_read = True
        notification.save(update_fields=["is_read"])

    if notification.url and url_has_allowed_host_and_scheme(
        notification.url,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        return redirect(notification.url)

    return redirect("teamlead_notifications")


# -----------------------------
# TEAM LEAD REMINDERS
# -----------------------------
import calendar
from datetime import datetime as reminder_datetime

from django.contrib import messages as reminder_messages
from django.utils import timezone as reminder_timezone
from django.utils.dateparse import parse_datetime as reminder_parse_datetime
from django.views.decorators.http import require_http_methods

from .models import TeamLeadReminder


def _parse_reminder_datetime(value):
    try:
        parsed = reminder_parse_datetime(value or "")

        if parsed and reminder_timezone.is_naive(parsed):
            parsed = reminder_timezone.make_aware(
                parsed,
                reminder_timezone.get_current_timezone(),
            )

        return parsed
    except (ValueError, TypeError, OverflowError):
        return None


@never_cache
@require_http_methods(["GET", "POST"])
def teamlead_reminders(request):
    team_lead = _notification_teamlead(request)

    if team_lead is None:
        return redirect("login_view")

    # All queries are restricted to this team lead.
    reminders = TeamLeadReminder.objects.filter(owner=team_lead)

    if request.method == "POST":
        action = request.POST.get("action", "")
        reminder = None

        if action in {"update", "delete", "complete", "reopen"}:
            reminder = get_object_or_404(
                reminders,
                pk=request.POST.get("reminder_id"),
            )

        if action == "delete":
            reminder.delete()
            reminder_messages.success(request, "Reminder deleted.")
            return redirect("teamlead_reminders")

        if action in {"complete", "reopen"}:
            reminder.is_completed = action == "complete"
            reminder.save(
                update_fields=["is_completed", "updated_at"]
            )

            response_message = (
                "Reminder completed."
                if reminder.is_completed
                else "Reminder reopened."
            )

            if request.headers.get("X-Requested-With") == "XMLHttpRequest":
               

                return JsonResponse({
                    "ok": True,
                    "id": reminder.pk,
                    "is_completed": reminder.is_completed,
                    "message": response_message,
                })

            reminder_messages.success(
                request,
                response_message,
            )
            return redirect("teamlead_reminders")

        if action not in {"create", "update"}:
            reminder_messages.error(request, "Invalid reminder action.")
            return redirect("teamlead_reminders")

        title = request.POST.get("title", "").strip()
        description = request.POST.get("description", "").strip()
        event_at = _parse_reminder_datetime(
            request.POST.get("event_at")
        )
        remind_at = _parse_reminder_datetime(
            request.POST.get("remind_at")
        )

        if not title or len(title) > 180:
            reminder_messages.error(
                request,
                "Enter a title with a maximum of 180 characters.",
            )
            return redirect("teamlead_reminders")

        if event_at is None or remind_at is None:
            reminder_messages.error(
                request,
                "Enter valid event and reminder dates and times.",
            )
            return redirect("teamlead_reminders")

        if remind_at > event_at:
            reminder_messages.error(
                request,
                "Reminder time must be on or before the event time.",
            )
            return redirect("teamlead_reminders")

        if remind_at <= reminder_timezone.now():
            reminder_messages.error(
                request,
                "Choose a future reminder time.",
            )
            return redirect("teamlead_reminders")

        if reminder is None:
            reminder = TeamLeadReminder(owner=team_lead)

        # Changing the dates allows a new notification to be scheduled.
        schedule_changed = (
            reminder.pk is None
            or reminder.event_at != event_at
            or reminder.remind_at != remind_at
        )

        reminder.title = title
        reminder.description = description
        reminder.event_at = event_at
        reminder.remind_at = remind_at

        if schedule_changed:
            reminder.notified_at = None
            reminder.notification = None
            reminder.is_completed = False

        reminder.save()

        reminder_messages.success(
            request,
            "Reminder created."
            if action == "create"
            else "Reminder updated.",
        )
        return redirect("teamlead_reminders")

    today = reminder_timezone.localdate()

    try:
        selected_month = reminder_datetime.strptime(
            request.GET.get("month", today.strftime("%Y-%m")),
            "%Y-%m",
        ).date()
    except (ValueError, TypeError):
        selected_month = today.replace(day=1)

    month_reminders = reminders.filter(
        event_at__year=selected_month.year,
        event_at__month=selected_month.month,
    ).order_by("event_at", "id")

    calendar_items = []

    for reminder in month_reminders:
        local_event = reminder_timezone.localtime(reminder.event_at)
        local_remind = reminder_timezone.localtime(reminder.remind_at)

        calendar_items.append({
            "id": reminder.pk,
            "title": reminder.title,
            "description": reminder.description,
            "date": local_event.strftime("%Y-%m-%d"),
            "event_at": local_event.strftime("%Y-%m-%dT%H:%M"),
            "remind_at": local_remind.strftime("%Y-%m-%dT%H:%M"),
            "is_completed": reminder.is_completed,
        })

    return render(
        request,
        "team_lead/teamlead_reminders.html",
        {
            "selected_month": selected_month.strftime("%Y-%m"),
            "month_label": selected_month.strftime("%B %Y"),
            "calendar_weeks": calendar.monthcalendar(
                selected_month.year,
                selected_month.month,
            ),
            "calendar_items": calendar_items,
            "today": today.isoformat(),
            "reminders": month_reminders,
        },
    )


from django.http import JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods
from .models import User, TeamLeadReminder
from django.http import JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_POST
from .models import User, TeamLeadNotification
from .member_notifications import sync_member_notifications

def _notification_teammember(request):
    if (
        not request.session.get("user_id")
        or request.session.get("position") != "team_member"
    ):
        return None

    return User.objects.filter(
        pk=request.session["user_id"]
    ).first()

@never_cache
def teammember_notifications(request):
    team_member = _notification_teammember(request)

    if team_member is None:
        return redirect("login_view")

    sync_member_notifications(team_member)

    notifications = TeamLeadNotification.objects.filter(
        recipient=team_member
    ).order_by("-created_at", "-id")

    if request.GET.get("counts") == "1":
        return JsonResponse({"ok": True, "unread_count": notifications.filter(is_read=False, is_archived=False).count()})

    return render(
        request,
        "team_member/teammember_notifications.html",
        {
            "user": team_member,
            "items": notifications,
            "unread_count": notifications.filter(
                is_read=False,
                is_archived=False,
            ).count(),
        },
    )

from django.http import JsonResponse
from django.views.decorators.http import require_POST

def _notify_member_operation(user, kind, title, message, route):
    return TeamLeadNotification.objects.create(
        recipient=user, kind=kind, title=title, message=message, url=reverse(route),
    )

@never_cache
def teammember_project(request):
    user_id = request.session.get("user_id")
    if not user_id:   # if session expired or user logged out
        return redirect("index")

    user = get_object_or_404(User, id=user_id)
    projects = ProjectAssign.objects.filter(assign_to=user).order_by("-assigned_date")

    if request.method == "POST":
        project_id = request.POST.get("project_id")
        status = request.POST.get("status")
        project = get_object_or_404(ProjectAssign, id=project_id, assign_to=user)
        if status not in dict(ProjectAssign.STATUS_CHOICES):
            messages.error(request, "Invalid status selected.")
            return redirect("teammember_project")
        # project.status = status
        # project.save()
        # messages.success(request, f"Project status updated to {status}.")
        # return redirect("teammember_project")
        
        old_status = project.status

        project.status = status
        project.save()

        # Notify the assigned Team Lead only when status changes.
        if old_status != status:
            team_lead = project.assigned_by

            if (
                team_lead
                and team_lead.pk != user.pk
                and (team_lead.job_Position or "").strip().lower() == "team lead"
            ):
                TeamLeadNotification.objects.create(
                    recipient=team_lead,
                    kind="project",
                    title="Project Status Updated",
                    message=(
                        f"{user.name} updated the project "
                        f"'{project.work_name}' from "
                        f"{old_status} to {status}."
                    ),
                    url=reverse("teamlead_project_assigning"),
                )

        messages.success(request, f"Project status updated to {status}.")
        return redirect("teammember_project")


    return render(request, "team_member/teammember_project.html", {"projects": projects})


@never_cache
def teammember_profile(request):
    user_id = request.session.get("user_id")
    if not user_id:
        return redirect("index")

    user = get_object_or_404(User, id=user_id)

    if request.method == "POST":
        action = request.POST.get("action")

        # Profile image AJAX upload
        if action == "edit_profile" and request.FILES.get("profile_image"):
            user.profile_image = request.FILES["profile_image"]
            user.save()
            _notify_member_operation(user, "profile", "Profile image updated", "Your profile image was updated.", "teammember_profile")

            # Check if AJAX
            if request.headers.get("x-requested-with") == "XMLHttpRequest":
                return JsonResponse({
                    "success": True,
                    "image_url": user.profile_image.url
                })

            # fallback redirect for normal form submit
            messages.success(request, "Profile image updated successfully!")
            return redirect("teammember_profile")

        # Normal form update (other details)
        elif action == "edit_profile":
            user.name = request.POST.get("name")
            user.email = request.POST.get("email")
            user.phone = request.POST.get("phone")
            user.work_location = request.POST.get("work_location")

            if "profile_image" in request.FILES:
                user.profile_image = request.FILES["profile_image"]

            user.save()
            _notify_member_operation(user, "profile", "Profile updated", "Your profile details were updated.", "teammember_profile")
            messages.success(request, "Profile updated successfully!")
            return redirect("teammember_profile")

    return render(request, "team_member/teammember_profile.html", {"user": user})


@never_cache
@require_POST
def update_task(request, task_id):
    user_id = request.session.get("user_id")
    if not user_id or request.session.get("position") != "team_member":
        return redirect("index")
    user = get_object_or_404(User, pk=user_id)
    # Members may update received tasks and their own personal tasks.
    permitted = Task.objects.filter(
        Q(assigned_to=user) |
        (Q(created_by=user) & Q(assigned_to__isnull=True))
    )
    task = get_object_or_404(permitted, pk=task_id)
    status = request.POST.get("status", "")
    progress_by_status = {"pending": 0, "in_progress": 50, "completed": 100}
    if status not in progress_by_status:
        messages.error(request, "Invalid task status.")
        return redirect("teammember_task")

    old_status = task.status

    task.status = status
    task.progress = progress_by_status[status]
    task.save(update_fields=["status", "progress"])

    # Notify Team Lead only when an assigned task changes status.
    if old_status != status and task.created_by_id != user.pk:
        team_lead = task.created_by

        if (
            team_lead
            and (team_lead.job_Position or "").strip().lower()
            == "team lead"
        ):
            status_labels = {
                "pending": "Pending",
                "in_progress": "In Progress",
                "completed": "Completed",
            }

            new_status = status_labels.get(status, status)

            notification_title = (
                "Task Completed"
                if status == "completed"
                else "Task Status Updated"
            )

            TeamLeadNotification.objects.create(
                recipient=team_lead,
                kind="task",
                title=notification_title,
                message=(
                    f"{user.name} updated the task "
                    f"'{task.title}' to {new_status}."
                ),
                url=reverse("teamlead_task"),
            )

    messages.success(request, "Task status updated successfully.")

    tab = "my" if task.created_by_id == user.pk else "assigned"
    return redirect(reverse("teammember_task") + "?tab=" + tab)


@never_cache
@require_http_methods(["GET", "POST"])
def teammember_reminders(request):
    user_id = request.session.get("user_id")
    if not user_id or request.session.get("position") != "team_member":
        return redirect("index")
    team_member = get_object_or_404(User, pk=user_id)

    # All reads and writes are restricted to this member.
    reminders = TeamLeadReminder.objects.filter(owner=team_member)

    # The shared Member sidebar polls this endpoint while the site is open.
    # It does not consume or change the background processor's notified_at field.
    if request.method == "GET" and request.GET.get("alerts") == "1":
        due = reminders.filter(
            is_completed=False, remind_at__lte=reminder_timezone.now()
        ).order_by("-remind_at", "-pk")[:100]
        return JsonResponse({"ok": True, "items": [{
            "id": item.pk,
            "title": item.title,
            "remind_at": item.remind_at.isoformat(),
            "event_at": reminder_timezone.localtime(item.event_at).strftime("%d %b %Y, %I:%M %p"),
        } for item in due]})

    if request.method == "POST":
        action = request.POST.get("action", "")
        reminder = None

        if action in {"update", "delete", "complete", "reopen"}:
            reminder = get_object_or_404(
                reminders,
                pk=request.POST.get("reminder_id"),
            )

        if action == "delete":
            reminder.delete()
            reminder_messages.success(request, "Reminder deleted.")
            return redirect("teammember_reminders")

        if action in {"complete", "reopen"}:
            reminder.is_completed = action == "complete"
            reminder.save(
                update_fields=["is_completed", "updated_at"]
            )

            response_message = (
                "Reminder completed."
                if reminder.is_completed
                else "Reminder reopened."
            )

            if request.headers.get("X-Requested-With") == "XMLHttpRequest":
                return JsonResponse({
                    "ok": True,
                    "id": reminder.pk,
                    "is_completed": reminder.is_completed,
                    "message": response_message,
                })

            reminder_messages.success(
                request,
                response_message,
            )
            return redirect("teammember_reminders")

        if action not in {"create", "update"}:
            reminder_messages.error(request, "Invalid reminder action.")
            return redirect("teammember_reminders")

        title = request.POST.get("title", "").strip()
        description = request.POST.get("description", "").strip()
        event_at = _parse_reminder_datetime(
            request.POST.get("event_at")
        )
        remind_at = _parse_reminder_datetime(
            request.POST.get("remind_at")
        )

        if not title or len(title) > 180:
            reminder_messages.error(
                request,
                "Enter a title with a maximum of 180 characters.",
            )
            return redirect("teammember_reminders")

        if event_at is None or remind_at is None:
            reminder_messages.error(
                request,
                "Enter valid event and reminder dates and times.",
            )
            return redirect("teammember_reminders")

        if remind_at > event_at:
            reminder_messages.error(
                request,
                "Reminder time must be on or before the event time.",
            )
            return redirect("teammember_reminders")

        if remind_at <= reminder_timezone.now():
            reminder_messages.error(
                request,
                "Choose a future reminder time.",
            )
            return redirect("teammember_reminders")

        if reminder is None:
            reminder = TeamLeadReminder(owner=team_member)

        # Changing the dates allows a new notification to be scheduled.
        schedule_changed = (
            reminder.pk is None
            or reminder.event_at != event_at
            or reminder.remind_at != remind_at
        )

        reminder.title = title
        reminder.description = description
        reminder.event_at = event_at
        reminder.remind_at = remind_at

        if schedule_changed:
            reminder.notified_at = None
            reminder.notification = None
            reminder.is_completed = False

        reminder.save()

        reminder_messages.success(
            request,
            "Reminder created."
            if action == "create"
            else "Reminder updated.",
        )
        return redirect("teammember_reminders")

    today = reminder_timezone.localdate()

    try:
        selected_month = reminder_datetime.strptime(
            request.GET.get("month", today.strftime("%Y-%m")),
            "%Y-%m",
        ).date()
    except (ValueError, TypeError):
        selected_month = today.replace(day=1)

    month_reminders = reminders.filter(
        event_at__year=selected_month.year,
        event_at__month=selected_month.month,
    ).order_by("event_at", "id")

    calendar_items = []

    for reminder in month_reminders:
        local_event = reminder_timezone.localtime(reminder.event_at)
        local_remind = reminder_timezone.localtime(reminder.remind_at)

        calendar_items.append({
            "id": reminder.pk,
            "title": reminder.title,
            "description": reminder.description,
            "date": local_event.strftime("%Y-%m-%d"),
            "event_at": local_event.strftime("%Y-%m-%dT%H:%M"),
            "remind_at": local_remind.strftime("%Y-%m-%dT%H:%M"),
            "is_completed": reminder.is_completed,
            "event_at_iso": reminder.event_at.isoformat(),
            "remind_at_iso": reminder.remind_at.isoformat(),
        })

    return render(
        request,
        "team_member/teammember_reminders.html",
        {
            "user": team_member,
            "selected_month": selected_month.strftime("%Y-%m"),
            "month_label": selected_month.strftime("%B %Y"),
            "calendar_weeks": calendar.monthcalendar(
                selected_month.year,
                selected_month.month,
            ),
            "calendar_items": calendar_items,
            "today": today.isoformat(),
            "reminders": month_reminders,
        },
    )


@require_POST
def teammember_notifications_read_all(request):
    team_member = _notification_teammember(request)

    if team_member is None:
        return redirect("login_view")

    TeamLeadNotification.objects.filter(
        recipient=team_member,
        is_read=False,
        is_archived=False,
    ).update(is_read=True)

    messages.success(request, "All notifications marked as read.")
    return redirect("teammember_notifications")


@require_POST
def teammember_notification_open(request, notification_id):
    team_member = _notification_teammember(request)

    if team_member is None:
        return redirect("login_view")

    notification = get_object_or_404(
        TeamLeadNotification,
        pk=notification_id,
        recipient=team_member,
    )

    action = request.POST.get("action", "open")

    changes = {
        "mark_read": {"is_read": True},
        "mark_unread": {"is_read": False},
        "archive": {"is_archived": True},
        "restore": {"is_archived": False},
    }

    if action in changes:
        TeamLeadNotification.objects.filter(
            pk=notification.pk,
            recipient=team_member,
        ).update(**changes[action])

        labels = {"mark_read": "Notification marked as read.", "mark_unread": "Notification marked as unread.", "archive": "Notification archived.", "restore": "Notification restored."}
        messages.success(request, labels[action])
        return redirect("teammember_notifications")

    if action != "open":
        return redirect("teammember_notifications")

    if not notification.is_read:
        notification.is_read = True
        notification.save(update_fields=["is_read"])

    # Older member notifications may contain Lead links: route them by category.
    routes = {
        "meeting": "teammember_google_meet", "profile": "teammember_profile",
        "announcement": "teammember_announcements", "project": "teammember_project",
        "task": "teammember_task", "note": "teammember_notepad",
        "resource": "teammember_repository", "report": "teammember_reports",
        "schedule": "teammember_dashboard", "reminder": "teammember_reminders",
    }
    route = routes.get(notification.kind)
    if route:
        target = reverse(route)
        # Preserve a reminder's date query only for the correct Member page.
        if notification.url and notification.url.split("?", 1)[0] == target:
            target = notification.url
        return redirect(target)
    return redirect("teammember_notifications")



"""Team Member chat pages, sharing Lead conversations and group messages."""
from django.db import transaction
from django.db.models import Q
from django.http import JsonResponse, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.views.decorators.cache import never_cache
from chat.models import Conversation
from .models import User, Group, GroupMember, GroupMessage, ExtraContact
from .member_chat_actions import chat_action


def _member_chat_member(request):
    position = str(request.session.get('position', '')).strip().lower().replace(' ', '_')
    if position != 'team_member' or not request.session.get('user_id'):
        return None
    return get_object_or_404(User, pk=request.session['user_id'])


def _member_chat_context(user):
    return {'current_user': user, 'user': user,
            'users': User.objects.exclude(pk=user.pk).order_by('name', 'pk'),
            'groups': Group.objects.filter(memberships__user=user).distinct().order_by('name'),
            'extra_contacts': ExtraContact.objects.all(), 'user_role': 'team_member'}


def _member_chat_error(message, status=400):
    return JsonResponse({'ok': False, 'message': message}, status=status)


def _member_chat_create_group(request, user , group_view_name="teammember_group_chat"):
    name = request.POST.get('group_name', '').strip()
    ids = request.POST.getlist('members')
    if not name or len(name) > 150:
        return _member_chat_error('Enter a group name of 150 characters or fewer.')
    try:
        ids = set(int(value) for value in ids) - {user.pk}
    except (ValueError, TypeError):
        return _member_chat_error('Select valid group members.')
    if not ids:
        return _member_chat_error('Select at least one other member (two people including you).')
    people = list(User.objects.filter(pk__in=ids))
    if len(people) != len(ids):
        return _member_chat_error('One of the selected members no longer exists.')
    with transaction.atomic():
        group = Group.objects.create(name=name, created_by=user)
        GroupMember.objects.create(group=group, user=user, role='admin')
        GroupMember.objects.bulk_create([GroupMember(group=group, user=person) for person in people])
    return JsonResponse({'ok': True, 'message': 'Group created successfully.',
                         'url': reverse(group_view_name, args=[group.pk])})


def _member_chat_payload(message, group=False):
    sender = message.sender
    return {'message_id': message.pk, 'message': message.message if group else message.content,
            'sender_id': sender.pk, 'sender': sender.name,
            'sender_profile': sender.profile_image.url if sender.profile_image else None,
            'timestamp': (message.timestamp if group else message.created_at).isoformat(),
            'edited_at': message.edited_at.isoformat() if message.edited_at else None,
            'deleted': message.deleted_for_everyone}


@never_cache
def teammember_chat(request):
    user = _member_chat_member(request)
    if user is None:
        return redirect('index')
    if request.method == 'POST':
        return _member_chat_create_group(request, user) if request.POST.get('action') == 'create_group' else _member_chat_error('Invalid action.')
    return render(request, 'team_member/teammember_chat.html', _member_chat_context(user))


@never_cache
def teammember_chat_room(request, user_id):
    user = _member_chat_member(request)
    if user is None:
        return redirect('index')
    other = get_object_or_404(User, pk=user_id)
    if user.pk == other.pk:
        return redirect('teammember_chat')
    room = Conversation.objects.filter(Q(user1=user, user2=other) | Q(user1=other, user2=user)).order_by('pk').first()
    if room is None:
        first, second = sorted([user, other], key=lambda person: person.pk)
        room, _ = Conversation.objects.get_or_create(user1=first, user2=second)
    if request.method == 'POST':
        return _member_chat_create_group(request, user) if request.POST.get('action') == 'create_group' else chat_action(request, user, room, False)
    messages_qs = room.messages.exclude(hidden_for=user).select_related('sender').order_by('created_at', 'pk')
    if request.GET.get('messages') == '1':
        return JsonResponse({'ok': True, 'items': [_member_chat_payload(message) for message in messages_qs]})
    context = _member_chat_context(user)
    context.update(room=room, other_user=other, messages=messages_qs)
    return render(request, 'team_member/teammember_chat_room.html', context)

@never_cache
def teammember_group_chat(request, group_id):
    from asgiref.sync import async_to_sync
    from channels.layers import get_channel_layer
    import logging

    user = _member_chat_member(request)
    if user is None:
        return redirect('index')
    group = get_object_or_404(Group, pk=group_id)
    membership = get_object_or_404(GroupMember, group=group, user=user)
    can_manage = group.created_by_id == user.pk or membership.role == 'admin'

    def member_state():
        current = GroupMember.objects.filter(group=group, user=user).first()
        permitted = bool(current and (group.created_by_id == user.pk or current.role == 'admin'))
        entries = list(GroupMember.objects.filter(group=group).select_related('user').order_by('user__name', 'pk'))
        return {
            'ok': True, 'can_manage': permitted, 'count': len(entries),
            'members': [{
                'id': entry.user_id, 'name': entry.user.name,
                'role': entry.role, 'is_creator': entry.user_id == group.created_by_id,
                'is_self': entry.user_id == user.pk,
                'profile': entry.user.profile_image.url if entry.user.profile_image else None,
            } for entry in entries],
            'all_users': [{'id': person.pk, 'name': person.name} for person in
                          User.objects.exclude(group_memberships__group=group).order_by('name', 'pk')] if permitted else [],
        }

    if request.method == 'POST':
        action = request.POST.get('action', '')
        if action == 'create_group':
            return _member_chat_create_group(request, user)
        if action in {'edit_message', 'delete_for_me', 'delete_for_everyone', 'clear_chat'}:
            return chat_action(request, user, group, True)
        if action not in {'add_member', 'remove_member', 'make_admin', 'remove_admin'}:
            return _member_chat_error('Invalid action.')
        if not can_manage:
            return _member_chat_error('Only the group creator or an admin can manage members.', 403)
        try:
            target_id = int(request.POST.get('user_id', ''))
        except (ValueError, TypeError):
            return _member_chat_error('Select a valid member.')
        with transaction.atomic():
            Group.objects.select_for_update().get(pk=group.pk)
            # Check permission again after obtaining the same lock used by all membership changes.
            membership = get_object_or_404(GroupMember, group=group, user=user)
            if group.created_by_id != user.pk and membership.role != 'admin':
                return _member_chat_error('You no longer have permission to manage this group.', 403)
            if action == 'add_member':
                target_user = get_object_or_404(User, pk=target_id)
                _, created = GroupMember.objects.get_or_create(group=group, user=target_user)
                notice = f'{target_user.name} added to the group.' if created else f'{target_user.name} is already a member.'
            else:
                target = get_object_or_404(GroupMember.objects.select_related('user'), group=group, user_id=target_id)
                if target_id == group.created_by_id:
                    return _member_chat_error('The group creator cannot be removed or demoted.')
                if target_id == user.pk:
                    return _member_chat_error('You cannot remove yourself or change your own admin role here.')
                if action == 'remove_member':
                    if GroupMember.objects.filter(group=group).count() <= 2:
                        return _member_chat_error('Keep at least two people in the group.')
                    name = target.user.name
                    target.delete()
                    notice = f'{name} removed from the group.'
                else:
                    target.role = 'admin' if action == 'make_admin' else 'member'
                    target.save(update_fields=['role'])
                    notice = f'{target.user.name} is now an admin.' if action == 'make_admin' else f'Admin access removed for {target.user.name}.'

            def notify_members():
                layer = get_channel_layer()
                if layer:
                    try:
                        async_to_sync(layer.group_send)(f'group_{group.pk}', {'type': 'chat_change'})
                    except Exception:
                        logging.getLogger(__name__).exception('Membership broadcast failed; member polling will reconcile.')
            transaction.on_commit(notify_members)
        return JsonResponse({**member_state(), 'message': notice})

    if request.GET.get('members') == '1':
        return JsonResponse(member_state())
    messages_qs = GroupMessage.objects.filter(group=group).exclude(hidden_for=user).select_related('sender').order_by('timestamp', 'pk')
    if request.GET.get('messages') == '1':
        return JsonResponse({'ok': True, 'items': [_member_chat_payload(message, True) for message in messages_qs]})
    context = _member_chat_context(user)
    context.update(group=group, messages=messages_qs,
                   members=GroupMember.objects.filter(group=group).select_related('user').order_by('user__name', 'pk'),
                   all_users=User.objects.exclude(group_memberships__group=group).order_by('name', 'pk'),
                   can_manage_group=can_manage)
    return render(request, 'team_member/teammember_group_chat.html', context)



# <----------------------------ACCOUNTS TEAM ()--------------------->

@never_cache
@require_GET
def accounts_dashboard(request):
    import calendar
    import re
    from datetime import date
    from decimal import Decimal
    from django.db.models import F, Sum
    from django.http import HttpResponseForbidden
    from django.shortcuts import redirect, render
    from django.urls import reverse
    from django.utils import timezone
    from .models import AccountsIncome, AccountsExpense, AccountsSale

    if not request.session.get("user_id"):
        return redirect("login_view")
    current_user = _accounts_user(request)
    if current_user is None:
        return HttpResponseForbidden("You do not have access to the Accounts dashboard.")

    today = timezone.localdate()
    period = request.GET.get("period", "monthly").strip()
    selected_date = request.GET.get("date", "").strip() or today.isoformat()
    selected_month = request.GET.get("month", "").strip() or today.strftime("%Y-%m")
    from_date = request.GET.get("start", "").strip()
    to_date = request.GET.get("end", "").strip()
    range_start = range_end = None
    filter_error = ""

    def read_date(value):
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ValueError("Choose a valid date.")
        return date.fromisoformat(value)

    try:
        if period == "monthly":
            if not re.fullmatch(r"\d{4}-\d{2}", selected_month):
                raise ValueError("Choose a valid month.")
            year, month = map(int, selected_month.split("-"))
            range_start = date(year, month, 1)
            range_end = date(year, month, calendar.monthrange(year, month)[1])
        elif period == "weekly":
            anchor = read_date(selected_date)
            ordinal = max(date.min.toordinal(), anchor.toordinal() - anchor.weekday())
            range_start = date.fromordinal(ordinal)
            range_end = date.fromordinal(min(date.max.toordinal(), ordinal + 6))
        elif period == "custom":
            if not from_date or not to_date:
                raise ValueError("Choose both From Date and To Date.")
            range_start, range_end = read_date(from_date), read_date(to_date)
            if range_start > range_end:
                raise ValueError("From Date must be on or before To Date.")
        elif period != "all":
            raise ValueError("Choose a valid period.")
    except (ValueError, OverflowError) as error:
        filter_error = str(error) if period == "custom" else "Choose a valid date, month or period."
        range_start = range_end = None

    income_qs = AccountsIncome.objects.select_related("created_by")
    expense_qs = AccountsExpense.objects.select_related("created_by")
    sales_qs = AccountsSale.objects.select_related("created_by")
    if filter_error:
        income_qs, expense_qs, sales_qs = income_qs.none(), expense_qs.none(), sales_qs.none()
    elif range_start is not None:
        income_qs = income_qs.filter(date__range=(range_start, range_end))
        expense_qs = expense_qs.filter(date__range=(range_start, range_end))
        sales_qs = sales_qs.filter(date__range=(range_start, range_end))

    def total(queryset, field="amount"):
        return queryset.aggregate(total=Sum(field))["total"] or Decimal("0.00")

    income_total, expense_total = total(income_qs), total(expense_qs)
    sales_totals = sales_qs.aggregate(amount=Sum("amount"), received=Sum("received_amount"))
    sales_total = sales_totals["amount"] or Decimal("0.00")
    received_total = sales_totals["received"] or Decimal("0.00")
    overdue_qs = sales_qs.filter(due_date__lt=today, received_amount__lt=F("amount"))
    overdue_totals = overdue_qs.aggregate(amount=Sum("amount"), received=Sum("received_amount"))
    overdue_balance = (overdue_totals["amount"] or Decimal("0.00")) - (overdue_totals["received"] or Decimal("0.00"))
    recent = []
    # Ten rows from each ledger suffice to find the latest ten overall.
    for queryset, kind, route in [(income_qs, "income", "accounts_income"),
                                  (expense_qs, "expense", "accounts_expenses")]:
        for entry in queryset.order_by("-date", "-created_at", "-pk")[:10]:
            recent.append({"date": entry.date, "created_at": entry.created_at, "pk": entry.pk,
                "kind": kind, "description": entry.description, "category": entry.category,
                "amount": entry.amount, "reference": entry.reference,
                "recorded_by": getattr(entry.created_by, "name", "") or "—",
                "url": reverse(route)})
    recent.sort(key=lambda item: (item["date"], item["created_at"], item["pk"], item["kind"]), reverse=True)
    period_label = "All dates" if period == "all" else (
        f"{range_start:%d %b %Y} – {range_end:%d %b %Y}" if range_start else "Invalid date filter")
    return render(request, "accounts/accounts_dashboard.html", {
        "current_user": current_user, "accounts_active": "dashboard",
        "period": period, "selected_date": selected_date, "selected_month": selected_month,
        "from_date": from_date, "to_date": to_date, "filter_error": filter_error,
        "period_label": period_label,
        "account_summary": {"income": income_total, "expenses": expense_total,
            "net_movement": income_total - expense_total, "sales": sales_total,
            "received": received_total, "outstanding": sales_total - received_total,
            "overdue": overdue_balance, "overdue_count": overdue_qs.count()},
        "recent_transactions": recent[:10],
        "overdue_sales": overdue_qs.order_by("due_date", "-pk")[:10],
        "income_count": income_qs.count(), "expense_count": expense_qs.count(),
        "sales_count": sales_qs.count(),
    }, status=400 if filter_error else 200)



@never_cache
@require_POST
def accounts_logout(request):
    user_id = request.session.get("user_id")
    session_role = str(request.session.get("position", "")).strip().lower().replace(" ", "_")
    if user_id and session_role not in {"accounts", "accounts_team"}:
        return HttpResponseForbidden("Use your own account's logout option.")

    if user_id:
        User.objects.filter(pk=user_id).update(
            status="inactive",
            last_logout_time=timezone.now(),
        )

    request.session.flush()
    return redirect("login_view")


from decimal import Decimal
from django.db.models import Sum
from .forms import AccountsIncomeForm
from .models import AccountsIncome
from django.views.decorators.http import require_http_methods


def _accounts_user(request):
    """Return an Accounts user only when both DB and session roles match."""
    user_id = request.session.get("user_id")
    if not user_id:
        return None
    user = User.objects.filter(pk=user_id).first()
    if user is None:
        return None
    roles = {"accounts", "accounts_team"}
    database_role = str(user.job_Position or "").strip().lower().replace(" ", "_")
    session_role = str(request.session.get("position", "")).strip().lower().replace(" ", "_")
    return user if database_role in roles and session_role in roles else None


def _create_accounts_notification(*, kind, title, message, actor=None):
    """Create a separate notification for every Accounts user."""
    from uuid import uuid4

    from django.db import transaction
    from django.db.models import Value
    from django.db.models.functions import Lower, Replace, Trim
    from django.urls import reverse

    from .models import AccountsNotification, User

    routes = {
        "income": "accounts_income",
        "expense": "accounts_expenses",
        "sale": "accounts_sales",
        "payment": "accounts_sales",
        "reminder": "accounts_reminders",
    }

    if kind not in routes:
        raise ValueError("Unsupported Accounts notification kind.")

    recipient_ids = list(
        User.objects.annotate(
            accounts_role=Replace(
                Lower(Trim("job_Position")),
                Value(" "),
                Value("_"),
            )
        )
        .filter(accounts_role__in=["accounts", "accounts_team"])
        .values_list("pk", flat=True)
    )

    if not recipient_ids:
        return 0

    event_id = uuid4().hex
    target_url = reverse(routes[kind])

    notifications = [
        AccountsNotification(
            recipient_id=recipient_id,
            actor=actor,
            kind=kind,
            title=title,
            message=message,
            url=target_url,
            event_key=f"accounts:{event_id}:{recipient_id}",
        )
        for recipient_id in recipient_ids
    ]

    with transaction.atomic():
        AccountsNotification.objects.bulk_create(notifications)

    return len(notifications)

from pathlib import Path
from django.db import transaction
from django.db.models import F, Q
from django.http import FileResponse, JsonResponse
from django.urls import reverse
from .models import AccountsIncome, AccountsIncomeHistory
from .forms import AccountsIncomeForm
from django.contrib import messages
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_POST, require_http_methods


def _income_snapshot(entry):
    return {
        "date": entry.date.isoformat(), "category": entry.category,
        "amount": str(entry.amount),
        "payment_method": entry.get_payment_method_display(),
        "description": entry.description, "reference": entry.reference,
        "receipt": Path(entry.receipt.name).name if entry.receipt else "",
    }


def _income_audit(entry, actor, action, before=None):
    before = before or {}
    after = {} if action == "deleted" else _income_snapshot(entry)
    changes = {key: {"before": before.get(key, ""), "after": after.get(key, "")}
               for key in (before.keys() | after.keys()) if before.get(key, "") != after.get(key, "")}
    if changes or action != "updated":
        AccountsIncomeHistory.objects.create(income=entry, income_number=entry.pk,
            actor=actor, actor_name=getattr(actor, "name", "") or str(actor),
            action=action, changes=changes)


def _income_duplicate(form, exclude_pk=None):
    values = form.cleaned_data
    matches = AccountsIncome.objects.all()
    if exclude_pk:
        matches = matches.exclude(pk=exclude_pk)
    if values.get("reference"):
        matches = matches.filter(reference__iexact=values["reference"])
    else:
        matches = matches.filter(date=values["date"], amount=values["amount"],
            category__iexact=values["category"])
    return matches.exists()


def _income_duplicate_response():
    return JsonResponse({"ok": False, "duplicate": True,
        "message": "A similar income or the same reference already exists. Review it before saving again."}, status=409)



@never_cache
@require_http_methods(["GET", "POST"])
def accounts_income(request):
    import calendar
    import re
    from datetime import date
    from decimal import Decimal
    from io import BytesIO

    from django.core.paginator import Paginator
    from django.db.models import Q, Sum
    from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
    from django.shortcuts import redirect, render
    from django.utils import timezone

    if not request.session.get("user_id"):
        return redirect("login_view")
    current_user = _accounts_user(request)
    if current_user is None:
        return HttpResponseForbidden("You do not have access to Accounts income.")

    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    form = AccountsIncomeForm(request.POST if request.method == "POST" else None,
        request.FILES if request.method == "POST" else None)
    if request.method == "POST" and form.is_valid():
        if _income_duplicate(form) and request.POST.get("confirm_duplicate") != "1":
            if is_ajax:
                return _income_duplicate_response()
            form.add_error(None, "A similar income already exists. Review the entry and use Save Anyway to confirm.")
        else:
            with transaction.atomic():
                income = form.save(commit=False)
                income.created_by = current_user
                income.save()
                _income_audit(income, current_user, "created")
                _create_accounts_notification(
                    kind="income",
                    title="Income added",
                    message=(
                        f"{current_user.name} added income "
                        f"#{income.pk}: {income.category}, "
                        f"₹{income.amount:,.2f} "
                        f"on {income.date:%d %b %Y}."
                    ),
                    actor=current_user,
                )
            if is_ajax:
                return JsonResponse({"ok": True, "message": "Income added successfully.", "id": income.pk}, status=201)
            messages.success(request, "Income added successfully.")
            return redirect("accounts_income")

    if request.method == "POST" and is_ajax:
        return JsonResponse({
            "ok": False,
            "message": "Please correct the highlighted fields.",
            "errors": form.errors.get_json_data(),
        }, status=400)

    today = timezone.localdate()
    month_start = today.replace(day=1)
    month_end = today.replace(day=calendar.monthrange(today.year, today.month)[1])
    all_entries = AccountsIncome.objects.select_related("created_by")
    entries = all_entries
    query = request.GET.get("q", "").strip()[:200]
    if query:
        entries = entries.filter(
            Q(category__icontains=query)
            | Q(description__icontains=query)
            | Q(reference__icontains=query)
        )

    category_filter = request.GET.get("category", "").strip()[:100]
    if category_filter:
        entries = entries.filter(category__iexact=category_filter)

    period = request.GET.get("period", "all").strip()
    selected_date = request.GET.get("date", "").strip() or today.isoformat()
    selected_month = request.GET.get("month", "").strip() or today.strftime("%Y-%m")
    from_date = request.GET.get("start", "").strip()
    to_date = request.GET.get("end", "").strip()
    range_start = range_end = None
    filter_error = ""

    def read_date(value):
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ValueError("Invalid date format.")
        return date.fromisoformat(value)

    try:
        if period == "weekly":
            anchor = read_date(selected_date)
            monday_ordinal = max(date.min.toordinal(), anchor.toordinal() - anchor.weekday())
            range_start = date.fromordinal(monday_ordinal)
            range_end = date.fromordinal(min(date.max.toordinal(), monday_ordinal + 6))
        elif period == "monthly":
            if not re.fullmatch(r"\d{4}-\d{2}", selected_month):
                raise ValueError("Invalid month format.")
            year, month = map(int, selected_month.split("-"))
            range_start = date(year, month, 1)
            range_end = date(year, month, calendar.monthrange(year, month)[1])
        elif period == "custom":
            if not from_date or not to_date:
                raise ValueError("Choose both From Date and To Date.")
            range_start, range_end = read_date(from_date), read_date(to_date)
            if range_start > range_end:
                raise ValueError("From Date must be on or before To Date.")
        elif period != "all":
            raise ValueError("Choose a valid period.")
    except (ValueError, OverflowError) as error:
        filter_error = (
            str(error) if period == "custom"
            else "Choose a valid date, month or period."
        )
        range_start = range_end = None
        entries = entries.none()

    if range_start is not None and not filter_error:
        entries = entries.filter(date__range=(range_start, range_end))
    entries = entries.order_by("-date", "-id")

    def income_total(queryset):
        return queryset.aggregate(total=Sum("amount"))["total"] or Decimal("0.00")

    # Export the entire filtered queryset before applying pagination.
    if request.method == "GET" and request.GET.get("export") == "xlsx" and not filter_error:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Income"
        headers = ["Date", "Category", "Description", "Payment Method", "Reference", "Recorded By", "Amount (INR)"]
        sheet.append(headers)
        total = Decimal("0.00")
        for entry in entries.iterator():
            creator = getattr(entry.created_by, "name", "") or "—"
            sheet.append([
                entry.date, entry.category, entry.description,
                entry.get_payment_method_display(), entry.reference or "—",
                creator, entry.amount,
            ])
            row = sheet.max_row
            # User-entered strings must remain text, including values starting with '='.
            for column in range(2, 7):
                cell = sheet.cell(row, column)
                cell.value = str(cell.value or "")
                cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.cell(row, 1).number_format = "dd mmm yyyy"
            for column in (7,):
                sheet.cell(row, column).number_format = "#,##0.00"
            total += entry.amount

        last_data_row = sheet.max_row
        sheet.append(["Total", None, None, None, None, None, total])
        total_row = sheet.max_row
        sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=6)
        for column in (7,):
            sheet.cell(total_row, column).number_format = "#,##0.00"
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="5B32A7")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for cell in sheet[total_row]:
            cell.fill = PatternFill("solid", fgColor="EDE9FE")
            cell.font = Font(bold=True)
        for column, width in enumerate([18, 25, 50, 22, 25, 25, 20], start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.row_dimensions[1].height = 26
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:G{last_data_row}"
        sheet.sheet_view.showGridLines = False
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="income_{today.isoformat()}.xlsx"'
        response["Cache-Control"] = "no-store"
        return response

    page_obj = Paginator(entries, 20).get_page(request.GET.get("page"))
    return render(request, "accounts/accounts_income.html", {
        "current_user": current_user,
        "accounts_active": "income",
        "form": form,
        "query": query,
        "period": period,
        "selected_date": selected_date,
        "selected_month": selected_month,
        "from_date": from_date,
        "to_date": to_date,
        "range_start": range_start,
        "range_end": range_end,
        "filter_error": filter_error,
        "has_filters": bool(query or period != "all" or category_filter),
        "category_filter": category_filter,
        "categories": all_entries.order_by("category").values_list("category", flat=True).distinct(),
        "page_obj": page_obj,
        "total_entries": entries.count(),
        "total_income": income_total(entries),
        "today_income": income_total(all_entries.filter(date=today)),
        "month_income": income_total(all_entries.filter(date__range=(month_start, month_end))),
        "month_label": month_start.strftime("%B %Y"),
    }, status=400 if request.method == "POST" or filter_error else 200)




@never_cache
@require_POST
def accounts_income_delete(request, pk):
    actor = _accounts_user(request)
    if actor is None:
        return JsonResponse({"ok": False, "message": "Please log in with an Accounts account."},
            status=403 if request.session.get("user_id") else 401)
    with transaction.atomic():
        entry = AccountsIncome.objects.select_for_update().filter(pk=pk).first()
        if entry is None:
            return JsonResponse({"ok": False, "message": "This income no longer exists."}, status=404)
        _income_audit(entry, actor, "deleted", _income_snapshot(entry))
        _create_accounts_notification(
            kind="income",
            title="Income deleted",
            message=(
                f"{actor.name} deleted income "
                f"#{entry.pk}: {entry.category}, "
                f"₹{entry.amount:,.2f} "
                f"on {entry.date:%d %b %Y}."
            ),
            actor=actor,
        )
        entry.delete()
    return JsonResponse({"ok": True, "message": "Income deleted successfully.", "id": pk})


@never_cache
@require_http_methods(["GET"])
def accounts_income_receipt(request, pk):
    if _accounts_user(request) is None:
        return JsonResponse({"ok": False, "message": "Accounts access required."}, status=403)
    entry = AccountsIncome.objects.filter(pk=pk).first()
    if entry is None or not entry.receipt:
        return JsonResponse({"ok": False, "message": "Receipt not available."}, status=404)
    try:
        handle = entry.receipt.open("rb")
    except (FileNotFoundError, OSError):
        return JsonResponse({"ok": False, "message": "Receipt file not available."}, status=404)
    suffix = Path(entry.receipt.name).suffix.lower()
    mime = {".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}.get(suffix, "application/octet-stream")
    response = FileResponse(handle, as_attachment=request.GET.get("download") == "1",
        filename=f"income_{pk}_receipt{suffix}", content_type=mime)
    response["X-Content-Type-Options"] = "nosniff"
    return response


@never_cache
@require_http_methods(["GET"])
def accounts_income_history(request, pk):
    if _accounts_user(request) is None:
        return JsonResponse({"ok": False, "message": "Accounts access required."}, status=403)
    if not AccountsIncome.objects.filter(pk=pk).exists():
        return JsonResponse({"ok": False, "message": "This income no longer exists."}, status=404)
    items = AccountsIncomeHistory.objects.filter(income_id=pk)[:50]
    return JsonResponse({"ok": True, "items": [{"action": item.get_action_display(),
        "actor": item.actor_name or "—", "time": timezone.localtime(item.created_at).strftime("%d %b %Y, %I:%M %p"),
        "changes": item.changes} for item in items]})

from decimal import Decimal
from django.db.models import Q, Sum
from django.core.paginator import Paginator
from django.http import JsonResponse, HttpResponseForbidden
from django.contrib import messages
from django.shortcuts import render, redirect
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods, require_POST
from .forms import AccountsExpenseForm
from .models import AccountsExpense
from pathlib import Path
from django.db import transaction
from django.db.models import F, Q
from django.http import FileResponse, JsonResponse
from django.urls import reverse
from .models import AccountsExpenseHistory


def _expense_snapshot(entry):
    return {
        "date": entry.date.isoformat(), "category": entry.category,
        "amount": str(entry.amount), "paid_amount": str(entry.effective_paid_amount),
        "payment_method": entry.get_payment_method_display(), "paid_to": entry.paid_to,
        "description": entry.description, "reference": entry.reference,
        "receipt": Path(entry.receipt.name).name if entry.receipt else "",
    }


def _expense_audit(entry, actor, action, before=None):
    before = before or {}
    after = {} if action == "deleted" else _expense_snapshot(entry)
    changes = {key: {"before": before.get(key, ""), "after": after.get(key, "")}
               for key in (before.keys() | after.keys()) if before.get(key, "") != after.get(key, "")}
    if changes or action != "updated":
        AccountsExpenseHistory.objects.create(expense=entry, expense_number=entry.pk,
            actor=actor, actor_name=getattr(actor, "name", "") or str(actor),
            action=action, changes=changes)


def _expense_duplicate(form, exclude_pk=None):
    values = form.cleaned_data
    matches = AccountsExpense.objects.all()
    if exclude_pk:
        matches = matches.exclude(pk=exclude_pk)
    if values.get("reference"):
        matches = matches.filter(reference__iexact=values["reference"])
    else:
        matches = matches.filter(date=values["date"], amount=values["amount"],
            category__iexact=values["category"], paid_to__iexact=values.get("paid_to", ""))
    return matches.exists()


def _expense_duplicate_response():
    return JsonResponse({"ok": False, "duplicate": True,
        "message": "A similar expense or the same reference already exists. Review it before saving again."}, status=409)



@never_cache
@require_http_methods(["GET", "POST"])
def accounts_expenses(request):
    import calendar
    import re
    from datetime import date
    from decimal import Decimal
    from io import BytesIO

    from django.core.paginator import Paginator
    from django.db.models import Q, Sum
    from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
    from django.shortcuts import redirect, render
    from django.utils import timezone

    if not request.session.get("user_id"):
        return redirect("login_view")
    current_user = _accounts_user(request)
    if current_user is None:
        return HttpResponseForbidden("You do not have access to Accounts expense.")

    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    form = AccountsExpenseForm(request.POST if request.method == "POST" else None,
        request.FILES if request.method == "POST" else None)
    if request.method == "POST" and form.is_valid():
        if _expense_duplicate(form) and request.POST.get("confirm_duplicate") != "1":
            if is_ajax:
                return _expense_duplicate_response()
            form.add_error(None, "A similar expense already exists. Review the entry and use Save Anyway to confirm.")
        else:
            with transaction.atomic():
                expense = form.save(commit=False)
                expense.created_by = current_user
                expense.save()
                _expense_audit(expense, current_user, "created")
                _create_accounts_notification(
                    kind="expense",
                    title="Expense added",
                    message=(
                        f"{current_user.name} added expense "
                        f"#{expense.pk}: {expense.category}, "
                        f"₹{expense.amount:,.2f} "
                        f"on {expense.date:%d %b %Y}."
                    ),
                    actor=current_user,
                )
            if is_ajax:
                return JsonResponse({"ok": True, "message": "Expense added successfully.", "id": expense.pk}, status=201)
            messages.success(request, "Expense added successfully.")
            return redirect("accounts_expenses")

    if request.method == "POST" and is_ajax:
        return JsonResponse({
            "ok": False,
            "message": "Please correct the highlighted fields.",
            "errors": form.errors.get_json_data(),
        }, status=400)

    today = timezone.localdate()
    month_start = today.replace(day=1)
    month_end = today.replace(day=calendar.monthrange(today.year, today.month)[1])
    all_entries = AccountsExpense.objects.select_related("created_by")
    entries = all_entries
    query = request.GET.get("q", "").strip()[:200]
    if query:
        entries = entries.filter(
            Q(category__icontains=query)
            | Q(description__icontains=query)
            | Q(reference__icontains=query)
            | Q(paid_to__icontains=query)
        )

    category_filter = request.GET.get("category", "").strip()[:100]
    payment_filter = request.GET.get("status", "").strip()
    if category_filter:
        entries = entries.filter(category__iexact=category_filter)
    if payment_filter == "paid":
        entries = entries.filter(Q(paid_amount__isnull=True) | Q(paid_amount=F("amount")))
    elif payment_filter == "pending":
        entries = entries.filter(paid_amount=0)
    elif payment_filter == "partial":
        entries = entries.filter(paid_amount__gt=0, paid_amount__lt=F("amount"))
    else:
        payment_filter = ""

    period = request.GET.get("period", "all").strip()
    selected_date = request.GET.get("date", "").strip() or today.isoformat()
    selected_month = request.GET.get("month", "").strip() or today.strftime("%Y-%m")
    from_date = request.GET.get("start", "").strip()
    to_date = request.GET.get("end", "").strip()
    range_start = range_end = None
    filter_error = ""

    def read_date(value):
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ValueError("Invalid date format.")
        return date.fromisoformat(value)

    try:
        if period == "weekly":
            anchor = read_date(selected_date)
            monday_ordinal = max(date.min.toordinal(), anchor.toordinal() - anchor.weekday())
            range_start = date.fromordinal(monday_ordinal)
            range_end = date.fromordinal(min(date.max.toordinal(), monday_ordinal + 6))
        elif period == "monthly":
            if not re.fullmatch(r"\d{4}-\d{2}", selected_month):
                raise ValueError("Invalid month format.")
            year, month = map(int, selected_month.split("-"))
            range_start = date(year, month, 1)
            range_end = date(year, month, calendar.monthrange(year, month)[1])
        elif period == "custom":
            if not from_date or not to_date:
                raise ValueError("Choose both From Date and To Date.")
            range_start, range_end = read_date(from_date), read_date(to_date)
            if range_start > range_end:
                raise ValueError("From Date must be on or before To Date.")
        elif period != "all":
            raise ValueError("Choose a valid period.")
    except (ValueError, OverflowError) as error:
        filter_error = (
            str(error) if period == "custom"
            else "Choose a valid date, month or period."
        )
        range_start = range_end = None
        entries = entries.none()

    if range_start is not None and not filter_error:
        entries = entries.filter(date__range=(range_start, range_end))
    entries = entries.order_by("-date", "-id")

    def expense_total(queryset):
        return queryset.aggregate(total=Sum("amount"))["total"] or Decimal("0.00")

    # Export the entire filtered queryset before applying pagination.
    if request.method == "GET" and request.GET.get("export") == "xlsx" and not filter_error:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Expenses"
        headers = ["Date", "Category", "Description", "Paid To", "Payment Method", "Reference", "Recorded By", "Total Amount (INR)", "Paid Amount (INR)", "Balance (INR)", "Status"]
        sheet.append(headers)
        total = paid_total = balance_total = Decimal("0.00")
        for entry in entries.iterator():
            creator = getattr(entry.created_by, "name", "") or "—"
            sheet.append([
                entry.date, entry.category, entry.description, entry.paid_to or "—",
                entry.get_payment_method_display(), entry.reference or "—",
                creator, entry.amount, entry.effective_paid_amount, entry.balance_amount, entry.payment_status_label,
            ])
            row = sheet.max_row
            # User-entered strings must remain text, including values starting with '='.
            for column in range(2, 8):
                cell = sheet.cell(row, column)
                cell.value = str(cell.value or "")
                cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.cell(row, 1).number_format = "dd mmm yyyy"
            for column in (8, 9, 10):
                sheet.cell(row, column).number_format = "#,##0.00"
            sheet.cell(row, 11).data_type = "s"
            paid_total += entry.effective_paid_amount
            balance_total += entry.balance_amount
            total += entry.amount

        last_data_row = sheet.max_row
        sheet.append(["Total", None, None, None, None, None, None, total, paid_total, balance_total, None])
        total_row = sheet.max_row
        sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=7)
        for column in (8, 9, 10):
            sheet.cell(total_row, column).number_format = "#,##0.00"
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="5B32A7")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for cell in sheet[total_row]:
            cell.fill = PatternFill("solid", fgColor="EDE9FE")
            cell.font = Font(bold=True)
        for column, width in enumerate([18, 25, 50, 25, 22, 25, 25, 20, 20, 20, 20], start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.row_dimensions[1].height = 26
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:K{last_data_row}"
        sheet.sheet_view.showGridLines = False
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="expense_{today.isoformat()}.xlsx"'
        response["Cache-Control"] = "no-store"
        return response

    page_obj = Paginator(entries, 20).get_page(request.GET.get("page"))
    return render(request, "accounts/accounts_expenses.html", {
        "current_user": current_user,
        "accounts_active": "expenses",
        "form": form,
        "query": query,
        "period": period,
        "selected_date": selected_date,
        "selected_month": selected_month,
        "from_date": from_date,
        "to_date": to_date,
        "range_start": range_start,
        "range_end": range_end,
        "filter_error": filter_error,
        "has_filters": bool(query or period != "all" or category_filter or payment_filter),
        "category_filter": category_filter,
        "payment_filter": payment_filter,
        "categories": all_entries.order_by("category").values_list("category", flat=True).distinct(),
        "page_obj": page_obj,
        "total_entries": entries.count(),
        "total_expense": expense_total(entries),
        "today_expense": expense_total(all_entries.filter(date=today)),
        "month_expense": expense_total(all_entries.filter(date__range=(month_start, month_end))),
        "month_label": month_start.strftime("%B %Y"),
    }, status=400 if request.method == "POST" or filter_error else 200)


@never_cache
@require_POST
def accounts_expense_delete(request, pk):
    actor = _accounts_user(request)
    if actor is None:
        return JsonResponse({"ok": False, "message": "Please log in with an Accounts account."},
            status=403 if request.session.get("user_id") else 401)
    with transaction.atomic():
        entry = AccountsExpense.objects.select_for_update().filter(pk=pk).first()
        if entry is None:
            return JsonResponse({"ok": False, "message": "This expense no longer exists."}, status=404)
        _expense_audit(entry, actor, "deleted", _expense_snapshot(entry))
        _create_accounts_notification(
            kind="expense",
            title="Expense deleted",
            message=(
                f"{actor.name} deleted expense "
                f"#{entry.pk}: {entry.category}, "
                f"₹{entry.amount:,.2f} "
                f"on {entry.date:%d %b %Y}."
            ),
            actor=actor,
        )
        entry.delete()
    return JsonResponse({"ok": True, "message": "Expense deleted successfully.", "id": pk})


@never_cache
@require_http_methods(["GET"])
def accounts_expense_receipt(request, pk):
    if _accounts_user(request) is None:
        return JsonResponse({"ok": False, "message": "Accounts access required."}, status=403)
    entry = AccountsExpense.objects.filter(pk=pk).first()
    if entry is None or not entry.receipt:
        return JsonResponse({"ok": False, "message": "Receipt not available."}, status=404)
    try:
        handle = entry.receipt.open("rb")
    except (FileNotFoundError, OSError):
        return JsonResponse({"ok": False, "message": "Receipt file not available."}, status=404)
    suffix = Path(entry.receipt.name).suffix.lower()
    mime = {".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}.get(suffix, "application/octet-stream")
    response = FileResponse(handle, as_attachment=request.GET.get("download") == "1",
        filename=f"expense_{pk}_receipt{suffix}", content_type=mime)
    response["X-Content-Type-Options"] = "nosniff"
    return response


@never_cache
@require_http_methods(["GET"])
def accounts_expense_history(request, pk):
    if _accounts_user(request) is None:
        return JsonResponse({"ok": False, "message": "Accounts access required."}, status=403)
    if not AccountsExpense.objects.filter(pk=pk).exists():
        return JsonResponse({"ok": False, "message": "This expense no longer exists."}, status=404)
    items = AccountsExpenseHistory.objects.filter(expense_id=pk)[:50]
    return JsonResponse({"ok": True, "items": [{"action": item.get_action_display(),
        "actor": item.actor_name or "—", "time": timezone.localtime(item.created_at).strftime("%d %b %Y, %I:%M %p"),
        "changes": item.changes} for item in items]})


from decimal import Decimal
from django.contrib import messages
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import F, Q, Sum
from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods, require_POST
from .models import AccountsSale
from .forms import AccountsSaleForm


def _sale_duplicate(form, exclude_pk=None):
    values = form.cleaned_data
    matches = AccountsSale.objects.all()
    if exclude_pk is not None:
        matches = matches.exclude(pk=exclude_pk)
    invoice = values.get("invoice_number", "").strip()
    if invoice:
        matches = matches.filter(invoice_number__iexact=invoice)
    else:
        matches = matches.filter(date=values["date"], amount=values["amount"],
            customer_name__iexact=values["customer_name"])
    return matches.exists()


def _sale_duplicate_response():
    return JsonResponse({"ok": False, "duplicate": True,
        "message": "A similar sale or the same invoice already exists. Review it before saving again."}, status=409)

from decimal import Decimal
from django.contrib import messages
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import F, Q, Sum
from django.http import FileResponse, HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods, require_POST
from django.urls import reverse
from pathlib import Path
from .models import AccountsSale, AccountsSalePayment
from .forms import AccountsSaleForm, AccountsSalePaymentForm


def _sale_duplicate(form, exclude_pk=None):
    values = form.cleaned_data
    matches = AccountsSale.objects.all()
    if exclude_pk is not None:
        matches = matches.exclude(pk=exclude_pk)
    invoice = values.get("invoice_number", "").strip()
    if invoice:
        matches = matches.filter(invoice_number__iexact=invoice)
    else:
        matches = matches.filter(date=values["date"], amount=values["amount"],
            customer_name__iexact=values["customer_name"])
    return matches.exists()


def _sale_duplicate_response():
    return JsonResponse({"ok": False, "duplicate": True,
        "message": "A similar sale or the same invoice already exists. Review it before saving again."}, status=409)


@never_cache
@require_http_methods(["GET", "POST"])
def accounts_sales(request):
    import calendar
    import re
    from datetime import date
    from io import BytesIO

    if not request.session.get("user_id"):
        return redirect("login_view")
    current_user = _accounts_user(request)
    if current_user is None:
        return HttpResponseForbidden("You do not have access to Accounts sales.")
    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    form = AccountsSaleForm(request.POST if request.method == "POST" else None, request.FILES if request.method == "POST" else None)
    if request.method == "POST" and form.is_valid():
        if _sale_duplicate(form) and request.POST.get("confirm_duplicate") != "1":
            if is_ajax:
                return _sale_duplicate_response()
            form.add_error(None, "A similar sale already exists. Review it and use Save Anyway to confirm.")
        else:
            with transaction.atomic():
                sale = form.save(commit=False)
                sale.created_by = current_user
                sale.save()
                if sale.received_amount > 0:
                    AccountsSalePayment.objects.create(
                        sale=sale,
                        date=sale.date,
                        amount=sale.received_amount,
                        created_by=current_user,
                        payment_method="other",
                        note="Initial received amount at sale creation.",
                    )

                _create_accounts_notification(
                    kind="sale",
                    title="Sale added",
                    message=(
                        f"{current_user.name} added sale "
                        f"#{sale.pk} for {sale.customer_name}: "
                        f"₹{sale.amount:,.2f}. "
                        f"Received: ₹{sale.received_amount:,.2f}."
                    ),
                    actor=current_user,
                )
            if is_ajax:
                return JsonResponse({"ok": True, "message": "Sale added successfully.", "id": sale.pk}, status=201)
            messages.success(request, "Sale added successfully.")
            return redirect("accounts_sales")
    if request.method == "POST" and is_ajax:
        return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
            "errors": form.errors.get_json_data()}, status=400)

    today = timezone.localdate()
    month_start = today.replace(day=1)
    month_end = today.replace(day=calendar.monthrange(today.year, today.month)[1])
    all_entries = AccountsSale.objects.select_related("created_by")
    entries = all_entries
    query = request.GET.get("q", "").strip()[:200]
    status_filter = request.GET.get("status", "").strip()
    if status_filter not in {"", "paid", "partial", "unpaid", "overdue"}:
        status_filter = ""
    if query:
        entries = entries.filter(Q(customer_name__icontains=query) |
            Q(invoice_number__icontains=query) | Q(description__icontains=query))
    if status_filter == "paid":
        entries = entries.filter(received_amount=F("amount"))
    elif status_filter == "partial":
        entries = entries.filter(received_amount__gt=0, received_amount__lt=F("amount"))
    elif status_filter == "unpaid":
        entries = entries.filter(received_amount=0)

    if status_filter == "overdue":
        entries = entries.filter(due_date__lt=today, received_amount__lt=F("amount"))

    period = request.GET.get("period", "all").strip()
    selected_date = request.GET.get("date", "").strip() or today.isoformat()
    selected_month = request.GET.get("month", "").strip() or today.strftime("%Y-%m")
    from_date = request.GET.get("start", "").strip()
    to_date = request.GET.get("end", "").strip()
    range_start = range_end = None
    filter_error = ""

    def read_date(value):
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ValueError("Invalid date format.")
        return date.fromisoformat(value)

    try:
        if period == "weekly":
            anchor = read_date(selected_date)
            monday_ordinal = max(date.min.toordinal(), anchor.toordinal() - anchor.weekday())
            range_start = date.fromordinal(monday_ordinal)
            range_end = date.fromordinal(min(date.max.toordinal(), monday_ordinal + 6))
        elif period == "monthly":
            if not re.fullmatch(r"\d{4}-\d{2}", selected_month):
                raise ValueError("Invalid month format.")
            year, month = map(int, selected_month.split("-"))
            range_start = date(year, month, 1)
            range_end = date(year, month, calendar.monthrange(year, month)[1])
        elif period == "custom":
            if not from_date or not to_date:
                raise ValueError("Choose both From Date and To Date.")
            range_start, range_end = read_date(from_date), read_date(to_date)
            if range_start > range_end:
                raise ValueError("From Date must be on or before To Date.")
        elif period != "all":
            raise ValueError("Choose a valid period.")
    except (ValueError, OverflowError) as error:
        filter_error = (
            str(error) if period == "custom"
            else "Choose a valid date, month or period."
        )
        range_start = range_end = None
        entries = entries.none()

    if range_start is not None and not filter_error:
        entries = entries.filter(date__range=(range_start, range_end))
    entries = entries.order_by("-date", "-id")

    totals = entries.aggregate(sales=Sum("amount"), received=Sum("received_amount"))
    total_sales = totals["sales"] or Decimal("0.00")
    total_received = totals["received"] or Decimal("0.00")
    total_balance = total_sales - total_received

    # Export all matching rows, before pagination.
    if request.method == "GET" and request.GET.get("export") == "xlsx" and not filter_error:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sales"
        sheet.append(["Date", "Customer", "Invoice", "Description", "Sale Amount (INR)",
            "Received (INR)", "Balance (INR)", "Status", "Recorded By", "Due Date", "Overdue"])
        for entry in entries.iterator():
            sheet.append([entry.date, entry.customer_name, entry.invoice_number or "—",
                entry.description, entry.amount, entry.received_amount, entry.balance_amount,
                entry.payment_status_label, getattr(entry.created_by, "name", "") or "—", entry.due_date, "Yes" if entry.is_overdue else "No"])
            row = sheet.max_row
            for column in (2, 3, 4, 8, 9, 11):
                cell = sheet.cell(row, column)
                cell.value = str(cell.value or "")
                cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.cell(row, 10).number_format = "dd mmm yyyy"
            sheet.cell(row, 1).number_format = "dd mmm yyyy"
            for column in (5, 6, 7):
                sheet.cell(row, column).number_format = "#,##0.00"
        last_data_row = sheet.max_row
        sheet.append(["Total", None, None, None, total_sales, total_received, total_balance, None, None, None, None])
        sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=4)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="5B32A7")
            cell.font = Font(color="FFFFFF", bold=True)
        for cell in sheet[sheet.max_row]:
            cell.fill = PatternFill("solid", fgColor="EDE9FE")
            cell.font = Font(bold=True)
        for column in (5, 6, 7):
            sheet.cell(sheet.max_row, column).number_format = "#,##0.00"
        for column, width in enumerate([18, 30, 25, 50, 22, 22, 22, 22, 28, 18, 14], start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.row_dimensions[1].height = 26
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:K{last_data_row}"
        sheet.sheet_view.showGridLines = False
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        response = HttpResponse(output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="sales_{today.isoformat()}.xlsx"'
        response["Cache-Control"] = "no-store"
        return response

    page_obj = Paginator(entries, 20).get_page(request.GET.get("page"))
    def sales_total(queryset):
        return queryset.aggregate(total=Sum("amount"))["total"] or Decimal("0.00")
    return render(request, "accounts/accounts_sales.html", {
        "current_user": current_user, "accounts_active": "sales", "form": form,
        "query": query, "status_filter": status_filter,
        "has_filters": bool(query or status_filter or period != "all"),
        "period": period, "selected_date": selected_date, "selected_month": selected_month,
        "from_date": from_date, "to_date": to_date, "range_start": range_start,
        "range_end": range_end, "filter_error": filter_error,
        "page_obj": page_obj, "total_entries": entries.count(), "total_sales": total_sales,
        "total_received": total_received, "total_balance": total_balance,
        "month_sales": sales_total(all_entries.filter(date__range=(month_start, month_end))),
        "today_sales": sales_total(all_entries.filter(date=today)),
        "month_label": month_start.strftime("%B %Y"),
    }, status=400 if request.method == "POST" or filter_error else 200)


# @never_cache
# @require_http_methods(["GET", "POST"])
# def accounts_sale_edit(request, pk):
#     actor = _accounts_user(request)
#     if actor is None:
#         return JsonResponse({"ok": False, "message": "Please log in with an Accounts account."},
#             status=403 if request.session.get("user_id") else 401)
#     with transaction.atomic():
#         sale = AccountsSale.objects.select_for_update().filter(pk=pk).first()
#         if sale is None:
#             return JsonResponse({"ok": False, "message": "This sale no longer exists."}, status=404)
#         if request.method == "GET":
#             return JsonResponse({"ok": True, "item": {
#                 "date": sale.date.isoformat(), "customer_name": sale.customer_name,
#                 "invoice_number": sale.invoice_number, "description": sale.description,
#                 "amount": str(sale.amount), "due_date": sale.due_date.isoformat() if sale.due_date else ""},
#                 "invoice_url": reverse("accounts_sale_invoice", args=[sale.pk]) if sale.invoice_file else "",
#                 "received_amount": str(sale.received_amount)})
#         before = {"date": sale.date.isoformat(), "customer_name": sale.customer_name,
#             "invoice_number": sale.invoice_number, "amount": str(sale.amount)}
#         form = AccountsSaleForm(request.POST, request.FILES, instance=sale)
#         if not form.is_valid():
#             return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
#                 "errors": form.errors.get_json_data()}, status=400)
#         if (any(str(form.cleaned_data.get(key) or "") != value for key, value in before.items())
#                 and _sale_duplicate(form, sale.pk) and request.POST.get("confirm_duplicate") != "1"):
#             return _sale_duplicate_response()
#         sale = form.save(commit=False)
#         if request.POST.get("remove_invoice") == "1" and not request.FILES.get("invoice_file"):
#             sale.invoice_file = ""
#             sale.save()
#         _create_accounts_notification(
#             kind="sale",
#             title="Sale updated",
#             message=(
#                 f"{actor.name} updated sale "
#                 f"#{sale.pk} for {sale.customer_name}: "
#                 f"₹{sale.amount:,.2f}."
#             ),
#             actor=actor,
#         )
#     return JsonResponse({
#         "ok": True,
#         "message": "Sale updated successfully.",
#         "id": sale.pk,
#     })

@never_cache
@require_POST
def accounts_sale_delete(request, pk):
    actor = _accounts_user(request)
    if actor is None:
        return JsonResponse({
            "ok": False,
            "message": "Please log in with an Accounts account.",
        }, status=403 if request.session.get("user_id") else 401)

    with transaction.atomic():
        sale = (
            AccountsSale.objects
            .select_for_update()
            .filter(pk=pk)
            .first()
        )
        if sale is None:
            return JsonResponse({
                "ok": False,
                "message": "This sale no longer exists.",
            }, status=404)

        _create_accounts_notification(
            kind="sale",
            title="Sale deleted",
            message=(
                f"{actor.name} deleted sale "
                f"#{sale.pk} for {sale.customer_name}: "
                f"₹{sale.amount:,.2f}."
            ),
            actor=actor,
        )
        sale.delete()

    return JsonResponse({
        "ok": True,
        "message": "Sale deleted successfully.",
        "id": pk,
    })

@never_cache
@require_http_methods(["GET"])
def accounts_sale_invoice(request, pk):
    if _accounts_user(request) is None:
        return JsonResponse({"ok": False, "message": "Accounts access required."}, status=403)
    sale = AccountsSale.objects.filter(pk=pk).first()
    if sale is None or not sale.invoice_file:
        return JsonResponse({"ok": False, "message": "Invoice not available."}, status=404)
    try:
        handle = sale.invoice_file.open("rb")
    except (OSError, FileNotFoundError):
        return JsonResponse({"ok": False, "message": "Invoice file not available."}, status=404)
    suffix = Path(sale.invoice_file.name).suffix.lower()
    mime = {".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}.get(suffix, "application/octet-stream")
    response = FileResponse(handle, as_attachment=request.GET.get("download") == "1",
        filename=f"sale_{pk}_invoice{suffix}", content_type=mime)
    response["X-Content-Type-Options"] = "nosniff"
    return response


@never_cache
@require_http_methods(["GET", "POST"])
def accounts_sale_payments(request, pk):
    actor = _accounts_user(request)
    if actor is None:
        return JsonResponse({"ok": False, "message": "Accounts access required."}, status=403)
    with transaction.atomic():
        sale = AccountsSale.objects.select_for_update().filter(pk=pk).first()
        if sale is None:
            return JsonResponse({"ok": False, "message": "This sale no longer exists."}, status=404)
        if request.method == "POST":
            form = AccountsSalePaymentForm(request.POST)
            if not form.is_valid():
                return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
                    "errors": form.errors.get_json_data()}, status=400)
            token = form.cleaned_data["request_token"]
            existing = AccountsSalePayment.objects.filter(request_token=token).first()
            if existing:
                if existing.sale_id != pk or any(getattr(existing, key) != form.cleaned_data[key]
                        for key in ("date", "amount", "payment_method", "reference", "note")):
                    return JsonResponse({"ok": False, "message": "This payment request changed. Close and reopen Payments."}, status=409)
                return JsonResponse({"ok": True, "message": "Payment already recorded.", "id": existing.pk})
            if form.cleaned_data["date"] < sale.date:
                form.add_error("date", "Payment date cannot be before the sale date.")
            if form.cleaned_data["date"] > timezone.localdate():
                form.add_error("date", "Use the date the payment was received; future payments cannot be recorded.")
            if form.cleaned_data["amount"] > sale.balance_amount:
                form.add_error("amount", "Payment cannot exceed the outstanding balance.")
            if form.errors:
                return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
                    "errors": form.errors.get_json_data()}, status=400)
            payment = form.save(commit=False)
            payment.sale, payment.created_by, payment.request_token = sale, actor, token
            changed = AccountsSale.objects.filter(pk=pk,
                received_amount__lte=F("amount") - payment.amount).update(
                    received_amount=F("received_amount") + payment.amount, updated_at=timezone.now())
            if not changed:
                return JsonResponse({"ok": False, "message": "Balance changed. Close and reopen Payments."}, status=409)
            payment.save()
            _create_accounts_notification(
                kind="payment",
                title="Sales payment received",
                message=(
                    f"{actor.name} recorded a payment of "
                    f"₹{payment.amount:,.2f} for "
                    f"{sale.customer_name}, sale #{sale.pk}, "
                    f"on {payment.date:%d %b %Y}."
                ),
                actor=actor,
            )
            return JsonResponse({
                "ok": True,
                "message": "Payment recorded successfully.",
                "id": payment.pk,
            })
        items = list(sale.payments.select_related("created_by").all())
        recorded = sum((item.amount for item in items), Decimal("0.00"))
        return JsonResponse({"ok": True, "customer": sale.customer_name,
            "received": str(sale.received_amount), "balance": str(sale.balance_amount),
            "opening_received": str(sale.received_amount - recorded),
            "today": timezone.localdate().isoformat(), "sale_date": sale.date.isoformat(),
            "items": [{"date": item.date.isoformat(), "amount": str(item.amount),
                "method": item.get_payment_method_display(), "reference": item.reference,
                "note": item.note, "actor": getattr(item.created_by, "name", "") or "—"} for item in items]})


from datetime import date as accounts_date, timedelta as accounts_timedelta
from decimal import Decimal
from django.db.models import Sum
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import render, redirect
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .models import AccountsIncome, AccountsExpense, AccountsSale

def _accounts_report_range(params, today):
    period = params.get("period", "monthly")
    if period not in {"weekly", "monthly", "custom"}:
        period = "monthly"
    anchor = today
    start = today.replace(day=1)
    end = today
    error = ""
    try:
        raw_anchor = params.get("date", "")
        anchor = accounts_date.fromisoformat(raw_anchor) if raw_anchor else today
        if period == "weekly":
            start = anchor - accounts_timedelta(days=anchor.weekday())
            end = start + accounts_timedelta(days=6)
        elif period == "monthly":
            start = anchor.replace(day=1)
            next_month = (
                start.replace(year=start.year + 1, month=1)
                if start.month == 12
                else start.replace(month=start.month + 1)
            )
            end = next_month - accounts_timedelta(days=1)
        else:
            start = accounts_date.fromisoformat(params.get("start", ""))
            end = accounts_date.fromisoformat(params.get("end", ""))
        if end < start:
            error = "End date must be on or after the start date."
        elif (end - start).days > 365:
            error = "Select a date range of 366 days or fewer."
    except (ValueError, TypeError, OverflowError):
        error = "Enter valid dates for the selected report period."
    return period, anchor, start, end, error


def _accounts_report_data(start, end):
    income = AccountsIncome.objects.filter(date__range=(start, end))
    expenses = AccountsExpense.objects.filter(date__range=(start, end))
    sales = AccountsSale.objects.filter(date__range=(start, end))
    zero = Decimal("0.00")
    income_total = income.aggregate(total=Sum("amount"))["total"] or zero
    expense_total = expenses.aggregate(total=Sum("amount"))["total"] or zero
    sale_totals = sales.aggregate(total=Sum("amount"), received=Sum("received_amount"))
    sales_total = sale_totals["total"] or zero
    received_total = sale_totals["received"] or zero

    income_days = {
        item["date"]: item["total"]
        for item in income.order_by().values("date").annotate(total=Sum("amount"))
    }
    expense_days = {
        item["date"]: item["total"]
        for item in expenses.order_by().values("date").annotate(total=Sum("amount"))
    }
    sale_days = {
        item["date"]: item
        for item in sales.order_by().values("date").annotate(
            total=Sum("amount"), received=Sum("received_amount")
        )
    }
    daily_rows = []
    for offset in range((end - start).days + 1):
        day = start + accounts_timedelta(days=offset)
        day_income = income_days.get(day, zero)
        day_expense = expense_days.get(day, zero)
        day_sale = sale_days.get(day, {})
        day_total = day_sale.get("total", zero)
        day_received = day_sale.get("received", zero)
        daily_rows.append({
            "date": day,
            "income": day_income,
            "expenses": day_expense,
            "net_movement": day_income - day_expense,
            "sales": day_total,
            "received": day_received,
            "outstanding": day_total - day_received,
        })
    return {
        "income_total": income_total,
        "expense_total": expense_total,
        "net_movement": income_total - expense_total,
        "sales_total": sales_total,
        "received_total": received_total,
        "outstanding_total": sales_total - received_total,
        "income_count": income.count(),
        "expense_count": expenses.count(),
        "sales_count": sales.count(),
        "daily_rows": daily_rows,
    }


def _accounts_report_excel(data, start, end):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Accounts Report", "Value"])
    summary.append(["Start Date", start])
    summary.append(["End Date", end])
    summary.append(["Income (INR)", data["income_total"]])
    summary.append(["Expenses (INR)", data["expense_total"]])
    summary.append(["Net Movement (INR)", data["net_movement"]])
    summary.append(["Sales (INR)", data["sales_total"]])
    summary.append(["Received Against Sales (INR)", data["received_total"]])
    summary.append(["Outstanding Against Sales (INR)", data["outstanding_total"]])
    summary.append(["Income Entries", data["income_count"]])
    summary.append(["Expense Entries", data["expense_count"]])
    summary.append(["Sales Entries", data["sales_count"]])
    summary.append(["Calculation", "Net Movement = Income - Expenses. Sales are shown separately."])
    summary.append(["Sales Payments", "Received/outstanding reflect current values of sales dated within this range."])
    summary["B2"].number_format = summary["B3"].number_format = "dd mmm yyyy"
    for row in range(4, 10):
        summary.cell(row, 2).number_format = '#,##0.00'

    daily = workbook.create_sheet("Daily Breakdown")
    daily.append(["Date", "Income (INR)", "Expenses (INR)", "Net Movement (INR)",
                  "Sales (INR)", "Received (INR)", "Outstanding (INR)"])
    for item in data["daily_rows"]:
        daily.append([item["date"], item["income"], item["expenses"], item["net_movement"],
                      item["sales"], item["received"], item["outstanding"]])
        daily.cell(daily.max_row, 1).number_format = "dd mmm yyyy"
        for column in range(2, 8):
            daily.cell(daily.max_row, column).number_format = '#,##0.00'

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="5B3B87")
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 25
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = (
                36 if sheet.title == "Summary" else 24
            )
    summary.column_dimensions["B"].width = 75
    summary["B13"].alignment = Alignment(wrap_text=True)
    summary["B14"].alignment = Alignment(wrap_text=True)
    summary.row_dimensions[13].height = summary.row_dimensions[14].height = 34
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="accounts_report_{start.isoformat()}_{end.isoformat()}.xlsx"'
    )
    workbook.save(response)
    return response


@never_cache
@require_GET
def accounts_reports(request):
    if not request.session.get("user_id"):
        return redirect("login_view")
    current_user = _accounts_user(request)
    if current_user is None:
        return HttpResponseForbidden("You do not have access to Accounts reports.")
    period, anchor, start, end, error = _accounts_report_range(
        request.GET, timezone.localdate()
    )
    data = None if error else _accounts_report_data(start, end)
    if request.GET.get("export") == "xlsx" and not error:
        return _accounts_report_excel(data, start, end)
    return render(request, "accounts/reports.html", {
        "current_user": current_user,
        "accounts_active": "reports",
        "period": period,
        "anchor_date": anchor,
        "start_date": start,
        "end_date": end,
        "report_error": error,
        "report": data,
    }, status=400 if error else 200)

@never_cache
@require_GET
def accounts_notifications(request):
    from django.core.paginator import Paginator
    from django.http import JsonResponse, HttpResponseForbidden
    from django.shortcuts import redirect, render
    from .models import AccountsNotification

    current_user = _accounts_user(request)
    if current_user is None:
        if not request.session.get("user_id"):
            return redirect("login_view")
        return HttpResponseForbidden("Accounts access required.")

    _process_accounts_reminders(current_user)

    inbox = AccountsNotification.objects.filter(
        recipient=current_user,
        is_archived=False,
    )
    unread_count = inbox.filter(is_read=False).count()

    selected_filter = request.GET.get("filter", "all")
    if selected_filter not in {"all", "unread", "read", "archived"}:
        selected_filter = "all"

    query = request.GET.get("q", "").strip()[:200]
    selected_kind = request.GET.get("kind", "").strip()
    if selected_kind not in {"", "income", "expense", "sale", "payment", "reminder", "note"}:
        selected_kind = ""
    archived = AccountsNotification.objects.filter(recipient=current_user, is_archived=True)
    items = archived if selected_filter == "archived" else inbox
    if query:
        from django.db.models import Q
        items = items.filter(Q(title__icontains=query) | Q(message__icontains=query))
    if selected_kind:
        items = items.filter(kind=selected_kind)
    if selected_filter == "unread":
        items = items.filter(is_read=False)
    elif selected_filter == "read":
        items = items.filter(is_read=True)

    page_obj = Paginator(items, 20).get_page(request.GET.get("page"))

    # Used for live sidebar badge and notification list updates.
    if request.GET.get("format") == "json":
        return JsonResponse({
            "ok": True,
            "unread_count": unread_count,
            "total_count": inbox.count(),
            "page": page_obj.number,
            "num_pages": page_obj.paginator.num_pages,
            "items": [
                {
                    "id": item.pk,
                    "kind": item.kind,
                    "title": item.title,
                    "message": item.message,
                    "is_read": item.is_read,
                    "created_at": item.created_at.isoformat(),
                }
                for item in page_obj
            ],
        })

    from urllib.parse import urlencode
    return render(request, "accounts/accounts_notifications.html", {
        "current_user": current_user,
        "accounts_active": "notifications",
        "page_obj": page_obj,
        "selected_filter": selected_filter,
        "query": query,
        "notification_query": urlencode({"filter": selected_filter, "q": query, "kind": selected_kind}),
        "selected_kind": selected_kind,
        "archived_count": archived.count(),
        "result_count": page_obj.paginator.count,
        "unread_count": unread_count,
        "total_count": inbox.count(),
    })



@never_cache
@require_POST
def accounts_notifications_read_all(request):
    from django.http import JsonResponse
    from .models import AccountsNotification

    current_user = _accounts_user(request)
    if current_user is None:
        return JsonResponse({
            "ok": False,
            "message": "Accounts access required.",
        }, status=403 if request.session.get("user_id") else 401)

    AccountsNotification.objects.filter(
        recipient=current_user,
        is_archived=False,
        is_read=False,
    ).update(is_read=True)

    return JsonResponse({
        "ok": True,
        "message": "All notifications marked as read.",
        "unread_count": 0,
    })
    from django.http import JsonResponse
    from .models import AccountsNotification

    current_user = _accounts_user(request)
    if current_user is None:
        return JsonResponse({
            "ok": False,
            "message": "Accounts access required.",
        }, status=403 if request.session.get("user_id") else 401)

    AccountsNotification.objects.filter(
        recipient=current_user,
        is_archived=False,
        is_read=False,
    ).update(is_read=True)

    return JsonResponse({
        "ok": True,
        "message": "All notifications marked as read.",
        "unread_count": 0,
    })

def _process_accounts_reminders(owner):
    from uuid import uuid4

    from django.db import transaction
    from django.urls import reverse
    from django.utils import timezone

    from .models import AccountsNotification, AccountsReminder

    now = timezone.now()

    due_ids = list(
        AccountsReminder.objects.filter(
            owner=owner,
            is_completed=False,
            notified_at__isnull=True,
            remind_at__lte=now,
        )
        .order_by("remind_at", "id")
        .values_list("pk", flat=True)[:200]
    )

    created_count = 0

    for reminder_id in due_ids:
        with transaction.atomic():
            reminder = (
                AccountsReminder.objects
                .select_for_update()
                .filter(
                    pk=reminder_id,
                    owner=owner,
                    is_completed=False,
                    notified_at__isnull=True,
                    remind_at__lte=now,
                )
                .first()
            )

            if reminder is None:
                continue

            claimed = AccountsReminder.objects.filter(
                pk=reminder.pk,
                is_completed=False,
                notified_at__isnull=True,
            ).update(notified_at=now)

            if not claimed:
                continue

            event_time = timezone.localtime(
                reminder.event_at
            ).strftime("%d %b %Y, %I:%M %p")

            notification = AccountsNotification.objects.create(
                recipient=owner,
                kind="reminder",
                title=reminder.title,
                message=(
                    f"Reminder: {reminder.title}\n"
                    f"Event: {event_time}"
                    + (
                        f"\n{reminder.description}"
                        if reminder.description else ""
                    )
                ),
                url=reverse("accounts_notifications"),
                event_key=(
                    f"accounts-reminder:"
                    f"{reminder.pk}:{uuid4().hex}"
                ),
            )

            AccountsReminder.objects.filter(
                pk=reminder.pk
            ).update(notification=notification)

            created_count += 1

    return created_count


@never_cache
@require_http_methods(["GET", "POST"])
def accounts_reminders(request):
    import calendar
    import re
    from datetime import date

    from django.contrib import messages
    from django.db import transaction
    from django.http import JsonResponse, HttpResponseForbidden
    from django.shortcuts import redirect, render
    from django.utils import timezone
    from django.utils.dateparse import parse_datetime

    from .models import AccountsNotification, AccountsReminder

    current_user = _accounts_user(request)

    if current_user is None:
        if not request.session.get("user_id"):
            return redirect("login_view")
        return HttpResponseForbidden("Accounts access required.")

    is_ajax = (
        request.headers.get("X-Requested-With")
        == "XMLHttpRequest"
    )

    def serialize(reminder):
        local_event = timezone.localtime(reminder.event_at)
        local_remind = timezone.localtime(reminder.remind_at)

        return {
            "id": reminder.pk,
            "title": reminder.title,
            "description": reminder.description,
            "date": local_event.strftime("%Y-%m-%d"),
            "event_at": local_event.strftime("%Y-%m-%dT%H:%M"),
            "remind_at": local_remind.strftime("%Y-%m-%dT%H:%M"),
            "is_completed": reminder.is_completed,
        }

    def parse_local_datetime(value):
        try:
            parsed = parse_datetime(value or "")
            if parsed and timezone.is_naive(parsed):
                parsed = timezone.make_aware(
                    parsed,
                    timezone.get_current_timezone(),
                )
            return parsed
        except (ValueError, TypeError, OverflowError):
            return None

    def failure(message, status=400):
        if is_ajax:
            return JsonResponse({
                "ok": False,
                "message": message,
            }, status=status)

        messages.error(request, message)
        return redirect("accounts_reminders")

    def success(message, reminder=None, deleted_id=None):
        if is_ajax:
            unread_count = AccountsNotification.objects.filter(
                recipient=current_user,
                is_archived=False,
                is_read=False,
            ).count()

            return JsonResponse({
                "ok": True,
                "message": message,
                "id": reminder.pk if reminder else deleted_id,
                "is_completed": (
                    reminder.is_completed if reminder else None
                ),
                "item": serialize(reminder) if reminder else None,
                "unread_count": unread_count,
            })

        messages.success(request, message)
        return redirect("accounts_reminders")

    reminders = AccountsReminder.objects.filter(
        owner=current_user
    )

    if request.method == "POST":
        action = request.POST.get("action", "")

        if action not in {
            "create", "update", "delete", "complete", "reopen"
        }:
            return failure("Invalid reminder action.")

        with transaction.atomic():
            reminder = None

            if action != "create":
                try:
                    reminder_id = int(
                        request.POST.get("reminder_id", "")
                    )
                except (ValueError, TypeError):
                    return failure("Invalid reminder ID.")

                reminder = (
                    reminders.select_for_update()
                    .filter(pk=reminder_id)
                    .first()
                )

                if reminder is None:
                    return failure(
                        "Reminder not found.",
                        status=404,
                    )

            if action == "delete":
                deleted_id = reminder.pk

                if reminder.notification_id:
                    AccountsNotification.objects.filter(
                        pk=reminder.notification_id,
                        recipient=current_user,
                    ).update(is_archived=True)

                reminder.delete()
                return success(
                    "Reminder deleted.",
                    deleted_id=deleted_id,
                )

            if action in {"complete", "reopen"}:
                reminder.is_completed = action == "complete"
                reminder.save(
                    update_fields=["is_completed", "updated_at"]
                )

                if (
                    reminder.is_completed
                    and reminder.notification_id
                ):
                    AccountsNotification.objects.filter(
                        pk=reminder.notification_id,
                        recipient=current_user,
                    ).update(is_read=True)

                return success(
                    "Reminder completed."
                    if reminder.is_completed
                    else "Reminder reopened.",
                    reminder,
                )

            title = request.POST.get("title", "").strip()
            description = request.POST.get(
                "description", ""
            ).strip()

            event_at = parse_local_datetime(
                request.POST.get("event_at")
            )
            remind_at = parse_local_datetime(
                request.POST.get("remind_at")
            )

            if not title or len(title) > 180:
                return failure(
                    "Enter a title with a maximum of 180 characters."
                )

            if event_at is None or remind_at is None:
                return failure(
                    "Enter valid event and reminder dates and times."
                )

            if remind_at > event_at:
                return failure(
                    "Reminder time must be on or before the event."
                )

            schedule_changed = (
                reminder is None
                or reminder.event_at != event_at
                or reminder.remind_at != remind_at
            )

            if schedule_changed and remind_at <= timezone.now():
                return failure("Choose a future reminder time.")

            if reminder is None:
                reminder = AccountsReminder(owner=current_user)

            if schedule_changed:
                if reminder.notification_id:
                    AccountsNotification.objects.filter(
                        pk=reminder.notification_id,
                        recipient=current_user,
                    ).update(is_archived=True)

                reminder.notified_at = None
                reminder.notification = None
                reminder.is_completed = False

            reminder.title = title
            reminder.description = description
            reminder.event_at = event_at
            reminder.remind_at = remind_at
            reminder.save()

            return success(
                "Reminder created."
                if action == "create"
                else "Reminder updated.",
                reminder,
            )

    _process_accounts_reminders(current_user)

    today = timezone.localdate()
    month_value = request.GET.get(
        "month", today.strftime("%Y-%m")
    )

    try:
        if not re.fullmatch(r"\d{4}-\d{2}", month_value):
            raise ValueError
        selected_month = date.fromisoformat(
            month_value + "-01"
        )
    except (ValueError, TypeError):
        selected_month = today.replace(day=1)

    month_reminders = reminders.filter(
        event_at__year=selected_month.year,
        event_at__month=selected_month.month,
    ).order_by("event_at", "id")

    calendar_items = [
        serialize(reminder)
        for reminder in month_reminders
    ]

    context = {
        "current_user": current_user,
        "accounts_active": "reminders",
        "selected_month": selected_month.strftime("%Y-%m"),
        "month_label": selected_month.strftime("%B %Y"),
        "calendar_weeks": calendar.Calendar(
            firstweekday=0
        ).monthdayscalendar(
            selected_month.year,
            selected_month.month,
        ),
        "calendar_items": calendar_items,
        "today": today.isoformat(),
        "reminders": month_reminders,
    }

    if request.GET.get("format") == "json":
        return JsonResponse({
            "ok": True,
            "selected_month": context["selected_month"],
            "month_label": context["month_label"],
            "calendar_weeks": context["calendar_weeks"],
            "items": calendar_items,
            "today": context["today"],
        })

    return render(
        request,
        "accounts/accounts_reminders.html",
        context,
    )

from urllib.parse import urlencode
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponseBadRequest, HttpResponseForbidden, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_POST, require_http_methods
from .models import Notepad


@never_cache
@require_http_methods(["GET", "POST"])
def accounts_notepad(request):
    if not request.session.get("user_id"):
        return redirect("login_view")
    user = _accounts_user(request)
    if user is None:
        return HttpResponseForbidden("Accounts access required.")
    owned = Notepad.objects.filter(user=user)
    params = request.POST if request.method == "POST" else request.GET
    query = params.get("q", "").strip()[:200]
    sort = params.get("sort", "updated")
    orderings = {"updated": ("-updated_at", "-id"), "created": ("-created_at", "-id"), "title": ("title", "id")}
    if sort not in orderings:
        sort = "updated"
    notes = owned
    if query:
        notes = notes.filter(Q(title__icontains=query) | Q(content__icontains=query))
    page_obj = Paginator(notes.order_by(*orderings[sort]), 4).get_page(params.get("page"))
    selected_id = request.POST.get("note_id") if request.method == "POST" else None
    note = None
    if selected_id:
        try:
            selected_id = int(selected_id)
        except (ValueError, TypeError):
            return HttpResponseBadRequest("Invalid note ID.")
        note = get_object_or_404(owned, pk=selected_id)
    error = ""
    title = note.title if note else ""
    content = (note.content or "") if note else ""
    if request.method == "POST":
        title = request.POST.get("title", "").strip() or "Untitled"
        content = request.POST.get("content", "")
        if len(title) > 255:
            error = "Title must be 255 characters or fewer."
        elif not content.strip():
            error = "Write some content before saving the note."
        else:
            if note:
                note.title, note.content = title, content
                note.save()
                _accounts_notepad_notification(user, note.title, "updated")               
                if request.headers.get("X-Requested-With") == "XMLHttpRequest":
                    return JsonResponse({
                        "ok": True, "id": note.pk, "title": note.title,
                        "content": note.content or "",
                        "updated_at": timezone.localtime(note.updated_at).strftime("%d %b %Y, %I:%M %p"),
                    })
                messages.success(request, "Note updated successfully.")
            else:
                note = Notepad.objects.create(
                    user=user,
                    title=title,
                    content=content,
                )
                _accounts_notepad_notification(user, note.title, "created")
                messages.success(request, "Note created successfully.")
            return redirect(reverse("accounts_notepad") + "?" + urlencode({
                "q": query if selected_id else "", "sort": sort if selected_id else "updated", "page": page_obj.number if selected_id else 1,
            }))
    if error and selected_id and request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return JsonResponse({"ok": False, "message": error}, status=400)
    response = render(request, "accounts/accounts_notepad.html", {
        "current_user": user, "accounts_active": "notepad",
        "note": None, "page_obj": page_obj, "query": query, "sort": sort,
        "form_title": title if not selected_id else "", "form_content": content if not selected_id else "", "note_error": error,
        "total_notes": owned.count(),
    })
    if error:
        response.status_code = 400
    return response


@never_cache
@require_POST
def accounts_notepad_delete(request, pk):
    if not request.session.get("user_id"):
        return redirect("login_view")
    user = _accounts_user(request)
    if user is None:
        return HttpResponseForbidden("Accounts access required.")
    note = get_object_or_404(Notepad, pk=pk, user=user)
    note_title = note.title
    note.delete()
    _accounts_notepad_notification(user, note_title, "deleted")
    messages.success(request, "Note deleted successfully.")
    sort = request.POST.get("sort", "updated")
    if sort not in {"updated", "created", "title"}:
        sort = "updated"
    return redirect(reverse("accounts_notepad") + "?" + urlencode({
        "q": request.POST.get("q", "").strip()[:200],
        "sort": sort, "page": request.POST.get("page", "1"),
    }))


def _accounts_notepad_notification(user, note_title, action):
    from uuid import uuid4
    from django.urls import reverse
    from .models import AccountsNotification

    labels = {
        "created": "Note created",
        "updated": "Note updated",
        "deleted": "Note deleted",
    }

    AccountsNotification.objects.create(
        recipient=user,
        kind="note",
        title=labels[action],
        message=f'{labels[action]}: "{note_title}"',
        url=reverse("accounts_notepad"),
        event_key=f"accounts-note:{uuid4().hex}",
    )

from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods

@never_cache
@require_http_methods(["GET", "POST"])
def accounts_profile(request):
    from django import forms
    from django.contrib import messages
    from django.db import IntegrityError, transaction
    from django.http import HttpResponseForbidden, JsonResponse
    from django.shortcuts import redirect, render
    from .models import User

    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    user = _accounts_user(request)
    if user is None:
        if is_ajax:
            return JsonResponse({"success": False, "message": "Please sign in with an Accounts account."}, status=403 if request.session.get("user_id") else 401)
        if not request.session.get("user_id"):
            return redirect("login_view")
        return HttpResponseForbidden("Accounts access required.")

    class ProfileForm(forms.ModelForm):
        class Meta:
            model = User
            fields = ["name", "email", "phone"]

        def clean_name(self):
            name = self.cleaned_data["name"].strip()
            if not name:
                raise forms.ValidationError("Enter your full name.")
            return name

    form = ProfileForm(instance=user)
    error = ""

    def fail(message, errors=None):
        if is_ajax:
            return JsonResponse({"success": False, "message": message, "errors": errors or {}}, status=400)
        messages.error(request, message)
        return redirect("accounts_profile")

    def result(message):
        if is_ajax:
            return JsonResponse({
                "success": True,
                "message": message,
                "name": user.name,
                "email": user.email,
                "phone": user.phone or "",
                "image_url": user.profile_image.url if user.profile_image else "",
                "initial": (user.name or user.username or "A")[:1].upper(),
            })
        messages.success(request, message)
        return redirect("accounts_profile")

    if request.method == "POST":
        action = request.POST.get("action")
        if action == "upload_photo":
            upload = request.FILES.get("profile_image")
            if not upload:
                return fail("Choose a profile photo.")
            if upload.size > 5 * 1024 * 1024:
                return fail("Choose an image smaller than 5 MB.")
            try:
                image = forms.ImageField().clean(upload)
                if getattr(image, "image", None).format not in {"JPEG", "PNG", "WEBP"}:
                    return fail("Use a JPG, PNG or WebP image.")
            except forms.ValidationError:
                return fail("Choose a valid JPG, PNG or WebP image.")
            user.profile_image = image
            user.save(update_fields=["profile_image"])
            return result("Profile photo updated successfully.")

        if action != "edit_profile":
            return fail("Invalid profile action.")

        form = ProfileForm(request.POST, instance=user)
        if form.is_valid():
            try:
                with transaction.atomic():
                    user = form.save(commit=False)
                    # Save only editable personal fields; retain employment fields.
                    user.save(update_fields=["name", "email", "phone"])
            except IntegrityError:
                return fail("This email is already in use. Choose another email.")
            return result("Profile updated successfully.")
        error = "Please correct the highlighted fields."
        if is_ajax:
            return fail(error, form.errors.get_json_data())

    return render(request, "accounts/accounts_profile.html", {
        "current_user": user,
        "profile_user": user,
        "accounts_active": "profile",
        "profile_form": form,
        "profile_error": error,
    }, status=400 if error else 200)


@never_cache
@require_http_methods(["GET", "POST"])
def accounts_income_edit(request, pk):
    actor = _accounts_user(request)
    if actor is None:
        return JsonResponse({"ok": False, "message": "Please log in with an Accounts account."},
            status=403 if request.session.get("user_id") else 401)
    with transaction.atomic():
        entry = AccountsIncome.objects.select_for_update().filter(pk=pk).first()
        if entry is None:
            return JsonResponse({"ok": False, "message": "This income no longer exists."}, status=404)
        if request.method == "GET":
            return JsonResponse({"ok": True, "item": {
                "date": entry.date.isoformat(), "category": entry.category,
                "amount": str(entry.amount),
                "payment_method": entry.payment_method,
                "description": entry.description, "reference": entry.reference,
            }, "receipt_url": reverse("accounts_income_receipt", args=[entry.pk]) if entry.receipt else ""})
        before = _income_snapshot(entry)
        form = AccountsIncomeForm(request.POST, request.FILES, instance=entry)
        if not form.is_valid():
            return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
                "errors": form.errors.get_json_data()}, status=400)
        # Recheck duplicates when identifying fields change.
        identifying = {"date", "category", "amount", "reference"}
        if any(str(form.cleaned_data.get(key) or "") != before.get(key, "") for key in identifying) and _income_duplicate(form, entry.pk) and request.POST.get("confirm_duplicate") != "1":
            return _income_duplicate_response()
        entry = form.save(commit=False)
        if request.POST.get("remove_receipt") == "1" and not request.FILES.get("receipt"):
            entry.receipt = ""
        entry.save()
        _income_audit(entry, actor, "updated", before)
        _create_accounts_notification(
            kind="income",
            title="Income updated",
            message=(
                f"{actor.name} updated income "
                f"#{entry.pk}: {entry.category}, "
                f"₹{entry.amount:,.2f} "
                f"on {entry.date:%d %b %Y}."
            ),
            actor=actor,
        )
    return JsonResponse({"ok": True, "message": "Income updated successfully.", "id": entry.pk})


@never_cache
@require_http_methods(["GET", "POST"])
def accounts_expense_edit(request, pk):
    actor = _accounts_user(request)
    if actor is None:
        return JsonResponse({"ok": False, "message": "Please log in with an Accounts account."},
            status=403 if request.session.get("user_id") else 401)
    with transaction.atomic():
        entry = AccountsExpense.objects.select_for_update().filter(pk=pk).first()
        if entry is None:
            return JsonResponse({"ok": False, "message": "This expense no longer exists."}, status=404)
        if request.method == "GET":
            return JsonResponse({"ok": True, "item": {
                "date": entry.date.isoformat(), "category": entry.category,
                "amount": str(entry.amount), "paid_amount": str(entry.effective_paid_amount),
                "payment_method": entry.payment_method, "paid_to": entry.paid_to,
                "description": entry.description, "reference": entry.reference,
            }, "receipt_url": reverse("accounts_expense_receipt", args=[entry.pk]) if entry.receipt else ""})
        before = _expense_snapshot(entry)
        form = AccountsExpenseForm(request.POST, request.FILES, instance=entry)
        if not form.is_valid():
            return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
                "errors": form.errors.get_json_data()}, status=400)
        # Recheck duplicates when identifying fields change; editing only paid amount is safe.
        identifying = {"date", "category", "amount", "paid_to", "reference"}
        if any(str(form.cleaned_data.get(key) or "") != before.get(key, "") for key in identifying) and _expense_duplicate(form, entry.pk) and request.POST.get("confirm_duplicate") != "1":
            return _expense_duplicate_response()
        entry = form.save(commit=False)
        if request.POST.get("remove_receipt") == "1" and not request.FILES.get("receipt"):
            entry.receipt = ""
        entry.save()
        _expense_audit(entry, actor, "updated", before)
        _create_accounts_notification(
            kind="expense",
            title="Expense updated",
            message=(
                f"{actor.name} updated expense "
                f"#{entry.pk}: {entry.category}, "
                f"₹{entry.amount:,.2f} "
                f"on {entry.date:%d %b %Y}."
            ),
            actor=actor,
        )
    return JsonResponse({"ok": True, "message": "Expense updated successfully.", "id": entry.pk})


@never_cache
@require_http_methods(["GET", "POST"])
def accounts_sale_edit(request, pk):
    actor = _accounts_user(request)
    if actor is None:
        return JsonResponse({"ok": False, "message": "Please log in with an Accounts account."},
            status=403 if request.session.get("user_id") else 401)
    with transaction.atomic():
        sale = AccountsSale.objects.select_for_update().filter(pk=pk).first()
        if sale is None:
            return JsonResponse({"ok": False, "message": "This sale no longer exists."}, status=404)
        if request.method == "GET":
            return JsonResponse({"ok": True, "item": {
                "date": sale.date.isoformat(), "customer_name": sale.customer_name,
                "invoice_number": sale.invoice_number, "description": sale.description,
                "amount": str(sale.amount), "due_date": sale.due_date.isoformat() if sale.due_date else ""},
                "invoice_url": reverse("accounts_sale_invoice", args=[sale.pk]) if sale.invoice_file else "",
                "received_amount": str(sale.received_amount)})
        before = {"date": sale.date.isoformat(), "customer_name": sale.customer_name,
            "invoice_number": sale.invoice_number, "amount": str(sale.amount)}
        form = AccountsSaleForm(request.POST, request.FILES, instance=sale)
        if not form.is_valid():
            return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
                "errors": form.errors.get_json_data()}, status=400)
        if (any(str(form.cleaned_data.get(key) or "") != value for key, value in before.items())
                and _sale_duplicate(form, sale.pk) and request.POST.get("confirm_duplicate") != "1"):
            return _sale_duplicate_response()
        sale = form.save(commit=False)
        if request.POST.get("remove_invoice") == "1" and not request.FILES.get("invoice_file"):
            sale.invoice_file = ""
        sale.save()
        _create_accounts_notification(
            kind="sale",
            title="Sale updated",
            message=(
                f"{actor.name} updated sale "
                f"#{sale.pk} for {sale.customer_name}: "
                f"₹{sale.amount:,.2f}."
            ),
            actor=actor,
        )
    return JsonResponse({
        "ok": True,
        "message": "Sale updated successfully.",
        "id": sale.pk,
    })





#<------------MASTER-ADMIN------------------>

def _master_admin_allowed(request):
    user = request.user
    return bool(user.is_authenticated and user.is_active and user.is_superuser)


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_login(request):
    if _master_admin_allowed(request):
        return redirect("master_admin_dashboard")
    if request.method == "POST":
        username = request.POST.get("username", "").strip()
        password = request.POST.get("password", "")
        auth_user = authenticate(request, username=username, password=password)
        if auth_user is not None and auth_user.is_active and auth_user.is_superuser:
            login(request, auth_user)
            # This portal uses Django authentication, not an employee identity.
            for key in ("user_id", "position", "login_time"):
                request.session.pop(key, None)
            return redirect("master_admin_dashboard")
        messages.error(request, "Invalid credentials or Master Admin access unavailable.")
    return render(request, "master_admin/master_admin_login.html")


@never_cache
@require_POST
def master_admin_logout(request):
    from django.http import HttpResponseForbidden
    if not _master_admin_allowed(request):
        return HttpResponseForbidden("Master Admin access required.")
    logout(request)
    return redirect("master_admin_login")


@never_cache
@require_GET
def master_admin_dashboard(request):
    from datetime import timedelta
    from decimal import Decimal
    from django.contrib.auth import get_user_model
    from django.db.models import Count, Q, Sum
    from django.db.models.functions import Lower, Trim
    from django.http import HttpResponseForbidden, JsonResponse
    from django.shortcuts import redirect, render
    from django.utils import timezone
    from .models import User as Employee, Department, Team
    from .models import AccountsIncome, AccountsExpense, AccountsSale

    json_request = request.GET.get("format") == "json"
    if not _master_admin_allowed(request):
        if json_request:
            return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
        if not request.user.is_authenticated:
            return redirect("master_admin_login")
        return HttpResponseForbidden("Master Admin access required.")
    now = timezone.now()
    today = timezone.localdate(now)
    month_start = today.replace(day=1)
    next_month = (month_start.replace(year=month_start.year + 1, month=1)
                  if month_start.month == 12 else month_start.replace(month=month_start.month + 1))
    # Master Admin's linked chat identity is not another employee registration.
    employees = Employee.objects.annotate(portal_role=Lower(Trim("job_Position"))).exclude(
        portal_role__in=["master admin", "master_admin"])
    counts = employees.aggregate(
        users=Count("pk"),
        admins=Count("pk", filter=Q(portal_role__in=["management", "admin"])),
        leads=Count("pk", filter=Q(portal_role__in=["team lead", "team_lead", "team leader", "teamlead"])),
        members=Count("pk", filter=Q(portal_role__in=["team member", "team_member", "team members", "teammember"])),
        accounts=Count("pk", filter=Q(portal_role__in=["accounts", "accounts team", "accounts_team"])),
        active=Count("pk", filter=Q(status="active", last_activity__gte=now-timedelta(minutes=10))),
    )
    counts["inactive"] = counts["users"] - counts["active"]
    counts["master_admins"] = get_user_model().objects.filter(is_superuser=True, is_active=True).count()
    counts["departments"] = Department.objects.count()
    counts["teams"] = Team.objects.count()

    def month_total(model):
        total = model.objects.filter(date__gte=month_start, date__lt=next_month).aggregate(total=Sum("amount"))["total"]
        return format(total if total is not None else Decimal("0.00"), ".2f")

    financial = {"income": month_total(AccountsIncome), "expenses": month_total(AccountsExpense), "sales": month_total(AccountsSale)}
    recent_users = [{"id": item.pk, "name": item.name, "employee_id": item.employee_id,
        "role": item.job_Position, "department": item.department.name if item.department else "Not assigned",
        "team": item.team.name if item.team else "Not assigned"}
        for item in employees.select_related("department", "team").order_by("-id")[:10]]
    payload = {"ok": True, "counts": counts, "financial": financial, "recent_users": recent_users,
        "month_label": month_start.strftime("%B %Y"),
        "updated_at": timezone.localtime(now).strftime("%d %b %Y, %I:%M:%S %p")}
    if json_request:
        return JsonResponse(payload)
    return render(request, "master_admin/master_admin_dashboard.html", {
        "master_user": request.user, "dashboard_data": payload, **payload,
    })




def _master_employee_form(data=None, files=None, instance=None):
    from django import forms
    from django.contrib.auth.hashers import make_password
    from django.contrib.auth.password_validation import validate_password
    from .models import User as Employee, Designation

    class EmployeeForm(forms.ModelForm):
        password_input = forms.CharField(required=False, strip=False)
        joining_date_input = forms.DateField(required=False)
        job_Position = forms.ChoiceField(choices=[
            ("Management", "Admin (Management)"), ("Team Lead", "Team Lead"),
            ("Team Member", "Team Member"), ("Accounts", "Accounts"),
        ])

        class Meta:
            model = Employee
            fields = ["employee_id", "name", "email", "phone", "department", "team",
                      "job_Position", "designation", "work_location", "username", "profile_image", "status"]

        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.fields["department"].required = False
            self.fields["designation"].required = False

        def clean(self):
            cleaned = super().clean()
            department, designation = cleaned.get("department"), cleaned.get("designation", "")
            if department and not Designation.objects.filter(department=department, name=designation).exists():
                self.add_error("designation", "Choose a designation belonging to the selected department.")
            elif not department:
                cleaned["designation"] = ""
            password = cleaned.get("password_input", "")
            if not self.instance.pk and not password:
                self.add_error("password_input", "Enter a password for the new user.")
            elif password:
                try:
                    validate_password(password, self.instance)
                except forms.ValidationError as error:
                    self.add_error("password_input", error)
            return cleaned

        def save(self, commit=True):
            employee = super().save(commit=False)
            password = self.cleaned_data.get("password_input")
            if password:
                employee.password = make_password(password)
            if commit:
                employee.save()
                # joining_date is auto_now_add: explicitly persist a chosen date after creation.
                chosen_date = self.cleaned_data.get("joining_date_input")
                if chosen_date:
                    Employee.objects.filter(pk=employee.pk).update(joining_date=chosen_date)
                    employee.joining_date = chosen_date
            return employee

    return EmployeeForm(data, files, instance=instance)


def _master_employee_save(request, instance=None):
    from django.db import IntegrityError, transaction
    from django.http import JsonResponse

    editing = instance is not None
    keys = {
        "employee_id": "edit_emp_id", "name": "edit_name", "email": "edit_email",
        "phone": "edit_phone", "department": "edit_department", "team": "edit_team",
        "job_Position": "edit_job_position", "designation": "edit_designation",
        "work_location": "edit_work_location", "username": "edit_username",
        "status": "edit_status",
    }
    data = {key: request.POST.get(source if editing else key, "") for key, source in keys.items()}
    data["password_input"] = request.POST.get("edit_password" if editing else "password", "")
    data["joining_date_input"] = request.POST.get("edit_joining_date" if editing else "joining_date", "")
    image = request.FILES.get("edit_profile_upload" if editing else "profile_image")
    files = {"profile_image": image} if image else None
    form = _master_employee_form(data, files, instance)
    ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"

    def result(ok, message, status=400, errors=None, employee=None):
        if ajax:
            return JsonResponse({"ok": ok, "message": message, "errors": errors or {},
                                 "id": employee.pk if employee else None}, status=(200 if editing else 201) if ok else status)
        (messages.success if ok else messages.error)(request, message)
        return redirect("master_admin_usermanagement")

    if not form.is_valid():
        details = " ".join(f"{form.fields[field].label or field.replace('_', ' ').title()}: {error}"
                           for field, errors in form.errors.items() for error in errors if field in form.fields)
        return result(False, details or "Please correct the form fields.", errors=form.errors.get_json_data())
    try:
        with transaction.atomic():
            employee = form.save()
    except IntegrityError:
        return result(False, "Employee ID, username or email is already in use.")
    return result(True, "User updated successfully." if editing else "User created successfully.", employee=employee)


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_usermanagement(request):
    from django.core.paginator import Paginator
    from django.db.models import Q
    from django.http import HttpResponseForbidden, JsonResponse
    from django.utils import timezone
    from .models import User as Employee, Department, Team, Designation

    if not _master_admin_allowed(request):
        if request.method == "POST" or request.headers.get("X-Requested-With") == "XMLHttpRequest":
            return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
        if not request.user.is_authenticated:
            return redirect("master_admin_login")
        return HttpResponseForbidden("Master Admin access required.")
    if request.method == "POST":
        return _master_employee_save(request)
    if "designation_department" in request.GET:
        try:
            department_id = int(request.GET.get("designation_department", ""))
        except (ValueError, TypeError):
            return JsonResponse({"ok": False, "message": "Invalid department."}, status=400)
        if not Department.objects.filter(pk=department_id).exists():
            return JsonResponse({"ok": False, "message": "Department not found."}, status=404)
        return JsonResponse({"ok": True, "items": list(Designation.objects.filter(
            department_id=department_id).order_by("name").values_list("name", flat=True))})

    users = Employee.objects.select_related("department", "team").order_by("-id")
    query = request.GET.get("search", "").strip()[:200]
    if query:
        users = users.filter(Q(name__icontains=query) | Q(email__icontains=query)
                             | Q(employee_id__icontains=query) | Q(username__icontains=query))
    department = request.GET.get("department", "")
    team = request.GET.get("team", "")
    position = request.GET.get("position", "")
    if department:
        users = users.filter(department__name=department)
    if team:
        users = users.filter(team__name=team)
    if position:
        users = users.filter(job_Position=position)
    total = users.count()
    page = Paginator(users, 25).get_page(request.GET.get("page", 1))
    # Display the stored status selected in Add/Edit.
    catalog = {}
    for department_id, name in Designation.objects.order_by("name").values_list("department_id", "name"):
        catalog.setdefault(str(department_id), []).append(name)
    roles = sorted(set(Employee.objects.exclude(job_Position="").values_list("job_Position", flat=True))
                   | {"Management", "Team Lead", "Team Member", "Accounts"})
    return render(request, "master_admin/master_admin_usermanagement.html", {
        "users": page.object_list, "page_obj": page, "total_filtered": total,
        "departments": Department.objects.order_by("name"), "teams": Team.objects.order_by("name"),
        "job_positions": roles, "designation_options": catalog,
    })


@never_cache
@require_POST
def master_admin_edit_user(request):
    from django.http import JsonResponse
    from .models import User as Employee
    if not _master_admin_allowed(request):
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    try:
        employee_id = int(request.POST.get("id", ""))
    except (ValueError, TypeError):
        return JsonResponse({"ok": False, "message": "Invalid user ID."}, status=400)
    employee = Employee.objects.filter(pk=employee_id).first()
    if employee is None:
        return JsonResponse({"ok": False, "message": "User not found."}, status=404)
    return _master_employee_save(request, employee)


@never_cache
@require_POST
def master_admin_delete_user(request, pk):
    from django.db import IntegrityError, transaction
    from django.db.models.deletion import ProtectedError, RestrictedError
    from django.http import JsonResponse
    from .models import User as Employee
    if not _master_admin_allowed(request):
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    employee = Employee.objects.filter(pk=pk).first()
    if employee is None:
        return JsonResponse({"ok": False, "message": "User not found."}, status=404)
    try:
        with transaction.atomic():
            employee.delete()
    except (ProtectedError, RestrictedError):
        return JsonResponse({"ok": False, "message": "This user has protected records and cannot be deleted."}, status=409)
    except IntegrityError:
        return JsonResponse({"ok": False, "message": "Linked records prevent deletion. No records were deleted."}, status=409)
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return JsonResponse({"ok": True, "message": "User deleted successfully.", "id": pk})
    messages.success(request, "User deleted successfully.")
    return redirect("master_admin_usermanagement")



# ADD this NEW function once to monitoringapp/views.py.
# Requires the existing _master_admin_allowed() helper. No migrations.

@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_departments(request):
    from django.contrib import messages
    from django.db import IntegrityError, transaction
    from django.db.models.deletion import ProtectedError, RestrictedError
    from django.http import JsonResponse, HttpResponseForbidden
    from django.shortcuts import redirect, render
    from .models import Department, Designation, User

    ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    if not _master_admin_allowed(request):
        if request.method == "POST" or ajax or request.GET.get("format") == "json":
            return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
        if not request.user.is_authenticated:
            return redirect("master_admin_login")
        return HttpResponseForbidden("Master Admin access required.")

    def finish(message, ok=True):
        if ajax:
            return JsonResponse({"ok": ok, "message": message}, status=200 if ok else 400)
        (messages.success if ok else messages.error)(request, message)
        return redirect("master_admin_departments")

    if request.method == "POST":
        names = [value.strip() for value in request.POST.getlist("designation_names") if value.strip()]
        # Discard duplicate names within this form, ignoring case.
        names = list({name.casefold(): name for name in names}.values())
        if any(len(name) > 100 for name in names):
            return finish("Designation names must be 100 characters or fewer.", False)
        try:
            with transaction.atomic():
                if "delete_department" in request.POST:
                    try:
                        department_id = int(request.POST.get("department_id", ""))
                    except (ValueError, TypeError):
                        return finish("Choose a valid department.", False)
                    department = Department.objects.select_for_update().filter(pk=department_id).first()
                    if department is None:
                        return finish("Department not found.", False)
                    if User.objects.filter(department=department).exists():
                        return finish("This department is assigned to employees. Reassign them before deleting it.", False)
                    for relation in Department._meta.related_objects:
                        if relation.related_model is Designation:
                            continue
                        if relation.related_model._default_manager.filter(**{relation.field.name: department}).exists():
                            return finish("This department is used by other records and cannot be deleted.", False)
                    department.delete()
                    return finish("Department deleted successfully.")
                if "edit_department" in request.POST:
                    try:
                        department_id = int(request.POST.get("department_id", ""))
                    except (ValueError, TypeError):
                        return finish("Choose a valid department.", False)
                    department = Department.objects.select_for_update().filter(pk=department_id).first()
                    if department is None:
                        return finish("Department not found.", False)
                    department_name = request.POST.get("department_name", "").strip()
                    if not department_name or len(department_name) > 100:
                        return finish("Enter a department name with up to 100 characters.", False)
                    if Department.objects.exclude(pk=department.pk).filter(name__iexact=department_name).exists():
                        return finish("This department name already exists.", False)
                    raw_names = request.POST.getlist("designation_names")
                    raw_ids = request.POST.getlist("designation_ids")
                    if len(raw_names) != len(raw_ids) or not raw_names:
                        return finish("Add at least one designation.", False)
                    existing = {str(item.pk): item for item in department.designations.select_for_update()}
                    changes = []; seen_names = set(); seen_ids = set()
                    for item_id, raw_name in zip(raw_ids, raw_names):
                        name = raw_name.strip()
                        if not name or len(name) > 100 or name.casefold() in seen_names:
                            return finish("Use unique designation names, each up to 100 characters.", False)
                        if item_id and (item_id not in existing or item_id in seen_ids):
                            return finish("Invalid designation selection.", False)
                        seen_names.add(name.casefold()); seen_ids.add(item_id)
                        changes.append((item_id, name))
                    removed = [item for key, item in existing.items() if key not in seen_ids]
                    for item in removed:
                        if User.objects.filter(department=department, designation__iexact=item.name).exists():
                            return finish(f"{item.name} is assigned to employees. Reassign them before removing it.", False)
                    # Temporarily rename retained rows to permit swapping two designation names.
                    from uuid import uuid4
                    for item_id, name in changes:
                        if item_id:
                            Designation.objects.filter(pk=existing[item_id].pk).update(name="tmp-" + uuid4().hex)
                    for item in removed:
                        item.delete()
                    # Match employees against original values before renaming to avoid swap collisions.
                    employee_changes = []
                    for item_id, name in changes:
                        if item_id:
                            item = existing[item_id]
                            employee_ids = list(User.objects.filter(department=department, designation=item.name).values_list("pk", flat=True))
                            employee_changes.append((employee_ids, name))
                            item.name = name; item.save(update_fields=["name"])
                        else:
                            Designation.objects.create(department=department, name=name)
                    for employee_ids, name in employee_changes:
                        User.objects.filter(pk__in=employee_ids).update(designation=name)
                    department.name = department_name; department.save(update_fields=["name"])
                    return finish("Department updated successfully.")
                if "add_department" in request.POST:
                    name = request.POST.get("department_name", "").strip()
                    if not name or len(name) > 100:
                        return finish("Enter a department name with up to 100 characters.", False)
                    if not names:
                        return finish("Add at least one designation.", False)
                    if Department.objects.filter(name__iexact=name).exists():
                        return finish("This department already exists. Add designations to its existing row.", False)
                    department = Department.objects.create(name=name)
                    for designation in names:
                        Designation.objects.create(department=department, name=designation)
                    return finish("Department and designations added successfully.")
                if "add_designation" in request.POST:
                    try:
                        department_id = int(request.POST.get("department_id", ""))
                    except (ValueError, TypeError):
                        return finish("Choose a valid department.", False)
                    department = Department.objects.filter(pk=department_id).first()
                    if department is None or not names:
                        return finish("Choose a department and enter a designation.", False)
                    for name in names:
                        if not Designation.objects.filter(department=department, name__iexact=name).exists():
                            Designation.objects.create(department=department, name=name)
                    return finish("Designations saved successfully.")
                return finish("Invalid department action.", False)
        except (ProtectedError, RestrictedError):
            return finish("This department has protected records and cannot be deleted.", False)
        except IntegrityError:
            return finish("This department or designation already exists.", False)

    from django.db.models import Count
    counts = dict(User.objects.exclude(department_id=None).values_list("department_id").annotate(total=Count("pk")))
    items = [{"id": item.pk, "name": item.name, "employees": counts.get(item.pk, 0),
              "designations": [{"id": value.pk, "name": value.name} for value in item.designations.all()]}
             for item in Department.objects.prefetch_related("designations").order_by("name", "id")]
    if request.GET.get("format") == "json":
        return JsonResponse({"ok": True, "items": items})
    return render(request, "master_admin/master_admin_departments.html", {"department_items": items})



@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_teams(request):
    from django.contrib import messages
    from django.db import IntegrityError, transaction
    from django.db.models import Count
    from django.db.models.deletion import ProtectedError, RestrictedError
    from django.http import JsonResponse, HttpResponseForbidden
    from django.shortcuts import redirect, render
    from .models import Team, User

    ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    if not _master_admin_allowed(request):
        if request.method == "POST" or ajax or request.GET.get("format") == "json":
            return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
        if not request.user.is_authenticated:
            return redirect("master_admin_login")
        return HttpResponseForbidden("Master Admin access required.")

    def finish(message, ok=True, status=400):
        if ajax:
            return JsonResponse({"ok": ok, "message": message}, status=200 if ok else status)
        (messages.success if ok else messages.error)(request, message)
        return redirect("master_admin_teams")

    if request.method == "POST":
        actions = [action for action in ("create", "update", "delete") if action in request.POST]
        if len(actions) != 1:
            return finish("Choose a valid team action.", False)
        action = actions[0]
        name = request.POST.get("team_name", "").strip()
        if action != "delete" and (not name or len(name) > 100):
            return finish("Enter a team name with up to 100 characters.", False)
        try:
            with transaction.atomic():
                team = None
                if action != "create":
                    try:
                        team_id = int(request.POST.get("team_id", ""))
                    except (ValueError, TypeError):
                        return finish("Invalid team ID.", False)
                    team = Team.objects.select_for_update().filter(pk=team_id).first()
                    if team is None:
                        return finish("Team not found.", False, 404)
                if action == "delete":
                    # Preserve employee assignments and other linked records.
                    for relation in Team._meta.related_objects:
                        if relation.related_model._default_manager.filter(**{relation.field.name: team}).exists():
                            return finish("This team is in use. Reassign its employees and linked records before deleting it.", False, 409)
                    team.delete()
                    return finish("Team deleted successfully.")
                duplicates = Team.objects.filter(name__iexact=name)
                if team is not None:
                    duplicates = duplicates.exclude(pk=team.pk)
                if duplicates.exists():
                    return finish("A team with this name already exists.", False)
                if action == "create":
                    Team.objects.create(name=name)
                    return finish("Team created successfully.")
                team.name = name
                team.save(update_fields=["name"])
                return finish("Team updated successfully.")
        except (ProtectedError, RestrictedError):
            return finish("This team has protected records and cannot be deleted.", False, 409)
        except IntegrityError:
            return finish("The change conflicts with an existing team or linked record. No changes were saved.", False, 409)

    counts = dict(User.objects.exclude(team_id=None).values_list("team_id").annotate(total=Count("pk")))
    department_names = {}
    for team_id, department_name in User.objects.exclude(team_id=None).exclude(department_id=None).values_list("team_id", "department__name").distinct():
        department_names.setdefault(team_id, set()).add(department_name)
    items = [{"id": team.pk, "name": team.name, "employees": counts.get(team.pk, 0),
              "departments": sorted(department_names.get(team.pk, set()))}
             for team in Team.objects.order_by("name", "id")]
    if request.GET.get("format") == "json":
        return JsonResponse({"ok": True, "items": items})
    return render(request, "master_admin/master_admin_teams.html", {"team_items": items})



@never_cache
@require_GET
def master_admin_reports(request):
    from datetime import datetime, time, timedelta
    from django.conf import settings
    from django.contrib import messages
    from django.db.models import Q
    from django.http import HttpResponse, HttpResponseForbidden
    from django.shortcuts import redirect, render
    from django.utils import timezone
    from django.utils.dateparse import parse_date
    from .models import MorningReport, EveningReport

    if not _master_admin_allowed(request):
        if not request.user.is_authenticated:
            return redirect("master_admin_login")
        return HttpResponseForbidden("Master Admin access required.")
    report_tab = request.GET.get("tab", "leads")
    if report_tab not in {"leads", "members"}:
        report_tab = "leads"
    role = "Team Lead" if report_tab == "leads" else "Team Member"
    show_all = request.GET.get("all") == "1"
    query = request.GET.get("q", "").strip()[:200]
    filter_date = None
    raw_date = request.GET.get("date", "").strip()
    if raw_date and raw_date != "None":
        try:
            selected_date = parse_date(raw_date)
        except ValueError:
            selected_date = None
        if selected_date:
            filter_date = selected_date.isoformat()
        else:
            messages.error(request, "Choose a valid report date.")

    morning = MorningReport.objects.filter(user__job_Position__iexact=role).select_related("user").order_by("-created_at", "-id")
    evening = EveningReport.objects.filter(user__job_Position__iexact=role).select_related("user").order_by("-created_at", "-id")
    if query:
        match = Q(user__name__icontains=query)
        morning, evening = morning.filter(match), evening.filter(match)
    day = parse_date(filter_date) if filter_date else (None if show_all else timezone.localdate())
    if day is not None:
        start = datetime.combine(day, time.min)
        end = datetime.combine(day + timedelta(days=1), time.min)
        if settings.USE_TZ:
            tz = timezone.get_current_timezone()
            start, end = timezone.make_aware(start, tz), timezone.make_aware(end, tz)
        morning = morning.filter(created_at__gte=start, created_at__lt=end)
        evening = evening.filter(created_at__gte=start, created_at__lt=end)

    if request.GET.get("export") == "1":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        wb = Workbook()
        ws = wb.active
        ws.title = "Team Reports"
        ws.append(["Report Type", "Date", "Time", "Name", "Team", "Department", "Report", "Status"])
        for rows, label in ((morning, "Morning"), (evening, "Evening")):
            for report in rows.iterator(chunk_size=500):
                stamp = report.created_at
                if timezone.is_aware(stamp):
                    stamp = timezone.localtime(stamp)
                values = [label, stamp.strftime("%Y-%m-%d"), stamp.strftime("%I:%M %p"),
                          report.user.name or report.user.username, report.team,
                          report.department, report.report_text, report.status]
                ws.append([ILLEGAL_CHARACTERS_RE.sub("", str(value or "")) for value in values])
                # Treat all employee-entered content as text, including strings starting with '='.
                for cell in ws[ws.max_row]:
                    cell.data_type = "s"
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="7351A4")
        for index, width in enumerate([15, 15, 12, 24, 20, 24, 65, 18], 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="master_admin_{report_tab}_reports.xlsx"'
        wb.save(response)
        return response

    # Return the complete filtered lists; the template scrolls within each table.
    return render(request, "master_admin/master_admin_reports.html", {
        "report_tab": report_tab,
        "show_all": show_all,
        "filter_date": filter_date,
        "query": query,
        "morning_reports": morning,
        "evening_reports": evening,
        "morning_count": morning.count(),
        "evening_count": evening.count(),
    })



from pathlib import Path
from decimal import Decimal
from django.contrib import messages
from django.db import transaction
from django.db.models import F, Q
from django.http import FileResponse, JsonResponse
from django.shortcuts import render, redirect
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods, require_POST, require_GET
from .models import AccountsIncome, AccountsIncomeHistory, AccountsExpense, AccountsExpenseHistory, AccountsSale, AccountsSalePayment, MasterAdminAccountsAudit
from .forms import AccountsIncomeForm, AccountsExpenseForm, AccountsSaleForm, AccountsSalePaymentForm


def _master_accounts_actor(request):
    return request.user if _master_admin_allowed(request) else None


def _master_accounts_actor_name(actor):
    return f"Master Admin: {actor.get_full_name() or actor.username} (#{actor.pk})"[:255]


def _master_accounts_audit(entry, actor, kind, action, before=None):
    snapshot = {}
    for field in entry._meta.concrete_fields:
        if field.name in {"created_by", "created_at", "updated_at"}:
            continue
        value = getattr(entry, field.attname)
        snapshot[field.name] = str(value) if value is not None else ""
    MasterAdminAccountsAudit.objects.create(actor=actor, actor_name=_master_accounts_actor_name(actor),
        kind=kind, entry_number=entry.pk, action=action,
        snapshot={"before": before or {}, "after": {} if action == "deleted" else snapshot, "deleted": snapshot if action == "deleted" else {}})
    if kind in {"income", "expense"}:
        before = before or {}
        after = {} if action == "deleted" else (_income_snapshot(entry) if kind == "income" else _expense_snapshot(entry))
        changes = {key: {"before": before.get(key, ""), "after": after.get(key, "")}
                   for key in before.keys() | after.keys() if before.get(key, "") != after.get(key, "")}
        model = AccountsIncomeHistory if kind == "income" else AccountsExpenseHistory
        model.objects.create(**{kind: entry, kind + "_number": entry.pk}, actor=None,
            actor_name=_master_accounts_actor_name(actor), action=action, changes=changes)


def _master_accounts_creators(page, kind):
    entries = list(page.object_list)
    names = {}
    for audit in MasterAdminAccountsAudit.objects.filter(kind=kind, action="created", entry_number__in=[item.pk for item in entries]).order_by("created_at", "pk"):
        names.setdefault(audit.entry_number, audit.actor_name)
    for item in entries:
        item.recorded_by_name = getattr(item.created_by, "name", "") or names.get(item.pk, "—")
    page.object_list = entries


def _master_accounts_notification(kind, title, message):
    # Existing Accounts notification model accepts actor=None for Django superusers.
    return _create_accounts_notification(kind=kind, title=title, message=message, actor=None)

@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_accounts_income(request):
    import calendar
    import re
    from datetime import date
    from decimal import Decimal
    from io import BytesIO

    from django.core.paginator import Paginator
    from django.db.models import Q, Sum
    from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
    from django.shortcuts import redirect, render
    from django.utils import timezone

    current_user = _master_accounts_actor(request)
    if current_user is None:
        return HttpResponseForbidden("You do not have access to Accounts income.")

    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    form = AccountsIncomeForm(request.POST if request.method == "POST" else None,
        request.FILES if request.method == "POST" else None)
    if request.method == "POST" and form.is_valid():
        if _income_duplicate(form) and request.POST.get("confirm_duplicate") != "1":
            if is_ajax:
                return _income_duplicate_response()
            form.add_error(None, "A similar income already exists. Review the entry and use Save Anyway to confirm.")
        else:
            with transaction.atomic():
                income = form.save(commit=False)
                income.created_by = None
                income.save()
                _master_accounts_audit(income, current_user, "income", "created")
                _master_accounts_notification(
                    kind="income",
                    title="Income added",
                    message=(
                        f"{_master_accounts_actor_name(current_user)} added income "
                        f"#{income.pk}: {income.category}, "
                        f"₹{income.amount:,.2f} "
                        f"on {income.date:%d %b %Y}."
                    ),

                )
            if is_ajax:
                return JsonResponse({"ok": True, "message": "Income added successfully.", "id": income.pk}, status=201)
            messages.success(request, "Income added successfully.")
            return redirect("master_admin_accounts_income")

    if request.method == "POST" and is_ajax:
        return JsonResponse({
            "ok": False,
            "message": "Please correct the highlighted fields.",
            "errors": form.errors.get_json_data(),
        }, status=400)

    today = timezone.localdate()
    month_start = today.replace(day=1)
    month_end = today.replace(day=calendar.monthrange(today.year, today.month)[1])
    all_entries = AccountsIncome.objects.select_related("created_by")
    entries = all_entries
    query = request.GET.get("q", "").strip()[:200]
    if query:
        entries = entries.filter(
            Q(category__icontains=query)
            | Q(description__icontains=query)
            | Q(reference__icontains=query)
        )

    category_filter = request.GET.get("category", "").strip()[:100]
    if category_filter:
        entries = entries.filter(category__iexact=category_filter)

    period = request.GET.get("period", "all").strip()
    selected_date = request.GET.get("date", "").strip() or today.isoformat()
    selected_month = request.GET.get("month", "").strip() or today.strftime("%Y-%m")
    from_date = request.GET.get("start", "").strip()
    to_date = request.GET.get("end", "").strip()
    range_start = range_end = None
    filter_error = ""

    def read_date(value):
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ValueError("Invalid date format.")
        return date.fromisoformat(value)

    try:
        if period == "weekly":
            anchor = read_date(selected_date)
            monday_ordinal = max(date.min.toordinal(), anchor.toordinal() - anchor.weekday())
            range_start = date.fromordinal(monday_ordinal)
            range_end = date.fromordinal(min(date.max.toordinal(), monday_ordinal + 6))
        elif period == "monthly":
            if not re.fullmatch(r"\d{4}-\d{2}", selected_month):
                raise ValueError("Invalid month format.")
            year, month = map(int, selected_month.split("-"))
            range_start = date(year, month, 1)
            range_end = date(year, month, calendar.monthrange(year, month)[1])
        elif period == "custom":
            if not from_date or not to_date:
                raise ValueError("Choose both From Date and To Date.")
            range_start, range_end = read_date(from_date), read_date(to_date)
            if range_start > range_end:
                raise ValueError("From Date must be on or before To Date.")
        elif period != "all":
            raise ValueError("Choose a valid period.")
    except (ValueError, OverflowError) as error:
        filter_error = (
            str(error) if period == "custom"
            else "Choose a valid date, month or period."
        )
        range_start = range_end = None
        entries = entries.none()

    if range_start is not None and not filter_error:
        entries = entries.filter(date__range=(range_start, range_end))
    entries = entries.order_by("-date", "-id")

    def income_total(queryset):
        return queryset.aggregate(total=Sum("amount"))["total"] or Decimal("0.00")

    # Export the entire filtered queryset before applying pagination.
    if request.method == "GET" and request.GET.get("export") == "xlsx" and not filter_error:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Income"
        headers = ["Date", "Category", "Description", "Payment Method", "Reference", "Recorded By", "Amount (INR)"]
        sheet.append(headers)
        total = Decimal("0.00")
        export_creators = dict(MasterAdminAccountsAudit.objects.filter(kind="income", action="created", entry_number__in=entries.values("pk")).values_list("entry_number", "actor_name"))
        for entry in entries.iterator():
            creator = getattr(entry.created_by, "name", "") or export_creators.get(entry.pk, "—")
            sheet.append([
                entry.date, entry.category, entry.description,
                entry.get_payment_method_display(), entry.reference or "—",
                creator, entry.amount,
            ])
            row = sheet.max_row
            # User-entered strings must remain text, including values starting with '='.
            for column in range(2, 7):
                cell = sheet.cell(row, column)
                cell.value = str(cell.value or "")
                cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.cell(row, 1).number_format = "dd mmm yyyy"
            for column in (7,):
                sheet.cell(row, column).number_format = "#,##0.00"
            total += entry.amount

        last_data_row = sheet.max_row
        sheet.append(["Total", None, None, None, None, None, total])
        total_row = sheet.max_row
        sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=6)
        for column in (7,):
            sheet.cell(total_row, column).number_format = "#,##0.00"
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="5B32A7")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for cell in sheet[total_row]:
            cell.fill = PatternFill("solid", fgColor="EDE9FE")
            cell.font = Font(bold=True)
        for column, width in enumerate([18, 25, 50, 22, 25, 25, 20], start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.row_dimensions[1].height = 26
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:G{last_data_row}"
        sheet.sheet_view.showGridLines = False
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="income_{today.isoformat()}.xlsx"'
        response["Cache-Control"] = "no-store"
        return response

    page_obj = Paginator(entries, 20).get_page(request.GET.get("page"))
    _master_accounts_creators(page_obj, "income")
    return render(request, "master_admin/master_admin_accounts_income.html", {
        "current_user": current_user,
        "accounts_active": "income",
        "form": form,
        "query": query,
        "period": period,
        "selected_date": selected_date,
        "selected_month": selected_month,
        "from_date": from_date,
        "to_date": to_date,
        "range_start": range_start,
        "range_end": range_end,
        "filter_error": filter_error,
        "has_filters": bool(query or period != "all" or category_filter),
        "category_filter": category_filter,
        "categories": all_entries.order_by("category").values_list("category", flat=True).distinct(),
        "page_obj": page_obj,
        "total_entries": entries.count(),
        "total_income": income_total(entries),
        "today_income": income_total(all_entries.filter(date=today)),
        "month_income": income_total(all_entries.filter(date__range=(month_start, month_end))),
        "month_label": month_start.strftime("%B %Y"),
    }, status=400 if request.method == "POST" or filter_error else 200)


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_accounts_income_edit(request, pk):
    actor = _master_accounts_actor(request)
    if actor is None:
        return JsonResponse({"ok": False, "message": "Master Admin access required."},
            status=403 if request.session.get("user_id") else 401)
    with transaction.atomic():
        entry = AccountsIncome.objects.select_for_update().filter(pk=pk).first()
        if entry is None:
            return JsonResponse({"ok": False, "message": "This income no longer exists."}, status=404)
        if request.method == "GET":
            return JsonResponse({"ok": True, "item": {
                "date": entry.date.isoformat(), "category": entry.category,
                "amount": str(entry.amount),
                "payment_method": entry.payment_method,
                "description": entry.description, "reference": entry.reference,
            }, "receipt_url": reverse("master_admin_accounts_income_receipt", args=[entry.pk]) if entry.receipt else ""})
        before = _income_snapshot(entry)
        form = AccountsIncomeForm(request.POST, request.FILES, instance=entry)
        if not form.is_valid():
            return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
                "errors": form.errors.get_json_data()}, status=400)
        # Recheck duplicates when identifying fields change.
        identifying = {"date", "category", "amount", "reference"}
        if any(str(form.cleaned_data.get(key) or "") != before.get(key, "") for key in identifying) and _income_duplicate(form, entry.pk) and request.POST.get("confirm_duplicate") != "1":
            return _income_duplicate_response()
        entry = form.save(commit=False)
        if request.POST.get("remove_receipt") == "1" and not request.FILES.get("receipt"):
            entry.receipt = ""
        entry.save()
        _master_accounts_audit(entry, actor, "income", "updated", before)
        _master_accounts_notification(
            kind="income",
            title="Income updated",
            message=(
                f"{_master_accounts_actor_name(actor)} updated income "
                f"#{entry.pk}: {entry.category}, "
                f"₹{entry.amount:,.2f} "
                f"on {entry.date:%d %b %Y}."
            ),

        )
    return JsonResponse({"ok": True, "message": "Income updated successfully.", "id": entry.pk})


@never_cache
@require_POST
def master_admin_accounts_income_delete(request, pk):
    actor = _master_accounts_actor(request)
    if actor is None:
        return JsonResponse({"ok": False, "message": "Master Admin access required."},
            status=403 if request.session.get("user_id") else 401)
    with transaction.atomic():
        entry = AccountsIncome.objects.select_for_update().filter(pk=pk).first()
        if entry is None:
            return JsonResponse({"ok": False, "message": "This income no longer exists."}, status=404)
        _master_accounts_audit(entry, actor, "income", "deleted", _income_snapshot(entry))
        _master_accounts_notification(
            kind="income",
            title="Income deleted",
            message=(
                f"{_master_accounts_actor_name(actor)} deleted income "
                f"#{entry.pk}: {entry.category}, "
                f"₹{entry.amount:,.2f} "
                f"on {entry.date:%d %b %Y}."
            ),

        )
        entry.delete()
    return JsonResponse({"ok": True, "message": "Income deleted successfully.", "id": pk})


@never_cache
@require_http_methods(["GET"])
def master_admin_accounts_income_receipt(request, pk):
    if _master_accounts_actor(request) is None:
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    entry = AccountsIncome.objects.filter(pk=pk).first()
    if entry is None or not entry.receipt:
        return JsonResponse({"ok": False, "message": "Receipt not available."}, status=404)
    try:
        handle = entry.receipt.open("rb")
    except (FileNotFoundError, OSError):
        return JsonResponse({"ok": False, "message": "Receipt file not available."}, status=404)
    suffix = Path(entry.receipt.name).suffix.lower()
    mime = {".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}.get(suffix, "application/octet-stream")
    response = FileResponse(handle, as_attachment=request.GET.get("download") == "1",
        filename=f"income_{pk}_receipt{suffix}", content_type=mime)
    response["X-Content-Type-Options"] = "nosniff"
    return response


@never_cache
@require_http_methods(["GET"])
def master_admin_accounts_income_history(request, pk):
    if _master_accounts_actor(request) is None:
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    if not AccountsIncome.objects.filter(pk=pk).exists():
        return JsonResponse({"ok": False, "message": "This income no longer exists."}, status=404)
    items = AccountsIncomeHistory.objects.filter(income_id=pk)[:50]
    return JsonResponse({"ok": True, "items": [{"action": item.get_action_display(),
        "actor": item.actor_name or "—", "time": timezone.localtime(item.created_at).strftime("%d %b %Y, %I:%M %p"),
        "changes": item.changes} for item in items]})


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_accounts_expenses(request):
    import calendar
    import re
    from datetime import date
    from decimal import Decimal
    from io import BytesIO

    from django.core.paginator import Paginator
    from django.db.models import Q, Sum
    from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
    from django.shortcuts import redirect, render
    from django.utils import timezone

    current_user = _master_accounts_actor(request)
    if current_user is None:
        return HttpResponseForbidden("You do not have access to Accounts expense.")

    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    form = AccountsExpenseForm(request.POST if request.method == "POST" else None,
        request.FILES if request.method == "POST" else None)
    if request.method == "POST" and form.is_valid():
        if _expense_duplicate(form) and request.POST.get("confirm_duplicate") != "1":
            if is_ajax:
                return _expense_duplicate_response()
            form.add_error(None, "A similar expense already exists. Review the entry and use Save Anyway to confirm.")
        else:
            with transaction.atomic():
                expense = form.save(commit=False)
                expense.created_by = None
                expense.save()
                _master_accounts_audit(expense, current_user, "expense", "created")
                _master_accounts_notification(
                    kind="expense",
                    title="Expense added",
                    message=(
                        f"{_master_accounts_actor_name(current_user)} added expense "
                        f"#{expense.pk}: {expense.category}, "
                        f"₹{expense.amount:,.2f} "
                        f"on {expense.date:%d %b %Y}."
                    ),

                )
            if is_ajax:
                return JsonResponse({"ok": True, "message": "Expense added successfully.", "id": expense.pk}, status=201)
            messages.success(request, "Expense added successfully.")
            return redirect("master_admin_accounts_expenses")

    if request.method == "POST" and is_ajax:
        return JsonResponse({
            "ok": False,
            "message": "Please correct the highlighted fields.",
            "errors": form.errors.get_json_data(),
        }, status=400)

    today = timezone.localdate()
    month_start = today.replace(day=1)
    month_end = today.replace(day=calendar.monthrange(today.year, today.month)[1])
    all_entries = AccountsExpense.objects.select_related("created_by")
    entries = all_entries
    query = request.GET.get("q", "").strip()[:200]
    if query:
        entries = entries.filter(
            Q(category__icontains=query)
            | Q(description__icontains=query)
            | Q(reference__icontains=query)
            | Q(paid_to__icontains=query)
        )

    category_filter = request.GET.get("category", "").strip()[:100]
    payment_filter = request.GET.get("status", "").strip()
    if category_filter:
        entries = entries.filter(category__iexact=category_filter)
    if payment_filter == "paid":
        entries = entries.filter(Q(paid_amount__isnull=True) | Q(paid_amount=F("amount")))
    elif payment_filter == "pending":
        entries = entries.filter(paid_amount=0)
    elif payment_filter == "partial":
        entries = entries.filter(paid_amount__gt=0, paid_amount__lt=F("amount"))
    else:
        payment_filter = ""

    period = request.GET.get("period", "all").strip()
    selected_date = request.GET.get("date", "").strip() or today.isoformat()
    selected_month = request.GET.get("month", "").strip() or today.strftime("%Y-%m")
    from_date = request.GET.get("start", "").strip()
    to_date = request.GET.get("end", "").strip()
    range_start = range_end = None
    filter_error = ""

    def read_date(value):
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ValueError("Invalid date format.")
        return date.fromisoformat(value)

    try:
        if period == "weekly":
            anchor = read_date(selected_date)
            monday_ordinal = max(date.min.toordinal(), anchor.toordinal() - anchor.weekday())
            range_start = date.fromordinal(monday_ordinal)
            range_end = date.fromordinal(min(date.max.toordinal(), monday_ordinal + 6))
        elif period == "monthly":
            if not re.fullmatch(r"\d{4}-\d{2}", selected_month):
                raise ValueError("Invalid month format.")
            year, month = map(int, selected_month.split("-"))
            range_start = date(year, month, 1)
            range_end = date(year, month, calendar.monthrange(year, month)[1])
        elif period == "custom":
            if not from_date or not to_date:
                raise ValueError("Choose both From Date and To Date.")
            range_start, range_end = read_date(from_date), read_date(to_date)
            if range_start > range_end:
                raise ValueError("From Date must be on or before To Date.")
        elif period != "all":
            raise ValueError("Choose a valid period.")
    except (ValueError, OverflowError) as error:
        filter_error = (
            str(error) if period == "custom"
            else "Choose a valid date, month or period."
        )
        range_start = range_end = None
        entries = entries.none()

    if range_start is not None and not filter_error:
        entries = entries.filter(date__range=(range_start, range_end))
    entries = entries.order_by("-date", "-id")

    def expense_total(queryset):
        return queryset.aggregate(total=Sum("amount"))["total"] or Decimal("0.00")

    # Export the entire filtered queryset before applying pagination.
    if request.method == "GET" and request.GET.get("export") == "xlsx" and not filter_error:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Expenses"
        headers = ["Date", "Category", "Description", "Paid To", "Payment Method", "Reference", "Recorded By", "Total Amount (INR)", "Paid Amount (INR)", "Balance (INR)", "Status"]
        sheet.append(headers)
        total = paid_total = balance_total = Decimal("0.00")
        export_creators = dict(MasterAdminAccountsAudit.objects.filter(kind="expense", action="created", entry_number__in=entries.values("pk")).values_list("entry_number", "actor_name"))
        for entry in entries.iterator():
            creator = getattr(entry.created_by, "name", "") or export_creators.get(entry.pk, "—")
            sheet.append([
                entry.date, entry.category, entry.description, entry.paid_to or "—",
                entry.get_payment_method_display(), entry.reference or "—",
                creator, entry.amount, entry.effective_paid_amount, entry.balance_amount, entry.payment_status_label,
            ])
            row = sheet.max_row
            # User-entered strings must remain text, including values starting with '='.
            for column in range(2, 8):
                cell = sheet.cell(row, column)
                cell.value = str(cell.value or "")
                cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.cell(row, 1).number_format = "dd mmm yyyy"
            for column in (8, 9, 10):
                sheet.cell(row, column).number_format = "#,##0.00"
            sheet.cell(row, 11).data_type = "s"
            paid_total += entry.effective_paid_amount
            balance_total += entry.balance_amount
            total += entry.amount

        last_data_row = sheet.max_row
        sheet.append(["Total", None, None, None, None, None, None, total, paid_total, balance_total, None])
        total_row = sheet.max_row
        sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=7)
        for column in (8, 9, 10):
            sheet.cell(total_row, column).number_format = "#,##0.00"
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="5B32A7")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for cell in sheet[total_row]:
            cell.fill = PatternFill("solid", fgColor="EDE9FE")
            cell.font = Font(bold=True)
        for column, width in enumerate([18, 25, 50, 25, 22, 25, 25, 20, 20, 20, 20], start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.row_dimensions[1].height = 26
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:K{last_data_row}"
        sheet.sheet_view.showGridLines = False
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="expense_{today.isoformat()}.xlsx"'
        response["Cache-Control"] = "no-store"
        return response

    page_obj = Paginator(entries, 20).get_page(request.GET.get("page"))
    _master_accounts_creators(page_obj, "expense")
    return render(request, "master_admin/master_admin_accounts_expenses.html", {
        "current_user": current_user,
        "accounts_active": "expenses",
        "form": form,
        "query": query,
        "period": period,
        "selected_date": selected_date,
        "selected_month": selected_month,
        "from_date": from_date,
        "to_date": to_date,
        "range_start": range_start,
        "range_end": range_end,
        "filter_error": filter_error,
        "has_filters": bool(query or period != "all" or category_filter or payment_filter),
        "category_filter": category_filter,
        "payment_filter": payment_filter,
        "categories": all_entries.order_by("category").values_list("category", flat=True).distinct(),
        "page_obj": page_obj,
        "total_entries": entries.count(),
        "total_expense": expense_total(entries),
        "today_expense": expense_total(all_entries.filter(date=today)),
        "month_expense": expense_total(all_entries.filter(date__range=(month_start, month_end))),
        "month_label": month_start.strftime("%B %Y"),
    }, status=400 if request.method == "POST" or filter_error else 200)


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_accounts_expense_edit(request, pk):
    actor = _master_accounts_actor(request)
    if actor is None:
        return JsonResponse({"ok": False, "message": "Master Admin access required."},
            status=403 if request.session.get("user_id") else 401)
    with transaction.atomic():
        entry = AccountsExpense.objects.select_for_update().filter(pk=pk).first()
        if entry is None:
            return JsonResponse({"ok": False, "message": "This expense no longer exists."}, status=404)
        if request.method == "GET":
            return JsonResponse({"ok": True, "item": {
                "date": entry.date.isoformat(), "category": entry.category,
                "amount": str(entry.amount), "paid_amount": str(entry.effective_paid_amount),
                "payment_method": entry.payment_method, "paid_to": entry.paid_to,
                "description": entry.description, "reference": entry.reference,
            }, "receipt_url": reverse("master_admin_accounts_expense_receipt", args=[entry.pk]) if entry.receipt else ""})
        before = _expense_snapshot(entry)
        form = AccountsExpenseForm(request.POST, request.FILES, instance=entry)
        if not form.is_valid():
            return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
                "errors": form.errors.get_json_data()}, status=400)
        # Recheck duplicates when identifying fields change; editing only paid amount is safe.
        identifying = {"date", "category", "amount", "paid_to", "reference"}
        if any(str(form.cleaned_data.get(key) or "") != before.get(key, "") for key in identifying) and _expense_duplicate(form, entry.pk) and request.POST.get("confirm_duplicate") != "1":
            return _expense_duplicate_response()
        entry = form.save(commit=False)
        if request.POST.get("remove_receipt") == "1" and not request.FILES.get("receipt"):
            entry.receipt = ""
        entry.save()
        _master_accounts_audit(entry, actor, "expense", "updated", before)
        _master_accounts_notification(
            kind="expense",
            title="Expense updated",
            message=(
                f"{_master_accounts_actor_name(actor)} updated expense "
                f"#{entry.pk}: {entry.category}, "
                f"₹{entry.amount:,.2f} "
                f"on {entry.date:%d %b %Y}."
            ),

        )
    return JsonResponse({"ok": True, "message": "Expense updated successfully.", "id": entry.pk})


@never_cache
@require_POST
def master_admin_accounts_expense_delete(request, pk):
    actor = _master_accounts_actor(request)
    if actor is None:
        return JsonResponse({"ok": False, "message": "Master Admin access required."},
            status=403 if request.session.get("user_id") else 401)
    with transaction.atomic():
        entry = AccountsExpense.objects.select_for_update().filter(pk=pk).first()
        if entry is None:
            return JsonResponse({"ok": False, "message": "This expense no longer exists."}, status=404)
        _master_accounts_audit(entry, actor, "expense", "deleted", _expense_snapshot(entry))
        _master_accounts_notification(
            kind="expense",
            title="Expense deleted",
            message=(
                f"{_master_accounts_actor_name(actor)} deleted expense "
                f"#{entry.pk}: {entry.category}, "
                f"₹{entry.amount:,.2f} "
                f"on {entry.date:%d %b %Y}."
            ),

        )
        entry.delete()
    return JsonResponse({"ok": True, "message": "Expense deleted successfully.", "id": pk})


@never_cache
@require_http_methods(["GET"])
def master_admin_accounts_expense_receipt(request, pk):
    if _master_accounts_actor(request) is None:
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    entry = AccountsExpense.objects.filter(pk=pk).first()
    if entry is None or not entry.receipt:
        return JsonResponse({"ok": False, "message": "Receipt not available."}, status=404)
    try:
        handle = entry.receipt.open("rb")
    except (FileNotFoundError, OSError):
        return JsonResponse({"ok": False, "message": "Receipt file not available."}, status=404)
    suffix = Path(entry.receipt.name).suffix.lower()
    mime = {".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}.get(suffix, "application/octet-stream")
    response = FileResponse(handle, as_attachment=request.GET.get("download") == "1",
        filename=f"expense_{pk}_receipt{suffix}", content_type=mime)
    response["X-Content-Type-Options"] = "nosniff"
    return response


@never_cache
@require_http_methods(["GET"])
def master_admin_accounts_expense_history(request, pk):
    if _master_accounts_actor(request) is None:
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    if not AccountsExpense.objects.filter(pk=pk).exists():
        return JsonResponse({"ok": False, "message": "This expense no longer exists."}, status=404)
    items = AccountsExpenseHistory.objects.filter(expense_id=pk)[:50]
    return JsonResponse({"ok": True, "items": [{"action": item.get_action_display(),
        "actor": item.actor_name or "—", "time": timezone.localtime(item.created_at).strftime("%d %b %Y, %I:%M %p"),
        "changes": item.changes} for item in items]})


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_accounts_sales(request):
    import calendar
    import re
    from datetime import date
    from io import BytesIO

    current_user = _master_accounts_actor(request)
    if current_user is None:
        return HttpResponseForbidden("You do not have access to Accounts sales.")
    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    form = AccountsSaleForm(request.POST if request.method == "POST" else None, request.FILES if request.method == "POST" else None)
    if request.method == "POST" and form.is_valid():
        if _sale_duplicate(form) and request.POST.get("confirm_duplicate") != "1":
            if is_ajax:
                return _sale_duplicate_response()
            form.add_error(None, "A similar sale already exists. Review it and use Save Anyway to confirm.")
        else:
            with transaction.atomic():
                sale = form.save(commit=False)
                sale.created_by = None
                sale.save()
                _master_accounts_audit(sale, current_user, "sale", "created")
                if sale.received_amount > 0:
                    initial_payment = AccountsSalePayment.objects.create(
                        sale=sale,
                        date=sale.date,
                        amount=sale.received_amount,
                        created_by=None,
                        payment_method="other",
                        note="Initial received amount at sale creation.",
                    )
                    _master_accounts_audit(initial_payment, current_user, "payment", "created")

                _master_accounts_notification(
                    kind="sale",
                    title="Sale added",
                    message=(
                        f"{_master_accounts_actor_name(current_user)} added sale "
                        f"#{sale.pk} for {sale.customer_name}: "
                        f"₹{sale.amount:,.2f}. "
                        f"Received: ₹{sale.received_amount:,.2f}."
                    ),

                )
            if is_ajax:
                return JsonResponse({"ok": True, "message": "Sale added successfully.", "id": sale.pk}, status=201)
            messages.success(request, "Sale added successfully.")
            return redirect("master_admin_accounts_sales")
    if request.method == "POST" and is_ajax:
        return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
            "errors": form.errors.get_json_data()}, status=400)

    today = timezone.localdate()
    month_start = today.replace(day=1)
    month_end = today.replace(day=calendar.monthrange(today.year, today.month)[1])
    all_entries = AccountsSale.objects.select_related("created_by")
    entries = all_entries
    query = request.GET.get("q", "").strip()[:200]
    status_filter = request.GET.get("status", "").strip()
    if status_filter not in {"", "paid", "partial", "unpaid", "overdue"}:
        status_filter = ""
    if query:
        entries = entries.filter(Q(customer_name__icontains=query) |
            Q(invoice_number__icontains=query) | Q(description__icontains=query))
    if status_filter == "paid":
        entries = entries.filter(received_amount=F("amount"))
    elif status_filter == "partial":
        entries = entries.filter(received_amount__gt=0, received_amount__lt=F("amount"))
    elif status_filter == "unpaid":
        entries = entries.filter(received_amount=0)

    if status_filter == "overdue":
        entries = entries.filter(due_date__lt=today, received_amount__lt=F("amount"))

    period = request.GET.get("period", "all").strip()
    selected_date = request.GET.get("date", "").strip() or today.isoformat()
    selected_month = request.GET.get("month", "").strip() or today.strftime("%Y-%m")
    from_date = request.GET.get("start", "").strip()
    to_date = request.GET.get("end", "").strip()
    range_start = range_end = None
    filter_error = ""

    def read_date(value):
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ValueError("Invalid date format.")
        return date.fromisoformat(value)

    try:
        if period == "weekly":
            anchor = read_date(selected_date)
            monday_ordinal = max(date.min.toordinal(), anchor.toordinal() - anchor.weekday())
            range_start = date.fromordinal(monday_ordinal)
            range_end = date.fromordinal(min(date.max.toordinal(), monday_ordinal + 6))
        elif period == "monthly":
            if not re.fullmatch(r"\d{4}-\d{2}", selected_month):
                raise ValueError("Invalid month format.")
            year, month = map(int, selected_month.split("-"))
            range_start = date(year, month, 1)
            range_end = date(year, month, calendar.monthrange(year, month)[1])
        elif period == "custom":
            if not from_date or not to_date:
                raise ValueError("Choose both From Date and To Date.")
            range_start, range_end = read_date(from_date), read_date(to_date)
            if range_start > range_end:
                raise ValueError("From Date must be on or before To Date.")
        elif period != "all":
            raise ValueError("Choose a valid period.")
    except (ValueError, OverflowError) as error:
        filter_error = (
            str(error) if period == "custom"
            else "Choose a valid date, month or period."
        )
        range_start = range_end = None
        entries = entries.none()

    if range_start is not None and not filter_error:
        entries = entries.filter(date__range=(range_start, range_end))
    entries = entries.order_by("-date", "-id")

    totals = entries.aggregate(sales=Sum("amount"), received=Sum("received_amount"))
    total_sales = totals["sales"] or Decimal("0.00")
    total_received = totals["received"] or Decimal("0.00")
    total_balance = total_sales - total_received

    # Export all matching rows, before pagination.
    if request.method == "GET" and request.GET.get("export") == "xlsx" and not filter_error:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sales"
        sheet.append(["Date", "Customer", "Invoice", "Description", "Sale Amount (INR)",
            "Received (INR)", "Balance (INR)", "Status", "Recorded By", "Due Date", "Overdue"])
        export_creators = dict(MasterAdminAccountsAudit.objects.filter(kind="sale", action="created", entry_number__in=entries.values("pk")).values_list("entry_number", "actor_name"))
        for entry in entries.iterator():
            sheet.append([entry.date, entry.customer_name, entry.invoice_number or "—",
                entry.description, entry.amount, entry.received_amount, entry.balance_amount,
                entry.payment_status_label, getattr(entry.created_by, "name", "") or export_creators.get(entry.pk, "—"), entry.due_date, "Yes" if entry.is_overdue else "No"])
            row = sheet.max_row
            for column in (2, 3, 4, 8, 9, 11):
                cell = sheet.cell(row, column)
                cell.value = str(cell.value or "")
                cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.cell(row, 10).number_format = "dd mmm yyyy"
            sheet.cell(row, 1).number_format = "dd mmm yyyy"
            for column in (5, 6, 7):
                sheet.cell(row, column).number_format = "#,##0.00"
        last_data_row = sheet.max_row
        sheet.append(["Total", None, None, None, total_sales, total_received, total_balance, None, None, None, None])
        sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=4)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="5B32A7")
            cell.font = Font(color="FFFFFF", bold=True)
        for cell in sheet[sheet.max_row]:
            cell.fill = PatternFill("solid", fgColor="EDE9FE")
            cell.font = Font(bold=True)
        for column in (5, 6, 7):
            sheet.cell(sheet.max_row, column).number_format = "#,##0.00"
        for column, width in enumerate([18, 30, 25, 50, 22, 22, 22, 22, 28, 18, 14], start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.row_dimensions[1].height = 26
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:K{last_data_row}"
        sheet.sheet_view.showGridLines = False
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        response = HttpResponse(output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="sales_{today.isoformat()}.xlsx"'
        response["Cache-Control"] = "no-store"
        return response

    page_obj = Paginator(entries, 20).get_page(request.GET.get("page"))
    def sales_total(queryset):
        return queryset.aggregate(total=Sum("amount"))["total"] or Decimal("0.00")
    _master_accounts_creators(page_obj, "sale")
    return render(request, "master_admin/master_admin_accounts_sales.html", {
        "current_user": current_user, "accounts_active": "sales", "form": form,
        "query": query, "status_filter": status_filter,
        "has_filters": bool(query or status_filter or period != "all"),
        "period": period, "selected_date": selected_date, "selected_month": selected_month,
        "from_date": from_date, "to_date": to_date, "range_start": range_start,
        "range_end": range_end, "filter_error": filter_error,
        "page_obj": page_obj, "total_entries": entries.count(), "total_sales": total_sales,
        "total_received": total_received, "total_balance": total_balance,
        "month_sales": sales_total(all_entries.filter(date__range=(month_start, month_end))),
        "today_sales": sales_total(all_entries.filter(date=today)),
        "month_label": month_start.strftime("%B %Y"),
    }, status=400 if request.method == "POST" or filter_error else 200)


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_accounts_sale_edit(request, pk):
    actor = _master_accounts_actor(request)
    if actor is None:
        return JsonResponse({"ok": False, "message": "Master Admin access required."},
            status=403 if request.session.get("user_id") else 401)
    with transaction.atomic():
        sale = AccountsSale.objects.select_for_update().filter(pk=pk).first()
        if sale is None:
            return JsonResponse({"ok": False, "message": "This sale no longer exists."}, status=404)
        if request.method == "GET":
            return JsonResponse({"ok": True, "item": {
                "date": sale.date.isoformat(), "customer_name": sale.customer_name,
                "invoice_number": sale.invoice_number, "description": sale.description,
                "amount": str(sale.amount), "due_date": sale.due_date.isoformat() if sale.due_date else ""},
                "invoice_url": reverse("master_admin_accounts_sale_invoice", args=[sale.pk]) if sale.invoice_file else "",
                "received_amount": str(sale.received_amount)})
        before = {"date": sale.date.isoformat(), "customer_name": sale.customer_name,
            "invoice_number": sale.invoice_number, "amount": str(sale.amount)}
        form = AccountsSaleForm(request.POST, request.FILES, instance=sale)
        if not form.is_valid():
            return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
                "errors": form.errors.get_json_data()}, status=400)
        if (any(str(form.cleaned_data.get(key) or "") != value for key, value in before.items())
                and _sale_duplicate(form, sale.pk) and request.POST.get("confirm_duplicate") != "1"):
            return _sale_duplicate_response()
        sale = form.save(commit=False)
        if request.POST.get("remove_invoice") == "1" and not request.FILES.get("invoice_file"):
            sale.invoice_file = ""
        sale.save()
        _master_accounts_audit(sale, actor, "sale", "updated", before)
        _master_accounts_notification(
            kind="sale",
            title="Sale updated",
            message=(
                f"{_master_accounts_actor_name(actor)} updated sale "
                f"#{sale.pk} for {sale.customer_name}: "
                f"₹{sale.amount:,.2f}."
            ),

        )
    return JsonResponse({
        "ok": True,
        "message": "Sale updated successfully.",
        "id": sale.pk,
    })


@never_cache
@require_POST
def master_admin_accounts_sale_delete(request, pk):
    actor = _master_accounts_actor(request)
    if actor is None:
        return JsonResponse({
            "ok": False,
            "message": "Master Admin access required.",
        }, status=403 if request.session.get("user_id") else 401)

    with transaction.atomic():
        sale = (
            AccountsSale.objects
            .select_for_update()
            .filter(pk=pk)
            .first()
        )
        if sale is None:
            return JsonResponse({
                "ok": False,
                "message": "This sale no longer exists.",
            }, status=404)

        _master_accounts_notification(
            kind="sale",
            title="Sale deleted",
            message=(
                f"{_master_accounts_actor_name(actor)} deleted sale "
                f"#{sale.pk} for {sale.customer_name}: "
                f"₹{sale.amount:,.2f}."
            ),

        )
        _master_accounts_audit(sale, actor, "sale", "deleted")
        sale.delete()

    return JsonResponse({
        "ok": True,
        "message": "Sale deleted successfully.",
        "id": pk,
    })


@never_cache
@require_http_methods(["GET"])
def master_admin_accounts_sale_invoice(request, pk):
    if _master_accounts_actor(request) is None:
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    sale = AccountsSale.objects.filter(pk=pk).first()
    if sale is None or not sale.invoice_file:
        return JsonResponse({"ok": False, "message": "Invoice not available."}, status=404)
    try:
        handle = sale.invoice_file.open("rb")
    except (OSError, FileNotFoundError):
        return JsonResponse({"ok": False, "message": "Invoice file not available."}, status=404)
    suffix = Path(sale.invoice_file.name).suffix.lower()
    mime = {".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}.get(suffix, "application/octet-stream")
    response = FileResponse(handle, as_attachment=request.GET.get("download") == "1",
        filename=f"sale_{pk}_invoice{suffix}", content_type=mime)
    response["X-Content-Type-Options"] = "nosniff"
    return response


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_accounts_sale_payments(request, pk):
    actor = _master_accounts_actor(request)
    if actor is None:
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    with transaction.atomic():
        sale = AccountsSale.objects.select_for_update().filter(pk=pk).first()
        if sale is None:
            return JsonResponse({"ok": False, "message": "This sale no longer exists."}, status=404)
        if request.method == "POST":
            form = AccountsSalePaymentForm(request.POST)
            if not form.is_valid():
                return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
                    "errors": form.errors.get_json_data()}, status=400)
            token = form.cleaned_data["request_token"]
            existing = AccountsSalePayment.objects.filter(request_token=token).first()
            if existing:
                if existing.sale_id != pk or any(getattr(existing, key) != form.cleaned_data[key]
                        for key in ("date", "amount", "payment_method", "reference", "note")):
                    return JsonResponse({"ok": False, "message": "This payment request changed. Close and reopen Payments."}, status=409)
                return JsonResponse({"ok": True, "message": "Payment already recorded.", "id": existing.pk})
            if form.cleaned_data["date"] < sale.date:
                form.add_error("date", "Payment date cannot be before the sale date.")
            if form.cleaned_data["date"] > timezone.localdate():
                form.add_error("date", "Use the date the payment was received; future payments cannot be recorded.")
            if form.cleaned_data["amount"] > sale.balance_amount:
                form.add_error("amount", "Payment cannot exceed the outstanding balance.")
            if form.errors:
                return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
                    "errors": form.errors.get_json_data()}, status=400)
            payment = form.save(commit=False)
            payment.sale, payment.created_by, payment.request_token = sale, None, token
            changed = AccountsSale.objects.filter(pk=pk,
                received_amount__lte=F("amount") - payment.amount).update(
                    received_amount=F("received_amount") + payment.amount, updated_at=timezone.now())
            if not changed:
                return JsonResponse({"ok": False, "message": "Balance changed. Close and reopen Payments."}, status=409)
            payment.save()
            _master_accounts_audit(payment, actor, "payment", "created")
            _master_accounts_notification(
                kind="payment",
                title="Sales payment received",
                message=(
                    f"{_master_accounts_actor_name(actor)} recorded a payment of "
                    f"₹{payment.amount:,.2f} for "
                    f"{sale.customer_name}, sale #{sale.pk}, "
                    f"on {payment.date:%d %b %Y}."
                ),

            )
            return JsonResponse({
                "ok": True,
                "message": "Payment recorded successfully.",
                "id": payment.pk,
            })
        items = list(sale.payments.select_related("created_by").all())
        master_payment_names = {item.entry_number: item.actor_name for item in MasterAdminAccountsAudit.objects.filter(
            kind="payment", action="created", entry_number__in=[item.pk for item in items])}
        recorded = sum((item.amount for item in items), Decimal("0.00"))
        return JsonResponse({"ok": True, "customer": sale.customer_name,
            "received": str(sale.received_amount), "balance": str(sale.balance_amount),
            "opening_received": str(sale.received_amount - recorded),
            "today": timezone.localdate().isoformat(), "sale_date": sale.date.isoformat(),
            "items": [{"date": item.date.isoformat(), "amount": str(item.amount),
                "method": item.get_payment_method_display(), "reference": item.reference,
                "note": item.note, "actor": getattr(item.created_by, "name", "") or master_payment_names.get(item.pk, "—")} for item in items]})



@never_cache
@require_GET
def master_admin_accounts_overview(request):
    from django.db.models import Sum
    from django.http import HttpResponseForbidden, JsonResponse
    if not _master_admin_allowed(request):
        if request.GET.get("format") == "json":
            return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
        if not request.user.is_authenticated:
            return redirect("master_admin_login")
        return HttpResponseForbidden("Master Admin access required.")
    period, anchor, start, end, error = _accounts_report_range(request.GET, timezone.localdate())
    income = AccountsIncome.objects.filter(date__range=(start, end))
    expense = AccountsExpense.objects.filter(date__range=(start, end))
    sales = AccountsSale.objects.filter(date__range=(start, end))
    if error:
        income, expense, sales = income.none(), expense.none(), sales.none()
    zero = Decimal("0.00")
    income_total = income.aggregate(total=Sum("amount"))["total"] or zero
    expense_total = expense.aggregate(total=Sum("amount"))["total"] or zero
    totals = sales.aggregate(amount=Sum("amount"), received=Sum("received_amount"))
    sale_total, received = totals["amount"] or zero, totals["received"] or zero
    audit = [{"actor": item.actor_name, "kind": item.get_kind_display(),
              "action": item.get_action_display(), "entry": item.entry_number,
              "time": timezone.localtime(item.created_at).strftime("%d %b %Y, %I:%M %p")}
             for item in MasterAdminAccountsAudit.objects.all()[:30]]
    data = {"ok": not bool(error), "message": error,
        "income": str(income_total), "expense": str(expense_total),
        "net": str(income_total - expense_total), "sales": str(sale_total),
        "received": str(received), "balance": str(sale_total - received), "audit": audit}
    if request.GET.get("format") == "json":
        return JsonResponse(data, status=400 if error else 200)
    return render(request, "master_admin/master_admin_accounts_overview.html", {
        "accounts_active": "overview", "period": period, "anchor": anchor,
        "range_start": start, "range_end": end, "filter_error": error, "data": data,
        "from_date": request.GET.get("start", ""), "to_date": request.GET.get("end", ""),
    })


from django import forms
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Q, Count
from django.db.models.functions import Lower, Trim
from django.http import JsonResponse, HttpResponseForbidden
from django.shortcuts import render, redirect
from django.urls import reverse
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods, require_POST
from .models import Task, User, Team, TeamLeadNotification, MasterAdminTaskAudit


def _master_task_people():
    return User.objects.annotate(task_role=Lower(Trim("job_Position"))).filter(
        task_role__in=["team lead", "team_lead", "team member", "team_member"]
    ).select_related("team", "department").order_by("name", "pk")


class MasterAdminTaskForm(forms.ModelForm):
    class Meta:
        model = Task
        fields = ["title", "description", "assigned_to", "status", "progress"]

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["assigned_to"].queryset = _master_task_people()
        self.fields["progress"].min_value = 0
        self.fields["progress"].max_value = 100

    def clean_title(self):
        title = self.cleaned_data["title"].strip()
        if not title:
            raise forms.ValidationError("Enter a task title.")
        return title

    def clean_progress(self):
        progress = self.cleaned_data["progress"]
        if not 0 <= progress <= 100:
            raise forms.ValidationError("Progress must be between 0 and 100.")
        return progress

    def clean(self):
        data = super().clean()
        if data.get("status") == "pending":
            data["progress"] = 0
        elif data.get("status") == "completed":
            data["progress"] = 100
        elif data.get("status") == "in_progress" and data.get("progress") == 100:
            self.add_error("progress", "Use Completed for 100% progress.")
        return data


def _master_task_snapshot(task):
    return {"title": task.title, "description": task.description or "",
        "assigned_to": task.assigned_to_id, "status": task.status, "progress": task.progress}


def _master_task_record(request, task, action, before=None):
    audit = MasterAdminTaskAudit.objects.create(
        actor=request.user, actor_name=(request.user.get_full_name() or request.user.username)[:255],
        task_number=task.pk, action=action, before=before or {},
        after={} if action == "deleted" else _master_task_snapshot(task))
    recipients = {task.assigned_to_id}
    if before and before.get("assigned_to"):
        recipients.add(before["assigned_to"])
    for person in _master_task_people().filter(pk__in=recipients):
        role = person.task_role
        TeamLeadNotification.objects.create(recipient=person, kind="task",
            title=f"Task {action}", message=f"Master Admin {audit.actor_name} {action} task #{task.pk}: {task.title}",
            url=reverse("teamlead_task" if role in {"team lead", "team_lead"} else "teammember_task"))


def _master_task_denied(request):
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    if not request.user.is_authenticated:
        return redirect("master_admin_login")
    return HttpResponseForbidden("Master Admin access required.")


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_tasks(request):
    if not _master_admin_allowed(request):
        return _master_task_denied(request)
    if request.method == "POST":
        form = MasterAdminTaskForm(request.POST)
        if not form.is_valid():
            return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
                "errors": form.errors.get_json_data()}, status=400)
        with transaction.atomic():
            task = form.save(commit=False)
            task.created_by = None
            task.save()
            _master_task_record(request, task, "created")
        return JsonResponse({"ok": True, "message": "Task assigned successfully.", "id": task.pk}, status=201)
    tasks = Task.objects.select_related("assigned_to__team", "assigned_to__department", "created_by").annotate(
        task_role=Lower(Trim("assigned_to__job_Position")))
    tab = request.GET.get("tab", "all")
    if tab not in {"all", "leads", "members"}:
        tab = "all"
    if tab == "leads":
        tasks = tasks.filter(task_role__in=["team lead", "team_lead"])
    elif tab == "members":
        tasks = tasks.filter(task_role__in=["team member", "team_member"])
    query = request.GET.get("q", "").strip()[:200]
    if query:
        tasks = tasks.filter(Q(title__icontains=query) | Q(description__icontains=query)
            | Q(assigned_to__name__icontains=query) | Q(created_by__name__icontains=query))
    team = request.GET.get("team", "")
    if team.isdigit():
        tasks = tasks.filter(assigned_to__team_id=int(team))
    else:
        team = ""
    state = request.GET.get("status", "")
    if state in dict(Task.STATUS_CHOICES):
        tasks = tasks.filter(status=state)
    else:
        state = ""
    counts = dict(tasks.values_list("status").annotate(total=Count("pk")))
    page = Paginator(tasks.order_by("-created_at", "-pk"), 20).get_page(request.GET.get("page"))
    items = list(page.object_list)
    creators = {}
    for audit in MasterAdminTaskAudit.objects.filter(action="created", task_number__in=[item.pk for item in items]).order_by("pk"):
        creators.setdefault(audit.task_number, audit.actor_name)
    data = []
    for item in items:
        item.creator_label = (item.created_by.name if item.created_by else
            "Master Admin: " + creators[item.pk] if item.pk in creators else "Not recorded")
        data.append({"id": item.pk, **_master_task_snapshot(item), "assignee": item.assigned_to.name,
            "creator": item.creator_label, "team": item.assigned_to.team.name if item.assigned_to.team else "No team"})
    page.object_list = items
    params = request.GET.copy()
    params.pop("page", None)
    params.pop("fragment", None)
    params.pop("tab", None)
    return render(request, "master_admin/master_admin_tasks.html", {
        "master_tasks_active": True, "task_tab": tab, "query": query, "selected_team": team,
        "selected_status": state, "teams": Team.objects.order_by("name"), "people": _master_task_people(),
        "page_obj": page, "task_items": data, "filter_params": params.urlencode(),
        "total_tasks": sum(counts.values()), "pending_tasks": counts.get("pending", 0),
        "in_progress_tasks": counts.get("in_progress", 0), "completed_tasks": counts.get("completed", 0),
    })


@never_cache
@require_POST
def master_admin_task_update(request, pk):
    if not _master_admin_allowed(request):
        return _master_task_denied(request)
    with transaction.atomic():
        task = Task.objects.select_for_update().filter(pk=pk).first()
        if task is None:
            return JsonResponse({"ok": False, "message": "Task no longer exists."}, status=404)
        before = _master_task_snapshot(task)
        data = request.POST.copy()
        # A row status change is supported without overwriting other fields.
        if data.get("mode") == "status":
            status = data.get("status")
            data = {**before, "status": status,
                "progress": {"pending": 0, "in_progress": 50, "completed": 100}.get(status, task.progress)}
        form = MasterAdminTaskForm(data, instance=task)
        if not form.is_valid():
            return JsonResponse({"ok": False, "message": "Please correct the highlighted fields.",
                "errors": form.errors.get_json_data()}, status=400)
        task = form.save()
        if before != _master_task_snapshot(task):
            _master_task_record(request, task, "updated", before)
    return JsonResponse({"ok": True, "message": "Task updated successfully."})


@never_cache
@require_POST
def master_admin_task_delete(request, pk):
    import logging
    import uuid
    from django.db import transaction
    from django.db.models.deletion import ProtectedError, RestrictedError
    from django.http import JsonResponse
    from .models import Task

    if not _master_admin_allowed(request):
        return JsonResponse({"ok": False, "message": "Master Admin access required. Sign in again."}, status=403)

    try:
        with transaction.atomic():
            task = Task.objects.select_for_update().filter(pk=pk).first()
            if task is None:
                return JsonResponse({"ok": False, "message": "Task no longer exists. Refresh the task list."}, status=404)
            _master_task_record(request, task, "deleted", _master_task_snapshot(task))
            task.delete()
    except (ProtectedError, RestrictedError):
        return JsonResponse({"ok": False, "message": "This task is linked to protected records and cannot be deleted."}, status=409)
    except Exception:
        reference = uuid.uuid4().hex[:8]
        logging.getLogger(__name__).exception("Master Admin task delete failed: task=%s reference=%s", pk, reference)
        return JsonResponse({"ok": False, "message": f"Task deletion failed. Check the Django terminal traceback (reference {reference})."}, status=500)

    return JsonResponse({"ok": True, "message": "Task deleted successfully.", "id": pk})


# ADD once at END of monitoringapp/views.py. Existing TL/TM functions stay unchanged.
from django import forms
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Q, Count
from django.db.models.functions import Lower, Trim
from django.http import JsonResponse, HttpResponseForbidden
from django.shortcuts import render, redirect
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods, require_POST
from .models import ProjectAssign, ProjectFile, ProjectImage, User, Team, Department, TeamLeadNotification, MasterAdminProjectAudit


def _master_project_people():
    return User.objects.annotate(project_role=Lower(Trim("job_Position"))).filter(
        project_role__in=["team lead", "team_lead", "team member", "team_member"]
    ).select_related("team", "department").order_by("name", "pk")


class MasterAdminProjectForm(forms.ModelForm):
    class Meta:
        model = ProjectAssign
        fields = ["team", "department", "assign_to", "assigned_by", "work_name", "work_type",
            "category", "description", "deadline", "additional_notes", "color_preference",
            "content_example", "priority", "status"]

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["assign_to"].queryset = _master_project_people()
        owner_filter = Q(project_role__in=["team lead", "team_lead"])
        if self.instance.pk:
            owner_filter |= Q(pk=self.instance.assigned_by_id)
        self.fields["assigned_by"].queryset = User.objects.annotate(
            project_role=Lower(Trim("job_Position"))).filter(owner_filter)

    def clean_work_name(self):
        value = self.cleaned_data["work_name"].strip()
        if not value:
            raise forms.ValidationError("Enter a project name.")
        return value

    def clean(self):
        data = super().clean()
        team, department = data.get("team"), data.get("department")
        assignee, owner = data.get("assign_to"), data.get("assigned_by")
        if team and assignee and assignee.team_id != team.pk:
            self.add_error("assign_to", "Choose someone from the selected team.")
        if department and assignee and assignee.department_id != department.pk:
            self.add_error("assign_to", "Choose someone from the selected department.")
        if team and owner and owner.team_id != team.pk:
            self.add_error("assigned_by", "Choose a responsible lead from the selected team.")
        if department and owner and owner.department_id != department.pk:
            self.add_error("assigned_by", "Choose a responsible lead from the selected department.")
        return data


def _master_project_denied(request):
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    return redirect("master_admin_login") if not request.user.is_authenticated else HttpResponseForbidden("Master Admin access required.")


def _master_project_snapshot(project):
    data = {}
    for name in MasterAdminProjectForm.Meta.fields:
        field = project._meta.get_field(name)
        value = getattr(project, field.attname)
        data[name] = value.isoformat() if hasattr(value, "isoformat") else value
    return data


def _master_project_audit(request, project, action, before=None):
    before = before or {}
    record = MasterAdminProjectAudit.objects.create(actor=request.user,
        actor_name=(request.user.get_full_name() or request.user.username)[:255],
        project_number=project.pk, action=action, before=before,
        after={} if action == "deleted" else _master_project_snapshot(project))
    recipients = {project.assign_to_id, project.assigned_by_id,
        before.get("assign_to"), before.get("assigned_by")} - {None}
    for person in _master_project_people().filter(pk__in=recipients):
        TeamLeadNotification.objects.create(recipient=person, kind="project",
            title=f"Project {action}", message=f"Master Admin {record.actor_name} {action} project #{project.pk}: {project.work_name}",
            url=reverse("teamlead_project_assigning" if person.project_role in {"team lead", "team_lead"} else "teammember_project"))


def _master_project_uploads(request):
    files = request.FILES.getlist("upload_file[]")
    images = request.FILES.getlist("upload_image[]")
    if len(files) + len(images) > 10:
        return None, None, "Upload a maximum of 10 attachments at a time."
    if any(item.size > 10 * 1024 * 1024 for item in files + images):
        return None, None, "Each attachment must be 10 MB or smaller."
    try:
        for item in images:
            forms.ImageField().clean(item)
            item.seek(0)
    except forms.ValidationError:
        return None, None, "Choose valid image files."
    return files, images, ""


def _master_project_save_files(project, files, images):
    for item in files:
        ProjectFile.objects.create(project=project, file=item)
    for item in images:
        ProjectImage.objects.create(project=project, image=item)


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_projects(request):
    if not _master_admin_allowed(request):
        return _master_project_denied(request)
    if request.method == "POST":
        form = MasterAdminProjectForm(request.POST)
        files, images, error = _master_project_uploads(request)
        if not form.is_valid() or error:
            return JsonResponse({"ok": False, "message": error or "Please correct the highlighted fields.",
                "errors": form.errors.get_json_data()}, status=400)
        with transaction.atomic():
            project = form.save()
            _master_project_save_files(project, files, images)
            _master_project_audit(request, project, "created")
        return JsonResponse({"ok": True, "message": "Project assigned successfully.", "id": project.pk}, status=201)
    projects = ProjectAssign.objects.select_related("team", "department", "assign_to", "assigned_by").prefetch_related("files", "images")
    query = request.GET.get("q", "").strip()[:200]
    if query:
        projects = projects.filter(Q(work_name__icontains=query) | Q(description__icontains=query)
            | Q(assign_to__name__icontains=query) | Q(assigned_by__name__icontains=query))
    team, department = request.GET.get("team", ""), request.GET.get("department", "")
    if team.isdigit():
        projects = projects.filter(team_id=int(team))
    else:
        team = ""
    if department.isdigit():
        projects = projects.filter(department_id=int(department))
    else:
        department = ""
    state = request.GET.get("status", "")
    if state in dict(ProjectAssign.STATUS_CHOICES):
        projects = projects.filter(status=state)
    else:
        state = ""
    priority = request.GET.get("priority", "")
    if priority in dict(ProjectAssign.PRIORITY_CHOICES):
        projects = projects.filter(priority=priority)
    else:
        priority = ""
    counts = dict(projects.order_by().values_list("status").annotate(total=Count("pk")))
    today = timezone.localdate()
    overdue = projects.filter(deadline__lt=today).exclude(status="Completed").count()
    page = Paginator(projects.order_by("-assigned_date", "-pk"), 20).get_page(request.GET.get("page"))
    items = list(page.object_list)
    authors = {}
    for audit in MasterAdminProjectAudit.objects.filter(action="created", project_number__in=[p.pk for p in items]).order_by("pk"):
        authors.setdefault(audit.project_number, audit.actor_name)
    data = []
    for item in items:
        item.is_overdue = bool(item.deadline and item.deadline < today and item.status != "Completed")
        item.recorded_by_name = "Master Admin: " + authors[item.pk] if item.pk in authors else item.assigned_by.name
        attachments = [{"name": f.file.name.split("/")[-1], "url": f.file.url} for f in item.files.all() if f.file]
        attachments += [{"name": f.image.name.split("/")[-1], "url": f.image.url} for f in item.images.all() if f.image]
        for field in [item.upload_file, item.upload_image]:
            if field:
                attachments.append({"name": field.name.split("/")[-1], "url": field.url})
        data.append({"id": item.pk, **_master_project_snapshot(item), "assignee_name": item.assign_to.name,
            "owner_name": item.assigned_by.name, "team_name": item.team.name, "department_name": item.department.name,
            "recorded_by": item.recorded_by_name, "attachments": attachments})
    page.object_list = items
    params = request.GET.copy()
    params.pop("page", None)
    people = list(_master_project_people())
    return render(request, "master_admin/master_admin_projects.html", {
        "master_projects_active": True, "page_obj": page, "project_items": data,
        "query": query, "selected_team": team, "selected_department": department,
        "selected_status": state, "selected_priority": priority, "filter_params": params.urlencode(),
        "teams": Team.objects.order_by("name"), "departments": Department.objects.order_by("name"),
        "people": people, "people_data": [{"id": p.pk, "name": p.name, "team": p.team_id,
            "department": p.department_id, "role": p.project_role} for p in people],
        "statuses": ProjectAssign.STATUS_CHOICES, "priorities": ProjectAssign.PRIORITY_CHOICES,
        "total_projects": sum(counts.values()), "ongoing_projects": counts.get("In Progress", 0),
        "completed_projects": counts.get("Completed", 0), "overdue_projects": overdue,
    })


@never_cache
@require_POST
def master_admin_project_update(request, pk):
    if not _master_admin_allowed(request):
        return _master_project_denied(request)
    with transaction.atomic():
        project = ProjectAssign.objects.select_for_update().filter(pk=pk).first()
        if project is None:
            return JsonResponse({"ok": False, "message": "Project no longer exists."}, status=404)
        before = _master_project_snapshot(project)
        data = request.POST
        if data.get("mode") == "status":
            status = data.get("status", "")
            if status not in dict(ProjectAssign.STATUS_CHOICES):
                return JsonResponse({"ok": False, "message": "Invalid project status."}, status=400)
            project.status = status
            project.save(update_fields=["status"])
            _master_project_audit(request, project, "updated", before)
            return JsonResponse({"ok": True, "message": "Project status updated successfully."})
        form = MasterAdminProjectForm(data, instance=project)
        files, images, error = _master_project_uploads(request)
        if not form.is_valid() or error:
            return JsonResponse({"ok": False, "message": error or "Please correct the highlighted fields.",
                "errors": form.errors.get_json_data()}, status=400)
        project = form.save()
        _master_project_save_files(project, files, images)
        _master_project_audit(request, project, "updated", before)
    return JsonResponse({"ok": True, "message": "Project updated successfully."})


@never_cache
@require_POST
def master_admin_project_delete(request, pk):
    if not _master_admin_allowed(request):
        return _master_project_denied(request)
    with transaction.atomic():
        project = ProjectAssign.objects.select_for_update().filter(pk=pk).first()
        if project is None:
            return JsonResponse({"ok": False, "message": "Project no longer exists."}, status=404)
        _master_project_audit(request, project, "deleted", _master_project_snapshot(project))
        project.delete()
    return JsonResponse({"ok": True, "message": "Project deleted successfully."})


# ADD once at END of monitoringapp/views.py. Existing TL/TM views stay unchanged.
from django.contrib.auth import get_user_model
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Q, Prefetch
from django.http import JsonResponse, HttpResponseForbidden, Http404
from django.shortcuts import render, redirect
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods, require_POST
from .models import Announcement, AnnouncementRecipient, AnnouncementAuthRecipient, User, AccountsNotification, AdminNotification, TeamLeadNotification


def _master_announcement_employees():
    return User.objects.select_related("team", "department").order_by("name", "pk")


def _master_announcement_admins():
    return get_user_model().objects.filter(is_active=True).filter(Q(is_staff=True) | Q(is_superuser=True)).order_by("username", "pk")


def _master_announcement_notify(announcement, *, edited=False):
    title = ("Announcement updated: " if edited else "New announcement: ") + announcement.title
    for row in announcement.recipients.select_related("recipient"):
        user = row.recipient
        role = (user.job_Position or "").strip().lower().replace("_", " ")
        key = f"master-announcement:{announcement.pk}:employee:{user.pk}"
        url = reverse("portal_announcement_detail", args=[announcement.pk])
        defaults = {"title": title[:180], "message": announcement.message, "url": url,
            "kind": "announcement", "is_archived": False, "is_read": False}
        if role in {"team lead", "team member"}:
            # Inbox pages already receive these shared AnnouncementRecipient rows.
            url = reverse("teamlead_announcements" if role == "team lead" else "teammember_announcements")
            if role == "team member":
                url += "?source=admin" if announcement.is_admin_announcement else "?source=lead"
            defaults["url"] = url
            TeamLeadNotification.objects.update_or_create(event_key=key, defaults={"recipient": user, **defaults})
        elif role in {"accounts", "account", "accounts team"}:
            AccountsNotification.objects.update_or_create(event_key=key, defaults={"recipient": user, "actor": None, **defaults})
        else:
            AdminNotification.objects.update_or_create(event_key=key, defaults={"management_owner": user, "auth_owner": None, **defaults})
    for row in announcement.auth_recipients.select_related("recipient"):
        AdminNotification.objects.update_or_create(
            event_key=f"master-announcement:{announcement.pk}:auth:{row.recipient_id}",
            defaults={"auth_owner": row.recipient, "management_owner": None, "kind": "announcement",
                "title": title[:180], "message": announcement.message,
                "url": reverse("portal_announcement_detail", args=[announcement.pk]), "is_read": False, "is_archived": False})


def _master_announcement_denied(request):
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    return redirect("master_admin_login") if not request.user.is_authenticated else HttpResponseForbidden("Master Admin access required.")


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_announcements(request):
    if not _master_admin_allowed(request):
        return _master_announcement_denied(request)
    employees = _master_announcement_employees()
    admins = _master_announcement_admins().exclude(pk=request.user.pk)
    if request.method == "POST":
        action = request.POST.get("action", "")
        if action not in {"create", "update", "delete"}:
            return JsonResponse({"ok": False, "message": "Invalid action."}, status=400)
        title = request.POST.get("title", "").strip()
        message = request.POST.get("message", "").strip()
        if action != "delete" and (not title or len(title) > 200 or not message):
            return JsonResponse({"ok": False, "message": "Enter a title (maximum 200 characters) and message."}, status=400)
        if action == "create":
            audience = request.POST.get("audience", "all")
            if audience == "selected":
                try:
                    employee_ids = set(map(int, request.POST.getlist("employee_ids")))
                    auth_ids = set(map(int, request.POST.getlist("auth_ids")))
                except (TypeError, ValueError):
                    return JsonResponse({"ok": False, "message": "Select valid recipients."}, status=400)
                employees = employees.filter(pk__in=employee_ids)
                admins = admins.filter(pk__in=auth_ids)
                if employees.count() != len(employee_ids) or admins.count() != len(auth_ids):
                    return JsonResponse({"ok": False, "message": "A selected recipient is no longer available."}, status=400)
            elif audience != "all":
                return JsonResponse({"ok": False, "message": "Invalid audience."}, status=400)
            employees, admins = list(employees), list(admins)
            if not employees and not admins:
                return JsonResponse({"ok": False, "message": "Choose at least one recipient."}, status=400)
            with transaction.atomic():
                announcement = Announcement.objects.create(title=title, message=message,
                    admin_created_by=request.user, is_admin_announcement=True, is_master_admin_announcement=True)
                AnnouncementRecipient.objects.bulk_create([AnnouncementRecipient(announcement=announcement, recipient=user) for user in employees])
                AnnouncementAuthRecipient.objects.bulk_create([AnnouncementAuthRecipient(announcement=announcement, recipient=user) for user in admins])
                _master_announcement_notify(announcement)
            return JsonResponse({"ok": True, "message": "Announcement sent successfully."}, status=201)
        try:
            pk = int(request.POST.get("announcement_id", ""))
        except (ValueError, TypeError):
            return JsonResponse({"ok": False, "message": "Invalid announcement ID."}, status=400)
        with transaction.atomic():
            announcement = Announcement.objects.select_for_update().filter(pk=pk, is_active=True).first()
            if announcement is None:
                return JsonResponse({"ok": False, "message": "Announcement no longer exists."}, status=404)
            # Superusers may moderate every announcement; original author stays unchanged.
            if action == "delete":
                announcement.is_active = False
                announcement.save(update_fields=["is_active", "updated_at"])
                prefix = f"master-announcement:{pk}:"
                for model in [TeamLeadNotification, AccountsNotification, AdminNotification]:
                    model.objects.filter(event_key__startswith=prefix).update(is_archived=True)
                return JsonResponse({"ok": True, "message": "Announcement deleted successfully."})
            announcement.title, announcement.message = title, message
            announcement.save(update_fields=["title", "message", "updated_at"])
            _master_announcement_notify(announcement, edited=True)
        return JsonResponse({"ok": True, "message": "Announcement updated successfully."})
    all_items = Announcement.objects.filter(is_active=True)
    source = request.GET.get("source", "master")
    if source not in {"master", "admin", "lead"}:
        source = "master"
    groups = {
        "master": all_items.filter(is_master_admin_announcement=True),
        "admin": all_items.filter(is_admin_announcement=True, is_master_admin_announcement=False),
        "lead": all_items.filter(is_admin_announcement=False),
    }
    query = request.GET.get("q", "").strip()[:200]
    items = groups[source]
    if query:
        items = items.filter(Q(title__icontains=query) | Q(message__icontains=query) | Q(created_by__name__icontains=query))
    items = items.select_related("created_by", "admin_created_by").prefetch_related(
        Prefetch("recipients", queryset=AnnouncementRecipient.objects.select_related("recipient").order_by("recipient__name", "pk")),
        Prefetch("auth_recipients", queryset=AnnouncementAuthRecipient.objects.select_related("recipient").order_by("recipient__username", "pk")))
    page = Paginator(items.order_by("-created_at", "-pk"), 15).get_page(request.GET.get("page"))
    data = []
    for item in page:
        auth = item.admin_created_by
        author = (auth.get_full_name() or auth.username) if auth else (item.created_by.name if item.created_by else "Not recorded")
        item.sender_label = ("Master Admin: " if item.is_master_admin_announcement else "Admin: " if item.is_admin_announcement else "Team Lead: ") + author
        recipients = [{"name": row.recipient.name, "role": row.recipient.job_Position,
            "read_at": timezone.localtime(row.read_at).strftime("%d %b %Y, %I:%M %p") if row.read_at else None} for row in item.recipients.all()]
        recipients += [{"name": row.recipient.get_full_name() or row.recipient.username, "role": "Admin",
            "read_at": timezone.localtime(row.read_at).strftime("%d %b %Y, %I:%M %p") if row.read_at else None} for row in item.auth_recipients.all()]
        item.recipient_total = len(recipients)
        item.seen_total = sum(bool(row["read_at"]) for row in recipients)
        data.append({"id": item.pk, "title": item.title, "message": item.message,
            "sender": item.sender_label, "created": timezone.localtime(item.created_at).strftime("%d %b %Y, %I:%M %p"),
            "recipients": recipients, "seen": item.seen_total, "total": item.recipient_total})
    params = request.GET.copy()
    params.pop("page", None)
    params.pop("source", None)
    return render(request, "master_admin/master_admin_announcements.html", {
        "master_announcements_active": True, "source": source, "query": query,
        "master_count": groups["master"].count(), "admin_count": groups["admin"].count(), "lead_count": groups["lead"].count(),
        "page_obj": page, "announcement_items": data, "employees": employees, "auth_admins": admins,
        "filter_params": params.urlencode(),
    })


def _portal_announcement_receipt(request, pk):
    employee_id = request.session.get("user_id")
    if employee_id:
        return AnnouncementRecipient.objects.filter(announcement_id=pk, recipient_id=employee_id, announcement__is_active=True).select_related("announcement").first()
    if request.user.is_authenticated and request.user.is_active:
        return AnnouncementAuthRecipient.objects.filter(announcement_id=pk, recipient=request.user, announcement__is_active=True).select_related("announcement").first()
    return None


@never_cache
@require_http_methods(["GET"])
def portal_announcement_detail(request, pk):
    # The display is now part of master_admin_announcements.html.
    # Keep the existing URL/function name so saved notifications still open.
    from django.template import engines

    receipt = _portal_announcement_receipt(request, pk)
    if receipt:
        announcement = receipt.announcement
    elif _master_admin_allowed(request):
        announcement = Announcement.objects.filter(pk=pk, is_active=True).first()
    else:
        return HttpResponseForbidden("This announcement is not addressed to you.")
    if announcement is None:
        raise Http404("Announcement no longer available.")

    detail_base = engines["django"].from_string(
        '<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        '<title>{% block title %}Announcement{% endblock %}</title>'
        '{% block extra_css %}{% endblock %}</head><body>'
        '{% block content %}{% endblock %}'
        '{% block extra_js %}{% endblock %}</body></html>'
    )
    return render(request, "master_admin/master_admin_announcements.html", {
        "announcement": announcement,
        "can_mark_seen": bool(receipt and receipt.read_at is None),
        "announcement_detail_mode": True,
        "announcement_base_template": detail_base,
    })


@never_cache
@require_POST
def portal_announcement_seen(request, pk):
    receipt = _portal_announcement_receipt(request, pk)
    if receipt is None:
        return JsonResponse({"ok": False, "message": "This announcement is not addressed to you."}, status=403)
    type(receipt).objects.filter(pk=receipt.pk, read_at__isnull=True).update(read_at=timezone.now())
    return JsonResponse({"ok": True})



@never_cache
@require_http_methods(["GET", "POST"])
def admin_announcements(request):
    owner = _admin_notepad_owner(request)
    if owner is None:
        return redirect("admin_login")
    recipients = MonitoringUser.objects.filter(
        Q(job_Position__iexact="Team Lead") | Q(job_Position__iexact="Team Member")
    ).order_by("name", "pk")
    if request.method == "POST":
        action = request.POST.get("action", "")
        if action in {"edit_announcement", "delete_announcement"}:
            try:
                pk = int(request.POST.get("announcement_id", ""))
            except (ValueError, TypeError):
                messages.error(request, "Invalid announcement ID.")
                return redirect("admin_announcements")
            announcement = get_object_or_404(Announcement, pk=pk, is_admin_announcement=True, is_master_admin_announcement=False, is_active=True)
            if action == "delete_announcement":
                announcement.is_active = False
                announcement.save(update_fields=["is_active", "updated_at"])
                messages.success(request, "Announcement deleted successfully.")
            else:
                message = request.POST.get("message", "").strip()
                if not message:
                    messages.error(request, "Announcement message cannot be empty.")
                    return redirect("admin_announcements")
                announcement.message = message
                announcement.save(update_fields=["message", "updated_at"])
                messages.success(request, "Announcement updated successfully.")
            return redirect("admin_announcements")
        if action != "create_announcement":
            messages.error(request, "Invalid announcement action.")
            return redirect("admin_announcements")
        message = request.POST.get("message", "").strip()
        if not message:
            messages.error(request, "Please enter an announcement message.")
            return redirect("admin_announcements")
        send_to = request.POST.get("send_to", "team")
        if send_to == "selected":
            try:
                ids = [int(value) for value in request.POST.getlist("recipient_ids")]
            except (ValueError, TypeError):
                messages.error(request, "Select valid recipients.")
                return redirect("admin_announcements")
            recipients = recipients.filter(pk__in=ids)
        elif send_to != "team":
            messages.error(request, "Invalid recipient option.")
            return redirect("admin_announcements")
        recipients = list(recipients)
        if not recipients:
            messages.error(request, "Select at least one Team Lead or Team Member.")
            return redirect("admin_announcements")
        sender = {"admin_created_by": owner["auth_owner"]} if "auth_owner" in owner else {"created_by": owner["management_owner"]}
        with transaction.atomic():
            announcement = Announcement.objects.create(
                title="Admin Announcement", message=message,
                is_admin_announcement=True, **sender,
            )
            AnnouncementRecipient.objects.bulk_create([
                AnnouncementRecipient(announcement=announcement, recipient=user)
                for user in recipients
            ])
        messages.success(request, "Announcement sent successfully.")
        return redirect("admin_announcements")
    response = render(request, "admin_panel/admin_announcements.html", {
        "announcements": Announcement.objects.filter(is_admin_announcement=True, is_master_admin_announcement=False, is_active=True).prefetch_related("recipients__recipient").order_by("-created_at", "-pk")[:50],
        "team_members": recipients,
        "lead_announcements": Announcement.objects.filter(
            is_admin_announcement=False, is_active=True,
            created_by__job_Position__iexact="Team Lead",
        ).select_related("created_by", "created_by__department").order_by("-created_at", "-pk"),
    })

    if request.method == "GET" and request.GET.get("format") == "json":
        return JsonResponse({"ok": True, "html": response.content.decode(response.charset)})
    return response


from pathlib import Path
from types import SimpleNamespace
from django.core.exceptions import ValidationError
from django.core.validators import URLValidator
from django.db import transaction
from django.http import JsonResponse, HttpResponseForbidden, Http404, FileResponse
from django.shortcuts import render, redirect
from django.urls import reverse
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods, require_POST
from .models import AdminKnowledge, Knowledge


def _master_repository_denied(request):
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    return redirect("master_admin_login") if not request.user.is_authenticated else HttpResponseForbidden("Master Admin access required.")


def _master_repository_model(source):
    if source == "admin":
        return AdminKnowledge
    if source == "team":
        return Knowledge
    raise Http404("Unknown resource type.")


def _master_repository_values(request):
    title = request.POST.get("title", "").strip()
    description = request.POST.get("description", "").strip()
    link = request.POST.get("link", "").strip()
    file = request.FILES.get("file")
    error = ""
    if not title or len(title) > 255:
        error = "Enter a title with a maximum of 255 characters."
    elif link:
        try:
            URLValidator(schemes=["http", "https"])(link)
            if len(link) > 200:
                raise ValidationError("Link is too long.")
        except ValidationError:
            error = "Enter a valid HTTP/HTTPS link with a maximum of 200 characters."
    if file and file.size > 10 * 1024 * 1024:
        error = "The file must be 10 MB or smaller."
    return {"title": title, "description": description, "link": link}, file, error


def _master_repository_card(request, resource, source):
    admin = source == "admin"
    uploader = (resource.auth_owner or resource.management_owner) if admin else resource.user
    own = bool(admin and resource.auth_owner_id == request.user.pk)
    category = "mine" if own else "admin" if admin else "lead" if (resource.user.job_Position or "").strip().lower().replace("_", " ") == "team lead" else "member"
    link = resource.link or ""
    if link:
        try:
            URLValidator(schemes=["http", "https"])(link)
        except ValidationError:
            link = ""
    file = SimpleNamespace(name=resource.file.name,
        url=reverse("master_admin_repository_file", args=[source, resource.pk])) if resource.file else None
    return SimpleNamespace(id=resource.pk, card_key=f"{source}-{resource.pk}",
        title=resource.title, description=resource.description or "", link=link,
        file=file, created_at=resource.created_at, user=uploader,
        owner_username=uploader.username if uploader else "Not recorded",
        source_label="Master Admin" if own else "Admin" if admin else resource.user.job_Position,
        is_admin_resource=admin, can_manage=True, category=category, resource_source=source,
        update_url=reverse("master_admin_repository_update", args=[source, resource.pk]))


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_repository(request):
    if not _master_admin_allowed(request):
        return _master_repository_denied(request)
    if request.method == "POST":
        values, file, error = _master_repository_values(request)
        if error:
            return JsonResponse({"ok": False, "message": error}, status=400)
        with transaction.atomic():
            resource = AdminKnowledge.objects.create(auth_owner=request.user, file=file, **values)
        return JsonResponse({"ok": True, "message": "Resource published to Team Leads.", "id": resource.pk}, status=201)
    admin_items = AdminKnowledge.objects.select_related("auth_owner", "management_owner")
    team_items = Knowledge.objects.select_related("user", "department")
    cards = [_master_repository_card(request, item, "admin") for item in admin_items]
    cards += [_master_repository_card(request, item, "team") for item in team_items]
    cards.sort(key=lambda item: (item.created_at, item.card_key), reverse=True)
    return render(request, "master_admin/master_admin_repository.html", {
        "master_repository_active": True, "knowledge_items": cards,
        "resource_edit_items": [{"key": item.card_key, "title": item.title,
            "description": item.description, "link": item.link, "url": item.update_url,
            "file_name": Path(item.file.name).name if item.file else ""} for item in cards],
    })


@never_cache
@require_POST
def master_admin_repository_update(request, source, pk):
    if not _master_admin_allowed(request):
        return _master_repository_denied(request)
    model = _master_repository_model(source)
    values, file, error = _master_repository_values(request)
    if error:
        return JsonResponse({"ok": False, "message": error}, status=400)
    remove_file = request.POST.get("remove_file") == "1"
    if remove_file and file:
        return JsonResponse({"ok": False, "message": "Choose either replacement upload or Remove file."}, status=400)
    with transaction.atomic():
        resource = model.objects.select_for_update().filter(pk=pk).first()
        if resource is None:
            return JsonResponse({"ok": False, "message": "Resource no longer exists."}, status=404)
        for name, value in values.items():
            setattr(resource, name, value)
        if file:
            resource.file = file
        elif remove_file:
            resource.file = ""
        # Original owner/department and sharing scope remain unchanged.
        resource.save()
    return JsonResponse({"ok": True, "message": "Resource updated successfully."})


@never_cache
@require_POST
def master_admin_repository_delete(request, source, pk):
    if not _master_admin_allowed(request):
        return _master_repository_denied(request)
    model = _master_repository_model(source)
    with transaction.atomic():
        resource = model.objects.select_for_update().filter(pk=pk).first()
        if resource is None:
            return JsonResponse({"ok": False, "message": "Resource no longer exists."}, status=404)
        resource.delete()
    return JsonResponse({"ok": True, "message": "Resource deleted successfully."})


@never_cache
@require_http_methods(["GET"])
def master_admin_repository_file(request, source, pk):
    if not _master_admin_allowed(request):
        raise Http404("Resource not found.")
    resource = _master_repository_model(source).objects.filter(pk=pk).first()
    if resource is None or not resource.file:
        raise Http404("File not found.")
    try:
        stream = resource.file.open("rb")
    except (FileNotFoundError, OSError):
        raise Http404("File not found.")
    return FileResponse(stream, as_attachment=True, filename=Path(resource.file.name).name)

# ADD once at END of monitoringapp/views.py. Requires _master_admin_allowed.
from datetime import datetime, timedelta
from urllib.parse import urlsplit
from types import SimpleNamespace
from django.contrib.auth import get_user_model
from django.db import transaction
from django.db.models import Q
from django.http import JsonResponse, HttpResponseForbidden
from django.shortcuts import redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_time
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods, require_POST
from .models import GoogleMeeting, User, Department, TeamLeadNotification, AccountsNotification, AdminNotification


def _master_meet_link(link):
    try:
        parsed = urlsplit(link)
        return bool(len(link) <= 200 and parsed.scheme == "https"
            and parsed.netloc == "meet.google.com" and parsed.path not in {"", "/"})
    except (ValueError, TypeError):
        return False


def _master_meet_denied(request):
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    return redirect("master_admin_login") if not request.user.is_authenticated else HttpResponseForbidden("Master Admin access required.")


def _master_meet_notify(meeting, edited=False):
    title = ("Meeting updated: " if edited else "Meeting scheduled: ") + meeting.purpose
    message = f"{meeting.purpose}\n{meeting.date:%d %b %Y} at {meeting.time:%I:%M %p}\n{meeting.meeting_link}"
    for user in meeting.participants.all():
        role = (user.job_Position or "").strip().lower().replace("_", " ")
        defaults = {"title": title[:180], "message": message, "kind": "meeting", "is_read": False, "is_archived": False}
        key = f"master-meeting:{meeting.pk}:employee:{user.pk}"
        if role in {"team lead", "team member"}:
            defaults["url"] = reverse("teamlead_google_meet" if role == "team lead" else "teammember_google_meet")
            TeamLeadNotification.objects.update_or_create(event_key=key, defaults={"recipient": user, **defaults})
        elif role in {"accounts", "account", "accounts team"}:
            AccountsNotification.objects.update_or_create(event_key=key, defaults={"recipient": user, "actor": None, "url": meeting.meeting_link, **defaults})
        else:
            AdminNotification.objects.update_or_create(event_key=key, defaults={"management_owner": user, "auth_owner": None, "url": meeting.meeting_link, **defaults})
    for user in meeting.auth_participants.all():
        AdminNotification.objects.update_or_create(event_key=f"master-meeting:{meeting.pk}:auth:{user.pk}",
            defaults={"management_owner": None, "auth_owner": user, "title": title[:180], "message": message,
                "kind": "meeting", "url": meeting.meeting_link, "is_read": False, "is_archived": False})


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_google_meet(request):
    if not _master_admin_allowed(request):
        return _master_meet_denied(request)
    eligible = User.objects.select_related("department").order_by("name", "pk")
    auth_users = get_user_model().objects.filter(is_active=True).filter(Q(is_staff=True) | Q(is_superuser=True)).order_by("username", "pk")
    def fail(message, status=400):
        return JsonResponse({"ok": False, "message": message}, status=status)
    if request.method == "POST":
        action = request.POST.get("action", "create_meeting")
        if action not in {"create_meeting", "edit_meeting"}:
            return fail("Invalid meeting action.")
        purpose = request.POST.get("purpose", "").strip()
        if not purpose or len(purpose) > 255:
            return fail("Enter a meeting purpose with a maximum of 255 characters.")
        try:
            date = parse_date(request.POST.get("date", ""))
            time = parse_time(request.POST.get("time", ""))
        except (ValueError, TypeError):
            return fail("Enter a valid date and time.")
        if date is None or time is None or time.tzinfo:
            return fail("Enter a valid date and time.")
        link = request.POST.get("meeting_link", "").strip()
        if not _master_meet_link(link):
            return fail("Enter a valid https://meet.google.com/... meeting link.")
        mode = request.POST.get("participant_type", "all")
        if mode == "all":
            employees, admins = list(eligible), list(auth_users)
        elif mode == "leads":
            employees, admins = list(eligible.filter(job_Position__iexact="Team Lead")), []
        elif mode == "department":
            try:
                department = int(request.POST.get("department_id", ""))
            except (TypeError, ValueError):
                return fail("Choose a valid department.")
            if not Department.objects.filter(pk=department).exists():
                return fail("Choose a valid department.")
            employees, admins = list(eligible.filter(department_id=department)), []
        elif mode == "selected":
            try:
                tokens = set(map(int, request.POST.getlist("participants")))
            except (TypeError, ValueError):
                return fail("Choose valid participants.")
            # Positive IDs refer to employee User; negative tokens refer to Django auth admins.
            employee_ids = {value for value in tokens if value > 0}
            auth_ids = {-value for value in tokens if value < 0}
            employees, admins = list(eligible.filter(pk__in=employee_ids)), list(auth_users.filter(pk__in=auth_ids))
            if 0 in tokens or len(employees) != len(employee_ids) or len(admins) != len(auth_ids):
                return fail("A selected participant is no longer available.")
        else:
            return fail("Invalid participant selection.")
        if not employees and not admins:
            return fail("Choose at least one participant.")
        scheduled = timezone.make_aware(datetime.combine(date, time), timezone.get_current_timezone())
        if action == "create_meeting" and scheduled <= timezone.now():
            return fail("Choose a future meeting date and time.")
        with transaction.atomic():
            if action == "edit_meeting":
                try:
                    pk = int(request.POST.get("meeting_id", ""))
                except (TypeError, ValueError):
                    return fail("Invalid meeting ID.")
                meeting = GoogleMeeting.objects.select_for_update().filter(pk=pk).first()
                if meeting is None:
                    return fail("Meeting no longer exists.", 404)
                if meeting.meeting_status == "ended":
                    return fail("An ended meeting cannot be edited.")
                if (meeting.date != date or meeting.time != time) and scheduled <= timezone.now():
                    return fail("Choose a future time when rescheduling.")
            else:
                meeting = GoogleMeeting(admin_created_by=request.user, is_admin_meeting=True, is_master_admin_meeting=True)
            # Preserve the original meeting's creator when moderating it.
            meeting.purpose, meeting.date, meeting.time, meeting.meeting_link = purpose, date, time, link
            meeting.save()
            meeting.participants.set(employees)
            meeting.auth_participants.set(admins)
            prefix = f"master-meeting:{meeting.pk}:"
            keys = {f"{prefix}employee:{u.pk}" for u in employees} | {f"{prefix}auth:{u.pk}" for u in admins}
            for model in [TeamLeadNotification, AccountsNotification, AdminNotification]:
                model.objects.filter(event_key__startswith=prefix).exclude(event_key__in=keys).update(is_archived=True)
            _master_meet_notify(meeting, edited=action == "edit_meeting")
        return JsonResponse({"ok": True, "message": "Meeting updated successfully." if action == "edit_meeting" else "Meeting scheduled successfully.", "id": meeting.pk}, status=200 if action == "edit_meeting" else 201)
    meetings = list(GoogleMeeting.objects.select_related("created_by", "admin_created_by").prefetch_related("participants", "auth_participants"))
    now = timezone.now()
    def priority(meeting):
        start = meeting.scheduled_datetime
        if start + timedelta(hours=1) <= now:
            return (2, -start.timestamp(), -meeting.pk)
        return (0 if start <= now else 1, start.timestamp(), -meeting.pk)
    meetings.sort(key=priority)
    for meeting in meetings:
        employees, admins = list(meeting.participants.all()), list(meeting.auth_participants.all())
        meeting.participant_ids = ",".join([str(u.pk) for u in employees] + [str(-u.pk) for u in admins])
        meeting.participant_names = ", ".join([u.name for u in employees] + [(u.get_full_name() or u.username) + " (Admin)" for u in admins])
        meeting.participant_total = len(employees) + len(admins)
    people = list(eligible) + [SimpleNamespace(pk=-u.pk, id=-u.pk, name=(u.get_full_name() or u.username) + " (Admin)", department=None, job_Position="Admin", designation="Admin") for u in auth_users]
    return render(request, "master_admin/master_admin_google_meet.html", {
        "master_meet_active": True, "meetings": meetings, "team_members": people,
        "departments": Department.objects.order_by("name"),
    })


@never_cache
@require_POST
def master_admin_google_meet_delete(request, pk):
    if not _master_admin_allowed(request):
        return _master_meet_denied(request)
    with transaction.atomic():
        meeting = GoogleMeeting.objects.select_for_update().filter(pk=pk).first()
        if meeting is None:
            return JsonResponse({"ok": False, "message": "Meeting no longer exists."}, status=404)
        prefix = f"master-meeting:{meeting.pk}:"
        for model in [TeamLeadNotification, AccountsNotification, AdminNotification]:
            model.objects.filter(event_key__startswith=prefix).update(is_archived=True)
        meeting.delete()
    return JsonResponse({"ok": True, "message": "Meeting removed from the portal."})



# REPLACE only these TWO existing functions in monitoringapp/views.py. Do not append duplicates.
# Admin: Master meetings cannot be edited in regular Admin portal.
# Accounts: meeting notifications open the validated Google Meet link.
# Previous announcement notification handling is preserved.

@never_cache
@require_http_methods(["GET", "POST"])
def admin_google_meet(request):
    owner = _admin_notepad_owner(request)
    if owner is None:
        return redirect("admin_login")

    eligible = User.objects.filter(
        Q(job_Position__iexact="Team Lead")
        | Q(job_Position__iexact="Team Member")
    ).select_related("department").order_by("name", "pk")
    meetings = GoogleMeeting.objects.filter(is_admin_meeting=True)

    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"

    def fail(message):
        if is_ajax:
            return JsonResponse({"ok": False, "message": message}, status=400)
        messages.error(request, message)
        return redirect("admin_google_meet")

    if request.method == "POST":
        action = request.POST.get("action", "create_meeting")
        if action not in {"create_meeting", "edit_meeting"}:
            return fail("Invalid meeting action.")
        purpose = request.POST.get("purpose", "").strip()
        if not purpose or len(purpose) > 255:
            return fail("Enter a meeting purpose with a maximum of 255 characters.")
        try:
            meeting_date = parse_date(request.POST.get("date", ""))
            meeting_time = parse_time(request.POST.get("time", ""))
        except (ValueError, TypeError):
            return fail("Enter a valid meeting date and time.")
        if meeting_date is None or meeting_time is None or meeting_time.tzinfo:
            return fail("Enter a valid meeting date and time.")
        link = request.POST.get("meeting_link", "").strip()
        try:
            parsed_link = urlsplit(link)
            valid_link = (
                parsed_link.scheme == "https"
                and parsed_link.netloc == "meet.google.com"
                and parsed_link.path not in {"", "/"}
                and len(link) <= 200
            )
        except ValueError:
            valid_link = False
        if not valid_link:
            return fail("Enter a valid https://meet.google.com/... meeting link.")

        mode = request.POST.get("participant_type", "all")
        if mode == "all":
            participants = eligible
        elif mode == "leads":
            participants = eligible.filter(job_Position__iexact="Team Lead")
        elif mode == "department":
            try:
                department_id = int(request.POST.get("department_id", ""))
            except (ValueError, TypeError):
                return fail("Choose a valid department.")
            if not Department.objects.filter(pk=department_id).exists():
                return fail("Choose a valid department.")
            participants = eligible.filter(department_id=department_id)
        elif mode == "selected":
            try:
                ids = set(int(value) for value in request.POST.getlist("participants"))
            except (ValueError, TypeError):
                return fail("Choose valid participants.")
            participants = eligible.filter(pk__in=ids)
            if not ids or participants.count() != len(ids):
                return fail("Choose valid Team Leads or Team Members.")
        else:
            return fail("Invalid participant selection.")
        participants = list(participants)
        if not participants:
            return fail("Select at least one participant.")

        scheduled = timezone.make_aware(
            datetime.combine(meeting_date, meeting_time),
            timezone.get_current_timezone(),
        )
        if action == "create_meeting" and scheduled <= timezone.now():
            return fail("Choose a future meeting date and time.")
        with transaction.atomic():
            if action == "edit_meeting":
                try:
                    meeting_id = int(request.POST.get("meeting_id", ""))
                except (ValueError, TypeError):
                    return fail("Invalid meeting ID.")
                meeting = meetings.select_for_update().filter(pk=meeting_id).first()
                if meeting is None:
                    return fail("Meeting no longer exists.")
                if meeting.is_master_admin_meeting:
                    return fail("This meeting is managed in the Master Admin portal.")
                if meeting.meeting_status == "ended":
                    return fail("An ended meeting cannot be edited.")
                if (meeting.date != meeting_date or meeting.time != meeting_time) and scheduled <= timezone.now():
                    return fail("Choose a future time when rescheduling a meeting.")
            else:
                sender = (
                    {"admin_created_by": owner["auth_owner"]}
                    if "auth_owner" in owner
                    else {"created_by": owner["management_owner"]}
                )
                meeting = GoogleMeeting(is_admin_meeting=True, **sender)
            meeting.purpose = purpose
            meeting.date = meeting_date
            meeting.time = meeting_time
            meeting.meeting_link = link
            meeting.save()
            meeting.participants.set(participants)
            
            # Windows notifications for invited Team Leads.
            teamlead_participants = [
                person for person in participants
                if (person.job_Position or "").strip().lower()
                == "team lead"
            ]

            notification_title = (
                "Google Meet Updated"
                if action == "edit_meeting"
                else "New Google Meet Scheduled"
            )

            for person in teamlead_participants:
                TeamLeadNotification.objects.create(
                    recipient=person,
                    kind="meeting",
                    title=notification_title,
                    message=(
                        f"{purpose}\n"
                        f"Date: {meeting_date:%d %b %Y}\n"
                        f"Time: {meeting_time:%I:%M %p}"
                    ),
                    url=reverse("teamlead_google_meet"),
                )


            
        message = "Meeting updated successfully." if action == "edit_meeting" else "Meeting scheduled successfully."
        if is_ajax:
            return JsonResponse({"ok": True, "message": message, "id": meeting.pk}, status=200 if action == "edit_meeting" else 201)
        messages.success(request, message)
        return redirect("admin_google_meet")

    now = timezone.now()
    meeting_list = list(meetings.prefetch_related("participants"))
    def priority(meeting):
        start = meeting.scheduled_datetime
        ended = start + timedelta(hours=1) <= now
        return (1 if ended else 0, -start.timestamp(), -meeting.pk)
    meeting_list.sort(key=priority)
    response = render(request, "admin_panel/admin_google_meet.html", {
        "meetings": meeting_list,
        "team_members": eligible,
        "departments": Department.objects.order_by("name"),
    })

    if request.GET.get("format") == "json":
        return JsonResponse({"ok": True, "html": response.content.decode(response.charset)})
    return response

@never_cache
@require_POST
def accounts_notification_action(request, pk):
    from django.http import JsonResponse
    from django.urls import reverse
    from .models import AccountsNotification

    current_user = _accounts_user(request)
    if current_user is None:
        return JsonResponse({
            "ok": False,
            "message": "Accounts access required.",
        }, status=403 if request.session.get("user_id") else 401)

    item = AccountsNotification.objects.filter(
        pk=pk,
        recipient=current_user,
    ).first()

    if item is None:
        return JsonResponse({
            "ok": False,
            "message": "Notification not found.",
        }, status=404)

    action = request.POST.get("action", "")
    target_url = ""

    if action in {"read", "open"}:
        item.is_read = True
        item.save(update_fields=["is_read"])

        if action == "open":
            routes = {
                "income": "accounts_income",
                "expense": "accounts_expenses",
                "sale": "accounts_sales",
                "payment": "accounts_sales",
                "reminder": "accounts_reminders",
                "note": "accounts_notepad",
            }
            target_url = reverse(routes.get(item.kind, "accounts_notifications"))
            if item.kind == "announcement":
                import re
                match = re.fullmatch(r"/portal/announcements/(\d+)/", item.url or "")
                if match:
                    target_url = reverse("portal_announcement_detail", args=[int(match.group(1))])
            elif item.kind == "meeting":
                from urllib.parse import urlsplit
                try:
                    parsed = urlsplit(item.url or "")
                    if parsed.scheme == "https" and parsed.netloc == "meet.google.com" and parsed.path not in {"", "/"}:
                        target_url = item.url
                except ValueError:
                    pass

    elif action == "unread":
        item.is_read = False
        item.save(update_fields=["is_read"])

    elif action == "archive":
        item.is_archived = True
        item.save(update_fields=["is_archived"])

    elif action == "restore":
        item.is_archived = False
        item.save(update_fields=["is_archived"])

    else:
        return JsonResponse({
            "ok": False,
            "message": "Invalid notification action.",
        }, status=400)

    unread_count = AccountsNotification.objects.filter(
        recipient=current_user,
        is_archived=False,
        is_read=False,
    ).count()

    return JsonResponse({
        "ok": True,
        "message": {
            "read": "Marked as read.",
            "open": "Notification opened.",
            "unread": "Marked as unread.",
            "archive": "Notification archived.",
            "restore": "Notification restored to inbox.",
        }[action],
        "unread_count": unread_count,
        "url": target_url,
    })




@never_cache
@require_GET
def master_admin_notifications(request):
    from django.core.paginator import Paginator
    from django.db.models import Count, Q
    from django.http import JsonResponse
    from django.urls import reverse
    from .models import AdminNotification
    if not _master_admin_allowed(request):
        if request.GET.get("format") == "json":
            return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
        return redirect("master_admin_login")
    _process_admin_reminders({"auth_owner": request.user})
    owned = AdminNotification.objects.filter(auth_owner=request.user).exclude(kind__in=["chat", "message", "group_chat"])
    totals = owned.aggregate(
        inbox=Count("pk", filter=Q(is_archived=False)),
        unread=Count("pk", filter=Q(is_read=False, is_archived=False)),
        archived=Count("pk", filter=Q(is_archived=True)),
    )
    folder_counts = {
        "all": totals["inbox"],
        "unread": totals["unread"],
        "archived": totals["archived"],
    }
    unread = totals["unread"]
    latest_notice = owned.filter(is_archived=False).order_by("-pk").values("pk", "title").first()
    latest_meta = {"latest_id": latest_notice["pk"] if latest_notice else 0, "latest_title": latest_notice["title"] if latest_notice else ""}
    if request.GET.get("counts") == "1":
        return JsonResponse({"ok": True, "unread_count": unread, "folder_counts": folder_counts, **latest_meta})
    tab = request.GET.get("tab", "all")
    if tab not in {"all", "unread", "archived"}:
        tab = "all"
    items = owned.filter(is_archived=tab == "archived")
    if tab == "unread":
        items = items.filter(is_read=False)
    query = request.GET.get("q", "").strip()[:200]
    if query:
        items = items.filter(Q(title__icontains=query) | Q(message__icontains=query))
    kind = request.GET.get("kind", "").strip()[:40]
    if kind:
        items = items.filter(kind=kind)
    page = Paginator(items.order_by("-created_at", "-pk"), 20).get_page(request.GET.get("page"))
    if request.GET.get("format") == "json":
        return JsonResponse({"ok": True, "unread_count": unread, "folder_counts": folder_counts, "page": page.number,
            "pages": page.paginator.num_pages, "total": page.paginator.count,
            **latest_meta, "items": [{"id": n.pk, "title": n.title, "message": n.message, "kind": n.kind,
                "created_at": timezone.localtime(n.created_at).strftime("%d %b %Y, %I:%M %p"),
                "is_read": n.is_read, "is_archived": n.is_archived,
                "url": n.url if n.url.startswith("/") and not n.url.startswith("//") and "\\" not in n.url else "",
                "action_url": reverse("master_admin_notification_action", args=[n.pk])} for n in page]})
    return render(request, "master_admin/master_admin_notifications.html", {"master_notifications_active": True})



def _master_notice_messages(request):
    from django.contrib import messages as message_api
    storage = message_api.get_messages(request)
    was_used = storage.used
    values = list(storage)
    storage.used = was_used
    return values


def _master_operation_notice(view, kind, route, label):
    from functools import wraps
    import json
    import logging

    @wraps(view)
    def wrapped(request, *args, **kwargs):
        if request.method != "POST" or not _master_admin_allowed(request):
            return view(request, *args, **kwargs)
        old_messages = {id(message) for message in _master_notice_messages(request)}
        previous_status = None
        if view.__name__ == "master_admin_edit_user":
            from .models import User as Employee
            previous_status = Employee.objects.filter(pk=request.POST.get("id") or None).values_list("status", flat=True).first() if str(request.POST.get("id", "")).isdigit() else None
        response = view(request, *args, **kwargs)
        data = None
        if "application/json" in response.get("Content-Type", ""):
            try:
                data = json.loads(response.content)
            except (ValueError, TypeError, UnicodeDecodeError):
                return response
            succeeded = response.status_code < 300 and data.get("ok") is True
            message = str(data.get("message") or label + " updated successfully.")
        elif response.status_code in {302, 303}:
            fresh = [m for m in _master_notice_messages(request) if id(m) not in old_messages]
            success = [m for m in fresh if m.level == 25]
            succeeded = bool(success) and not any(m.level >= 40 for m in fresh)
            message = str(success[-1]) if success else ""
        else:
            succeeded = False
            message = ""
        if not succeeded:
            return response
        from django.db import transaction
        from django.urls import reverse
        from .models import AdminNotification
        detail = message
        # Include the account status selected in User Management, rather than presence/idle status.
        status = request.POST.get("edit_status", request.POST.get("status", "")).strip()
        name = request.POST.get("edit_name", request.POST.get("name", "")).strip()
        if kind == "user":
            if name:
                detail += "\nUser: " + name[:150]
            if status:
                detail += "\nAccount status: " + status[:40]
                if previous_status is not None and str(previous_status) != status:
                    detail += " (previously " + str(previous_status)[:40] + ")"
        actor = request.user
        # Existing recipient notifications remain untouched. This is the acting Master's inbox entry.
        try:
            with transaction.atomic():
                AdminNotification.objects.create(auth_owner=actor, kind=kind, title=message[:180], message=detail, url=reverse(route))
            if data is not None:
                data["unread_count"] = AdminNotification.objects.filter(auth_owner=actor, is_read=False, is_archived=False).exclude(kind__in=["chat", "message", "group_chat"]).count()
                data["notification_saved"] = True
        except Exception:
            logging.getLogger(__name__).exception("Master Admin operation succeeded but its notification could not be saved")
            if data is not None:
                data["notification_saved"] = False
        if data is not None:
            response.content = json.dumps(data)
        return response
    wrapped._master_operation_notice = True
    return wrapped


# Only successful mutation views are registered. GETs, validation failures, Chat,
# notification read/archive actions, and passive activity heartbeats never create entries.
_MASTER_NOTICE_VIEWS = {
    "master_admin_usermanagement": ("user", "master_admin_usermanagement", "User"),
    "master_admin_edit_user": ("user", "master_admin_usermanagement", "User"),
    "master_admin_delete_user": ("user", "master_admin_usermanagement", "User"),
    "master_admin_departments": ("department", "master_admin_departments", "Department"),
    "master_admin_teams": ("team", "master_admin_teams", "Team"),
    "master_admin_accounts_income": ("income", "master_admin_accounts_income", "Income"),
    "master_admin_accounts_income_edit": ("income", "master_admin_accounts_income", "Income"),
    "master_admin_accounts_income_delete": ("income", "master_admin_accounts_income", "Income"),
    "master_admin_accounts_expenses": ("expense", "master_admin_accounts_expenses", "Expense"),
    "master_admin_accounts_expense_edit": ("expense", "master_admin_accounts_expenses", "Expense"),
    "master_admin_accounts_expense_delete": ("expense", "master_admin_accounts_expenses", "Expense"),
    "master_admin_accounts_sales": ("sale", "master_admin_accounts_sales", "Sale"),
    "master_admin_accounts_sale_edit": ("sale", "master_admin_accounts_sales", "Sale"),
    "master_admin_accounts_sale_delete": ("sale", "master_admin_accounts_sales", "Sale"),
    "master_admin_accounts_sale_payments": ("sale", "master_admin_accounts_sales", "Sale payment"),
    "master_admin_tasks": ("task", "master_admin_tasks", "Task"),
    "master_admin_task_update": ("task", "master_admin_tasks", "Task"),
    "master_admin_task_delete": ("task", "master_admin_tasks", "Task"),
    "master_admin_projects": ("project", "master_admin_projects", "Project"),
    "master_admin_project_update": ("project", "master_admin_projects", "Project"),
    "master_admin_project_delete": ("project", "master_admin_projects", "Project"),
    "master_admin_announcements": ("announcement", "master_admin_announcements", "Announcement"),
    "master_admin_repository": ("resource", "master_admin_repository", "Resource"),
    "master_admin_repository_update": ("resource", "master_admin_repository", "Resource"),
    "master_admin_repository_delete": ("resource", "master_admin_repository", "Resource"),
    "master_admin_google_meet": ("meeting", "master_admin_google_meet", "Meeting"),
    "master_admin_google_meet_delete": ("meeting", "master_admin_google_meet", "Meeting"),
    "master_admin_reminders": ("reminder", "master_admin_reminders", "Reminder"),
}


def _master_report_notice(sender, instance, created=None, **kwargs):
    # Report submissions happen on Team Lead/Member pages, so cover them at the model event.
    if kwargs.get("raw"):
        return
    from django.contrib.auth import get_user_model
    from django.db import transaction
    from django.urls import reverse
    from uuid import uuid4
    import logging
    from .models import AdminNotification
    action = "deleted" if created is None else "submitted" if created else "updated"
    try:
        who = getattr(instance.user, "name", "") or "Employee"
    except Exception:
        who = "Employee"
    title = f"{sender._meta.verbose_name.title()} {action}"
    message = f"{who}: report #{instance.pk} {action}."
    key = "master-report:" + uuid4().hex
    def publish():
        try:
            for account_id in get_user_model().objects.filter(is_active=True, is_superuser=True).values_list("pk", flat=True):
                AdminNotification.objects.get_or_create(event_key=f"{key}:{account_id}", defaults={"auth_owner_id": account_id, "kind": "report", "title": title[:180], "message": message, "url": reverse("master_admin_reports")})
        except Exception:
            logging.getLogger(__name__).exception("Unable to deliver Master Admin report notification")
    transaction.on_commit(publish)


from django.db.models.signals import post_save as _master_notice_post_save, post_delete as _master_notice_post_delete
from .models import MorningReport as _MasterMorningReport, EveningReport as _MasterEveningReport
for _notice_report_model in (_MasterMorningReport, _MasterEveningReport):
    _notice_key = _notice_report_model._meta.label_lower
    _master_notice_post_save.connect(_master_report_notice, sender=_notice_report_model, dispatch_uid="master-notice-save:" + _notice_key, weak=False)
    _master_notice_post_delete.connect(_master_report_notice, sender=_notice_report_model, dispatch_uid="master-notice-delete:" + _notice_key, weak=False)


@never_cache
@require_POST
def master_admin_notification_action(request, pk):
    from django.http import JsonResponse
    from .models import AdminNotification
    if not _master_admin_allowed(request):
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    owned = AdminNotification.objects.filter(auth_owner=request.user)
    actions = {"read": {"is_read": True}, "unread": {"is_read": False},
               "archive": {"is_archived": True}, "restore": {"is_archived": False}}
    action = request.POST.get("action")
    if action == "read_all":
        owned.filter(is_archived=False, is_read=False).update(is_read=True)
    elif action not in actions:
        return JsonResponse({"ok": False, "message": "Invalid action."}, status=400)
    elif not owned.filter(pk=pk).update(**actions[action]):
        return JsonResponse({"ok": False, "message": "Notification not found."}, status=404)
    return JsonResponse({"ok": True, "message": "Notifications updated.",
        "unread_count": owned.filter(is_archived=False, is_read=False).count()})


# Append to monitoringapp/views.py. Keep existing reminder helpers and Admin view.
@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_reminders(request):
    import calendar
    from datetime import date, datetime, time
    from django.db import transaction
    from django.http import JsonResponse
    from django.utils import timezone
    from .models import AdminReminder, AdminNotification
    owner = {"auth_owner": request.user} if _master_admin_allowed(request) else None
    ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    if owner is None:
        if ajax or request.GET.get("format") == "json":
            return JsonResponse({"ok": False, "message": "Admin access required."}, status=403)
        return redirect("master_admin_login")
    reminders = AdminReminder.objects.filter(**owner)

    def failure(message, status=400):
        if ajax:
            return JsonResponse({"ok": False, "message": message}, status=status)
        messages.error(request, message)
        return redirect("master_admin_reminders")

    def success(message, reminder=None, deleted_id=None):
        if ajax:
            return JsonResponse({"ok": True, "message": message, "id": reminder.pk if reminder else deleted_id,
                                 "item": _admin_reminder_item(reminder) if reminder else None,
                                 "is_completed": reminder.is_completed if reminder else None,
                                 "unread_count": AdminNotification.objects.filter(**owner, is_read=False, is_archived=False).count()})
        messages.success(request, message)
        return redirect("master_admin_reminders")

    if request.method == "POST":
        action = request.POST.get("action", "")
        if action not in {"create", "update", "delete", "complete", "reopen"}:
            return failure("Invalid reminder action.")
        with transaction.atomic():
            reminder = None
            if action != "create":
                try:
                    reminder_id = int(request.POST.get("reminder_id", ""))
                except (TypeError, ValueError):
                    return failure("Invalid reminder ID.")
                reminder = reminders.select_for_update().filter(pk=reminder_id).first()
                if reminder is None:
                    return failure("Reminder not found.", 404)
            if action == "delete":
                deleted_id = reminder.pk
                if reminder.notification_id:
                    AdminNotification.objects.filter(pk=reminder.notification_id, **owner).update(is_archived=True)
                reminder.delete()
                return success("Reminder deleted.", deleted_id=deleted_id)
            if action in {"complete", "reopen"}:
                reminder.is_completed = action == "complete"
                reminder.save(update_fields=["is_completed", "updated_at"])
                if reminder.notification_id and reminder.is_completed:
                    AdminNotification.objects.filter(pk=reminder.notification_id, **owner).update(is_read=True)
                # Reopening keeps an already-delivered alert delivered; reschedule to get a new alert.
                return success("Reminder completed." if reminder.is_completed else "Reminder reopened.", reminder)
            title = request.POST.get("title", "").strip()
            description = request.POST.get("description", "").strip()
            event_at = _admin_reminder_time(request.POST.get("event_at"))
            remind_at = _admin_reminder_time(request.POST.get("remind_at"))
            if not title or len(title) > 180:
                return failure("Enter a title with a maximum of 180 characters.")
            if event_at is None or remind_at is None:
                return failure("Enter valid event and reminder dates and times.")
            if remind_at > event_at:
                return failure("Reminder must be on or before the event.")
            changed = reminder is None or reminder.event_at != event_at or reminder.remind_at != remind_at
            if changed and remind_at <= timezone.now():
                return failure("Choose a future reminder time.")
            if reminder is None:
                reminder = AdminReminder(**owner)
            elif changed:
                if reminder.notification_id:
                    AdminNotification.objects.filter(pk=reminder.notification_id, **owner).update(is_archived=True)
                reminder.schedule_version += 1
            if changed:
                reminder.notified_at = None
                reminder.notification = None
                reminder.is_completed = False
            reminder.title, reminder.description = title, description
            reminder.event_at, reminder.remind_at = event_at, remind_at
            reminder.save()
            return success("Reminder created." if action == "create" else "Reminder updated.", reminder)

    _process_admin_reminders(owner)
    today = timezone.localdate()
    try:
        selected = date.fromisoformat(request.GET.get("month", today.strftime("%Y-%m")) + "-01")
        next_month = selected.replace(year=selected.year + 1, month=1) if selected.month == 12 else selected.replace(month=selected.month + 1)
    except (TypeError, ValueError, OverflowError):
        selected = today.replace(day=1)
        next_month = selected.replace(year=selected.year + 1, month=1) if selected.month == 12 else selected.replace(month=selected.month + 1)
    tz = timezone.get_current_timezone()
    start = timezone.make_aware(datetime.combine(selected, time.min), tz)
    end = timezone.make_aware(datetime.combine(next_month, time.min), tz)
    monthly = reminders.filter(event_at__gte=start, event_at__lt=end).order_by("event_at", "id")
    context = {"master_reminders_active": True, "selected_month": selected.strftime("%Y-%m"), "month_label": selected.strftime("%B %Y"),
               "calendar_weeks": calendar.Calendar(firstweekday=0).monthdayscalendar(selected.year, selected.month),
               "calendar_items": [_admin_reminder_item(r) for r in monthly], "today": today.isoformat(), "reminders": monthly}
    if request.GET.get("format") == "json":
        return JsonResponse({"ok": True, "items": context["calendar_items"], "selected_month": context["selected_month"],
                             "month_label": context["month_label"], "calendar_weeks": context["calendar_weeks"], "today": context["today"]})
    return render(request, "master_admin/master_admin_reminders.html", context)



# Append these NEW functions to monitoringapp/views.py. Keep Admin/Team Lead views.
@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_notepad(request):
    from django.core.paginator import Paginator
    from django.db.models import Q
    from django.http import JsonResponse
    from django.shortcuts import render, redirect
    from django.urls import reverse
    from django.utils import timezone
    from .models import AdminNotepad, AdminNotification
    ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    if not _master_admin_allowed(request):
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403) if ajax else redirect("master_admin_login")
    owned = AdminNotepad.objects.filter(auth_owner=request.user)
    title, content, error = "", "", ""
    if request.method == "POST":
        title = request.POST.get("title", "").strip() or "Untitled"
        content = request.POST.get("content", "")
        selected_id = request.POST.get("note_id")
        note = None
        if selected_id:
            try:
                note = owned.filter(pk=int(selected_id)).first()
            except (TypeError, ValueError, OverflowError):
                return JsonResponse({"ok": False, "message": "Invalid note ID."}, status=400)
            if note is None:
                return JsonResponse({"ok": False, "message": "Note not found."}, status=404)
        if len(title) > 255:
            error = "Title must be 255 characters or fewer."
        elif not content.strip():
            error = "Write some content before saving."
        if error and ajax:
            return JsonResponse({"ok": False, "message": error}, status=400)
        if not error:
            updated = note is not None
            if note is None:
                note = AdminNotepad(auth_owner=request.user)
            note.title, note.content = title, content
            note.save()
            message = "Note updated successfully." if updated else "Note created successfully."
            _create_admin_notification({"auth_owner": request.user}, kind="note", title=message,
                message="Your personal notepad was updated.", url=reverse("master_admin_notepad"))
            if ajax:
                return JsonResponse({"ok": True, "message": message, "id": note.pk,
                    "title": note.title, "content": note.content,
                    "updated_at": timezone.localtime(note.updated_at).strftime("%d %b %Y, %I:%M %p"),
                    "unread_count": AdminNotification.objects.filter(auth_owner=request.user, is_read=False, is_archived=False).count()})
            messages.success(request, message)
            return redirect("master_admin_notepad")
    params = request.GET
    query = params.get("q", "").strip()[:200]
    orderings = {"updated": ("-updated_at", "-pk"), "created": ("-created_at", "-pk"), "title": ("title", "pk")}
    sort = params.get("sort", "updated")
    if sort not in orderings:
        sort = "updated"
    notes = owned.filter(Q(title__icontains=query) | Q(content__icontains=query)) if query else owned
    page_obj = Paginator(notes.order_by(*orderings[sort]), 4).get_page(params.get("page"))
    return render(request, "master_admin/master_admin_notepad.html", {
        "master_notepad_active": True, "note": None, "page_obj": page_obj,
        "query": query, "sort": sort, "form_title": title, "form_content": content,
        "note_error": error, "total_notes": owned.count(),
    }, status=400 if error else 200)


@never_cache
@require_POST
def master_admin_notepad_delete(request, pk):
    from django.http import JsonResponse
    from django.urls import reverse
    from .models import AdminNotepad, AdminNotification
    if not _master_admin_allowed(request):
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403)
    note = AdminNotepad.objects.filter(auth_owner=request.user, pk=pk).first()
    if note is None:
        return JsonResponse({"ok": False, "message": "Note not found."}, status=404)
    note.delete()
    _create_admin_notification({"auth_owner": request.user}, kind="note", title="Note deleted successfully.",
        message="A personal note was deleted.", url=reverse("master_admin_notepad"))
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        return JsonResponse({"ok": True, "message": "Note deleted successfully.", "id": pk,
            "unread_count": AdminNotification.objects.filter(auth_owner=request.user, is_read=False, is_archived=False).count()})
    messages.success(request, "Note deleted successfully.")
    return redirect("master_admin_notepad")


# Append to monitoringapp/views.py. Keep existing Admin/Team Lead profile functions.
@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_profile(request):
    from io import BytesIO
    from uuid import uuid4
    from types import SimpleNamespace
    from PIL import Image, ImageOps, UnidentifiedImageError
    from django.contrib.auth import update_session_auth_hash
    from django.contrib.auth.forms import PasswordChangeForm
    from django.core.exceptions import ValidationError
    from django.core.files.base import ContentFile
    from django.core.validators import validate_email
    from django.db import transaction
    from django.http import JsonResponse
    from django.urls import reverse
    from .models import MasterAdminProfile, AdminNotification
    ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    if not _master_admin_allowed(request):
        return JsonResponse({"ok": False, "message": "Master Admin access required."}, status=403) if ajax else redirect("master_admin_login")
    account = request.user
    profile, _ = MasterAdminProfile.objects.get_or_create(user=account)

    def fail(message):
        return JsonResponse({"ok": False, "message": message}, status=400)

    if request.method == "POST":
        action = request.POST.get("action", "")
        if action == "change_password":
            password_form = PasswordChangeForm(account, request.POST)
            if not password_form.is_valid():
                errors = [str(error) for values in password_form.errors.values() for error in values]
                return fail(" ".join(errors))
            account = password_form.save()
            update_session_auth_hash(request, account)
            message = "Password changed successfully."
        elif action == "edit_profile" and request.FILES.get("profile_image"):
            upload = request.FILES["profile_image"]
            if upload.size > 5 * 1024 * 1024:
                return fail("Choose an image smaller than 5 MB.")
            try:
                with Image.open(upload) as source:
                    if source.format not in {"JPEG", "PNG", "WEBP"} or source.width * source.height > 20000000:
                        return fail("Choose a JPG, PNG or WebP image up to 20 megapixels.")
                    source.load()
                    image = ImageOps.fit(ImageOps.exif_transpose(source).convert("RGB"), (512, 512))
                    buffer = BytesIO()
                    image.save(buffer, format="JPEG", quality=90)
            except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError):
                return fail("Choose a valid JPG, PNG or WebP image.")
            profile.profile_image.save(uuid4().hex + ".jpg", ContentFile(buffer.getvalue()), save=True)
            message = "Profile photo updated successfully."
        elif action == "edit_profile":
            name = request.POST.get("name", "").strip()
            email = request.POST.get("email", "").strip()
            phone = request.POST.get("phone", "").strip()
            if not name or len(name) > 150:
                return fail("Enter a name with a maximum of 150 characters.")
            if len(phone) > 30:
                return fail("Phone must be 30 characters or fewer.")
            try:
                validate_email(email)
            except ValidationError:
                return fail("Enter a valid email address.")
            if len(email) > account._meta.get_field("email").max_length:
                return fail("Email address is too long.")
            account.first_name, _, account.last_name = name.partition(" ")
            account.email = email
            profile.phone = phone
            with transaction.atomic():
                account.save(update_fields=["first_name", "last_name", "email"])
                profile.save(update_fields=["phone"])
            message = "Profile updated successfully."
        else:
            return fail("Invalid profile action.")
        _create_admin_notification({"auth_owner": account}, kind="profile", title=message,
            message="Your Master Admin account was updated.", url=reverse("master_admin_profile"))
        if ajax:
            return JsonResponse({"ok": True, "success": True, "message": message,
                "image_url": profile.profile_image.url if profile.profile_image else "",
                "unread_count": AdminNotification.objects.filter(auth_owner=account, is_read=False, is_archived=False).count()})
        messages.success(request, message)
        return redirect("master_admin_profile")
    display_user = SimpleNamespace(name=account.get_full_name() or account.username,
        username=account.username, email=account.email, phone=profile.phone,
        profile_image=profile.profile_image, joined=account.date_joined,
        account_status="Active" if account.is_active else "Inactive")
    return render(request, "master_admin/master_admin_profile.html", {
        "user": display_user, "master_profile_active": True,
    })


# Append these NEW functions to monitoringapp/views.py. Keep all existing chat views.

def _master_chat_user(request):
    from .master_chat_identity import master_chat_user
    return master_chat_user(request.user, create=True, touch=True)


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_chat(request):
    current_user = _master_chat_user(request)
    if current_user is None:
        return redirect("master_admin_login")

    if request.method == "POST":
        action = request.POST.get("action")

        if action == "create_group":
            return _member_chat_create_group(
                request,
                current_user,
                group_view_name="master_admin_group_chat",
            )

        return _member_chat_error("Invalid action.")

    context = _member_chat_context(current_user)
    context["user_role"] = "master_admin"

    context["master_chat_active"] = True
    return render(
        request,
        "master_admin/master_admin_chat.html",
        context,
    )


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_chat_room(request, user_id):
    from .member_chat_actions import chat_action

    current_user = _master_chat_user(request)
    if current_user is None:
        return redirect("master_admin_login")

    other_user = get_object_or_404(
        User,
        pk=user_id,
    )

    if current_user.pk == other_user.pk:
        return redirect("master_admin_chat")

    # Preserve existing conversations in either user ordering.
    room = (
        Conversation.objects.filter(
            Q(user1=current_user, user2=other_user)
            | Q(user1=other_user, user2=current_user)
        )
        .order_by("pk")
        .first()
    )

    if room is None:
        first, second = sorted(
            [current_user, other_user],
            key=lambda person: person.pk,
        )

        room, _ = Conversation.objects.get_or_create(
            user1=first,
            user2=second,
        )

    if request.method == "POST":
        action = request.POST.get("action")

        if action == "create_group":
            return _member_chat_create_group(
                request,
                current_user,
                group_view_name="master_admin_group_chat",
            )

        return chat_action(
            request,
            current_user,
            room,
            group=False,
        )

    messages_qs = (
        room.messages
        .exclude(hidden_for=current_user)
        .select_related("sender")
        .order_by("created_at", "pk")
    )

    # Used by the chat UI to reconcile messages without refreshing.
    if request.GET.get("messages") == "1":
        return JsonResponse({
            "ok": True,
            "items": [
                _member_chat_payload(message)
                for message in messages_qs
            ],
        })

    context = _member_chat_context(current_user)

    context.update({
        "user_role": "master_admin",
        "room": room,
        "other_user": other_user,
        "messages": messages_qs,
    })

    context["master_chat_active"] = True
    return render(
        request,
        "master_admin/master_admin_chat_room.html",
        context,
    )


@never_cache
@require_http_methods(["GET", "POST"])
def master_admin_group_chat(request, group_id):
    import logging

    from asgiref.sync import async_to_sync
    from channels.layers import get_channel_layer

    from .chat_receipts import broadcast_chat_activity    
    from .member_chat_actions import chat_action

    current_user = _master_chat_user(request)
    if current_user is None:
        return redirect("master_admin_login")
    group = get_object_or_404(Group, pk=group_id)

    membership = get_object_or_404(
        GroupMember,
        group=group,
        user=current_user,
    )

    can_manage = (
        group.created_by_id == current_user.pk
        or membership.role == "admin"
    )

    def member_state():
        current_membership = GroupMember.objects.filter(
            group=group,
            user=current_user,
        ).first()

        permitted = bool(
            current_membership
            and (
                group.created_by_id == current_user.pk
                or current_membership.role == "admin"
            )
        )

        entries = list(
            GroupMember.objects.filter(group=group)
            .select_related("user")
            .order_by("user__name", "pk")
        )

        available_users = []

        if permitted:
            available_users = [
                {
                    "id": person.pk,
                    "name": person.name,
                }
                for person in User.objects.exclude(
                    group_memberships__group=group
                ).order_by("name", "pk")
            ]

        return {
            "ok": True,
            "can_manage": permitted,
            "count": len(entries),
            "members": [
                {
                    "id": entry.user_id,
                    "name": entry.user.name,
                    "role": entry.role,
                    "is_creator": (
                        entry.user_id == group.created_by_id
                    ),
                    "is_self": (
                        entry.user_id == current_user.pk
                    ),
                    "profile": (
                        entry.user.profile_image.url
                        if entry.user.profile_image
                        else None
                    ),
                }
                for entry in entries
            ],
            "all_users": available_users,
        }

    def notify_members():
        layer = get_channel_layer()

        if layer:
            try:
                async_to_sync(layer.group_send)(
                    f"group_{group.pk}",
                    {"type": "chat_change"},
                )
            except Exception:
                logging.getLogger(__name__).exception(
                    "Group membership update delivery failed."
                )

        async_to_sync(broadcast_chat_activity)(
            True,
            group.pk,
        )

    if request.method == "POST":
        action = request.POST.get("action", "")

        if action == "create_group":
            return _member_chat_create_group(
                request,
                current_user,
                group_view_name="master_admin_group_chat",
            )

        if action in {
            "edit_message",
            "delete_for_me",
            "delete_for_everyone",
            "clear_chat",
        }:
            return chat_action(
                request,
                current_user,
                group,
                group=True,
            )

        if action not in {
            "add_member",
            "remove_member",
            "make_admin",
            "remove_admin",
        }:
            return _member_chat_error("Invalid action.")

        if not can_manage:
            return _member_chat_error(
                "Only the group creator or an admin "
                "can manage members.",
                403,
            )

        try:
            target_id = int(
                request.POST.get("user_id", "")
            )
        except (ValueError, TypeError):
            return _member_chat_error(
                "Select a valid member."
            )

        with transaction.atomic():
            # Serialize membership changes for this group.
            locked_group = (
                Group.objects.select_for_update()
                .get(pk=group.pk)
            )

            membership = get_object_or_404(
                GroupMember,
                group=group,
                user=current_user,
            )

            if (
                locked_group.created_by_id != current_user.pk
                and membership.role != "admin"
            ):
                return _member_chat_error(
                    "You no longer have permission "
                    "to manage this group.",
                    403,
                )

            if action == "add_member":
                target_user = get_object_or_404(
                    User,
                    pk=target_id,
                )

                _, created = (
                    GroupMember.objects.get_or_create(
                        group=group,
                        user=target_user,
                    )
                )

                notice = (
                    f"{target_user.name} added to the group."
                    if created
                    else f"{target_user.name} is already a member."
                )

            else:
                target = get_object_or_404(
                    GroupMember.objects.select_related("user"),
                    group=group,
                    user_id=target_id,
                )

                if target_id == locked_group.created_by_id:
                    return _member_chat_error(
                        "The group creator cannot be "
                        "removed or demoted."
                    )

                if target_id == current_user.pk:
                    return _member_chat_error(
                        "You cannot remove yourself or "
                        "change your own admin role here."
                    )

                if action == "remove_member":
                    member_count = (
                        GroupMember.objects.filter(
                            group=group
                        ).count()
                    )

                    if member_count <= 2:
                        return _member_chat_error(
                            "Keep at least two people "
                            "in the group."
                        )

                    name = target.user.name
                    target.delete()

                    notice = (
                        f"{name} removed from the group."
                    )

                else:
                    target.role = (
                        "admin"
                        if action == "make_admin"
                        else "member"
                    )

                    target.save(update_fields=["role"])

                    notice = (
                        f"{target.user.name} is now an admin."
                        if action == "make_admin"
                        else (
                            "Admin access removed for "
                            f"{target.user.name}."
                        )
                    )

            transaction.on_commit(notify_members)

        return JsonResponse({
            **member_state(),
            "message": notice,
        })

    if request.GET.get("members") == "1":
        return JsonResponse(member_state())

    messages_qs = (
        GroupMessage.objects.filter(group=group)
        .exclude(hidden_for=current_user)
        .select_related("sender")
        .order_by("timestamp", "pk")
    )

    if request.GET.get("messages") == "1":
        return JsonResponse({
            "ok": True,
            "items": [
                _member_chat_payload(message, True)
                for message in messages_qs
            ],
        })

    context = _member_chat_context(current_user)

    context.update({
        "user_role": "master_admin",
        "group": group,
        "messages": messages_qs,
        "members": (
            GroupMember.objects.filter(group=group)
            .select_related("user")
            .order_by("user__name", "pk")
        ),
        "all_users": (
            User.objects.exclude(
                group_memberships__group=group
            ).order_by("name", "pk")
        ),
        "can_manage_group": can_manage,
    })

    context["master_chat_active"] = True
    return render(
        request,
        "master_admin/master_admin_group_chat.html",
        context,
    )


@never_cache
@require_GET
def master_admin_chat_meta(request):
    from .chat_receipts import chat_meta

    if not _master_admin_allowed(request):
        return JsonResponse({
            "ok": False,
            "message": "Master Admin access required.",
        }, status=403)

    current_user = _master_chat_user(request)
    if current_user is None:
        return JsonResponse({
            "ok": False,
            "message": "Master Admin chat identity is unavailable.",
        }, status=401)

    return chat_meta(request, actor=current_user)


@never_cache
@require_POST
def master_admin_chat_read(request):
    from .chat_receipts import chat_mark_read

    if not _master_admin_allowed(request):
        return JsonResponse({
            "ok": False,
            "message": "Master Admin access required.",
        }, status=403)

    current_user = _master_chat_user(request)
    if current_user is None:
        return JsonResponse({
            "ok": False,
            "message": "Master Admin chat identity is unavailable.",
        }, status=401)

    return chat_mark_read(request, actor=current_user)



for _notice_view_name, _notice_config in _MASTER_NOTICE_VIEWS.items():
    _notice_original = globals()[_notice_view_name]
    if not getattr(_notice_original, "_master_operation_notice", False):
        globals()[_notice_view_name] = _master_operation_notice(_notice_original, *_notice_config)
# Notepad/profile already create the Master's own inbox entry in the supplied views.
